from contextlib import nullcontext
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from toolbox.runtime import ToolPaths, load_module

ROOT = Path(__file__).resolve().parents[2]


class ReviewDataGuardsTests(unittest.TestCase):
    def startup(self, day_coal, day_gas, rt_coal, rt_gas):
        module = load_module("review_fix_group", ROOT / "集团每日上传/upload_daily_report.py")
        def post(*args, **kwargs):
            groups = json.loads(kwargs["data"])["spotClearEleGroupList"]
            values = (day_coal, day_gas) if groups[0].startswith("DAY_AHEAD") else (rt_coal, rt_gas)
            return {"data": [dict(date="2026-10-26", dataItemGroup=group, unitCount=curve)
                             for group, curve in zip(groups, values)]}
        with patch.object(module, "parse_api_payload", side_effect=lambda response, _: response), patch.object(module, "log"):
            return module.fetch_startup_max(SimpleNamespace(request=SimpleNamespace(post=post)), "2026-10-26", "044")

    def test_startup_sums_by_hour_and_uses_only_complete_fallback(self):
        coal, gas = [0, 10] + [0] * 22, [1] * 24
        self.assertEqual(self.startup(coal, gas, [2] * 24, [3] * 24), 11)
        for invalid in (None, True, -1, 1.5, "nan", "inf", "-"):
            with self.subTest(invalid=invalid):
                broken = [invalid] + coal[1:]
                self.assertEqual(self.startup(broken, gas, [4] * 24, [3] * 24), 7)

    def test_startup_rejects_two_incomplete_curves(self):
        for invalid in ([None, 10] + [0] * 22, [], [1] * 23, [1] * 25):
            with self.subTest(invalid=invalid), self.assertRaisesRegex(RuntimeError, "完整燃煤/燃气"):
                self.startup(invalid, [1] * 24, [1] * 24, invalid)

    def test_group_excel_readback_uses_aligned_count_and_failed_retry_keeps_output(self):
        module = load_module("review_fix_group_integration", ROOT / "集团每日上传/upload_daily_report.py")
        broken = [False]
        def post(*args, **kwargs):
            groups = json.loads(kwargs["data"])["spotClearEleGroupList"]
            curves = [[None if broken[0] else 0, 10] + [0] * 22, [1] * 24]
            return {"data": [dict(date="2026-10-26", dataItemGroup=g, unitCount=c) for g, c in zip(groups, curves)]}
        with tempfile.TemporaryDirectory() as folder:
            energy = Path(folder) / "广东-能销数据-20261026.xlsx"
            energy.write_bytes(b"energy sentinel")
            template_bytes = module.BUILTIN_PROVINCE_TEMPLATE.read_bytes()
            day, uploads = module.prepare_upload_files([energy])
            with patch.object(module, "parse_api_payload", side_effect=lambda response, _: response), patch.object(
                module, "fetch_capacity_average", return_value=100
            ), patch.object(module, "log"):
                context = SimpleNamespace(request=SimpleNamespace(post=post))
                updated = module.update_province_excel_for_upload(context, day, uploads)
                target = next(item.path for item in updated if item.kind == "省内")
                book = load_workbook(target)
                template = load_workbook(module.BUILTIN_PROVINCE_TEMPLATE)
                try:
                    self.assertEqual(book.active["B2"].value, 11)
                    self.assertEqual(book.active["B2"].number_format, template.active["B2"].number_format)
                    self.assertEqual(book.active["A2"].value.date().isoformat(), day)
                finally:
                    book.close()
                    template.close()
                saved = target.read_bytes()
                broken[0] = True
                with self.assertRaisesRegex(RuntimeError, "完整燃煤/燃气"):
                    module.update_province_excel_for_upload(context, day, uploads)
                self.assertEqual(target.read_bytes(), saved)
            self.assertEqual(module.BUILTIN_PROVINCE_TEMPLATE.read_bytes(), template_bytes)
            self.assertEqual(energy.read_bytes(), b"energy sentinel")

    def export_online(self, module, values, folder):
        context = SimpleNamespace(close=Mock())
        overrides = dict(
            sync_playwright=lambda: nullcontext(None), launch_context=lambda *a, **k: context,
            ensure_login=lambda *a: None, log=lambda *a: None,
            get_companies=lambda *a: [dict(tenantId="test", name="海门电厂")],
            switch_company=lambda *a: None, get_units=lambda *a: [dict(id="u", name="#1")],
            is_renewable_company=lambda *a: False,
            fetch_fire_rows=lambda *a: [dict(unitId="u", onlineEle=values)],
        )
        with patch.multiple(module, **overrides), patch.object(module.time, "sleep"):
            try:
                return module.export_online_energy("2026-10-26", module.Config(output_dir=folder))
            finally:
                context.close.assert_called_once()

    def test_incomplete_online_export_preserves_existing_file(self):
        module = load_module("review_fix_online", ROOT / "上网电量抓取/export_online_energy.py")
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            target = folder / "上网电量汇总_2026-10-26.xlsx"
            target.write_bytes(b"existing user file")
            cases = [None, [], [1, 2, 3], [0] * 25]
            cases += [[value] + [0] * 23 for value in (None, "", "-", "nan", "inf", True)]
            for values in cases:
                with self.subTest(values=values), self.assertRaises(ValueError):
                    self.export_online(module, values, folder)
                self.assertEqual(target.read_bytes(), b"existing user file")

    def test_complete_online_export_readback_keeps_zero_and_negative_values(self):
        module = load_module("review_fix_online", ROOT / "上网电量抓取/export_online_energy.py")
        with tempfile.TemporaryDirectory() as folder:
            path = self.export_online(module, [0, "-1.25"] + [2] * 22, Path(folder))
            book = load_workbook(path)
            try:
                self.assertEqual([book.active[f"A{row}"].value for row in (3, 4, 26)], [0, -1.25, 2])
                self.assertEqual(book.active["A4"].number_format, "General")
            finally:
                book.close()

    def source(self, path, market="日前", unit=1):
        book = Workbook()
        sheet = book.active
        sheet.title = f"汕头#{unit}机组"
        sheet.append(["时刻", "电量", "出清电价"])
        for hour in range(24):
            sheet.append([f"{hour:02d}:00", unit * 10, 100])
        book.save(path)
        book.close()

    def test_duplicate_workbook_sources_stop_before_replacing_summary(self):
        module = load_module("review_fix_summary", ROOT / "电力工具脚本/电量汇总工具/summarize_511_excel.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            station = root / "汕头"
            station.mkdir()
            paths = [station / f"日前交易结果查询({day}).xlsx" for day in ("2026-10-25", "2026-10-26")]
            for path in paths:
                self.source(path)
            output = root / "汇总.xlsx"
            output.write_bytes(b"previous report")
            with self.assertRaisesRegex(ValueError, "来源冲突") as error:
                module.run_summary(root, output)
            for path in paths:
                self.assertIn(str(path), str(error.exception))
            self.assertEqual(output.read_bytes(), b"previous report")

    def test_distinct_units_and_market_dates_still_export(self):
        module = load_module("review_fix_summary", ROOT / "电力工具脚本/电量汇总工具/summarize_511_excel.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            station = root / "汕头"
            station.mkdir()
            self.source(station / "1号日前交易结果查询(2026-10-26).xlsx")
            self.source(station / "1号实时交易结果查询(2026-10-24).xlsx")
            self.source(station / "2号日前交易结果查询(2026-10-26).xlsx", unit=2)
            path = module.run_summary(root, root / "汇总.xlsx")
            book = load_workbook(path)
            try:
                sheet = book["汕头"]
                self.assertEqual([sheet[f"{col}4"].value for col in ("B", "D", "G", "I")], [10, 20, 10, None])
            finally:
                book.close()

    def test_batch_collisions_block_all_writes_including_case_and_unicode(self):
        module = load_module("review_fix_trade", ROOT / "电力工具脚本/电力交易分析工具/generate_electricity_report.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / "out"
            output.mkdir()
            sentinel = output / "report_分析报告.html"
            sentinel.write_text("previous report")
            for first, second in (("report", "report"), ("Report", "report"), ("é", "e\u0301")):
                with self.subTest(first=first), patch.object(module, "generate_report") as generate:
                    with self.assertRaisesRegex(ValueError, "文件名冲突"):
                        module.generate_batch([root / "day1" / f"{first}.xlsx", root / "day2" / f"{second}.xlsx"], output)
                    generate.assert_not_called()
            self.assertEqual(sentinel.read_text(), "previous report")

    def test_service_validation_and_standalone_gui_share_collision_guard(self):
        from toolbox import workflows
        module = load_module("review_fix_trade", ROOT / "电力工具脚本/电力交易分析工具/generate_electricity_report.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            files = [root / "a/report.xlsx", root / "b/report.xlsx"]
            for source in files:
                source.parent.mkdir()
                source.touch()
            with patch.object(workflows, "load_module", return_value=module), patch.object(module, "generate_report") as generate:
                with self.assertRaisesRegex(ValueError, "文件名冲突"):
                    workflows.validate(ToolPaths(ROOT), "trade-analysis", "run", dict(files=[str(p) for p in files], output_dir=str(root / "out")))
                app = SimpleNamespace(files=files, output_dir=SimpleNamespace(get=lambda: str(root / "out")),
                                      root=SimpleNamespace(after=Mock()), append_log=Mock(), _finish_batch=Mock())
                module.ReportApp._run_batch_worker(app)
                generate.assert_not_called()
                self.assertEqual(app.root.after.call_count, 2)
