"""WPS configuration format shared by desktop and Tk; no browser or GUI imports."""
from __future__ import annotations
import json
from datetime import datetime
from pathlib import Path
from typing import Any
CONFIG_EXPORT_SCHEMA = "wps_excel_to_kdocs_config_export"
CONFIG_EXPORT_VERSION = 1

def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [item.strip() for item in value if isinstance(item, str) and item.strip()]

def _region_config_from_dict(row: dict[str, Any]) -> dict[str, str]:
    return {
        "source_sheet": str(row.get("source_sheet", "")),
        "source_start": str(row.get("source_start", "")),
        "source_end": str(row.get("source_end", "")),
        "target_sheet": str(row.get("target_sheet", "")),
        "target_start": str(row.get("target_start", "")),
        "target_end": str(row.get("target_end", "")),
    }

def _document_config_from_dict(row: dict[str, Any]) -> dict[str, Any]:
    regions = [
        _region_config_from_dict(region)
        for region in row.get("regions", [])
        if isinstance(region, dict)
    ]
    name = str(row.get("name", "")).strip() or Path(str(row.get("local_file", ""))).stem or "未命名配置"
    source_type = str(row.get("source_type", "excel")).strip() or "excel"
    if source_type not in {"excel", "kdocs"}:
        source_type = "excel"
    return {
        "name": name,
        "source_type": source_type,
        "source_url": str(row.get("source_url", "")),
        "kdocs_url": str(row.get("kdocs_url", "")),
        "local_file": str(row.get("local_file", "")),
        "regions": regions,
    }

def normalize_runtime_config(data: dict[str, Any]) -> dict[str, Any]:
    documents = data.get("configs")
    if not isinstance(documents, list):
        documents = legacy_documents(data)
    config = {
        "recent_urls": _string_list(data.get("recent_urls", [])),
        "browser_mode": str(data.get("browser_mode", "persistent") or "persistent"),
        "cdp_url": str(data.get("cdp_url", "http://127.0.0.1:9222") or "http://127.0.0.1:9222"),
        "configs": [
            _document_config_from_dict(row)
            for row in documents
            if isinstance(row, dict)
        ],
    }
    titles = data.get("document_titles")
    if isinstance(titles, dict):
        titles = {url.strip(): title.strip() for url, title in titles.items()
                  if isinstance(url, str) and url.strip() and isinstance(title, str) and title.strip()}
        if titles:
            config["document_titles"] = titles
    return config


def legacy_documents(data: dict[str, Any]) -> list[dict[str, Any]]:
    """Migrate the original flat mappings in memory; never rewrite the source."""
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    grouped = {}
    for mapping in data.get("mappings", []):
        if not isinstance(mapping, dict):
            continue
        region = _region_config_from_dict(mapping)
        if "source_range" in mapping and ("source_start" not in mapping or "source_end" not in mapping):
            left, top, right, bottom = range_boundaries(mapping["source_range"])
            region["source_start"] = f"{get_column_letter(left)}{top}"
            region["source_end"] = f"{get_column_letter(right)}{bottom}"
        if not region["target_end"] and region["target_start"] and "source_range" in mapping:
            left, top, right, bottom = range_boundaries(mapping["source_range"])
            target_left, target_top, _, _ = range_boundaries(region["target_start"])
            region["target_end"] = f"{get_column_letter(target_left + right - left)}{target_top + bottom - top}"
        key = (mapping.get("kdocs_url") or data.get("kdocs_url", ""), mapping.get("local_file") or data.get("excel_path", ""),
               mapping.get("source_type", "excel"), mapping.get("source_url", ""))
        if key not in grouped:
            grouped[key] = dict(name=Path(key[1]).stem or f"配置 {len(grouped) + 1}", kdocs_url=key[0], local_file=key[1],
                                source_type=key[2], source_url=key[3], regions=[])
        grouped[key]["regions"].append(region)
    return list(grouped.values())

def build_config_export(runtime_config: dict[str, Any], exported_at: str | None = None) -> dict[str, Any]:
    export_time = exported_at or datetime.now().replace(microsecond=0).isoformat()
    return {
        "schema": CONFIG_EXPORT_SCHEMA,
        "version": CONFIG_EXPORT_VERSION,
        "exported_at": export_time,
        **normalize_runtime_config(runtime_config),
    }

def runtime_config_from_export_payload(payload: dict[str, Any]) -> dict[str, Any]:
    if payload.get("schema") != CONFIG_EXPORT_SCHEMA:
        raise ValueError("无法识别此配置文件，请选择由本工具导出的配置。")
    if payload.get("version") != CONFIG_EXPORT_VERSION:
        raise ValueError(f"暂不支持此配置文件版本：{payload.get('version')!r}")
    runtime_config = normalize_runtime_config(payload)
    if not runtime_config["configs"]:
        raise ValueError("配置文件中没有可导入的写入配置。")
    return runtime_config

def load_config_export(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"配置文件格式有误，无法读取：{exc}") from exc
    if not isinstance(payload, dict):
        raise ValueError("配置文件结构有误，请重新导出后再导入。")
    return runtime_config_from_export_payload(payload)

def save_config_export(path: Path, runtime_config: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(build_config_export(runtime_config), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
