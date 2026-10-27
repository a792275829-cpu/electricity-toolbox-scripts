from __future__ import annotations
import importlib

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from toolbox.runtime import load_module


class WpsBatchReadTests(unittest.TestCase):
    def test_multiple_ranges_use_one_open_and_keep_values_in_mapping_order(self) -> None:
        root = Path(__file__).resolve().parents[2]
        module = importlib.import_module('toolbox_shared.wps_service')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "数据"
            first["A2"] = datetime(2026, 9, 24, 8, 30)
            first["B2"] = 0.125
            first["A5"] = 7
            second = workbook.create_sheet("补充")
            second["C3"] = "完成"
            workbook.save(source)
            workbook.close()

            def mapping(sheet: str, start: str, end: str, target: str):
                return module.MappingRow(
                    kdocs_url="https://example.invalid/table",
                    local_file=str(source),
                    source_sheet=sheet,
                    source_start=start,
                    source_end=end,
                    target_sheet="目标",
                    target_start=target,
                    target_end=target if start == end else "B" + target[1:],
                )

            mappings = [
                mapping("数据", "A2", "B2", "A10"),
                mapping("补充", "C3", "C3", "A20"),
                mapping("数据", "A5", "B5", "A30"),
            ]
            with mock.patch.object(
                module, "load_excel_workbook", wraps=module.load_excel_workbook
            ) as opened:
                prepared = module.prepare_mappings(mappings)

            self.assertEqual(1, opened.call_count)
            self.assertEqual([1, 2, 3], [item.index for item in prepared])
            self.assertEqual([["2026-09-24 08:30:00", 0.125]], prepared[0].values)
            self.assertEqual([["完成"]], prepared[1].values)
            self.assertEqual([[7, ""]], prepared[2].values)
            self.assertEqual([1, 1, 1], [item.rows for item in prepared])

    def test_online_and_excel_sources_keep_original_mapping_order(self) -> None:
        root = Path(__file__).resolve().parents[2]
        module = importlib.import_module('toolbox_shared.wps_service')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = Workbook()
            workbook.active.title = "数据"
            workbook.active["A1"] = 42
            workbook.save(source)
            workbook.close()
            mappings = [
                module.MappingRow(
                    "https://example.invalid/target", str(source), "数据", "A1", "A1", "目标", "A1", "A1"
                ),
                module.MappingRow(
                    "https://example.invalid/target", "", "在线", "B2", "B2", "目标", "A2", "A2",
                    source_type="kdocs", source_url="https://example.invalid/source",
                ),
                module.MappingRow(
                    "https://example.invalid/target", str(source), "数据", "A1", "A1", "目标", "A3", "A3"
                ),
            ]
            with mock.patch.object(module, "read_online_ranges", return_value=[[["在线值"]]]) as online:
                prepared = module.prepare_mappings(mappings)

            self.assertEqual([[[42]], [["在线值"]], [[42]]], [item.values for item in prepared])
            online.assert_called_once()

    def test_online_regions_share_context_and_preserve_order(self) -> None:
        root = Path(__file__).resolve().parents[2]
        module = importlib.import_module('toolbox_shared.wps_service')
        requests = [
            ("https://example.invalid/a", "数据", "B2", 1, 1),
            ("https://example.invalid/b", "补充", "A1:B1", 1, 2),
            ("https://example.invalid/a", "数据", "C3", 1, 1),
        ]
        context = mock.Mock()
        with mock.patch.object(module, "sync_playwright"), \
             mock.patch.object(module, "open_browser_context", return_value=(None, context, True)) as opened, \
             mock.patch.object(module, "find_or_open_page") as find, \
             mock.patch.object(module, "evaluate_wps_operation", side_effect=[None, [0.125, "完成"], 42]):
            result = module.read_online_ranges(requests, "persistent", "", lambda _: None)
        self.assertEqual(result, [[[""]], [[0.125, "完成"]], [[42]]])
        opened.assert_called_once()
        self.assertEqual(find.call_args_list, [mock.call(context, request[0]) for request in requests])
        context.close.assert_called_once()

    def test_online_failure_closes_owned_context_and_stops_remaining_reads(self) -> None:
        root = Path(__file__).resolve().parents[2]
        module = importlib.import_module('toolbox_shared.wps_service')
        requests = [("https://example.invalid/a", "数据", "A1", 1, 1)] * 3
        for owned in (False, True):
            with self.subTest(owned=owned):
                context = mock.Mock()
                with mock.patch.object(module, "sync_playwright"), \
                     mock.patch.object(module, "open_browser_context", return_value=(None, context, owned)), \
                     mock.patch.object(module, "find_or_open_page"), \
                     mock.patch.object(module, "evaluate_wps_operation", side_effect=[42, RuntimeError("source failed")]) as evaluate:
                    with self.assertRaisesRegex(RuntimeError, "source failed"):
                        module.read_online_ranges(requests, "cdp", "", lambda _: None)
                self.assertEqual(evaluate.call_count, 2)
                self.assertEqual(context.close.call_count, int(owned))


if __name__ == "__main__":
    unittest.main()
