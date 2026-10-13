from copy import copy
from datetime import date, timedelta
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string

from toolbox.runtime import load_module


class MarketTableMonthFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.module = load_module(
            "market_table_month_format_test",
            Path(__file__).resolve().parents[2] / "市场表更新" / "update_market_table.py",
        )

    def test_all_mappings_use_last_calendar_day_across_month_separator(self):
        market = "2026年市场情况"
        cost = "2026年运行方式及成本"
        mappings = [
            ("write_day_ahead_load_values", market, "B", 20, False),
            ("write_day_ahead_market_values", market, "V", 12, False),
            ("write_actual_load_values", market, "AH", 20, False),
            ("write_realtime_market_values", market, "BB", 4, False),
            ("write_generation_settlement_values", market, "BF", 12, True),
            ("write_user_settlement_values", market, "BR", 7, True),
            ("write_unit_cost_price_values", cost, "N", 11, False),
            ("write_market_board_online_unit_values", cost, ("AA", "AE"), 2, False),
            ("write_spot_clearing_unit_count_values", cost, "AB", 2, False),
            ("copy_previous_day_unit_operation_mode", cost, "B", 12, False),
        ]
        for target_date in (date(2026, 9, 1), date(2026, 3, 1), date(2026, 1, 1)):
            for name, sheet_name, start, count, overwrite in mappings:
                with self.subTest(date=target_date, mapping=name), tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "market.xlsx"
                    wb = Workbook()
                    ws = wb.active
                    ws.title = sheet_name
                    ws["A4"] = target_date - timedelta(days=2)
                    ws["A5"] = "=A4+1"
                    ws["A6"] = f"{target_date.month}月"
                    ws["A7"] = "=A5+1"
                    columns = (
                        [column_index_from_string(c) for c in start]
                        if isinstance(start, tuple)
                        else list(range(column_index_from_string(start), column_index_from_string(start) + count))
                    )
                    expected_styles = []
                    for index, column in enumerate(columns):
                        source = ws.cell(5, column, index + 10)
                        source.font = Font(name="Arial", size=11, color="000000")
                        source.fill = PatternFill("solid", fgColor="FFFFFF")
                        source.border = Border(bottom=Side(style="thin", color="000000"))
                        source.alignment = Alignment(horizontal="right")
                        source.number_format = "0.00%" if index % 2 else "0.0000"
                        expected_styles.append(copy(source._style))
                        separator = ws.cell(6, column, 0)
                        separator.fill = PatternFill("solid", fgColor="FFFF00")
                        separator.font = Font(bold=True)
                        separator.number_format = "0"
                    ws.cell(7, columns[0], 0)
                    ws.cell(7, columns[0]).fill = PatternFill("solid", fgColor="00FF00")
                    kept_style = copy(ws.cell(7, columns[0])._style)
                    wb.save(path)
                    wb.close()

                    values = list(range(10, 10 + count))
                    writer = getattr(self.module, name)
                    summary = (
                        writer(path, target_date)
                        if name == "copy_previous_day_unit_operation_mode"
                        else writer(path, target_date, values)
                    )
                    updated = load_workbook(path)
                    try:
                        sheet = updated[sheet_name]
                        for index, column in enumerate(columns):
                            skipped = index == 0 and not overwrite
                            cell = sheet.cell(7, column)
                            self.assertEqual(cell.value, 0 if skipped else values[index])
                            self.assertEqual(cell._style, kept_style if skipped else expected_styles[index])
                            self.assertEqual(sheet.cell(6, column).fill.fgColor.rgb, "00FFFF00")
                            self.assertEqual(sheet.cell(6, column).value, 0)
                        self.assertEqual(sheet["A7"].value, "=A5+1")
                        self.assertEqual(summary["count"], count if overwrite else count - 1)
                        self.assertEqual(summary["skipped"], 0 if overwrite else 1)
                    finally:
                        updated.close()

    def test_missing_previous_month_keeps_target_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A4"] = "9月"
            ws["A5"] = date(2026, 9, 1)
            for column in (28, 29):
                ws.cell(4, column).fill = PatternFill("solid", fgColor="FFFF00")
                ws.cell(5, column).number_format = "0.00"
            wb.save(path)
            wb.close()
            self.module.write_spot_clearing_unit_count_values(path, date(2026, 9, 1), [121, 46])
            updated = load_workbook(path)
            try:
                for cell in (updated.active["AB5"], updated.active["AC5"]):
                    self.assertEqual(cell.number_format, "0.00")
                    self.assertNotEqual(cell.fill.fgColor.rgb, "00FFFF00")
            finally:
                updated.close()
