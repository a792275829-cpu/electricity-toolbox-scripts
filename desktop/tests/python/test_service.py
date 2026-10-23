"""Exercise real JSON transport and owned business processes, without live accounts."""
from __future__ import annotations
import json
import os
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[3]
sys.path[:0] = [str(ROOT), str(ROOT / "电力工具箱")]
from toolbox.service import ToolboxService
from toolbox.runtime import ToolPaths
from toolbox.workflows import validate, plan


class HeadlessProtocolTests(unittest.TestCase):
    def test_real_transport_has_no_tk_and_recovers_from_invalid_json(self):
        script = '''import sys, json
from pathlib import Path
sys.path[:0] = [str(Path.cwd()), str(Path.cwd() / "电力工具箱"), str(Path.cwd() / "desktop/python")]
from service_bridge import serve
assert "tkinter" not in sys.modules
serve(Path.cwd())
assert "tkinter" not in sys.modules
'''
        proc = subprocess.run([sys.executable, "-B", "-c", script], cwd=ROOT,
            input='bad\n{"id":1,"method":"ping"}\n{"id":2,"method":"catalog"}\n{"id":3,"method":"shutdown"}\n',
            text=True, capture_output=True, timeout=15)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        rows = [json.loads(v) for v in proc.stdout.splitlines()]
        replies = {row['id']: row for row in rows if 'id' in row}
        self.assertIn('error', replies[None])
        self.assertEqual(replies[1]['result']['protocol'], 3)
        self.assertEqual(len(replies[2]['result']), 15)
        self.assertEqual(rows[-1]['event']['kind'], 'service-closed')


class ServiceTaskTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        script = self.root / '上网电量抓取/export_online_energy.py'
        script.parent.mkdir()
        script.write_text('import sys\nprint("sample-output",flush=True)\n')
        self.service = ToolboxService(self.root)
        self.params = dict(tool_id='online-energy', action='run', request_id='one', values=dict(run_date='2026-10-20'))
    def tearDown(self):
        self.service.registry.shutdown(timeout=5)
        self.temp.cleanup()
    def finish(self, task_id):
        deadline = time.monotonic() + 10
        while self.service.registry.has_running_tasks() and time.monotonic() < deadline:
            self.service.collect_events(); time.sleep(.02)
        self.service.collect_events()
        return self.service.registry.snapshot(task_id)
    def test_real_worker_streams_logs_and_deduplicates_request(self):
        task_id = self.service.call('start', self.params)
        self.assertEqual(self.service.call('start', self.params), task_id)
        snapshot = self.finish(task_id)
        self.assertEqual(snapshot.state.value, 'succeeded', snapshot.error)
        self.assertTrue(snapshot.result)
        batch = self.service.call('events', dict(after=0))
        self.assertTrue(any('sample-output' in e['message'] for e in batch['events']))
        self.assertFalse(self.service.call('events', dict(after=batch['cursor']))['events'])
        self.assertEqual(self.service.call('start', self.params), task_id)
        with self.assertRaisesRegex(ValueError, '不同参数'):
            self.service.call('start', {**self.params, 'values': {'run_date': '2026-10-21'}})
    def test_business_failure_retains_first_exception_and_exit_code(self):
        self.service.paths.online_energy.write_text('raise RuntimeError("first-business-failure")\n')
        task_id = self.service.call('start', self.params)
        snapshot = self.finish(task_id)
        self.assertEqual(snapshot.state.value, 'failed')
        self.assertIn('退出码 1', str(snapshot.error))
        self.assertIn('first-business-failure', str(snapshot.error))
        text = ''.join(e['message'] for e in self.service.call('events', {})['events'])
        self.assertIn('first-business-failure', text)
    def test_noninterruptible_write_blocks_close(self):
        release = threading.Event()
        self.service.registry.start_callable('write', lambda token, emit: release.wait(5), cancellable=False)
        try:
            with self.assertRaisesRegex(ValueError, '不可安全中断'):
                self.service.call('shutdown', {'cancel': True})
            self.assertFalse(self.service.closed.is_set())
        finally:
            release.set()
    def test_cancel_terminates_worker_and_owned_child(self):
        self.service.paths.online_energy.write_text(
            'import os,time\n'
            'print(f"started:{os.getpid()}:{os.getppid()}",flush=True)\n'
            'time.sleep(60)\n')
        task_id = self.service.call('start', self.params)
        deadline = time.monotonic() + 10
        owned_pids = None
        while time.monotonic() < deadline:
            events = self.service.call('events', {})['events']
            ready = next((e['message'].strip() for e in events
                          if e['message'].startswith('started:')), None)
            if ready:
                owned_pids = tuple(int(pid) for pid in ready.split(':')[1:])
                break
            snapshot = self.service.registry.snapshot(task_id)
            self.assertIn(snapshot.state.value, {'created', 'running'},
                          f'{snapshot!r}; events={events!r}')
            time.sleep(.01)
        self.assertIsNotNone(owned_pids, f'worker not ready; events={events!r}')
        self.assertTrue(self.service.call('cancel', {'task_id': task_id}))
        snapshot = self.finish(task_id)
        self.assertEqual(snapshot.state.value, 'cancelled',
                         f'{snapshot!r}; events={self.service.call("events", {})!r}')
        self.assertFalse(self.service.registry.has_running_tasks())
        for pid in owned_pids:
            deadline = time.monotonic() + 3
            while time.monotonic() < deadline:
                try:
                    os.kill(pid, 0)
                except ProcessLookupError:
                    break
                time.sleep(.01)
            else:
                self.fail(f'owned process {pid} survived cancellation')
    def test_cancel_before_worker_launch_does_not_run_business_script(self):
        from toolbox.tasks import ProcessRunner

        marker = self.root / 'unexpected-business-start'
        self.service.paths.online_energy.write_text(
            f'from pathlib import Path\nPath({str(marker)!r}).touch()\n')
        entered, release = threading.Event(), threading.Event()
        original_run = ProcessRunner.run
        def delayed_run(runner, *args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise TimeoutError('test did not release worker launch')
            return original_run(runner, *args, **kwargs)

        with patch.object(ProcessRunner, 'run', delayed_run):
            task_id = self.service.call('start', self.params)
            try:
                self.assertTrue(entered.wait(2))
                self.assertTrue(self.service.call('cancel', {'task_id': task_id}))
            finally:
                release.set()
            snapshot = self.finish(task_id)
        self.assertEqual(snapshot.state.value, 'cancelled', repr(snapshot))
        self.assertFalse(marker.exists())
    def test_unknown_action_and_shell_input_are_rejected(self):
        with self.assertRaises(ValueError):
            self.service.call('start', {**self.params, 'action': 'shell'})
        with self.assertRaises(ValueError):
            self.service.call('start', {**self.params, 'values': {'run_date':'2026-10-20', 'command':'echo wrong'}})
    def test_credentials_api_never_returns_password(self):
        with patch('toolbox.service.CredentialStore') as store:
            store.return_value.load.return_value = dict(username='sample', password='private')
            result = self.service.call('credentials.get', {})
            self.assertEqual(result, dict(username='sample', configured=True))
            self.assertNotIn('private', json.dumps(result))
    def test_nonfinite_threshold_rejected_before_task_creation(self):
        source = self.root / 'source.xlsx'; source.touch()
        topology = self.root / 'topology.md'; topology.touch()
        with self.assertRaises(ValueError):
            self.service.call('prepare_task', dict(tool_id='section-analysis', values=dict(files=[str(source)], topology_path=str(topology), output_root=str(self.root/'output'), threshold='nan')))
        self.assertFalse(self.service.registry.snapshots())
    def test_event_history_is_bounded_and_reports_a_gap(self):
        from toolbox.tasks import TaskEvent, TaskEventKind
        for index in range(3000):
            self.service.registry.events.publish(TaskEvent('sample', TaskEventKind.LOG, message=str(index)))
        self.assertEqual(self.service.registry.events.pending_count, 2000)
        batch = self.service.call('events', dict(after=0))
        self.assertTrue(batch['truncated'])
        self.assertEqual(len(batch['events']), 500)
    def test_ui_retry_after_an_uncertain_response_reuses_the_accepted_task(self):
        from toolbox.ui_tasks import RemoteTaskRegistry
        from toolbox.tasks import CancellationToken
        service = self.service
        class UncertainTransport:
            first = True
            def call(self, method, params=None):
                result = service.call(method, params or {})
                if method == 'start' and self.first:
                    self.first = False
                    raise TimeoutError('response lost after acceptance')
                return result
        registry = RemoteTaskRegistry(UncertainTransport())
        values = dict(run_date='2026-10-20')
        with self.assertRaises(TimeoutError):
            registry.execute_remote('online-energy', 'run', values, CancellationToken(), lambda _: None)
        task_id = service.registry.snapshots()[0].task_id
        self.finish(task_id)
        result = registry.execute_remote('online-energy', 'run', values, CancellationToken(), lambda _: None)
        self.assertTrue(result)
        self.assertEqual(len(service.registry.snapshots()), 1)
    def test_wps_inspection_forwards_the_original_range_requests(self):
        from toolbox.workflows import execute_wps
        from unittest.mock import Mock
        module = Mock()
        module.read_online_ranges.return_value = [[[1, 2]]]
        requests = [['https://example.invalid/doc', 'Sheet1', 'A1:B1', 1, 2]]
        with patch('toolbox.workflows.load_module', return_value=module):
            result = execute_wps(ToolPaths(ROOT), 'inspect', dict(operation=json.dumps(dict(method='read_online_ranges', arguments=dict(requests=requests, browser_mode='cdp', cdp_url='http://127.0.0.1:9222')))), lambda _: None)
        self.assertEqual(result, [[[1, 2]]])
        self.assertEqual(module.read_online_ranges.call_args.kwargs['requests'], requests)
    def test_real_summary_workbook_values_and_formats_are_read_back(self):
        from openpyxl import Workbook, load_workbook
        source = self.root / 'review'
        station = source / '海门'
        station.mkdir(parents=True)
        for label, day, price in [('日前', '2026-10-20', '711.47'), ('实时', '2026-10-18', '625.33')]:
            book = Workbook()
            sheet = book.active
            sheet.title = '海门#1机组'
            sheet.append(['时刻', '电量(MWh)', '标签', '出清电价', '结算电价'])
            for hour in range(24):
                sheet.append([f'{hour:02d}:00', f'{hour + 1}.125', '', '1.00', price])
            book.save(station / f'1号机{label}交易结果查询({day}).xlsx')
            book.close()
        self.service = ToolboxService(ROOT)
        output = source / '汇总.xlsx'
        task_id = self.service.call('start', dict(tool_id='summary-upload', action='run',
            request_id='summary-sample', confirmed=True, values=dict(input_dir=str(source), output_file=str(output), mode='summary')))
        snapshot = self.finish(task_id)
        self.assertEqual(snapshot.state.value, 'succeeded', snapshot.error)
        book = load_workbook(output, data_only=True)
        try:
            sheet = book['海门']
            self.assertEqual(sheet['B4'].value, 1.125)
            self.assertEqual(sheet['C4'].value, 711.47)
            realtime_energy = next(cell.column for cell in sheet[3] if cell.value == '实时电量\n（MWh）')
            realtime_price = next(cell.column for cell in sheet[3] if cell.value == '实时电价\n（元/MWh）')
            self.assertEqual(sheet.cell(4, realtime_energy).value, 1.125)
            self.assertEqual(sheet.cell(4, realtime_price).value, 625.33)
            self.assertEqual(sheet['B4'].number_format, '0.000')
            self.assertEqual(sheet['C4'].number_format, '0.00')
        finally:
            book.close()


class CommandContracts(unittest.TestCase):
    def test_transaction_optional_date_and_backup_flags_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory)/'book.xlsx'; workbook.touch()
            paths = ToolPaths(ROOT)
            values = validate(paths, 'transaction-info', 'run', dict(workbook=str(workbook), start_date='', end_date='2026-10-20', create_backup=False))
            command, _ = plan(paths, 'transaction-info', 'run', values)[0]
            self.assertIn('--no-backup', command)
            self.assertNotIn('--start-date', command)
            self.assertEqual(command[command.index('--end-date')+1], '2026-10-20')
    def test_only_selected_spot_workbooks_are_required(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory)/'book.xlsx'; workbook.touch()
            paths = ToolPaths(ROOT)
            values = validate(paths, 'spot-price', 'run', dict(guihu_workbook=str(workbook), guihu_selected=True, tuolian_selected=False, haifeng_selected=False, overwrite=False, create_backup=True))
            command, _ = plan(paths, 'spot-price', 'run', values)[0]
            self.assertEqual(command[-1], '--skip-existing')
            self.assertNotIn('--tuolian-workbook', command)
            self.assertNotIn('--no-backup', command)

if __name__ == '__main__': unittest.main()
