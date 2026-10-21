"""Measure local update/report stages against an explicit pre-change snapshot."""
import argparse
from dataclasses import asdict
from datetime import date
import hashlib
import json
import os
from pathlib import Path
import statistics
import sys
import tempfile
import time
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / '电力工具箱'), str(ROOT / '广东电价预测'), str(ROOT / '电力工具脚本/电力交易分析工具')]
from toolbox.runtime import load_module


def digest(path):
    book = load_workbook(path, read_only=True, data_only=False)
    result = hashlib.sha256()
    formats = {}
    try:
        for sheet in book:
            result.update(repr((sheet.title, sheet.max_row, sheet.max_column)).encode())
            for row in sheet:
                for cell in row:
                    style_id = getattr(cell, '_style_id', 0)
                    if style_id not in formats:
                        formats[style_id] = repr((str(cell.font),str(cell.fill),str(cell.border),str(cell.alignment),str(cell.protection)))
                    result.update(repr((cell.value, cell.data_type, cell.number_format, formats[style_id])).encode())
    finally:
        book.close()
    return result.hexdigest()


def run(args):
    files = {'forecast':'广东电价预测/update_gdfire_data.py',
             'report':'电力工具脚本/电力交易分析工具/generate_electricity_report.py',
             'deviation':'电力工具脚本/电力交易分析工具/find_output_range_deviation_days.py',
             'benefit':'功率调整收益统计表更新/update_power_benefit.py',
             'transaction':'交易系统信息更新/update_transaction_info.py'}
    modules = {mode:{key:load_module(f'bench_update_{mode}_{key}', base/relative) for key,relative in files.items()}
               for mode,base in [('before',args.baseline_dir),('after',ROOT)]}
    modules['before']['deviation'].ReportBuilder = modules['before']['report'].ReportBuilder
    refs = {}; times = {}; counts = {}
    def measure(mode, stage, function, verify=None):
        start=time.perf_counter(); value=function(); elapsed=time.perf_counter()-start
        times.setdefault(stage,{}).setdefault(mode,[]).append(elapsed)
        print(f'{mode} {stage}: {elapsed:.4f}s',flush=True)
        if verify: verify(value)
        return value
    def equal(key,value):
        if key not in refs: refs[key]=value
        elif value!=refs[key]: raise AssertionError(f'{key} differs')
    def frame_equal(value):
        if 'frame' not in refs:refs['frame']=value
        else:assert_frame_equal(refs['frame'],value,check_exact=True)
    with tempfile.TemporaryDirectory(prefix='toolbox-update-bench-') as folder:
        folder=Path(folder)
        for repeat in range(args.repeats):
            for mode in (('before','after') if repeat%2==0 else ('after','before')):
                m=modules[mode]
                os.environ['TOOLBOX_XLSX_BACKEND']='python' if mode=='before' else 'rust'
                frame=measure(mode,'forecast_read',lambda:m['forecast'].read_data_workbook(args.forecast),frame_equal)
                counts['forecast_shape']=list(frame.shape)
                output=folder/f'{mode}.xlsx'
                measure(mode,'forecast_write_with_layout',lambda:m['forecast'].write_large_xlsx(frame,output,'整合数据',format_source=args.forecast))
                if repeat==0:
                    equal('forecast_output',digest(output))
                report=folder/'report.html'
                measure(mode,'trade_report',lambda:m['report'].ReportBuilder(args.report,report).build(),lambda p:equal('report_html',p.read_bytes()))
                measure(mode,'trade_deviations',lambda:m['deviation'].workbook_hour_deviations(args.report),lambda rows:equal('deviations',[asdict(row) for row in rows]))
                measure(mode,'benefit_inspect',lambda:m['benefit'].inspect_workbook(args.benefit,end_date=date(2026,10,20)),lambda plan:equal('benefit',asdict(plan)))
                measure(mode,'transaction_tail',lambda:m['transaction'].inspect_workbook_tail(args.transaction),lambda tail:equal('transaction',asdict(tail)))
    result={'repeats':args.repeats,'counts':counts,'inputs':{k:getattr(args,k).name for k in ['forecast','report','benefit','transaction']},'stages':{},'parity':'exact dataframe values/dtypes; full exported cell value/type/format/style digest; report HTML bytes; deviation rows; benefit target plan; transaction tail', 'notes':'Local stages only; imports and verification excluded; process startup, JSON transport, temporary spool and compression included; alternating order.'}
    for stage,modes in times.items():
        result['stages'][stage]={mode:{'seconds':values,'median':statistics.median(values)} for mode,values in modes.items()}
        result['stages'][stage]['speedup']=statistics.median(modes['before'])/statistics.median(modes['after'])
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for key in ['baseline-dir','forecast','report','benefit','transaction','output']:parser.add_argument('--'+key,type=Path,required=True)
    parser.add_argument('--repeats',type=int,default=3)
    run(parser.parse_args())
