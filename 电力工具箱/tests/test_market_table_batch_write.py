from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from toolbox.runtime import load_module


class MarketTableBatchWriteTests(unittest.TestCase):
    def test_external_save_during_session_is_preserved(self) -> None:
        module = load_module("market_table_external_save_test",
            Path(__file__).resolve().parents[2] / "市场表更新/update_market_table.py")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "market.xlsx"
            book = Workbook()
            book.active["A1"] = "原值"
            book.save(path)
            book.close()
            session = module.WorkbookWriteSession(path)
            updated = session.get_workbook()
            updated.active["B1"] = 42
            session.record(updated.active["B1"])
            external = load_workbook(path)
            external.active["A1"] = "外部保存"
            external.save(path)
            external.close()
            original = path.read_bytes()
            try:
                with self.assertRaisesRegex(RuntimeError, "工作簿被修改"):
                    session.commit()
            finally:
                updated.close()
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [path])

    def test_readback_format_mismatch_preserves_original(self) -> None:
        module = load_module(
            "market_table_stream_verification_test",
            Path(__file__).resolve().parents[2] / "市场表更新/update_market_table.py",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "market.xlsx"
            workbook = Workbook()
            workbook.active.title = module.MARKET_SHEET_NAME
            workbook.active["V2"] = 0.125
            workbook.active["V2"].number_format = "0.00%"
            workbook.save(path)
            original = path.read_bytes()
            session = module.WorkbookWriteSession(path)
            updated = session.get_workbook()
            updated[module.MARKET_SHEET_NAME]["V2"] = 0.25
            session.record(updated[module.MARKET_SHEET_NAME]["V2"])
            save = updated.save

            def corrupt_format(target):
                save(target)
                changed = load_workbook(target)
                changed[module.MARKET_SHEET_NAME]["V2"].number_format = "0.000"
                changed.save(target)
                changed.close()

            try:
                with mock.patch.object(updated, "save", side_effect=corrupt_format):
                    with self.assertRaisesRegex(RuntimeError, "市场表保存后回读不一致"):
                        session.commit()
            finally:
                updated.close()
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(path.parent.glob(".market.*.xlsx")), [])

    def test_numeric_write_ignores_date_format_on_previous_blank_row(self) -> None:
        module = load_module(
            "market_table_numeric_format_test",
            Path(__file__).resolve().parents[2] / "市场表更新" / "update_market_table.py",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "market.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = module.COST_SHEET_NAME
            sheet["A2"] = date(2026, 9, 23)
            sheet["A3"] = date(2026, 9, 24)
            sheet["AA2"].number_format = "[$-409]yyyy/mm/dd;@"
            sheet["AA3"].number_format = "0_ "
            workbook.save(path)

            with module.workbook_write_session(path):
                module.write_market_board_online_unit_values(
                    path, date(2026, 9, 24), [504, 9000]
                )

            saved = load_workbook(path)
            try:
                self.assertEqual(saved[module.COST_SHEET_NAME]["AA3"].value, 504)
                self.assertEqual(saved[module.COST_SHEET_NAME]["AA3"].number_format, "0_ ")
            finally:
                saved.close()

    def test_failed_update_keeps_original_workbook(self) -> None:
        module = load_module(
            "market_table_batch_failure_test",
            Path(__file__).resolve().parents[2] / "市场表更新" / "update_market_table.py",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "market.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = module.MARKET_SHEET_NAME
            sheet["A2"] = date(2026, 7, 2)
            workbook.save(path)
            original_bytes = path.read_bytes()

            with self.assertRaisesRegex(RuntimeError, "模拟后续接口失败"):
                with module.workbook_write_session(path):
                    module.write_day_ahead_market_values(
                        path, date(2026, 7, 2), list(range(12))
                    )
                    raise RuntimeError("模拟后续接口失败")

            self.assertEqual(path.read_bytes(), original_bytes)

    def test_one_workbook_load_and_save_preserves_values_and_formats(self) -> None:
        module = load_module(
            "market_table_batch_write_test",
            Path(__file__).resolve().parents[2] / "市场表更新" / "update_market_table.py",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "market.xlsx"
            workbook = Workbook()
            market = workbook.active
            market.title = module.MARKET_SHEET_NAME
            cost = workbook.create_sheet(module.COST_SHEET_NAME)
            for sheet in (market, cost):
                sheet["A2"] = date(2026, 7, 1)
                sheet["A3"] = date(2026, 7, 2)
            market["V2"].number_format = "0.00%"
            cost["N2"].number_format = "0.00"
            cost["B2"] = "运行"
            cost["B3"] = "已有"
            workbook.save(path)

            original_load = module.load_workbook_checked
            original_loads = []

            def counted_load(workbook_path, **kwargs):
                if Path(workbook_path).resolve() == path.resolve():
                    original_loads.append(kwargs)
                return original_load(workbook_path, **kwargs)

            with mock.patch.object(module, "load_workbook_checked", side_effect=counted_load):
                with module.workbook_write_session(path):
                    market_summary = module.write_day_ahead_market_values(
                        path, date(2026, 7, 2), list(range(12))
                    )
                    cost_summary = module.write_unit_cost_price_values(
                        path, date(2026, 7, 2), list(range(11))
                    )
                    operation_summary = module.copy_previous_day_unit_operation_mode(
                        path, date(2026, 7, 2)
                    )

            self.assertEqual(len(original_loads), 1)
            self.assertEqual(market_summary["range"], "V3:AG3")
            self.assertEqual(cost_summary["range"], "N3:X3")
            self.assertEqual(operation_summary["skipped"], 1)
            saved = load_workbook(path, data_only=False)
            try:
                self.assertEqual(saved[module.MARKET_SHEET_NAME]["V3"].value, 0)
                self.assertEqual(saved[module.MARKET_SHEET_NAME]["V3"].number_format, "0.00%")
                self.assertEqual(saved[module.COST_SHEET_NAME]["N3"].value, 0)
                self.assertEqual(saved[module.COST_SHEET_NAME]["N3"].number_format, "0.00")
                self.assertEqual(saved[module.COST_SHEET_NAME]["B3"].value, "已有")
                self.assertEqual(saved[module.COST_SHEET_NAME]["C3"].value, None)
            finally:
                saved.close()


if __name__ == "__main__":
    unittest.main()
