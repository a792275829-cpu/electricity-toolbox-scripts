from contextlib import nullcontext
from datetime import date
import io
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from PIL import Image

from toolbox.runtime import load_module

ROOT = Path(__file__).resolve().parents[2]


class PerformanceBoundaryTests(unittest.TestCase):
    def test_date_formats_use_one_scan_and_keep_styled_blank_inheritance(self):
        module = load_module("six_spot", ROOT / "现货价格更新/update_spot_prices.py")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "dates.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["A2"] = date(2026, 10, 1)
            sheet["A2"].number_format = "yyyy/mm/dd"
            from openpyxl.styles import Font
            sheet["A4"].font = Font(bold=True)  # Styled General resets inheritance.
            sheet["A7"].number_format = "m/d/yy"
            sheet["B10"] = "tail"
            book.save(path)
            book.close()
            book = load_workbook(path, read_only=True)
            try:
                sheet = book.active
                with patch.object(sheet, "cell", side_effect=AssertionError("random access")), patch.object(
                    sheet, "iter_rows", wraps=sheet.iter_rows
                ) as scan:
                    result = module.effective_number_formats(sheet, [1, 3, 5, 8, 31], 1)
                self.assertEqual(result, {1: "General", 3: "yyyy/mm/dd", 5: "General", 8: "m/d/yy", 31: "m/d/yy"})
                scan.assert_called_once()
            finally:
                book.close()

    def test_early_summaries_share_closed_value_snapshot_and_invalidate_on_save(self):
        module = load_module("six_early", ROOT / "每日生产经营情况汇报自动生成工具/scripts/generate_red_marked_report.py")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "快报（日前）"
            for row in range(1, 41):
                sheet.append([None] * 5 + [row, row + .25, -row])
            for row, label in [(20, "海上风电"), (25, "鮀莲"), (30, "归湖")]:
                sheet.cell(row, 1, label)
            basic = book.create_sheet("基础信息")
            basic["AA64"], basic["AC64"] = 12.5, 0
            book.save(path)
            original = path.read_bytes()
            from openpyxl.worksheet._read_only import ReadOnlyWorksheet
            opened = []
            def load(*args, **kwargs):
                value = load_workbook(*args, **kwargs)
                opened.append(value)
                return value
            with patch.object(module, "load_workbook", side_effect=load) as loader, patch.object(
                ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")
            ), module.clearing_input_cache():
                ahead = module.load_day_ahead_clearing_summary("2026-10-24", path)
                prices = module.load_online_price_summary("2026-10-23", path)
                self.assertEqual(ahead["totalEnergy"], 19.5)
                self.assertEqual(ahead["companies"]["清远"]["energy"], 12.5)
                self.assertEqual(prices["companies"]["归湖"], -30)
                self.assertEqual(loader.call_count, 1)
                self.assertIsNone(opened[0]._archive.fp)
                sheet["G7"] = 888
                book.save(path)
                self.assertEqual(module.load_online_price_summary("2026-10-23", path)["averagePrice"], 888)
                self.assertEqual(loader.call_count, 2)
            self.assertIsNone(module._clearing_inputs.get())
            book.close()
            self.assertNotEqual(path.read_bytes(), original)  # Only our explicit save changed it.

    def test_sparse_existing_and_omitted_cells_keep_unrelated_xml(self):
        module = load_module("six_benefit", ROOT / "功率调整收益统计表更新/update_power_benefit.py")
        original = b'<worksheet><sheetData><row r="2"><c r="A2"><v>1</v></c><c r="C2" s="3"/><c r="N2"><v>9</v></c><c r="Q2"><f>C2+J2</f><v>123</v></c></row></sheetData></worksheet>'
        result = module.patch_sparse_sheet(original, {"C2": -2.5, "J2": 0})
        self.assertEqual(result, original.replace(b'<c r="C2" s="3"/>', b'<c r="C2" s="3"><v>-2.5</v></c>').replace(
            b'<c r="N2">', b'<c r="J2"><v>0</v></c><c r="N2">'
        ))

    def test_pdf_qualified_embedded_skips_crop_and_keeps_fallback_boundaries(self):
        sys.path.insert(0, str(ROOT / "广东运行日报PDF解析"))
        self.addCleanup(sys.path.remove, str(ROOT / "广东运行日报PDF解析"))
        from gd_pdf_tool import parser
        image = Image.new("RGB", (8, 8), "red")
        data = io.BytesIO()
        image.save(data, format="PNG")
        page = SimpleNamespace(width=600, height=800, images=[dict(
            width=400, top=100, stream=SimpleNamespace(get_data=lambda: data.getvalue()))])
        crop = Image.new("RGB", (8, 8), "blue")
        for score, fallback_score, expected in [(50, 60, "embedded-image"), (49, 60, "rendered-page-crop"), (10, 4, "embedded-image")]:
            with self.subTest(score=score), patch.object(parser, "_open_pdf", return_value=nullcontext(SimpleNamespace(pages=[None, None, page]))), patch.object(
                parser, "render_quote_chart_crop", return_value=crop
            ) as render, patch.object(parser, "_score_quote_chart_candidate", side_effect=lambda img: fallback_score if img is crop else score):
                selected, kind = parser.extract_quote_chart_image(Path("unused.pdf"))
                self.assertEqual(kind, expected)
                self.assertEqual(selected.tobytes(), (image if expected == "embedded-image" else crop).tobytes())
                self.assertEqual(render.call_count, int(score < 50))
        with patch.object(parser, "_open_pdf", side_effect=OSError("broken PDF")), patch.object(parser, "render_quote_chart_crop", side_effect=OSError):
            with self.assertRaises(parser.ParseError):
                parser.extract_quote_chart_image(Path("unused.pdf"))

    def test_forecast_reuses_only_single_day_features_with_identical_predictions(self):
        sys.path.insert(0, str(ROOT / "广东电价预测"))
        self.addCleanup(sys.path.remove, str(ROOT / "广东电价预测"))
        import train_d1_forecast as forecast
        times = pd.date_range("2026-06-01", periods=20 * 96, freq="15min")
        frame = pd.DataFrame({"日期": times.normalize(), "时刻": times.strftime("%H:%M"), "全省日前平均电价": np.arange(len(times), dtype=float)})
        frame.loc[frame["日期"].eq(frame["日期"].max()), forecast.TARGET] = np.nan
        options = dict(exogenous_mode="none")
        X, _, identity = forecast.build_features(frame, **options)
        day = frame["日期"].max()
        with patch.object(forecast, "predict_model", side_effect=lambda _, values: values["price_lag_1d"].to_numpy() * .9):
            reference = forecast._recursive_prediction(frame, object(), .7, [day], options)
            with patch.object(forecast, "build_features", side_effect=AssertionError("duplicate features")):
                reused = forecast._recursive_prediction(frame, object(), .7, [day], options, (X, identity))
            pd.testing.assert_frame_equal(reference, reused, check_exact=True)
            with patch.object(forecast, "build_features", wraps=forecast.build_features) as build:
                forecast._recursive_prediction(frame, object(), .7, [day - pd.Timedelta(days=1), day], options, (X, identity))
                self.assertEqual(build.call_count, 2)
