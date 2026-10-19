from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook
from pandas.testing import assert_frame_equal

from toolbox.runtime import load_module


ROOT = Path(__file__).resolve().parents[2]


class OptimizationRegressionTests(unittest.TestCase):
    def test_streamed_summary_preserves_formats_and_last_hour_value(self):
        module = load_module(
            "summary_stream_regression",
            ROOT / "电力工具脚本/电量汇总工具/summarize_511_excel.py",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "summary_input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "汕头1号机"
            sheet.append(["时刻", "电量", "结算电价"])
            sheet.append(["00:00", 1.25, 300.1234])
            sheet["B2"].number_format = "0.00"
            sheet["C2"].number_format = "0.0000"
            sheet.append([None, None, None])
            sheet.append(["01:00", 2, 200])
            sheet.append(["01:00", 3, 250])
            workbook.create_sheet("说明")["A1"] = "无匹配字段"
            workbook.save(path)
            parsed = module.parse_workbook(path)
        self.assertEqual(len(parsed), 1)
        _, energy, price = parsed[0]
        self.assertEqual(len(energy), 24)
        self.assertEqual(energy[0], module.NumericCell(1.25, "0.00"))
        self.assertEqual(price[0], module.NumericCell(300.1234, "0.0000"))
        self.assertEqual(energy[1].value, 3)
        self.assertEqual(price[1].value, 250)
        self.assertEqual(energy[2].value, 0)

    def test_commit_restores_all_files_when_post_replace_verification_fails(self):
        commit_workbooks = load_module(
            "workbook_commit_regression", ROOT / "toolbox_workbook_commit.py"
        ).commit_workbooks
        with tempfile.TemporaryDirectory() as directory:
            prepared = []
            for index in range(2):
                original = Path(directory) / f"{index}.xlsx"
                temporary = Path(directory) / f"{index}.tmp.xlsx"
                original.write_bytes(b"original")
                temporary.write_bytes(b"updated")
                prepared.append(SimpleNamespace(original=original, temporary=temporary))

            def verify(path, item):
                if item is prepared[1]:
                    raise RuntimeError("verification failed")

            with self.assertRaisesRegex(RuntimeError, "verification failed"):
                commit_workbooks(tuple(prepared), verify, create_backup=False)
            for item in prepared:
                self.assertEqual(item.original.read_bytes(), b"original")
                self.assertFalse(item.temporary.exists())
            self.assertEqual(list(Path(directory).glob("*.rollback.xlsx")), [])

    def test_forecast_merge_preserves_nonempty_values_and_input_frames(self):
        module = load_module("forecast_merge_regression", ROOT / "广东电价预测/update_gdfire_data.py")
        base = pd.DataFrame({"日期": ["2026-09-24", "2026-09-24"],
                             "时刻": ["00:00", "00:15"], "数值": [1.25, 2.5]})
        increment = pd.DataFrame({"日期": ["2026-09-24", "2026-09-25"],
                                  "时刻": ["00:00", "00:00"], "数值": [None, 3.75]})
        original_base, original_increment = base.copy(), increment.copy()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.xlsx"
            base.to_excel(path, index=False)
            legacy_result = module.merge_with_base(path, increment)
            result = module.merge_data_frames(base, increment)
        assert_frame_equal(result, legacy_result)
        self.assertEqual(result["数值"].tolist(), [1.25, 2.5, 3.75])
        assert_frame_equal(base, original_base)
        assert_frame_equal(increment, original_increment)

    def test_report_workbook_cache_is_scoped_and_closes_on_failure(self):
        sys.path.insert(0, str(ROOT / "南方分公司电力市场日报"))
        try:
            from gd_power_report import workbook_reader as reader
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "input.xlsx"
                Workbook().save(path)
                opened = mock.Mock()

                @reader.reuse_workbook_inputs
                def parse():
                    self.assertIs(reader.load_workbook(path, data_only=True), opened)
                    self.assertIs(reader.load_workbook(path, data_only=True), opened)
                    raise ValueError("parse failed")

                with mock.patch.object(reader, "_load_workbook", return_value=opened) as loader:
                    with self.assertRaisesRegex(ValueError, "parse failed"):
                        parse()
                    self.assertEqual(loader.call_count, 1)
                    opened.close.assert_called_once()
                    with self.assertRaises(ValueError):
                        parse()
                    self.assertEqual(loader.call_count, 2)
        finally:
            sys.path.remove(str(ROOT / "南方分公司电力市场日报"))

    def test_full_day_table_reuses_input_and_keeps_date_and_formats(self):
        from datetime import date
        from openpyxl.styles import PatternFill
        sys.path.insert(0, str(ROOT / "南方分公司电力市场日报"))
        try:
            from gd_power_report import workbook_reader as reader
            from gd_power_report.full_day_clearing_xlsx import read_full_day_clearing_table
            target_date = date(2026, 9, 20)
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "full_day.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "新"
                sheet["A6"] = target_date
                sheet["A10"] = "分公司全天出清明细（2026/9/20）"
                sheet["A11"] = "时刻"
                sheet.merge_cells("A11:A13")
                sheet["B11"] = "分公司"
                sheet.merge_cells("B11:E11")
                for address, label in (("B12", "日前出清"), ("D12", "实时出清")):
                    sheet[address] = label
                sheet.merge_cells("B12:C12")
                sheet.merge_cells("D12:E12")
                for column, label in enumerate(("电量", "电价", "电量", "电价"), 2):
                    sheet.cell(13, column, label)
                for row in range(14, 39):
                    sheet.cell(row, 1, f"{row - 14}:00" if row < 38 else "合计")
                    for column in range(2, 6):
                        sheet.cell(row, column, 0.125 if column % 2 == 0 else 321.5)
                        sheet.cell(row, column).number_format = "0.000"
                sheet["B14"].fill = PatternFill("solid", fgColor="00FF00")
                workbook.save(path)

                @reader.reuse_workbook_inputs
                def parse():
                    cached = reader.load_workbook(path, data_only=True, read_only=False)
                    result = read_full_day_clearing_table(path, target_date)
                    self.assertIs(reader.load_workbook(path, data_only=True, read_only=False), cached)
                    wrong_date = read_full_day_clearing_table(path, date(2026, 9, 21))
                    self.assertTrue(wrong_date["pending"])
                    return result

                with mock.patch.object(reader, "_load_workbook", wraps=reader._load_workbook) as opened:
                    result = parse()
                self.assertEqual(opened.call_count, 1)
                self.assertFalse(result["pending"])
                self.assertEqual(result["source_range"], "A11:E38")
                self.assertEqual(result["raw_values"], [[0.125, 321.5, 0.125, 321.5]] * 25)
                self.assertIn("0.125", str(result["table"]))
                self.assertIn("#00FF00", str(result["table"]).upper())
        finally:
            sys.path.remove(str(ROOT / "南方分公司电力市场日报"))


if __name__ == "__main__":
    unittest.main()
