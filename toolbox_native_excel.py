"""Shared optional Rust value/number-format reader for immutable XLSX inputs.

This is deliberately not a replacement for styled template editing.
TOOLBOX_XLSX_BACKEND=auto|python|rust controls new consumers.
"""
from __future__ import annotations

from collections import OrderedDict
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass
import json
import os
from pathlib import Path
import subprocess
import tempfile
import threading
import warnings

from openpyxl import load_workbook as python_load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel, from_ISO8601
from openpyxl.worksheet._reader import _cast_number

ROOT = Path(__file__).resolve().parent
_prefetched = ContextVar("native_excel_prefetched", default=None)
_excel_session = ContextVar("native_excel_session", default=None)


def native_binary():
    for path in (ROOT / "500kV断面分析/bin/toolbox-price-xlsx",
                 ROOT / "native/price-xlsx/target/release/toolbox-price-xlsx"):
        if path.is_file() and os.access(path, os.X_OK):
            return path
    return None


def _oneshot_call(binary, operation, path, **params):
    request = json.dumps(dict(protocol=1, operation=operation, path=str(path), **params),
                         ensure_ascii=False, allow_nan=False)
    transport = os.environ.get("TOOLBOX_FRAME_TRANSPORT", "file") if operation == "read_frame" else "pipe"
    if transport not in {"pipe", "file"}:
        raise ValueError(f"未知大表传输方式: {transport}")
    if transport == "file":
        with tempfile.TemporaryFile() as output:
            subprocess.run([str(binary)], input=request, stdout=output, stderr=subprocess.PIPE,
                           text=True, encoding="utf-8", check=True)
            output.seek(0)
            response = json.load(output)
    else:
        result = subprocess.run([str(binary)], input=request, capture_output=True,
                                text=True, encoding="utf-8", check=True)
        response = json.loads(result.stdout)
    if not isinstance(response, dict) or response.get("protocol") != 1:
        raise ValueError("Rust Excel 协议不匹配")
    return response


class _ExcelSession:
    def __init__(self, cache):
        self.process = None
        self.binary = None
        self.supported = None
        self.cache = OrderedDict()
        self.cache_bytes = 0
        self.enable_cache = cache
        self.lock = threading.RLock()

    def request(self, binary, operation, path, params):
        with self.lock:
            if self.supported is None:
                try:
                    response = _oneshot_call(binary, "capabilities", "")
                    self.supported = response.get("stream") is True
                except (OSError, ValueError, subprocess.SubprocessError):
                    self.supported = False
            if not self.supported:
                return _oneshot_call(binary, operation, path, **params)
            if self.process is None:
                self.binary = str(binary)
                self.process = subprocess.Popen([str(binary), "--serve"], stdin=subprocess.PIPE,
                    stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, encoding="utf-8")
            if str(binary) != self.binary:
                return _oneshot_call(binary, operation, path, **params)
            key = None
            if self.enable_cache and operation in {"read", "read_cells"}:
                try:
                    source = Path(path).resolve()
                    stat = source.stat()
                    identity = (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)
                    key = (operation, str(source), identity, json.dumps(params, sort_keys=True))
                except OSError:
                    pass
            if key in self.cache:
                self.cache.move_to_end(key)
                return json.loads(self.cache[key][0])
            request = dict(protocol=1, operation=operation, path=str(path), **params)
            self.process.stdin.write(json.dumps(request, ensure_ascii=False, allow_nan=False) + "\n")
            self.process.stdin.flush()
            line = self.process.stdout.readline()
            if not line:
                raise ValueError("Rust Excel 会话意外退出")
            response = json.loads(line)
            if not isinstance(response, dict):
                raise ValueError("Rust Excel 会话协议不匹配")
            if response.get("protocol") != 1 or "error" in response:
                raise ValueError(response.get("error", "Rust Excel 会话协议不匹配"))
            if key is not None:
                stat = source.stat()
                after = (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)
                size = len(line.encode("utf-8"))
                if identity == after and size <= 16 * 1024 * 1024:
                    while self.cache and (len(self.cache) >= 4 or self.cache_bytes + size > 16 * 1024 * 1024):
                        _, (_, removed) = self.cache.popitem(last=False)
                        self.cache_bytes -= removed
                    self.cache[key] = (line, size)
                    self.cache_bytes += size
            return response

    def close(self):
        self.cache.clear()
        if self.process is not None:
            process = self.process
            try:
                process.stdin.close()
            except OSError:
                pass
            try:
                process.wait(timeout=2)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
            process.stdout.close()


@contextmanager
def excel_session(*, cache=True, enabled=True):
    """A lazy, task-local native process and bounded immutable-response cache."""
    mode = os.environ.get("TOOLBOX_EXCEL_SESSION", "auto")
    if mode not in {"auto", "oneshot", "session"}:
        raise ValueError(f"未知 Excel 会话模式：{mode}")
    if mode == "oneshot" or (mode == "auto" and not enabled) or _excel_session.get() is not None or backend_mode() == "python":
        yield
        return
    session = _ExcelSession(cache)
    token = _excel_session.set(session)
    try:
        yield
    finally:
        _excel_session.reset(token)
        session.close()


def call(binary, operation, path, **params):
    # Large frames keep their selected standalone transport; small batch reads
    # can reuse a process without retaining data beyond the explicit scope.
    session = _excel_session.get()
    if session is not None and operation != "read_frame":
        return session.request(binary, operation, path, params)
    return _oneshot_call(binary, operation, path, **params)


def backend_mode():
    mode = os.environ.get("TOOLBOX_XLSX_BACKEND", "auto")
    if mode not in {"auto", "python", "rust"}:
        raise ValueError(f"未知 Excel 后端: {mode}")
    return mode


def fallback_warning(exc):
    detail = exc.stderr.strip() if isinstance(exc, subprocess.CalledProcessError) else str(exc)
    warnings.warn(f"Rust Excel 读取不可用，改用 Python: {detail}", RuntimeWarning)


@dataclass(frozen=True, slots=True)
class ValueCell:
    value: object
    number_format: str | None
    coordinate: str
    data_type: str = "n"


class ValueSheet:
    def __init__(self, data, book, read_only):
        self.title = data["name"]
        self._data = data
        self._formats = book.formats
        self._date_formats = book.date_formats
        self._duration_formats = book.duration_formats
        self._epoch = book.epoch
        self._styles = data["styles"]
        self._read_only = read_only
        self._merges = [range_boundaries(ref) for ref in data["merges"]] if not read_only else []
        self.max_row = data["max_row"]
        self.max_column = data["max_column"]
        if not read_only:
            self.max_row = max([data["actual_max_row"], *[r[3] for r in self._merges]], default=1)
            self.max_column = max([data["actual_max_column"], *[r[2] for r in self._merges]], default=1)
        if any(index >= len(book.formats) for index in self._styles.values()):
            raise ValueError("无效的 Excel 样式索引")

    def reset_dimensions(self):
        self.max_row = self._data["actual_max_row"]
        self.max_column = self._data["actual_max_column"]

    def cell(self, row, column):
        if row < 1 or column < 1:
            raise ValueError("Row or column values must be at least 1")
        coordinate = f"{get_column_letter(column)}{row}"
        style = self._styles.get(coordinate, 0)
        fmt = self._formats[style]
        if self._read_only and coordinate not in self._styles:
            fmt = None
        rr, cc = row - 1 - self._data["row_offset"], column - 1 - self._data["col_offset"]
        rows = self._data["rows"]
        value = rows[rr][cc] if 0 <= rr < len(rows) and 0 <= cc < len(rows[rr]) else None
        if coordinate in self._data["raw_numbers"]:
            value = _cast_number(self._data["raw_numbers"][coordinate])
        kind = self._data["types"].get(coordinate, "n")
        if isinstance(value, dict):
            if "error" in value:
                value, kind = value["error"], "e"
            elif "iso_date" in value:
                value, kind = from_ISO8601(value["iso_date"]), "d"
            elif "excel_date" in value:
                value = value["excel_date"]
            else:
                raise ValueError("不支持的 Rust Excel 单元格类型")
        if type(value) in (int, float) and self._date_formats[style]:
            value = from_excel(value, epoch=self._epoch, timedelta=self._duration_formats[style])
            kind = "d"
        elif isinstance(value, bool):
            kind = "b"
        elif isinstance(value, str) and kind != "e":
            kind = "s"
        for left, top, right, bottom in self._merges:
            if left <= column <= right and top <= row <= bottom and (row, column) != (top, left):
                value, fmt, kind = None, "General", "n"
                break
        return ValueCell(value, fmt, coordinate, kind)

    def __getitem__(self, coordinate):
        row, column = coordinate_to_tuple(coordinate)
        return self.cell(row, column)

    def iter_rows(self, min_row=1, max_row=None, min_col=1, max_col=None, values_only=False):
        max_row = self.max_row if max_row is None else max_row
        max_col = self.max_column if max_col is None else max_col
        for row in range(min_row, max_row + 1):
            cells = tuple(self.cell(row, col) for col in range(min_col, max_col + 1))
            yield tuple(cell.value for cell in cells) if values_only else cells

    @property
    def rows(self):
        return self.iter_rows()


class ValueWorkbook:
    def __init__(self, response, read_only=True):
        self.read_only = read_only
        self.epoch = CALENDAR_MAC_1904 if response["epoch1904"] else CALENDAR_WINDOWS_1900
        custom = response["custom_formats"]
        self.formats = [custom.get(str(fid), BUILTIN_FORMATS.get(fid, "General"))
                        for fid in response["format_ids"]]
        if any(fid >= 164 and str(fid) not in custom for fid in response["format_ids"]):
            raise ValueError("缺失的 Excel 数字格式")
        self.date_formats = [is_date_format(fmt) for fmt in self.formats]
        self.duration_formats = [is_timedelta_format(fmt) for fmt in self.formats]
        self.worksheets = [ValueSheet(data, self, read_only) for data in response["sheets"]]

    @property
    def sheetnames(self):
        return [sheet.title for sheet in self.worksheets]

    @property
    def active(self):
        return self.worksheets[0]

    def __getitem__(self, name):
        for sheet in self.worksheets:
            if sheet.title == name:
                return sheet
        raise KeyError(f"Worksheet {name} does not exist.")

    def close(self):
        pass


def read_model_frame(path, preferred="整合数据"):
    """Retain pandas' own Excel header/NA/dtype inference with native values."""
    import pandas as pd
    if backend_mode() == "python" or (backend_mode() == "auto" and native_binary() is None):
        with pd.ExcelFile(path) as excel:
            return pd.read_excel(excel, sheet_name=preferred if preferred in excel.sheet_names else excel.sheet_names[0])
    from pandas.io.excel._openpyxl import OpenpyxlReader
    from pandas.io.common import IOHandles
    class NativeReader(OpenpyxlReader):
        def __init__(self, source):
            binary = native_binary()
            try:
                self.book = ValueWorkbook(call(binary, "read_frame", source, selected=[preferred]))
            except (OSError, ValueError, KeyError, subprocess.SubprocessError) as exc:
                if backend_mode() == "rust":
                    raise
                fallback_warning(exc)
                self.book = python_load_workbook(source, read_only=True, data_only=True)
            self.handles = IOHandles(self.book, compression={"method": None})
        def get_sheet_data(self, sheet, file_rows_needed=None):
            if not isinstance(sheet, ValueSheet):
                return super().get_sheet_data(sheet, file_rows_needed)
            import numpy as np
            data = sheet._data
            special = {}
            for coordinate, style in data["styles"].items():
                if sheet._date_formats[style]:
                    special[coordinate_to_tuple(coordinate)] = True
            overrides = {coordinate_to_tuple(c): _cast_number(v) for c,v in data["raw_numbers"].items()}
            if data.get("prepared_frame"):
                # The helper has already normalized ordinary numbers/blank cells.
                # Only date/error/ISO cells and exact numeric overrides need Python.
                result = [[] for _ in range(data["row_offset"])]
                # This private reader owns the read_frame response. Reuse its
                # row lists instead of retaining two additional full matrices.
                for values in data["rows"]:
                    result.append([""] * data["col_offset"] + values if data["col_offset"] else values)
                    if file_rows_needed is not None and len(result) >= file_rows_needed:
                        break
                positions = set(special) | {tuple(pos) for pos in data.get("object_cells", [])} | set(overrides)
                for rr, cc in positions:
                    if not (0 < rr <= len(result) and 0 < cc <= len(result[rr - 1])):
                        continue
                    value = overrides.get((rr, cc), result[rr - 1][cc - 1])
                    if (rr, cc) in special or isinstance(value, dict):
                        cell = sheet.cell(rr, cc)
                        value = np.nan if cell.data_type == "e" else cell.value
                    if value is None:
                        value = ""
                    elif type(value) in (int, float) and not pd.isna(value):
                        integer = int(value)
                        value = integer if integer == value else float(value)
                    result[rr - 1][cc - 1] = value
                for row in result:
                    while row and row[-1] == "":
                        row.pop()
                while result and not result[-1]:
                    result.pop()
                width = max(map(len, result), default=0)
                for row in result:
                    row.extend([""] * (width - len(row)))
                return result
            result = [[] for _ in range(data["row_offset"])]
            for rr, values in enumerate(data["rows"], 1 + data["row_offset"]):
                row = [""] * data["col_offset"]
                for cc, value in enumerate(values, 1 + data["col_offset"]):
                    value = overrides.get((rr,cc), value)
                    if (rr,cc) in special or isinstance(value, dict):
                        cell = sheet.cell(rr,cc)
                        value = np.nan if cell.data_type == "e" else cell.value
                    if value is None:
                        value = ""
                    elif type(value) in (int,float) and not pd.isna(value):
                        integer = int(value)
                        value = integer if integer == value else float(value)
                    row.append(value)
                while row and row[-1] == "":
                    row.pop()
                result.append(row)
                if file_rows_needed is not None and len(result) >= file_rows_needed:
                    break
            while result and not result[-1]:
                result.pop()
            width = max(map(len, result), default=0)
            return [row + [""] * (width-len(row)) for row in result]
    reader = NativeReader(path)
    try:
        sheet = preferred if preferred in reader.sheet_names else reader.sheet_names[0]
        return reader.parse(sheet_name=sheet)
    finally:
        reader.close()


def load_value_workbook(path, *, data_only=True, read_only=True, selected=(), native_in_auto=True):
    if not data_only:
        raise ValueError("数值读取器仅接受 data_only=True")
    mode = backend_mode()
    # Some existing streaming workflows outperform whole-sheet native transport.
    # Keep their measured fast path in auto; strict rust still permits comparison.
    use_native = mode == "rust" or (mode == "auto" and native_in_auto)
    binary = native_binary() if use_native else None
    if mode == "rust" and binary is None:
        raise RuntimeError("未安装 Rust Excel 后端，请运行 native/price-xlsx/build.py")
    if binary is not None:
        try:
            cached = _prefetched.get()
            key = Path(path).resolve()
            if cached is not None and key in cached:
                response = cached[key]
                if response is not None:
                    return ValueWorkbook(response, read_only)
                return python_load_workbook(path, data_only=True, read_only=read_only)
            response = call(binary, "read_cells", path, selected=list(selected))
            return ValueWorkbook(response, read_only)
        except (OSError, ValueError, KeyError, IndexError, TypeError, subprocess.SubprocessError) as exc:
            if mode == "rust":
                raise
            fallback_warning(exc)
    return python_load_workbook(path, data_only=True, read_only=read_only)


@contextmanager
def prefetch_workbooks(paths):
    mode = backend_mode()
    binary = native_binary() if mode != "python" else None
    books = {}
    if binary is not None:
        paths = list(dict.fromkeys(Path(path).resolve() for path in paths))
        if paths:
            try:
                response = call(binary, "read_cells_many", "", paths=[str(p) for p in paths])
                if len(response["books"]) != len(paths):
                    raise ValueError("Rust Excel 批量读取结果数量不匹配")
                for path, item in zip(paths, response["books"]):
                    if "error" in item:
                        raise ValueError(item["error"])
                    books[path] = item
            except (OSError, ValueError, KeyError, TypeError, subprocess.SubprocessError) as exc:
                if mode == "rust":
                    raise
                fallback_warning(exc)
                books = {path: None for path in paths}
    token = _prefetched.set(books)
    try:
        yield
    finally:
        _prefetched.reset(token)
