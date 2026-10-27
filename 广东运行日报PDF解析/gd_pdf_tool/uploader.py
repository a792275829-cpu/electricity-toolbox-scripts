from __future__ import annotations

import math
import re
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable
from urllib.parse import urlsplit

from openpyxl import load_workbook

from .excel_writer import HEADERS
from .parser import QuoteRow, extract_report_date, price_intervals


TOOL_DIR = Path(__file__).resolve().parents[1]
WORKSPACE = TOOL_DIR.parent
BASE_URL = "https://xhxt.chng.com.cn"
UPLOAD_URL = f"{BASE_URL}/datacenter/GDDataManagement/SpotUpload"
CONFIG_PATH = TOOL_DIR / "config.json"
STATUS_PATH = "/datacenter/gd/api/pub/public/data/list/status"


@dataclass(frozen=True)
class UploadPlan:
    pdf_path: Path
    excel_path: Path
    report_date: date


def verify_quote_workbook(path: Path, expected_rows: list[QuoteRow] | None = None) -> int:
    """Read back the exact upload schema and, after parsing, every saved value."""
    workbook = load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("报价容量 Excel 必须只有 Sheet1 工作表。")
        sheet = workbook["Sheet1"]
        if tuple(cell.value for cell in sheet[1]) != HEADERS:
            raise ValueError(f"报价容量 Excel 表头必须为：{'、'.join(HEADERS)}。")
        count = sheet.max_row - 1
        if count < 1 or (expected_rows is not None and len(expected_rows) != count):
            raise ValueError("报价容量 Excel 数据行数不正确。")
        for index, interval in enumerate(price_intervals(count), 2):
            if sheet.cell(index, 1).value != interval:
                raise ValueError(f"报价容量 Excel 第 {index} 行报价区间顺序不正确。")
            for column in range(2, 5):
                cell = sheet.cell(index, column)
                value = cell.value
                if isinstance(value, bool) or not isinstance(value, (int, float)):
                    raise ValueError(f"报价容量 Excel {cell.coordinate} 必须为数值。")
                if not math.isfinite(value) or value < 0 or cell.number_format != "0.0":
                    raise ValueError(f"报价容量 Excel {cell.coordinate} 的容量或数字格式无效。")
                if expected_rows is not None:
                    row = expected_rows[index - 2]
                    expected = (row.coal, row.gas, row.renewable)[column - 2]
                    if value != expected:
                        raise ValueError(f"报价容量 Excel {cell.coordinate} 回读值与解析结果不一致。")
        return count
    finally:
        workbook.close()


def build_upload_plan(pdf_path: str | Path, excel_path: str | Path) -> UploadPlan:
    pdf = Path(pdf_path).expanduser().resolve()
    excel = Path(excel_path).expanduser().resolve()
    for path, suffix in ((pdf, ".pdf"), (excel, ".xlsx")):
        if not path.is_file() or path.suffix.lower() != suffix:
            raise ValueError(f"上传文件无效：{path}")
    stamp = extract_report_date(pdf)
    if not re.fullmatch(r"20\d{6}", stamp):
        raise ValueError("无法确定 PDF 报告日期，已停止上传。请检查源文件日期。")
    report_date = datetime.strptime(stamp, "%Y%m%d").date()
    filename_date = re.search(r"报价容量区间(20\d{6})", excel.stem)
    if filename_date and filename_date.group(1) != stamp:
        raise ValueError("解析 Excel 文件名日期与 PDF 报告日期不一致，已停止上传。")
    verify_quote_workbook(excel)
    return UploadPlan(pdf, excel, report_date)


def load_auth_module():
    from toolbox_shared.auth_loader import load_login_module

    return load_login_module(
        WORKSPACE, module_name="daily_pdf_upload_auth", state_dir=TOOL_DIR,
        config_path=CONFIG_PATH, app_url=UPLOAD_URL,
    )


def load_upload_config():
    from toolbox_shared.credentials import read_login_config
    from types import SimpleNamespace

    raw = read_login_config(CONFIG_PATH, WORKSPACE / "上网电量抓取" / "config.json")
    return SimpleNamespace(
        username=str(raw.get("username") or ""), password=str(raw.get("password") or ""),
        headless=bool(raw.get("headless", True)),
    )


def open_upload_page(context, auth):
    response = context.request.get(
        f"{BASE_URL}/usercenter/web/pf/tenant/user/application", timeout=20000,
    )
    payload = auth.parse_json_response(response, "读取数据中心授权")
    tenants = [
        tenant for tenant in payload.get("data", [])
        if "广东分公司" in str(tenant.get("name") or "")
        and tenant.get("tenantId")
        and any(app.get("name") == "数据中心" for app in tenant.get("applications", []))
    ]
    if len(tenants) != 1:
        raise RuntimeError("未能唯一确定有数据中心权限的广东分公司，已停止上传。")
    auth.switch_company(context, str(tenants[0]["tenantId"]))
    page = context.new_page()
    page.goto(UPLOAD_URL, wait_until="domcontentloaded", timeout=60000)
    page.get_by_role("tab", name="现货交易日报", exact=True).wait_for(timeout=30000)
    return page


def active_panel(page):
    # This site renders its table beside empty Ant Design tabpanels.
    # Scope to the visible date filter's nearest table-containing container.
    return page.locator("#runDates:visible").locator("xpath=ancestor::div[.//table][1]")


def set_report_date(page, report_date: date) -> None:
    """Select the same start/end date through the website's range picker."""
    panel = active_panel(page)
    panel.locator("#runDates").click()
    calendar = page.locator(".ant-picker-dropdown:visible")
    stamp = report_date.isoformat()
    for _ in range(120):
        dates = calendar.locator("td.ant-picker-cell-in-view[title]")
        dates.first.wait_for()
        current = date.fromisoformat(dates.first.get_attribute("title"))
        difference = (report_date.year - current.year) * 12 + report_date.month - current.month
        if difference == 0:
            break
        direction = "prev" if difference < 0 else "next"
        step = "super-" if abs(difference) >= 12 else ""
        calendar.locator(f".ant-picker-header-{step}{direction}-btn:visible").first.click()
    else:
        raise RuntimeError("无法在日期选择器中定位报告月份。")
    cell = calendar.locator(f'td.ant-picker-cell-in-view[title="{stamp}"]').first
    cell.click()
    with page.expect_response(lambda response: STATUS_PATH in response.url, timeout=30000):
        calendar.locator(f'td.ant-picker-cell-in-view[title="{stamp}"]').first.click()
    values = panel.locator(".ant-picker input").evaluate_all("es => es.map(e => e.value)")
    if values != [stamp, stamp]:
        raise RuntimeError(f"网站数据日期未正确切换到 {stamp}，已停止上传。")


def date_row(page, report_date: date):
    stamp = report_date.isoformat()
    row = active_panel(page).locator("tbody tr").filter(
        has=page.get_by_role("cell", name=stamp, exact=True),
    )
    row.wait_for(state="visible", timeout=30000)
    if row.count() != 1:
        raise RuntimeError(f"网站日期 {stamp} 未唯一匹配，已停止上传。")
    return row


def verify_remote_status(context, auth, *, data_type: str, plan: UploadPlan, path: Path) -> dict:
    stamp = plan.report_date.isoformat()
    response = context.request.get(
        f"{BASE_URL}{STATUS_PATH}", params={
            "dataType": data_type, "provinceAreaId": "044", "publicDataItemList": "",
            "startTime": stamp, "endTime": stamp,
        }, timeout=30000,
    )
    payload = auth.parse_json_response(response, "回读上传状态")
    records = payload.get("data")
    if not isinstance(records, list):
        raise RuntimeError("网站上传状态返回格式无效。")
    matching = [item for item in records if str(item.get("date", ""))[:10] == stamp]
    if len(matching) != 1:
        raise RuntimeError(f"网站未返回唯一的 {stamp} 上传记录。")
    record = matching[0]
    filename = record.get("importFileName") or ""
    # The quotation import stores parsed data and returns an empty filename.
    # Its selected filename was checked in the UI before the successful POST.
    quotation = data_type == "SPOT_RUN_REPORT_QUOTE_CAPACITY"
    filename_matches = filename == path.name or (
        quotation and not filename and record.get("dataItem") == "报价容量数据"
    )
    if record.get("importStatus") not in (1, "1") or not filename_matches:
        raise RuntimeError(f"网站 {stamp} 的上传状态或文件名回读不一致：{path.name}")
    if not record.get("updateTime"):
        raise RuntimeError(f"网站 {stamp} 未返回上传完成时间。")
    return record


def wait_upload_response(page, action, *, endpoint: str, label: str, logger,
                         submitted: bool = False, timeout_ms: int = 120000):
    """Report progress and request failures without retrying a possible write."""
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

    responses, failures = [], []

    def matches(request):
        return request.method == "POST" and urlsplit(request.url).path == endpoint

    def on_response(response):
        if matches(response.request):
            responses.append(response)

    def on_failure(request):
        if matches(request):
            failures.append(request.failure or "连接中断")

    page.on("response", on_response)
    page.on("requestfailed", on_failure)
    started = time.monotonic()
    next_progress = 15.0
    seconds = math.ceil(timeout_ms / 1000)
    detail = ("上传请求可能已提交，请核对网站状态后再重试。" if submitted
              else "文件校验未完成，尚未执行本文件的全部上传；可稍后使用“仅上传”重试。")
    try:
        logger(f"{label}：等待网站接口返回，最多 {seconds} 秒")
        action()
        while not responses:
            if failures:
                raise RuntimeError(f"{label}网络请求失败：{failures[0]}；{detail}本地 PDF 和 Excel 已保留。")
            elapsed = time.monotonic() - started
            if elapsed * 1000 >= timeout_ms:
                cause = PlaywrightTimeoutError(
                    f'Timeout {timeout_ms}ms exceeded while waiting for POST {endpoint}'
                )
                raise RuntimeError(
                    f"{label}超时：{seconds} 秒内网站接口未返回。"
                    f"{detail}本地 PDF 和 Excel 已保留。"
                ) from cause
            if elapsed >= next_progress:
                logger(f"{label}：已等待 {int(elapsed)} 秒，网站仍未返回结果")
                next_progress += 15
            page.wait_for_timeout(min(250, max(1, timeout_ms - elapsed * 1000)))
        return responses[0]
    finally:
        page.remove_listener("response", on_response)
        page.remove_listener("requestfailed", on_failure)


def upload_one(page, context, auth, plan: UploadPlan, path: Path, *, label: str, data_type: str, logger):
    from playwright.sync_api import expect

    page.get_by_role("tab", name=label, exact=True).click()
    set_report_date(page, plan.report_date)
    panel = active_panel(page)
    # A fresh page must never submit selections belonging to other dates.
    if panel.locator('input[type="file"]').evaluate_all("es => es.some(e => e.files.length)"):
        raise RuntimeError("页面有其他待上传文件，已停止，避免全部上传包含无关文件。")
    row = date_row(page, plan.report_date)
    logger(f"{label}：{plan.report_date.isoformat()} 选择文件 {path.name}")
    prefix = "/datacenter/gd/api/pub/data/report/import" if data_type == "SPOT_RUN_REPORT_DATA" else "/datacenter/gd/api/pub/data/import"
    validating = wait_upload_response(
        page, lambda: row.locator('input[type="file"]').set_input_files(str(path)),
        endpoint=prefix + "/valid/multi", label=f"{label}文件校验", logger=logger,
    )
    auth.parse_json_response(validating, f"{label}文件校验")
    button = panel.get_by_role("button", name="全部上传", exact=True)
    try:
        expect(button).to_be_enabled(timeout=30000)
        expect(row).to_contain_text(path.name, timeout=30000)
    except AssertionError as exc:
        raise RuntimeError(f"{label}文件未通过网站校验，未执行全部上传。") from exc
    logger(f"{label}：执行全部上传")
    uploading = wait_upload_response(
        page, button.click, endpoint=prefix + "/create/multi",
        label=f"{label}全部上传", logger=logger, submitted=True,
    )
    payload = auth.parse_json_response(uploading, f"{label}全部上传")
    if payload.get("data") != []:
        raise RuntimeError(f"{label}服务器返回导入异常，请检查网站导入结果。")
    expect(button).to_be_disabled(timeout=30000)
    record = verify_remote_status(context, auth, data_type=data_type, plan=plan, path=path)
    expect(date_row(page, plan.report_date)).to_contain_text(record["updateTime"], timeout=30000)
    logger(f"{label}上传成功：{plan.report_date.isoformat()}，网站完成时间 {record['updateTime']}")
    return record


def upload_report_files(plan: UploadPlan, *, logger: Callable[[str], None] = print, config=None) -> None:
    from playwright.sync_api import sync_playwright

    config = config or load_upload_config()
    auth = load_auth_module()
    completed: list[str] = []
    with auth.AuthStateLock():
        with sync_playwright() as playwright:
            context = auth.launch_context(playwright, headless=config.headless)
            try:
                logger("正在检查清能网站登录状态...")
                auth.ensure_login(context, config)
                logger("正在打开广东数据中心日报上传页面...")
                page = open_upload_page(context, auth)
                for path, label, data_type in (
                    (plan.pdf_path, "现货交易日报", "SPOT_RUN_REPORT_DATA"),
                    (plan.excel_path, "报价容量数据", "SPOT_RUN_REPORT_QUOTE_CAPACITY"),
                ):
                    upload_one(page, context, auth, plan, path, label=label, data_type=data_type, logger=logger)
                    completed.append(label)
                auth.save_auth_state(context)
            except Exception:
                if completed:
                    logger(f"已完成：{'、'.join(completed)}；后续上传失败或状态未确认，本地 PDF 和 Excel 已保留。")
                raise
            finally:
                context.close()
