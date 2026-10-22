from __future__ import annotations

import tempfile
import subprocess
import sys
import unittest
from unittest import mock
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from toolbox.runtime import ToolPaths, load_module


class TransactionInfoUpdateTests(unittest.TestCase):
    def test_commit_rejects_external_save_and_preserves_other_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "transaction.xlsx"
            self._build_fixture(path)
            day = date(2026, 9, 11)
            pair = self.module.build_daily_load_pair(day, self._payload(day), self._payload(day - timedelta(days=3)))
            prepared = self.module.prepare_workbook(path, (pair,))
            book = openpyxl.load_workbook(path)
            book["保留工作表"]["A1"] = "外部保存"
            book["保留工作表"]["B1"] = 0.125
            book["保留工作表"]["B1"].number_format = "0.00%"
            book.save(path)
            book.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(RuntimeError, "工作簿被修改"):
                self.module.commit_prepared_workbook(prepared, create_backup=False)
            self.assertEqual(path.read_bytes(), original)
            saved = openpyxl.load_workbook(path)
            try:
                self.assertEqual(saved["保留工作表"]["A1"].value, "外部保存")
                self.assertEqual(saved["保留工作表"]["B1"].value, 0.125)
                self.assertEqual(saved["保留工作表"]["B1"].number_format, "0.00%")
                self.assertIsNone(saved["Sheet1"]["C8"].value)
            finally:
                saved.close()
            self.assertEqual(list(Path(directory).iterdir()), [path])

    def test_tail_streaming_keeps_latest_group_and_rejects_truncated_labels(self) -> None:
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            self._build_fixture(path)
            workbook = openpyxl.load_workbook(path)
            sheet = workbook["Sheet1"]
            for row in range(1, 7):
                sheet.cell(row + 100, 4, sheet.cell(row, 4).value)
                sheet.cell(row + 100, 11, sheet.cell(row, 11).value)
            sheet["C101"] = date(2026, 9, 11)
            sheet["J101"] = date(2026, 9, 8)
            workbook.save(path)
            with mock.patch.object(ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")):
                tail = self.module.inspect_workbook_tail(path)
            self.assertEqual(tail.start_row, 101)
            self.assertEqual(tail.day_ahead_date, date(2026, 9, 11))
            sheet["K106"] = None
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(RuntimeError, "最后一组字段顺序异常"):
                self.module.inspect_workbook_tail(path)

    @classmethod
    def setUpClass(cls) -> None:
        cls.paths = ToolPaths(Path(__file__).resolve().parents[2])
        cls.module = load_module(
            "transaction_info_update_tests", cls.paths.transaction_info_update
        )

    @staticmethod
    def _payload(target_date: date) -> dict:
        rows = []
        for index, load_type in enumerate(("1", "6", "4", "2", "3", "5"), start=1):
            values = {
                "date": target_date.isoformat(),
                "max": 1000 * index,
                "min": 100 * index,
                "avg": 500.125 * index,
                "sum": 0.12003 * index,
            }
            rows.append(
                {
                    "loadType": load_type,
                    "forecastPeriodList": [values],
                    "actualPeriodList": [values],
                }
            )
        return {"data": {"dataNetLoadDTOList": rows}}

    @staticmethod
    def _build_fixture(path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        day_date = date(2026, 9, 10)
        actual_date = day_date - timedelta(days=3)
        categories = ("统调", "B类", "西电", "A类", "地方电", "粤港联络线")
        for offset, category in enumerate(categories, start=1):
            if offset == 1:
                sheet.cell(offset, 3, datetime.combine(day_date, datetime.min.time()))
                sheet.cell(offset, 10, datetime.combine(actual_date, datetime.min.time()))
                sheet.cell(offset, 3).number_format = "yyyy/m/d"
                sheet.cell(offset, 10).number_format = "yyyy/m/d"
            sheet.cell(offset, 4, category)
            sheet.cell(offset, 11, category)
            if category != "西电":
                for column in range(5, 9):
                    sheet.cell(offset, column, 10 * offset + column)
                    sheet.cell(offset, column).number_format = "0.0000"
            for column in range(12, 16):
                sheet.cell(offset, column, 20 * offset + column)
                sheet.cell(offset, column).number_format = "0.0000"
        workbook.create_sheet("保留工作表")["A1"] = "不应改动"
        workbook.save(path)
        workbook.close()

    def test_parse_load_dataset_distinguishes_day_ahead_and_actual(self) -> None:
        target_date = date(2026, 9, 11)
        payload = self._payload(target_date)
        day_ahead = self.module.parse_load_dataset(
            payload, target_date, day_ahead=True
        )
        actual = self.module.parse_load_dataset(payload, target_date, day_ahead=False)

        self.assertNotIn("西电", day_ahead)
        self.assertEqual(day_ahead["统调"], (1000, 100, 500.125, 12003))
        self.assertEqual(actual["西电"], (3000, 300, 1500.375, 36009))
        self.assertEqual(actual["粤港联络线"], (6000, 600, 3000.75, 72018))

    def test_build_pair_keeps_day_ahead_b_without_subtracting_actual_west(self) -> None:
        day_ahead_date = date(2026, 9, 11)
        actual_date = day_ahead_date - timedelta(days=3)

        pair = self.module.build_daily_load_pair(
            day_ahead_date,
            self._payload(day_ahead_date),
            self._payload(actual_date),
        )

        self.assertEqual(pair.actual_date, actual_date)
        self.assertEqual(pair.day_ahead["B类"], (2000, 200, 1000.25, 24006))
        self.assertEqual(pair.actual["西电"], (3000, 300, 1500.375, 36009))
        self.assertEqual(pair.day_ahead["统调"], (1000, 100, 500.125, 12003))

    def test_prepare_and_commit_append_group_preserve_styles_and_other_parts(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "交易系统信息收集.xlsx"
            self._build_fixture(workbook_path)
            target_date = date(2026, 9, 11)
            actual_date = target_date - timedelta(days=3)
            pair = self.module.build_daily_load_pair(
                target_date,
                self._payload(target_date),
                self._payload(actual_date),
            )

            prepared = self.module.prepare_workbook(workbook_path, (pair,))
            backup = self.module.commit_prepared_workbook(
                prepared, create_backup=True
            )

            workbook = openpyxl.load_workbook(workbook_path, data_only=True)
            try:
                sheet = workbook["Sheet1"]
                self.assertEqual(sheet["C8"].value.date(), target_date)
                self.assertEqual(sheet["J8"].value.date(), actual_date)
                self.assertEqual(sheet["D10"].value, "西电")
                self.assertIsNone(sheet["E10"].value)
                self.assertIsNone(sheet["H10"].value)
                self.assertEqual(sheet["L10"].value, 3000)
                self.assertEqual(sheet["O10"].value, 36009)
                self.assertEqual(sheet["E9"].value, 2000)
                self.assertEqual(sheet["H9"].value, 24006)
                self.assertEqual(sheet["E8"].number_format, sheet["E1"].number_format)
                self.assertEqual(workbook["保留工作表"]["A1"].value, "不应改动")
            finally:
                workbook.close()
            self.assertIsNotNone(backup)
            self.assertTrue(backup.is_file())

    def test_resolve_start_date_requires_contiguous_append(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "fixture.xlsx"
            self._build_fixture(workbook_path)
            self.assertEqual(
                self.module.resolve_start_date(workbook_path, None),
                date(2026, 9, 11),
            )
            with self.assertRaisesRegex(ValueError, "开始日期必须是 2026-09-11"):
                self.module.resolve_start_date(workbook_path, date(2026, 9, 12))

    def test_already_updated_date_is_successful_no_op(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "fixture.xlsx"
            self._build_fixture(workbook_path)
            original = workbook_path.read_bytes()
            result = subprocess.run(
                [
                    sys.executable,
                    str(self.paths.transaction_info_update),
                    "--workbook",
                    str(workbook_path),
                    "--end-date",
                    "2026-09-10",
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("无需更新", result.stdout)
            self.assertEqual(workbook_path.read_bytes(), original)
            self.assertEqual(list(workbook_path.parent.glob("*.备份-*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
