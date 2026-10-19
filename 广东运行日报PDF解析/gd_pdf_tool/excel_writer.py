from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .parser import QuoteRow


HEADERS = ("报价区间", "燃煤", "燃气", "新能源")


def write_quote_excel(rows: Iterable[QuoteRow], output_path: str | Path) -> Path:
    """将解析结果写成固定结构的单工作表 Excel。"""
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append((row.interval, row.coal, row.gas, row.renewable))

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cells in sheet.iter_rows(min_row=2, min_col=2, max_col=4):
        for cell in cells:
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="right")
    for cell in sheet["A"]:
        cell.alignment = Alignment(horizontal="center")

    for column, width in {"A": 16, "B": 12, "C": 12, "D": 12}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(output)
    return output
