from __future__ import annotations

import tempfile
import tkinter as tk
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Font, PatternFill

from toolbox.runtime import ToolPaths, load_module


class PowerBenefitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.paths = ToolPaths(Path(__file__).resolve().parents[2])
        cls.module = load_module("power_benefit_tests", cls.paths.power_benefit_update)

    def make_template(self, folder, name="归湖"):
        path = Path(folder) / f"功率预测调整收益统计-{name}.xlsx"
        w = openpyxl.Workbook()
        s = w.active
        s.title = "10月"
        for col, label in self.module.HEADERS.items():
            s[f"{col}2"] = label
        for start, day in ((4, datetime(2026, 10, 1)), (29, "=A4+1"), (54, "=A29+1")):
            s.cell(start, 1, day).number_format = "mm-dd-yy"
            for hour in range(24):
                row = start + hour
                s.cell(row, 2, f"{hour}:00")
                for col in self.module.FIELDS:
                    c = s[f"{col}{row}"]
                    c.font = Font(name="Arial", size=11)
                    c.fill = PatternFill("solid", fgColor="FFF2CC")
                    c.number_format = "0.00_ "
                s[f"Q{row}"] = f"=(M{row}-K{row})*(N{row}-L{row})"
            s.cell(start+24, 2, "合计")
            s[f"C{start+24}"] = f"=SUM(C{start}:C{start+23})"
        s["C4"] = 0
        s["J5"] = "=1+1"
        s["K6"] = 123.45
        w.save(path)
        w.close()
        return path

    def test_filename_and_exact_company_mapping(self):
        m = self.module
        self.assertEqual(m.station_from_filename(Path("功率预测调整收益统计-鮀莲(2)(2).xlsx")), "鮀莲")
        for name in ("无场站.xlsx", "归湖鮀莲.xlsx"):
            with self.assertRaises(ValueError): m.station_from_filename(Path(name))
        for spec in m.shared_spot().default_targets()[:2]:
            company = m.shared_spot().resolve_tenant([
                {"tenantId":"a", "name":"华能（潮州潮安）新能源有限责任公司"},
                {"tenantId":"b", "name":"华能（汕头金平）新能源有限责任公司"}], spec)
            self.assertEqual(company["tenantId"], "a" if spec.key == "归湖" else "b")

    def test_date_formulas_cutoff_and_nonempty_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_template(folder)
            plan = self.module.inspect_workbook(path, end_date=date(2026,10,2))
            self.assertEqual(len(plan.cells), 24*6*2-3)
            self.assertEqual({c.run_date for c in plan.cells}, {date(2026,10,1), date(2026,10,2)})
            self.assertFalse({"C4","J5","K6"} & {c.coordinate for c in plan.cells})
            plan = self.module.inspect_workbook(path, start_date=date(2026,10,2), end_date=date(2026,10,2))
            self.assertEqual(len(plan.cells), 144)

    def test_out_of_month_placeholder_is_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_template(folder)
            w = openpyxl.load_workbook(path)
            s = w.copy_worksheet(w["10月"]); s.title = "9月"; s["A4"] = datetime(2026,9,30)
            w.save(path); w.close()
            plan = self.module.inspect_workbook(path, end_date=date(2026,10,2))
            self.assertEqual(sum(c.sheet == "9月" for c in plan.cells), 141)

    def test_forecast_uses_original_and_declared_not_reversed(self):
        payload = {"data":[{"unitId":"u", "date":"2026-02-01",
                            "forecastPower24":[10]*24,"stPredictPowerList24":[20]*24}]}
        self.assertEqual(self.module.parse_forecast_payload(payload,unit_id="u"), {"C":(10,)*24,"J":(20,)*24})
        payload["data"][0]["rawShortTermPower24"] = [30]*24
        self.assertEqual(self.module.parse_forecast_payload(payload,unit_id="u",fallback=True)["J"], (30,)*24)

    def test_trade_units_date_and_unit_filter_and_negative_prices(self):
        row={"unitId":"u","date":"2026-10-02","electricity":[2.761]*24,"acPrice":[-25.89]*24,"dataType":"1"}
        data=self.module.parse_trade_payload({"data":{"unitDateList":[dict(row,unitId="-1"),row]}},unit_id="u",expected_dates={date(2026,10,2)},data_type=1)
        self.assertAlmostEqual(data[date(2026,10,2)]["K"][12],27.61)
        self.assertEqual(data[date(2026,10,2)]["L"][12],-25.89)
        with self.assertRaisesRegex(ValueError,"重复"):
            self.module.parse_trade_payload({"data":{"unitDateList":[row,row]}},unit_id="u",expected_dates={date(2026,10,2)},data_type=1)

    def test_unpublished_is_not_zero_and_invalid_points_fail(self):
        m=self.module
        for raw in (None, [], [None]*24, [0]*23+[None]): self.assertIsNone(m.points24(raw,label="test"))
        self.assertEqual(m.points24([0]*24,label="test"),(0,)*24)
        for raw in ([0]*23, [0]*23+[float("nan")], [False]*24):
            with self.assertRaises(ValueError): m.points24(raw,label="test")

    def test_sparse_write_preserves_other_values_formulas_styles_and_package(self):
        m=self.module
        with tempfile.TemporaryDirectory() as folder:
            path=self.make_template(folder)
            plan=m.inspect_workbook(path,end_date=date(2026,10,2))
            series={day:{col:tuple(index+hour/100 for hour in range(24)) for index,col in enumerate(m.FIELDS,1)}
                    for day in (date(2026,10,1),date(2026,10,2))}
            prepared=m.prepare_workbook(plan,series)
            original=openpyxl.load_workbook(path); result=openpyxl.load_workbook(prepared.temporary)
            changed={(c.sheet,c.coordinate):v for c,v in prepared.expectations}
            for sheet in original:
                for row in sheet:
                    for cell in row:
                        actual=result[sheet.title][cell.coordinate]
                        key = (sheet.title,cell.coordinate)
                        if key in changed:
                            self.assertAlmostEqual(actual.value, changed[key], places=12)
                        else:
                            self.assertEqual(actual.value, cell.value)
                        self.assertEqual(actual._style,cell._style)
            self.assertTrue(result.calculation.fullCalcOnLoad)
            original.close(); result.close()
            with zipfile.ZipFile(path) as a,zipfile.ZipFile(prepared.temporary) as b:
                self.assertEqual(a.namelist(),b.namelist())
                for name in a.namelist():
                    if name not in {"xl/workbook.xml","xl/worksheets/sheet1.xml"}: self.assertEqual(a.read(name),b.read(name))
            prepared.temporary.unlink()

    def test_concurrent_changes_block_preparation(self):
        with tempfile.TemporaryDirectory() as folder:
            path=self.make_template(folder); m=self.module
            plan=m.inspect_workbook(path,end_date=date(2026,10,2))
            with path.open("ab") as stream: stream.write(b"changed")
            with self.assertRaisesRegex(RuntimeError,"被修改"): m.prepare_workbook(plan,{})

    def test_blank_cell_omitted_from_xml_can_be_inserted_in_column_order(self):
        m = self.module
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_template(folder)
            w = openpyxl.load_workbook(path)
            w["10月"]["L4"]._style = None
            w.save(path); w.close()
            plan = m.inspect_workbook(path, end_date=date(2026,10,1))
            item = m.prepare_workbook(plan, {date(2026,10,1): {"L": (300,)*24}})
            w = openpyxl.load_workbook(item.temporary)
            self.assertEqual(w["10月"]["L4"].value, 300)
            self.assertEqual(w["10月"]["L4"].number_format, "General")
            w.close(); item.temporary.unlink()

    def test_adjacent_self_closing_cells_do_not_consume_next_cell_or_formula(self):
        source = b'<worksheet><sheetData><row r="4"><c r="K4" s="41"/><c r="L4" s="42"/><c r="M4" s="43"/><c r="O4" s="39"><f>M4-K4</f><v>0</v></c></row></sheetData></worksheet>'
        result = self.module.patch_sparse_sheet(source, {"K4": 10, "L4": 300, "M4": 12})
        self.assertIn(b'<c r="K4" s="41"><v>10</v></c>', result)
        self.assertIn(b'<c r="L4" s="42"><v>300</v></c>', result)
        self.assertIn(b'<c r="M4" s="43"><v>12</v></c>', result)
        self.assertIn(b'<c r="O4" s="39"><f>M4-K4</f><v>0</v></c>', result)

    def test_two_workbooks_rollback_on_second_verification_failure(self):
        m=self.module
        with tempfile.TemporaryDirectory() as folder:
            paths=[self.make_template(folder,name) for name in ("归湖","鮀莲")]
            original=[p.read_bytes() for p in paths]
            items=tuple(m.prepare_workbook(m.inspect_workbook(p,end_date=date(2026,10,1)),
                      {date(2026,10,1):{"C":(5,)*24}}) for p in paths)
            calls=[]
            def verify(path,item):
                calls.append(path)
                if len(calls)==2: raise RuntimeError("test verification failure")
                m.verify_workbook(path,item)
            with self.assertRaisesRegex(RuntimeError,"verification failure"):
                m.commit_workbooks(items,verify,create_backup=False)
            self.assertEqual([p.read_bytes() for p in paths],original)
            self.assertFalse(list(Path(folder).glob(".*.xlsx")))

    def test_auth_profile_and_lock_are_independent(self):
        from toolbox_shared.storage import data_path
        m=self.module; auth=m.load_auth_module()
        self.assertEqual(auth.AUTH_STATE_PATH,data_path(m.SCRIPT_DIR/"auth_state.json"))
        self.assertEqual(auth.AUTH_LOCK_PATH,data_path(m.SCRIPT_DIR/"auth_state.json.lock"))
        self.assertEqual(auth.PROFILE_DIR,data_path(m.SCRIPT_DIR/".browser-profile"))
        self.assertNotEqual(auth.AUTH_STATE_PATH,data_path(self.paths.online_energy_dir/"auth_state.json"))

    def haifeng_template(self, folder):
        path = Path(folder) / "海风博弈收益统计-新.xlsx"
        w = openpyxl.Workbook(); s = w.active; s.title = "10月份"
        for column, label in self.module.HAIFENG_HEADERS.items():
            s[f"{column}2"] = label
        for first, day in ((4, datetime(2026,10,16)), (29, "=A4+1")):
            s.cell(first, 1, day).number_format = "yyyy-mm-dd"
            for hour in range(24):
                row = first + hour
                s.cell(row, 2, f"{hour}:00")
                for col in self.module.HAIFENG_LAYOUT:
                    s[f"{col}{row}"] = 99
                    s[f"{col}{row}"].number_format = "0.0000"
                for col, source in (("M","T"),("N","U"),("O","V")):
                    s[f"{col}{row}"] = f"={source}{row}"
                for col in ("J","K","L","S","T","U","V"):
                    s[f"{col}{row}"] = 777
        w.save(path); w.close()
        return path

    def test_haifeng_overwrites_only_five_columns_in_selected_dates(self):
        m = self.module
        with tempfile.TemporaryDirectory() as folder:
            path = self.haifeng_template(folder)
            with self.assertRaisesRegex(ValueError,"开始日期"):
                m.inspect_workbook(path,end_date=date(2026,10,17))
            plan = m.inspect_workbook(path,start_date=date(2026,10,16),end_date=date(2026,10,16))
            for backend in ("stream", "full"):
                with patch.dict("os.environ", {"TOOLBOX_BENEFIT_INSPECT_BACKEND": backend}):
                    self.assertEqual(m.inspect_workbook(path,start_date=date(2026,10,16),end_date=date(2026,10,16)), plan)
            self.assertEqual(len(plan.cells),120)
            self.assertEqual({c.field for c in plan.cells},{"C","J","M","L","N"})
            data = {date(2026,10,16): {col: (float(index),)*24 for index,col in enumerate(("C","J","M","L","N"), 1)}}
            item = m.prepare_workbook(plan,data)
            w = openpyxl.load_workbook(item.temporary); s=w["10月份"]
            for col,expected in (("C",1),("R",2),("N",3),("M",4),("O",5)):
                self.assertEqual(s[f"{col}4"].value,expected)
                self.assertEqual(s[f"{col}4"].number_format,"0.0000")
            for col in ("J","K","L","S","T","U","V"):
                self.assertEqual(s[f"{col}4"].value,777)
            self.assertEqual(s["C29"].value,99)
            self.assertEqual(s["M29"].value,"=T29")
            w.close();item.temporary.unlink()

    def test_partial_shared_formula_overwrite_preserves_unavailable_followers(self):
        source = b'<worksheet><row r="4"><c r="M4" s="1"><f t="shared" ref="M4:O4" si="2">T4</f><v>1</v></c><c r="N4" s="1"><f t="shared" si="2"/><v>2</v></c><c r="O4" s="1"><f t="shared" si="2"/><v>3</v></c></row></worksheet>'
        result = self.module.patch_sparse_sheet(source,{"M4":300})
        self.assertIn(b'<c r="M4" s="1"><v>300</v></c>',result)
        self.assertIn(b'<f>U4</f><v>2</v>',result)
        self.assertIn(b'<f>V4</f><v>3</v>',result)


class PowerBenefitPageTests(unittest.TestCase):
    def test_page_factory_and_command(self):
        from toolbox.catalog import default_catalog
        from toolbox.runtime import TaskRegistry
        from toolbox.power_benefit_page import PowerBenefitUpdatePage
        root=tk.Tk(); root.withdraw()
        try:
            paths=ToolPaths(Path(__file__).resolve().parents[2])
            descriptor=next(d for d in default_catalog() if d.tool_id=="power-benefit")
            page=descriptor.factory()(root,paths,TaskRegistry())
            self.assertIsInstance(page,PowerBenefitUpdatePage)
            command=page.build_update_command([Path("/tmp/归湖.xlsx"),Path("/tmp/鮀莲.xlsx")],
                    date(2026,10,12),start_date=None,create_backup=True)
            self.assertEqual(command.count("--workbook"),2)
            self.assertNotIn("--no-backup",command)
            self.assertNotIn("--start-date",command)
            self.assertEqual(len(page.workbooks),3)
            self.assertIn("海风博弈收益统计-新.xlsx",page.workbooks[2].get())
            self.assertFalse(page.selected[2].get())
            page.destroy()
        finally: root.destroy()


class PowerBenefitDesktopTests(unittest.TestCase):
    def test_service_routes_haifeng_and_requires_explicit_overwrite_dates(self):
        from toolbox.service_catalog import desktop_catalog
        from toolbox.workflows import validate, plan
        paths = ToolPaths(Path(__file__).resolve().parents[2])
        tool = next(t for t in desktop_catalog(paths) if t["id"] == "power-benefit")
        self.assertIn("haifeng_workbook", {f["key"] for f in tool["fields"]})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"海风博弈收益统计-新.xlsx";path.touch()
            values = dict(guihu_selected=False,tuolian_selected=False,haifeng_selected=True,
                          haifeng_workbook=str(path),end_date="2026-10-16")
            with self.assertRaisesRegex(ValueError,"开始日期"):
                validate(paths,"power-benefit","run",values)
            values["start_date"] = "2026-10-16"
            params = validate(paths,"power-benefit","run",values)
            command, _ = plan(paths,"power-benefit","run",params)[0]
            self.assertIn(str(path.resolve()),command)
            self.assertEqual(command.count("--workbook"),1)


if __name__=="__main__": unittest.main()
