from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

import numpy as np
import pandas as pd
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[2] / '广东电价预测'


class ProductionForecastTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(PROJECT))

    @classmethod
    def tearDownClass(cls):
        if sys.path[0] == str(PROJECT):
            sys.path.pop(0)

    def sample(self):
        from gd_price_forecast.features import FORECAST_COLUMNS
        end = pd.Timestamp.now(tz='Asia/Shanghai').normalize().tz_localize(None)
        times = pd.date_range(end-pd.Timedelta(days=49), periods=50*96, freq='15min')
        frame = pd.DataFrame({'日期':times.normalize(), '时刻':times.strftime('%H:%M'),
                              '全省日前平均电价':300+20*np.sin(np.arange(len(times))/96)})
        for number, column in enumerate(FORECAST_COLUMNS):
            frame[column] = 10000.+number*100+np.arange(len(times))
        target = end+pd.Timedelta(days=1)
        snap = frame.iloc[-96:].copy(); snap['日期'] = target
        snap['全省日前平均电价'] = 999999.
        cutoff = pd.Timestamp.now(tz='Asia/Shanghai')-pd.Timedelta(minutes=1)
        receipt = dict(target_date=str(target.date()), cutoff=cutoff.isoformat(),
                       publications=[dict(field=c,published_at=(cutoff-pd.Timedelta(minutes=1)).isoformat())
                                     for c in FORECAST_COLUMNS])
        return frame, snap, receipt

    def run_forecast(self, frame, output, **kwargs):
        from train_d1_forecast import run_pipeline
        with mock.patch('train_d1_forecast.load_model_data',return_value=frame):
            return run_pipeline('unused.xlsx',output,
                configs=[dict(name='smoke',kind='hist_gradient_boosting',params=dict(max_iter=5))],
                **kwargs)

    def test_complete_snapshot_uses_supply_model_and_preserves_output_contract(self):
        frame, snap, receipt = self.sample()
        with tempfile.TemporaryDirectory() as folder:
            summary = self.run_forecast(frame,folder,future_snapshot=snap,snapshot_receipt=receipt)
            self.assertEqual('当日供需_12月',summary['strategy'])
            self.assertEqual(12,summary['training_months'])
            self.assertEqual(1.,summary['model_weight'])
            self.assertIsNone(summary['fallback_reason'])
            self.assertFalse(summary['historical_availability_verified'])
            self.assertEqual(receipt['cutoff'],summary['snapshot_cutoff'])
            book = load_workbook(Path(folder)/'D+1日前电价.xlsx',read_only=True,data_only=True)
            try:
                sheet = book.active
                self.assertEqual((97,3),(sheet.max_row,sheet.max_column))
                self.assertEqual(['日期','时刻','日前电价'],[c.value for c in sheet[1]])
                self.assertEqual(pd.Timestamp(receipt['target_date']),sheet['A2'].value)
                self.assertEqual('23:45',sheet['B97'].value)
                self.assertEqual('0.00',sheet['C2'].number_format)
                self.assertTrue(np.isfinite(sheet['C2'].value))
                self.assertLess(sheet['C2'].value,999999.)
            finally:
                book.close()

    def test_partial_snapshot_falls_back_without_training_on_missing_current_inputs(self):
        frame, snap, receipt = self.sample()
        snap.loc[snap.index[0],'日前统调负荷(MW)'] = np.nan
        with tempfile.TemporaryDirectory() as folder:
            summary = self.run_forecast(frame,folder,future_snapshot=snap,snapshot_receipt=receipt)
        self.assertEqual(3,summary['training_months'])
        self.assertEqual(.25,summary['model_weight'])
        self.assertIn('非有限值',summary['fallback_reason'])
        self.assertIsNone(summary['snapshot_cutoff'])

    def test_fetch_failure_still_exports_historical_forecast(self):
        frame, _, _ = self.sample()
        with tempfile.TemporaryDirectory() as folder, mock.patch(
            'train_d1_forecast.fetch_forecast_snapshot',side_effect=TimeoutError('network timeout')
        ) as fetch:
            summary = self.run_forecast(frame,folder)
            fetch.assert_called_once()
            self.assertTrue((Path(folder)/'D+1日前电价.xlsx').is_file())
        self.assertEqual(.25,summary['model_weight'])
        self.assertEqual('TimeoutError',summary['snapshot_error_type'])

    def test_historical_only_never_calls_remote_api(self):
        frame, _, _ = self.sample()
        with tempfile.TemporaryDirectory() as folder, mock.patch('train_d1_forecast.fetch_forecast_snapshot') as fetch:
            summary = self.run_forecast(frame,folder,use_future_inputs=False)
            fetch.assert_not_called()
        self.assertEqual(.25,summary['model_weight'])

    def test_snapshot_requires_publications_for_every_field_and_canonical_slots(self):
        from gd_price_forecast.future_inputs import validate_forecast_snapshot
        _, snap, receipt = self.sample()
        receipt['publications'].pop()
        with self.assertRaisesRegex(ValueError,'发布时间记录未齐全'):
            validate_forecast_snapshot(snap,receipt,receipt['target_date'])
        snap.loc[snap.index[0],'时刻'] = '00:07'
        with self.assertRaisesRegex(ValueError,'96个唯一时点'):
            validate_forecast_snapshot(snap,receipt,receipt['target_date'])

    def test_post_cutoff_publication_cannot_activate_supply_model(self):
        from gd_price_forecast.future_inputs import validate_forecast_snapshot
        _, snap, receipt = self.sample()
        receipt['publications'][0]['published_at'] = (pd.Timestamp(receipt['cutoff'])+pd.Timedelta(seconds=1)).isoformat()
        with self.assertRaisesRegex(ValueError,'晚于快照截止时间'):
            validate_forecast_snapshot(snap,receipt,receipt['target_date'])

    def test_snapshot_captured_after_historical_target_day_is_rejected(self):
        from gd_price_forecast.future_inputs import validate_forecast_snapshot
        _, snap, receipt = self.sample()
        target = pd.Timestamp(receipt['target_date'])-pd.Timedelta(days=1)
        snap['日期'] = target
        receipt['target_date'] = str(target.date())
        with self.assertRaisesRegex(ValueError,'预测日开始之前'):
            validate_forecast_snapshot(snap,receipt,receipt['target_date'])


if __name__ == '__main__':
    unittest.main()
