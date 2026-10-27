#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import sys
import uuid
import zipfile
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel, to_excel
from playwright.sync_api import sync_playwright


SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SCRIPT_DIR.parent
if str(WORKSPACE_DIR) not in sys.path:
    sys.path.insert(0, str(WORKSPACE_DIR))
from toolbox_shared.auth_loader import load_login_module, standalone_login
from toolbox_shared.dates import parse_date_input
from toolbox_shared.credentials import apply_credentials
from toolbox_shared.excel_read import read_cell_snapshots
from toolbox_shared.native_excel import load_value_workbook
from toolbox_shared.workbook_commit import (
    FileFingerprint, assert_workbook_unchanged, commit_workbooks, workbook_fingerprint,
)

BASE_URL = "https://xhxt.chng.com.cn"
APP_URL = f"{BASE_URL}/gdfire/PublicDataManage/SpotReport"
LOAD_API_URL = f"{BASE_URL}/gdfire/api/data/net/load"
from toolbox_shared.storage import data_path

PROFILE_DIR = data_path(SCRIPT_DIR / ".browser-profile")
AUTH_STATE_PATH = data_path(SCRIPT_DIR / "auth_state.json")
AUTH_LOCK_PATH = Path(str(AUTH_STATE_PATH) + ".lock")
CONFIG_PATH = data_path(SCRIPT_DIR / "config.json")
FALLBACK_CONFIG_PATH = data_path(WORKSPACE_DIR / "上网电量抓取" / "config.json")

DEFAULT_WORKBOOK = (
    Path.home() / "Desktop" / "日常工作" / "交易系统信息收集20260901.xlsx"
)
SHEET_NAME = "Sheet1"
DEFAULT_PROVINCE_AREA_ID = "044"
DATE_OFFSET_DAYS = 3
GROUP_HEIGHT = 7
CATEGORY_ORDER = ("统调", "B类", "西电", "A类", "地方电", "粤港联络线")
LOAD_TYPE_BY_CATEGORY = {
    "统调": "1",
    "B类": "6",
    "西电": "4",
    "A类": "2",
    "地方电": "3",
    "粤港联络线": "5",
}
DAY_AHEAD_REQUIRED = ("统调", "B类", "A类", "地方电", "粤港联络线")


@dataclass(frozen=True, slots=True)
class Config:
    username: str = ""
    password: str = ""
    headless: bool = True


@dataclass(frozen=True, slots=True)
class DailyLoadPair:
    day_ahead_date: date
    actual_date: date
    day_ahead: dict[str, tuple[int | float, ...]]
    actual: dict[str, tuple[int | float, ...]]


@dataclass(frozen=True, slots=True)
class WorkbookTail:
    start_row: int
    day_ahead_date: date
    actual_date: date


@dataclass(frozen=True, slots=True)
class PreparedWorkbook:
    original: Path
    temporary: Path
    source_fingerprint: FileFingerprint
    first_row: int
    last_row: int
    pairs: tuple[DailyLoadPair, ...]


def log(message: str) -> None:
    print(message, flush=True)


def _load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload if isinstance(payload, dict) else {}


def load_config() -> Config:
    raw = apply_credentials({**_load_json(FALLBACK_CONFIG_PATH), **_load_json(CONFIG_PATH)})
    return Config(
        username=str(raw.get("username") or ""),
        password=str(raw.get("password") or ""),
        headless=bool(raw.get("headless", True)),
    )


def load_auth_module():
    return load_login_module(
        WORKSPACE_DIR,
        module_name="transaction_info_auth",
        state_dir=SCRIPT_DIR,
        config_path=CONFIG_PATH,
        app_url=APP_URL,
    )


def validate_date(value: str) -> date:
    return parse_date_input(value)


def date_range(start_date: date, end_date: date) -> tuple[date, ...]:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期。")
    return tuple(
        start_date + timedelta(days=offset)
        for offset in range((end_date - start_date).days + 1)
    )


def _number(value: Any, *, label: str) -> int | float:
    if value in (None, "", "-") or isinstance(value, bool):
        raise RuntimeError(f"{label}为空。")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"{label}不是有效数值：{value!r}") from exc
    if not math.isfinite(number):
        raise RuntimeError(f"{label}不是有限数值：{value!r}")
    rounded = round(number, 10)
    if rounded.is_integer():
        return int(rounded)
    return rounded


def _payload_rows(dataset: dict[str, Any]) -> list[dict[str, Any]]:
    data = dataset.get("data")
    if isinstance(data, dict):
        rows = data.get("dataNetLoadDTOList") or []
    else:
        rows = []
    if not isinstance(rows, list):
        raise RuntimeError("市场负荷接口返回格式异常。")
    return [row for row in rows if isinstance(row, dict)]


def _period_date(value: Any) -> date | None:
    text = str(value or "")[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None


def summarize_load_type(
    dataset: dict[str, Any],
    load_type: str,
    run_date: date,
    *,
    period_list_key: str,
    label: str,
) -> tuple[int | float, ...]:
    series = next(
        (row for row in _payload_rows(dataset) if str(row.get("loadType") or "") == load_type),
        None,
    )
    if series is None:
        available = "、".join(
            sorted({str(row.get("loadType")) for row in _payload_rows(dataset) if row.get("loadType") is not None})
        )
        raise RuntimeError(
            f"{run_date.isoformat()} {label}缺少 loadType={load_type}，"
            f"可用类型：{available or '无'}"
        )
    periods = series.get(period_list_key) or []
    if not isinstance(periods, list):
        periods = []
    summary = next(
        (row for row in periods if isinstance(row, dict) and _period_date(row.get("date")) == run_date),
        None,
    )
    if summary is None and len(periods) == 1 and isinstance(periods[0], dict):
        summary = periods[0]
    if summary is None:
        raise RuntimeError(f"{run_date.isoformat()} loadType={load_type}没有返回{label}数据。")
    values = tuple(
        _number(summary.get(field), label=f"{run_date.isoformat()} {label} loadType={load_type} {field}")
        for field in ("max", "min", "avg", "sum")
    )
    # 网页的总电量单位是亿 kWh，目标工作簿按 MWh 存储。
    return (*values[:3], _number(values[3] * 100_000, label=f"{run_date.isoformat()} {label} 总电量"))


def parse_load_dataset(
    dataset: dict[str, Any],
    run_date: date,
    *,
    day_ahead: bool,
) -> dict[str, tuple[int | float, ...]]:
    categories = DAY_AHEAD_REQUIRED if day_ahead else CATEGORY_ORDER
    period_key = "forecastPeriodList" if day_ahead else "actualPeriodList"
    label = "日前" if day_ahead else "实际"
    return {
        category: summarize_load_type(
            dataset,
            LOAD_TYPE_BY_CATEGORY[category],
            run_date,
            period_list_key=period_key,
            label=label,
        )
        for category in categories
    }


def build_daily_load_pair(
    day_ahead_date: date,
    day_payload: dict[str, Any],
    actual_payload: dict[str, Any],
) -> DailyLoadPair:
    actual_date = day_ahead_date - timedelta(days=DATE_OFFSET_DAYS)
    actual = parse_load_dataset(actual_payload, actual_date, day_ahead=False)
    day_ahead = parse_load_dataset(day_payload, day_ahead_date, day_ahead=True)
    return DailyLoadPair(
        day_ahead_date=day_ahead_date,
        actual_date=actual_date,
        day_ahead=day_ahead,
        actual=actual,
    )


def fetch_load_dataset(context, auth, run_date: date, *, day_ahead: bool) -> dict[str, Any]:
    params: dict[str, Any] = {
        "startDate": run_date.isoformat(),
        "endDate": run_date.isoformat(),
        "provinceAreaId": DEFAULT_PROVINCE_AREA_ID,
        "timeSegment": 2,
    }
    if day_ahead:
        params["type"] = 1
    response = context.request.get(LOAD_API_URL, params=params, timeout=60_000)
    return auth.parse_json_response(
        response,
        f"读取{'日前' if day_ahead else '实际'}市场负荷 {run_date.isoformat()}",
    )


def collect_pairs(
    start_date: date,
    end_date: date,
    *,
    config: Config | None = None,
) -> tuple[DailyLoadPair, ...]:
    config = config or load_config()
    auth = load_auth_module()
    pairs: list[DailyLoadPair] = []
    with auth.AuthStateLock():
        with sync_playwright() as playwright:
            context = auth.launch_context(playwright, headless=config.headless)
            try:
                auth.ensure_login(context, config)
                for day_ahead_date in date_range(start_date, end_date):
                    actual_date = day_ahead_date - timedelta(days=DATE_OFFSET_DAYS)
                    log(
                        f"正在读取：日前 {day_ahead_date.isoformat()} / "
                        f"实际 {actual_date.isoformat()}"
                    )
                    day_payload = fetch_load_dataset(
                        context, auth, day_ahead_date, day_ahead=True
                    )
                    actual_payload = fetch_load_dataset(
                        context, auth, actual_date, day_ahead=False
                    )
                    pairs.append(
                        build_daily_load_pair(
                            day_ahead_date, day_payload, actual_payload
                        )
                    )
                auth.save_auth_state(context)
            finally:
                context.close()
    return tuple(pairs)


def _as_date(value: Any, *, epoch) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
        except (TypeError, ValueError, OverflowError):
            return None
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def inspect_workbook_tail(path: Path) -> WorkbookTail:
    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"仅支持 .xlsx 工作簿：{path}")
    workbook = load_value_workbook(path, data_only=True, read_only=True, native_in_auto=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(f"工作簿中没有工作表 {SHEET_NAME}。")
        sheet = workbook[SHEET_NAME]
        tail = None
        labels, right_labels = [], []
        for row, values in enumerate(sheet.iter_rows(
            min_row=1, max_row=sheet.max_row, min_col=3, max_col=11, values_only=True,
        ), start=1):
            if values[1] == "统调" and values[8] == "统调":
                day_ahead_date = _as_date(values[0], epoch=workbook.epoch)
                actual_date = _as_date(values[7], epoch=workbook.epoch)
                if day_ahead_date and actual_date:
                    tail = WorkbookTail(row, day_ahead_date, actual_date)
                    labels, right_labels = [], []
            if tail is not None and row < tail.start_row + len(CATEGORY_ORDER):
                labels.append(values[1])
                right_labels.append(values[8])
        if tail is None:
            raise RuntimeError(f"{SHEET_NAME} 中未找到可接续的日前/实际数据组。")
        if tail.day_ahead_date - tail.actual_date != timedelta(days=DATE_OFFSET_DAYS):
            raise RuntimeError(
                f"最后一组日期不符合 D / D-{DATE_OFFSET_DAYS}："
                f"日前 {tail.day_ahead_date} / 实际 {tail.actual_date}"
            )
        labels = tuple(labels + [None] * (len(CATEGORY_ORDER) - len(labels)))
        right_labels = tuple(right_labels + [None] * (len(CATEGORY_ORDER) - len(right_labels)))
        if labels != CATEGORY_ORDER or right_labels != CATEGORY_ORDER:
            raise RuntimeError(
                f"最后一组字段顺序异常：日前 {labels} / 实际 {right_labels}"
            )
        return tail
    finally:
        workbook.close()


def resolve_start_date(path: Path, requested: date | None) -> date:
    expected = inspect_workbook_tail(path).day_ahead_date + timedelta(days=1)
    if requested is not None and requested != expected:
        raise ValueError(
            f"为保持日期连续，开始日期必须是 {expected.isoformat()}；"
            f"当前输入为 {requested.isoformat()}。"
        )
    return expected


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def xlsx_sheet_paths(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    result: dict[str, str] = {}
    sheets = workbook_root.find(f"{{{MAIN_NS}}}sheets")
    if sheets is None:
        return result
    for sheet in sheets.findall(f"{{{MAIN_NS}}}sheet"):
        rel_id = sheet.attrib.get(f"{{{REL_NS}}}id")
        target = relationships.get(str(rel_id), "").lstrip("/")
        if target:
            result[str(sheet.attrib.get("name") or "")] = (
                target if target.startswith("xl/") else f"xl/{target}"
            )
    return result


def _row_xml(xml_text: str, row_number: int) -> str:
    pattern = re.compile(
        rf'<row\b(?=[^>]*\br="{row_number}")[^>]*>.*?</row>', re.DOTALL
    )
    match = pattern.search(xml_text)
    if match is None:
        raise RuntimeError(f"工作表 XML 未找到模板第 {row_number} 行。")
    return match.group(0)


def _renumber_row(row_xml: str, source_row: int, target_row: int) -> str:
    row_xml = re.sub(
        rf'(<row\b[^>]*\br="){source_row}("[^>]*>)',
        rf'\g<1>{target_row}\2',
        row_xml,
        count=1,
    )
    return re.sub(rf'([A-Z]+){source_row}(?=")', rf'\g<1>{target_row}', row_xml)


def _cell_pattern(cell_ref: str) -> re.Pattern[str]:
    return re.compile(
        rf'<c\b(?=[^>]*\br="{re.escape(cell_ref)}")[^>]*(?:>.*?</c>|/>)',
        re.DOTALL,
    )


def _replace_numeric_cell(row_xml: str, cell_ref: str, value: int | float) -> str:
    pattern = _cell_pattern(cell_ref)
    match = pattern.search(row_xml)
    if match is None:
        raise RuntimeError(f"模板行缺少单元格 {cell_ref}。")
    style_match = re.search(r'\bs="([^"]+)"', match.group(0))
    style = f' s="{style_match.group(1)}"' if style_match else ""
    replacement = f'<c r="{cell_ref}"{style}><v>{format(float(value), ".15g")}</v></c>'
    return row_xml[: match.start()] + replacement + row_xml[match.end() :]


def _build_group_rows(
    xml_text: str,
    *,
    template_start: int,
    target_start: int,
    pair: DailyLoadPair,
    epoch,
) -> str:
    rows: list[str] = []
    for offset, category in enumerate(CATEGORY_ORDER):
        source_row = template_start + offset
        target_row = target_start + offset
        row_xml = _renumber_row(_row_xml(xml_text, source_row), source_row, target_row)
        if offset == 0:
            row_xml = _replace_numeric_cell(
                row_xml, f"C{target_row}", to_excel(pair.day_ahead_date, epoch)
            )
            row_xml = _replace_numeric_cell(
                row_xml, f"J{target_row}", to_excel(pair.actual_date, epoch)
            )
        if category != "西电":
            for column, value in zip(("E", "F", "G", "H"), pair.day_ahead[category]):
                row_xml = _replace_numeric_cell(row_xml, f"{column}{target_row}", value)
        for column, value in zip(("L", "M", "N", "O"), pair.actual[category]):
            row_xml = _replace_numeric_cell(row_xml, f"{column}{target_row}", value)
        rows.append(row_xml)
    return "".join(rows)


def patch_sheet_xml(
    xml_bytes: bytes,
    *,
    template_start: int,
    pairs: tuple[DailyLoadPair, ...],
    epoch,
) -> tuple[bytes, int, int]:
    xml_text = xml_bytes.decode("utf-8")
    first_row = template_start + GROUP_HEIGHT
    appended: list[str] = []
    for index, pair in enumerate(pairs):
        target_start = first_row + index * GROUP_HEIGHT
        appended.append(
            _build_group_rows(
                xml_text,
                template_start=template_start,
                target_start=target_start,
                pair=pair,
                epoch=epoch,
            )
        )
    if "</sheetData>" not in xml_text:
        raise RuntimeError("工作表 XML 缺少 sheetData 结束标记。")
    xml_text = xml_text.replace("</sheetData>", "".join(appended) + "</sheetData>", 1)
    last_row = first_row + (len(pairs) - 1) * GROUP_HEIGHT + 5
    dimension = re.search(r'<dimension\s+ref="([^"]+)"', xml_text)
    if dimension is None:
        raise RuntimeError("工作表 XML 缺少 dimension。")
    old_ref = dimension.group(1)
    new_ref = re.sub(r"\d+$", str(last_row), old_ref)
    xml_text = xml_text[: dimension.start(1)] + new_ref + xml_text[dimension.end(1) :]
    return xml_text.encode("utf-8"), first_row, last_row


def verify_package_preservation(original: Path, patched: Path, changed_part: str) -> None:
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(patched) as result:
        if source.namelist() != result.namelist():
            raise RuntimeError(f"XLSX 包结构发生变化：{original.name}")
        for name in source.namelist():
            if name == changed_part:
                continue
            if source.read(name) != result.read(name):
                raise RuntimeError(f"XLSX 非目标部件被改动：{original.name} / {name}")


def _write_patched_workbook(
    original: Path,
    temporary: Path,
    *,
    template_start: int,
    pairs: tuple[DailyLoadPair, ...],
) -> tuple[int, int]:
    sheet_part = xlsx_sheet_paths(original).get(SHEET_NAME)
    if not sheet_part:
        raise RuntimeError(f"XLSX 包内未找到工作表：{SHEET_NAME}")
    workbook = openpyxl.load_workbook(original, data_only=False, read_only=True)
    try:
        epoch = workbook.epoch
    finally:
        workbook.close()
    bounds: tuple[int, int] | None = None
    with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(temporary, "w") as output:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == sheet_part:
                data, first_row, last_row = patch_sheet_xml(
                    data,
                    template_start=template_start,
                    pairs=pairs,
                    epoch=epoch,
                )
                bounds = (first_row, last_row)
            output.writestr(info, data)
    if bounds is None:
        raise RuntimeError(f"未改动目标工作表：{SHEET_NAME}")
    verify_package_preservation(original, temporary, sheet_part)
    return bounds


def _values_equal(actual: Any, expected: int | float) -> bool:
    if isinstance(actual, bool):
        return False
    try:
        return math.isclose(float(actual), float(expected), rel_tol=0, abs_tol=1e-7)
    except (TypeError, ValueError):
        return False


def verify_workbook(path: Path, prepared: PreparedWorkbook) -> None:
    workbook = load_value_workbook(path, data_only=True, read_only=True, native_in_auto=False)
    try:
        columns = (3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 14, 15)
        snapshots = read_cell_snapshots(workbook, (
            (SHEET_NAME, f"{get_column_letter(column)}{row}")
            for row in range(prepared.first_row, prepared.first_row + len(prepared.pairs) * GROUP_HEIGHT)
            for column in columns
        ))

        def value(row, column):
            return snapshots[(SHEET_NAME, f"{get_column_letter(column)}{row}")][0]

        for index, pair in enumerate(prepared.pairs):
            start = prepared.first_row + index * GROUP_HEIGHT
            day_date = _as_date(value(start, 3), epoch=workbook.epoch)
            actual_date = _as_date(value(start, 10), epoch=workbook.epoch)
            if day_date != pair.day_ahead_date or actual_date != pair.actual_date:
                raise RuntimeError(
                    f"日期回读失败：第 {start} 行，"
                    f"日前 {day_date} / 实际 {actual_date}"
                )
            for offset, category in enumerate(CATEGORY_ORDER):
                row = start + offset
                if value(row, 4) != category or value(row, 11) != category:
                    raise RuntimeError(f"字段顺序回读失败：第 {row} 行。")
                left = tuple(value(row, column) for column in range(5, 9))
                if category == "西电":
                    if any(value not in (None, "") for value in left):
                        raise RuntimeError(f"日前西电应留空：第 {row} 行。")
                elif not all(
                    _values_equal(actual, expected)
                    for actual, expected in zip(left, pair.day_ahead[category])
                ):
                    raise RuntimeError(f"日前数值回读失败：第 {row} 行。")
                right = tuple(value(row, column) for column in range(12, 16))
                if not all(
                    _values_equal(actual, expected)
                    for actual, expected in zip(right, pair.actual[category])
                ):
                    raise RuntimeError(f"实际数值回读失败：第 {row} 行。")
    finally:
        workbook.close()


def prepare_workbook(path: Path, pairs: Iterable[DailyLoadPair]) -> PreparedWorkbook:
    original = Path(path).expanduser().resolve()
    pair_tuple = tuple(pairs)
    if not pair_tuple:
        raise ValueError("没有可写入的日期数据。")
    source_fingerprint = workbook_fingerprint(original)
    tail = inspect_workbook_tail(original)
    expected_start = tail.day_ahead_date + timedelta(days=1)
    if pair_tuple[0].day_ahead_date != expected_start:
        raise ValueError(
            f"新数据必须从 {expected_start.isoformat()} 开始，"
            f"实际为 {pair_tuple[0].day_ahead_date.isoformat()}。"
        )
    for previous, current in zip(pair_tuple, pair_tuple[1:]):
        if current.day_ahead_date != previous.day_ahead_date + timedelta(days=1):
            raise ValueError("待写入的日前日期必须连续。")
    for pair in pair_tuple:
        if pair.day_ahead_date - pair.actual_date != timedelta(days=DATE_OFFSET_DAYS):
            raise ValueError("每组数据必须满足日前 D、实际 D-3。")
    temporary = original.with_name(f".{original.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        first_row, last_row = _write_patched_workbook(
            original,
            temporary,
            template_start=tail.start_row,
            pairs=pair_tuple,
        )
        prepared = PreparedWorkbook(
            original=original,
            temporary=temporary,
            source_fingerprint=source_fingerprint,
            first_row=first_row,
            last_row=last_row,
            pairs=pair_tuple,
        )
        verify_workbook(temporary, prepared)
        assert_workbook_unchanged(original, source_fingerprint)
        return prepared
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def commit_prepared_workbook(
    prepared: PreparedWorkbook,
    *,
    create_backup: bool = True,
) -> Path | None:
    return commit_workbooks(
        (prepared,), lambda path, item: verify_workbook(path, item),
        create_backup=create_backup,
    )[prepared.original]


def update_workbook(
    workbook_path: Path,
    start_date: date | None,
    end_date: date,
    *,
    config: Config | None = None,
    create_backup: bool = True,
) -> PreparedWorkbook:
    start = resolve_start_date(workbook_path, start_date)
    if end_date < start:
        raise ValueError(
            f"结束日期 {end_date.isoformat()} 早于待接续日期 "
            f"{start.isoformat()}，无需更新。"
        )
    pairs = collect_pairs(start, end_date, config=config)
    prepared = prepare_workbook(workbook_path, pairs)
    backup = commit_prepared_workbook(prepared, create_backup=create_backup)
    log(
        f"已更新 {prepared.original}：日前 {start.isoformat()} 至 "
        f"{end_date.isoformat()}，实际 {(start - timedelta(days=3)).isoformat()} 至 "
        f"{(end_date - timedelta(days=3)).isoformat()}；写入第 {prepared.first_row}-{prepared.last_row} 行。"
    )
    if backup:
        log(f"备份：{backup}")
    return prepared


def interactive_login(config: Config | None = None) -> None:
    standalone_login(
        load_auth_module(),
        config or load_config(),
        app_url=APP_URL,
        log=log,
        label="交易系统信息更新模块的独立",
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="抓取现货运行日报市场负荷信息并接续更新交易系统信息收集表"
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--start-date", help="日前开始日期；不填时自动从工作簿最后日期+1接续")
    parser.add_argument(
        "--end-date",
        default=(date.today() - timedelta(days=1)).isoformat(),
        help="日前结束日期 YYYY-MM-DD，默认昨日",
    )
    parser.add_argument("--no-backup", action="store_true", help="不保留时间戳备份")
    parser.add_argument("--headed", action="store_true", help="显示抓取浏览器")
    parser.add_argument("--dry-run", action="store_true", help="只抓取并校验，不写入 Excel")
    parser.add_argument("--login", action="store_true", help="登录并保存本模块独立会话")
    args = parser.parse_args()
    try:
        config = replace(load_config(), headless=not args.headed)
        if args.login:
            interactive_login(config)
            return 0
        requested_start = validate_date(args.start_date) if args.start_date else None
        start_date = resolve_start_date(args.workbook, requested_start)
        end_date = validate_date(args.end_date)
        if end_date < start_date:
            log(
                f"无需更新：工作簿已接续至日前 "
                f"{(start_date - timedelta(days=1)).isoformat()}，"
                f"所选结束日期为 {end_date.isoformat()}。"
            )
            return 0
        if args.dry_run:
            pairs = collect_pairs(start_date, end_date, config=config)
            log(
                f"校验通过：{len(pairs)} 组，日前 {start_date} 至 {end_date}，"
                f"实际 {start_date - timedelta(days=3)} 至 {end_date - timedelta(days=3)}。"
            )
            return 0
        update_workbook(
            args.workbook,
            start_date,
            end_date,
            config=config,
            create_backup=not args.no_backup,
        )
        return 0
    except KeyboardInterrupt:
        print("已取消。", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
