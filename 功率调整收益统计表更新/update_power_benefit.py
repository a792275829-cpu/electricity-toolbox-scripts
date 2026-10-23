#!/usr/bin/env python3
"""归湖、鮀莲补齐六列；海风按指定日期覆盖 C/R/N/M/O，保留其他 XLSX 部件。"""
from __future__ import annotations

import argparse
import importlib.util
import math
import os
import warnings
import re
import stat
import sys
import uuid
import zipfile
from html import escape, unescape
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import from_excel
from playwright.sync_api import sync_playwright

SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE = SCRIPT_DIR.parent
if str(WORKSPACE) not in sys.path:
    sys.path.insert(0, str(WORKSPACE))
from toolbox_auth_loader import load_login_module
from toolbox_credentials import read_login_config
from toolbox_dates import parse_date_input
from toolbox_excel_read import read_cell_snapshots
from toolbox_native_excel import load_value_workbook
from toolbox_workbook_inspection import load_inspection_pair, UnsupportedInspection
from toolbox_workbook_commit import (
    FileFingerprint, assert_workbook_unchanged, commit_workbooks, workbook_fingerprint,
)

BASE_URL = "https://xhxt.chng.com.cn"
APP_URL = f"{BASE_URL}/gdfire/PrivateDataManageCE/SpotBasicInfoSubmit"
FORECAST_API = f"{BASE_URL}/gdfire/api/data/personal/newEnergy/productRun/"
TRADE_API = f"{BASE_URL}/gdfire/api/data/spot/trade/unit"
from toolbox_storage import data_path

CONFIG_PATH = data_path(SCRIPT_DIR / "config.json")
DEFAULT_DIRECTORY = Path.home() / "Desktop" / "日常工作"
DEFAULT_WORKBOOKS = (
    DEFAULT_DIRECTORY / "功率预测调整收益统计-归湖.xlsx",
    DEFAULT_DIRECTORY / "功率预测调整收益统计-鮀莲(2)(2).xlsx",
)
FIELDS = {"C": "原始功率", "J": "申报功率", "K": "日前电量",
          "L": "日前电价", "M": "上网电量", "N": "实时电价"}
HAIFENG_WORKBOOK = DEFAULT_DIRECTORY / "海风博弈收益统计-新.xlsx"
# 实际目标列 -> 既有采集字段；避免将海风 N 上网电量与原模板 N 实时电价混淆。
HAIFENG_LAYOUT = {"C": "C", "R": "J", "N": "M", "M": "L", "O": "N"}
HAIFENG_HEADERS = {"C": "功率预测", "R": "功率预测", "N": "上网电量", "M": "日前电价", "O": "实时电价"}


def column_layout(station: str) -> dict[str, str]:
    return HAIFENG_LAYOUT if station == "海风" else {col: col for col in FIELDS}


HEADERS = {"C": "功率预测", "J": "功率预测", "K": "日前电量",
           "L": "日前电价", "M": "上网电量", "N": "实时电价"}


@dataclass(frozen=True)
class CellTarget:
    sheet: str
    coordinate: str
    run_date: date
    hour: int
    field: str
    number_format: str
    style_id: int


@dataclass(frozen=True)
class WorkbookPlan:
    path: Path
    station: str
    cells: tuple[CellTarget, ...]
    fingerprint: tuple[int, int]


@dataclass(frozen=True)
class PreparedWorkbook:
    original: Path
    temporary: Path
    expectations: tuple[tuple[CellTarget, float], ...]
    source_fingerprint: FileFingerprint


def log(message: str) -> None:
    print(message, flush=True)


@lru_cache(maxsize=1)
def shared_spot():
    """复用已有公司的场站匹配与 XLSX 包校验逻辑。"""
    spec = importlib.util.spec_from_file_location(
        "power_benefit_shared_spot", WORKSPACE / "现货价格更新" / "update_spot_prices.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法加载现货价格更新模块。")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_auth_module():
    return load_login_module(WORKSPACE, module_name="power_benefit_auth",
                             state_dir=SCRIPT_DIR, config_path=CONFIG_PATH, app_url=APP_URL)


def station_from_filename(path: Path) -> str:
    stations = [key for key in ("归湖", "鮀莲", "海风") if key in path.stem]
    if len(stations) != 1:
        raise ValueError(f"文件名必须包含且仅包含归湖、鮀莲或海风中的一个名称：{path.name}")
    return stations[0]


def fingerprint(path: Path) -> tuple[int, int]:
    info = path.stat()
    return info.st_size, info.st_mtime_ns


def make_writable(path: Path) -> None:
    info = path.stat()
    immutable = getattr(stat, "UF_IMMUTABLE", 0)
    if getattr(info, "st_flags", 0) & immutable:
        path.chflags(info.st_flags & ~immutable)
    path.chmod(stat.S_IMODE(info.st_mode) | stat.S_IWUSR)


def as_date(value: Any, epoch) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
        return value.date() if isinstance(value, datetime) else None
    if isinstance(value, str) and not value.startswith("="):
        try:
            return parse_date_input(value.strip()[:10])
        except ValueError:
            pass
    return None


def date_at(sheet, cached, row: int, epoch, seen: set[int] | None = None) -> date | None:
    """模板 A 列合并日期可含 =A上一日期行+1，缓存缺失也能解析。"""
    value = sheet.cell(row, 1).value
    if isinstance(value, str) and value.startswith("="):
        match = re.fullmatch(r"=\$?A\$?(\d+)\s*\+\s*(\d+)", value, re.I)
        if match:
            seen = set() if seen is None else seen
            if row in seen:
                raise ValueError(f"{sheet.title}!A{row} 的日期公式循环引用。")
            seen.add(row)
            base = date_at(sheet, cached, int(match[1]), epoch, seen)
            return base + timedelta(days=int(match[2])) if base else None
        result = as_date(cached.cell(row, 1).value, epoch)
        if result is None:
            raise ValueError(f"不能解析日期公式：{sheet.title}!A{row} {value}")
        return result
    return as_date(value, epoch)


def hour_of(value: Any) -> int | None:
    if isinstance(value, (datetime, time)):
        return value.hour if value.minute == value.second == value.microsecond == 0 else None
    match = re.fullmatch(r"(\d{1,2}):00(?::00)?", str(value).strip())
    if match and 0 <= int(match[1]) < 24:
        return int(match[1])
    if isinstance(value, (int, float)) and 0 <= value < 1:
        hour = round(value * 24)
        if hour < 24 and math.isclose(value * 24, hour, abs_tol=1e-8):
            return hour
    return None


def inspect_workbook(path: Path, *, end_date: date, start_date: date | None = None) -> WorkbookPlan:
    path = path.expanduser().resolve(strict=True)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"只支持 .xlsx：{path}")
    station = station_from_filename(path)
    if station == "海风" and start_date is None:
        raise ValueError("海风覆盖更新必须指定开始日期，避免覆盖未选定的历史数据。")
    layout = column_layout(station)
    headers = HAIFENG_HEADERS if station == "海风" else HEADERS
    target_columns = tuple(openpyxl.utils.column_index_from_string(col) for col in layout)
    before = fingerprint(path)
    mode = os.environ.get("TOOLBOX_BENEFIT_INSPECT_BACKEND", "auto")
    if mode not in {"auto", "full", "stream"}:
        raise ValueError(f"未知收益检查后端：{mode}")
    workbook = cached = None
    if mode != "full":
        try:
            workbook, cached = load_inspection_pair(
                path, columns=(1, 2, *target_columns),
                target_columns=target_columns,
                sheet_filter=lambda name: re.fullmatch(r"(\d{1,2})月(?:份)?", name),
            )
        except (UnsupportedInspection, OSError, ValueError, KeyError, IndexError, TypeError) as exc:
            if mode == "stream":
                raise
            warnings.warn(f"局部收益检查不可用，使用完整工作簿：{exc}", RuntimeWarning)
    if workbook is None:
        workbook = openpyxl.load_workbook(path)
    cells: list[CellTarget] = []
    seen_days: set[date] = set()
    try:
        if cached is None:
            cached = load_value_workbook(path, data_only=True, read_only=False)
        for sheet in workbook:
            month_match = re.fullmatch(r"(\d{1,2})月(?:份)?", sheet.title)
            if not month_match:
                continue
            header_row = next((r for r in range(1, min(sheet.max_row, 20) + 1)
                               if all(re.sub(r"\s+", "", str(sheet[f"{col}{r}"].value or "")) == label
                                      for col, label in headers.items())), None)
            if header_row is None:
                raise ValueError(f"{path.name}/{sheet.title} 目标列模板表头不匹配。")
            groups: dict[date, dict[int, int]] = {}
            current: date | None = None
            for row in range(header_row + 1, sheet.max_row + 1):
                if sheet.cell(row, 1).value is not None:
                    current = date_at(sheet, cached[sheet.title], row, workbook.epoch)
                hour = hour_of(sheet.cell(row, 2).value)
                if current is None or hour is None:
                    continue
                # 某些30天月份的模板带第31天占位，归属下月的块不能写入本月。
                if current.month != int(month_match[1]) or current > end_date or (start_date and current < start_date):
                    continue
                hours = groups.setdefault(current, {})
                if hour in hours:
                    raise ValueError(f"{path.name}/{sheet.title} {current} {hour}:00 重复。")
                hours[hour] = row
            for day, hours in sorted(groups.items()):
                if set(hours) != set(range(24)) or day in seen_days:
                    raise ValueError(f"{path.name} {day} 日期重复或不是完整的0–23点。")
                seen_days.add(day)
                for hour, row in sorted(hours.items()):
                    for column, field in layout.items():
                        cell = sheet[f"{column}{row}"]
                        if station == "海风" or cell.value in (None, ""):
                            cells.append(CellTarget(sheet.title, cell.coordinate, day, hour,
                                                    field, cell.number_format, cell.style_id))
        if not seen_days:
            raise ValueError(f"{path.name} 模板内未找到指定日期范围的完整小时块。")
    finally:
        workbook.close()
        if cached is not None:
            cached.close()
    if fingerprint(path) != before:
        raise RuntimeError(f"检查时工作簿发生变化，请关闭 Excel/WPS 后重试：{path}")
    return WorkbookPlan(path, station, tuple(cells), before)


def points24(raw: Any, *, label: str, scale: float = 1) -> tuple[float, ...] | None:
    if raw is None or raw == []:
        return None
    if not isinstance(raw, list) or len(raw) != 24:
        raise ValueError(f"{label} 应有24个数，实际为{len(raw) if isinstance(raw, list) else '非数组'}。")
    if any(value in (None, "", "-") for value in raw):
        return None
    return tuple(shared_spot()._number(value, label=label) * scale for value in raw)


def unique_unit_row(rows: Any, unit_id: str, *, label: str) -> dict[str, Any] | None:
    if rows is None:
        return None
    if not isinstance(rows, list):
        raise ValueError(f"{label} 的 data 不是数组。")
    matched = [row for row in rows if isinstance(row, dict) and str(row.get("unitId")) == unit_id]
    if len(matched) > 1:
        raise ValueError(f"{label} 同一场站返回多条记录，不能确定数据。")
    return matched[0] if matched else None


def parse_forecast_payload(payload: dict, *, unit_id: str, fallback: bool = False) -> dict[str, tuple[float, ...]]:
    row = unique_unit_row(payload.get("data"), unit_id, label="功率预测")
    if row is None:
        return {}
    # queryLatest 的 date 是运行参数配置生效日，非查询运行日；实际功率按请求 date 返回。
    fields = {"C": "forecastPower24", "J": "rawShortTermPower24" if fallback else "stPredictPowerList24"}
    result = {}
    for column, source in fields.items():
        values = points24(row.get(source), label=FIELDS[column])
        if values is not None:
            result[column] = values
    return result


def parse_trade_payload(payload: dict, *, unit_id: str, expected_dates: set[date],
                        data_type: int) -> dict[date, dict[str, tuple[float, ...]]]:
    data = payload.get("data") or {}
    rows = data.get("unitDateList", []) if isinstance(data, dict) else data
    if not isinstance(rows, list):
        raise ValueError("出清结果 unitDateList 不是数组。")
    result: dict[date, dict[str, tuple[float, ...]]] = {}
    for row in rows:
        if str(row.get("unitId")) != unit_id:
            continue
        day = as_date(row.get("date"), openpyxl.utils.datetime.WINDOWS_EPOCH)
        if day not in expected_dates:
            continue
        if str(row.get("dataType", data_type)) != str(data_type):
            raise ValueError(f"{day} 出清结果日前/实时类型不匹配。")
        if day in result:
            raise ValueError(f"{day} 同一场站出清记录重复。")
        values = {}
        fields = {"K": ("electricity", 10), "L": ("acPrice", 1)} if data_type == 1 else {"N": ("acPrice", 1)}
        for column, (field, scale) in fields.items():
            points = points24(row.get(field), label=f"{day} {FIELDS[column]}", scale=scale)
            if points is not None:
                values[column] = points
        result[day] = values
    return result


def collect_data(plans: tuple[WorkbookPlan, ...], *, headed: bool = False,
                 login_only: bool = False) -> dict[str, dict[date, dict[str, tuple[float, ...]]]]:
    raw = read_login_config(CONFIG_PATH, data_path(WORKSPACE / "上网电量抓取" / "config.json"))
    auth = load_auth_module()
    config = auth.Config(username=str(raw.get("username") or ""), password=str(raw.get("password") or ""))
    collected: dict[str, dict[date, dict[str, tuple[float, ...]]]] = {}
    with auth.AuthStateLock():
        with sync_playwright() as playwright:
            context = auth.launch_context(playwright, headless=not (headed or login_only))
            try:
                if login_only:
                    auth.interactive_login(context, config)
                    return collected
                auth.ensure_login(context, config)
                companies = auth.get_companies(context)
                for plan in plans:
                    if not plan.cells:
                        continue
                    spot = shared_spot()
                    spec = next(target for target in spot.default_targets() if target.key == plan.station)
                    tenant = spot.resolve_tenant(companies, spec)
                    auth.switch_company(context, str(tenant["tenantId"]))
                    unit = spot.resolve_unit(auth.get_units(context), spec, auth)
                    unit_id = str(unit["id"])
                    log(f"{plan.station}：{tenant['name']} / {auth.unit_display_name(unit)}")
                    needs: dict[date, set[str]] = {}
                    for cell in plan.cells:
                        needs.setdefault(cell.run_date, set()).add(cell.field)
                    series: dict[date, dict[str, tuple[float, ...]]] = {day: {} for day in needs}
                    for data_type, fields in ((1, {"K", "L"}), (2, {"N"})):
                        days = {day for day, needed in needs.items() if needed & fields}
                        if not days:
                            continue
                        payload = auth.request_json_with_retry(context, "get", TRADE_API, "读取出清结果",
                            params={"dataType": data_type, "startDate": min(days).isoformat(),
                                    "endDate": max(days).isoformat(), "unitIds": unit_id})
                        for day, values in parse_trade_payload(payload, unit_id=unit_id,
                                                               expected_dates=days, data_type=data_type).items():
                            series[day].update(values)
                    for day, needed in sorted(needs.items()):
                        log(f"{plan.station} {day}：读取 {'、'.join(FIELDS[f] for f in FIELDS if f in needed)}")
                        if needed & {"C", "J"}:
                            params = {"date": day.isoformat(), "orgId": tenant["tenantId"], "unitIds": unit_id}
                            payload = auth.request_json_with_retry(context, "get", FORECAST_API + "queryLatest",
                                                                   "读取功率预测", params=params)
                            powers = parse_forecast_payload(payload, unit_id=unit_id)
                            # 与网站相同：未保存申报功率时读取 rawShortTermPower24 的默认申报值。
                            row = unique_unit_row(payload.get("data"), unit_id, label="功率预测")
                            if row is None or (not row.get("stPredictPowerList24") and not row.get("stPredictPowerList96")):
                                fallback = auth.request_json_with_retry(context, "get", FORECAST_API + "queryRawShortTermPower",
                                                                        "读取默认功率预测", params=params)
                                powers.update(parse_forecast_payload(fallback, unit_id=unit_id, fallback=True))
                            series[day].update(powers)
                        if "M" in needed:
                            rows = auth.fetch_energy_rows(context, day.isoformat(), str(tenant["tenantId"]), [unit])
                            row = unique_unit_row(rows, unit_id, label="上网电量")
                            if row:
                                if as_date(row.get("runDate"), openpyxl.utils.datetime.WINDOWS_EPOCH) != day:
                                    raise ValueError(f"上网电量返回日期与请求 {day} 不符。")
                                values = points24(row.get("onlineEle"), label=f"{day} 上网电量")
                                if values is not None:
                                    series[day]["M"] = values
                        missing = needed - series[day].keys()
                        if missing:
                            log(f"待发布/不完整，保留{'原值' if plan.station == '海风' else '空白'}：{plan.station} {day} {'、'.join(FIELDS[f] for f in FIELDS if f in missing)}")
                    collected[plan.station] = series
                auth.save_auth_state(context)
            finally:
                context.close()
    return collected


def expand_replaced_shared_formulas(text: str, changes: dict[str, float]) -> str:
    """覆盖共享公式主单元格前，将同组公式独立化，保留未发布字段的引用。"""
    cells = re.compile(r'<c\b(?=[^>]*\br="([A-Z]+\d+)")[^>]*?(?:/>|>.*?</c>)', re.S)
    masters = {}
    for cell in cells.finditer(text):
        formula = re.search(r'<f\b([^>]*)(?<!/)>((?:(?!</f>).)*)</f>', cell[0], re.S)
        if formula and cell[1] in changes and re.search(r'\bt="shared"', formula[1]):
            shared_id = re.search(r'\bsi="(\d+)"', formula[1])
            if shared_id:
                masters[shared_id[1]] = (cell[1], "=" + unescape(formula[2]))
    if not masters:
        return text
    def expand(cell):
        formula = re.search(r'<f\b[^>]*?(?:/>|>.*?</f>)', cell[0], re.S)
        if not formula or not re.search(r'\bt="shared"', formula[0]):
            return cell[0]
        shared_id = re.search(r'\bsi="(\d+)"', formula[0])
        if shared_id is None or shared_id[1] not in masters:
            return cell[0]
        origin, expression = masters[shared_id[1]]
        translated = Translator(expression, origin=origin).translate_formula(cell[1])[1:]
        return cell[0][:formula.start()] + "<f>" + escape(translated, quote=False) + "</f>" + cell[0][formula.end():]
    return cells.sub(expand, text)


def patch_sparse_sheet(data: bytes, changes: dict[str, float]) -> bytes:
    spot = shared_spot()
    text = expand_replaced_shared_formulas(data.decode("utf-8"), changes)
    original_text = text
    # 自闭合空单元格必须优先匹配，不能吞入后面的单元格或公式。
    cell_pattern = re.compile(r'<c\b(?=[^>]*\br="([A-Z]+\d+)")[^>]*?(?:/>|>.*?</c>)', re.S)
    replaced = set()
    def replace_existing(match):
        coordinate = match[1]
        if coordinate not in changes or coordinate in replaced:
            return match[0]
        replaced.add(coordinate)
        return spot._numeric_cell_xml(coordinate, changes[coordinate], spot._style_from_cell_xml(match[0]))
    text = cell_pattern.sub(replace_existing, text)
    for coordinate, value in changes.items():
        if coordinate in replaced:
            continue
        pattern = re.compile(rf'<c\b(?=[^>]*\br="{re.escape(coordinate)}")[^>]*?(?:/>|>.*?</c>)', re.S)
        match = pattern.search(text)
        if match is None:
            # Excel 可省略完全空白的 <c>，插入时保留其默认/列继承样式。
            row_number, column_number = openpyxl.utils.cell.coordinate_to_tuple(coordinate)
            row = re.search(rf'(<row\b(?=[^>]*\br="{row_number}")[^>]*>)(.*?)(</row>)', text, re.S)
            if row is None:
                raise ValueError(f"模板缺少目标行 {row_number}，停止写入。")
            content = row[2]
            insertion = len(content)
            for existing in cell_pattern.finditer(content):
                if openpyxl.utils.cell.coordinate_to_tuple(existing[1])[1] > column_number:
                    insertion = existing.start()
                    break
            replacement = spot._numeric_cell_xml(coordinate, value, None)
            content = content[:insertion] + replacement + content[insertion:]
            text = text[:row.start()] + row[1] + content + row[3] + text[row.end():]
            continue
        # 只替换输入单元格，样式索引沿用原值；公式、图表、缓存及其他部件原样保留。
        replacement = spot._numeric_cell_xml(coordinate, value, spot._style_from_cell_xml(match[0]))
        text = text[:match.start()] + replacement + text[match.end():]
    def omit_inputs(match):
        return "" if match[1] in changes else match[0]
    if cell_pattern.sub(omit_inputs, original_text) != cell_pattern.sub(omit_inputs, text):
        raise RuntimeError("工作表非目标单元格或公式发生变化，停止写入。")
    return text.encode("utf-8")


def request_recalculation(data: bytes) -> bytes:
    text = data.decode("utf-8")
    match = re.search(r"<calcPr\b[^>]*/>", text)
    if match:
        tag = match[0]
        for key in ("fullCalcOnLoad", "forceFullCalc", "calcOnSave"):
            tag = re.sub(rf'\s{key}="[^"]*"', "", tag)
            tag = tag[:-2] + f' {key}="1"/>'
        text = text[:match.start()] + tag + text[match.end():]
    else:
        text = text.replace("</workbook>", '<calcPr fullCalcOnLoad="1" forceFullCalc="1" calcOnSave="1"/></workbook>')
    return text.encode("utf-8")


def verify_workbook(path: Path, prepared: PreparedWorkbook) -> None:
    mode = os.environ.get("TOOLBOX_BENEFIT_VERIFY_BACKEND", "auto")
    if mode not in {"auto", "full", "stream"}:
        raise ValueError(f"未知收益回读后端：{mode}")
    workbook = None
    if mode != "full":
        try:
            requests = {}
            for cell, _ in prepared.expectations:
                requests.setdefault(cell.sheet, set()).add(openpyxl.utils.cell.coordinate_to_tuple(cell.coordinate))
            columns = tuple(sorted({column for cells in requests.values() for _, column in cells})) or (3, 10, 11, 12, 13, 14)
            workbook, _ = load_inspection_pair(
                path, columns=columns, target_columns=columns,
                sheet_filter=lambda name: re.fullmatch(r"(\d{1,2})月(?:份)?", name),
                requests=requests, capture_cache=False,
            )
        except (UnsupportedInspection, OSError, ValueError, KeyError, IndexError, TypeError) as exc:
            if mode == "stream":
                raise
            warnings.warn(f"局部收益回读不可用，使用完整工作簿：{exc}", RuntimeWarning)
    if workbook is None:
        workbook = openpyxl.load_workbook(path)
    try:
        snapshots = read_cell_snapshots(workbook, ((cell.sheet, cell.coordinate) for cell, _ in prepared.expectations))
        for cell, expected in prepared.expectations:
            actual, number_format = snapshots[(cell.sheet, cell.coordinate)]
            if isinstance(actual, bool) or not isinstance(actual, (int, float)) or not math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-9):
                raise RuntimeError(f"回读不符：{cell.sheet}!{cell.coordinate}，期望{expected}，实际{actual}")
            if number_format != cell.number_format:
                raise RuntimeError(f"数字格式改变：{cell.sheet}!{cell.coordinate}")
            if workbook[cell.sheet][cell.coordinate].style_id != cell.style_id:
                raise RuntimeError(f"单元格样式改变：{cell.sheet}!{cell.coordinate}")
    finally:
        workbook.close()


def prepare_workbook(plan: WorkbookPlan, series: dict[date, dict[str, tuple[float, ...]]]) -> PreparedWorkbook:
    source_fingerprint = workbook_fingerprint(plan.path)
    if fingerprint(plan.path) != plan.fingerprint:
        raise RuntimeError(f"抓取期间工作簿被修改，停止写入：{plan.path}")
    expectations = []
    changes: dict[str, dict[str, float]] = {}
    for cell in plan.cells:
        values = series.get(cell.run_date, {}).get(cell.field)
        if values is None:
            continue
        values = points24(list(values), label=f"{cell.run_date} {FIELDS[cell.field]}")
        if values is None:
            continue
        value = values[cell.hour]
        expectations.append((cell, value))
        changes.setdefault(cell.sheet, {})[cell.coordinate] = value
    temporary = plan.path.with_name(f".{plan.path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    prepared = PreparedWorkbook(plan.path, temporary, tuple(expectations), source_fingerprint)
    try:
        sheet_paths = shared_spot().xlsx_sheet_paths(plan.path)
        part_changes = {sheet_paths[sheet]: values for sheet, values in changes.items()}
        with zipfile.ZipFile(plan.path) as source, zipfile.ZipFile(temporary, "w") as output:
            output.comment = source.comment
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in part_changes:
                    data = patch_sparse_sheet(data, part_changes[info.filename])
                elif info.filename == "xl/workbook.xml" and expectations:
                    data = request_recalculation(data)
                output.writestr(info, data)
        changed_parts = set(part_changes) | ({"xl/workbook.xml"} if expectations else set())
        shared_spot().verify_package_preservation(plan.path, temporary, changed_parts)
        verify_workbook(temporary, prepared)
        assert_workbook_unchanged(plan.path, source_fingerprint)
        return prepared
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def run_update(paths: tuple[Path, ...], *, end_date: date, start_date: date | None = None,
               create_backup: bool = True, headed: bool = False) -> tuple[PreparedWorkbook, ...]:
    if start_date and end_date < start_date:
        raise ValueError("截止日期不能早于开始日期。")
    plans = tuple(inspect_workbook(path, end_date=end_date, start_date=start_date) for path in paths)
    if len({plan.station for plan in plans}) != len(plans):
        raise ValueError("每个场站一次只能选择一个工作簿。")
    for plan in plans:
        make_writable(plan.path)
        log(f"{plan.path.name}：已解除文件只读；截至 {end_date} 待补 {len(plan.cells)} 个单元格。")
    data = collect_data(plans, headed=headed) if any(plan.cells for plan in plans) else {}
    prepared: list[PreparedWorkbook] = []
    try:
        for plan in plans:
            item = prepare_workbook(plan, data.get(plan.station, {}))
            if item.expectations:
                prepared.append(item)
            else:
                item.temporary.unlink(missing_ok=True)
                log(f"{plan.station}：本次无可写入数据。")
        # 准备完成后再检查一次，避免另一个程序在两张表准备期间修改原文件。
        for plan in plans:
            if fingerprint(plan.path) != plan.fingerprint:
                raise RuntimeError(f"提交前工作簿被修改，停止写入：{plan.path}")
        backups = commit_workbooks(tuple(prepared), verify_workbook, create_backup=create_backup)
        for item in prepared:
            counts = {column: sum(re.match(r"[A-Z]+", cell.coordinate)[0] == column for cell, _ in item.expectations)
                      for column in column_layout(station_from_filename(item.original))}
            log(f"完成并回读：{item.original.name}，写入{len(item.expectations)}格；" +
                "，".join(f"{column} {count}" for column, count in counts.items()))
            if backups[item.original]:
                log(f"备份：{backups[item.original]}")
        return tuple(prepared)
    finally:
        for item in prepared:
            item.temporary.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", action="append", type=Path, help="可重复传入；按文件名识别归湖/鮀莲/海风")
    parser.add_argument("--start-date", type=parse_date_input, help="海风覆盖必填；其他场站留空检查历史空白")
    parser.add_argument("--end-date", type=parse_date_input, default=date.today(), help="截止日期 D（含），默认今天")
    parser.add_argument("--no-backup", action="store_true")
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--login", action="store_true")
    args = parser.parse_args()
    try:
        if args.login:
            collect_data((), login_only=True)
        else:
            run_update(tuple(args.workbook or DEFAULT_WORKBOOKS), end_date=args.end_date,
                       start_date=args.start_date, create_backup=not args.no_backup, headed=args.headed)
        return 0
    except Exception as exc:
        log(f"更新失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
