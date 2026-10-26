from __future__ import annotations

import sys
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook
import numpy as np
from PIL import Image


WORKSPACE = Path(__file__).resolve().parents[2]
TOOL_DIR = WORKSPACE / "广东运行日报PDF解析"
if str(TOOL_DIR) not in sys.path:
    sys.path.insert(0, str(TOOL_DIR))

from gd_pdf_tool.excel_writer import HEADERS, write_quote_excel
from gd_pdf_tool.parser import (
    DeclarationInfo, DeclarationRow, ParseError, QuoteRow, analyze_chart_image,
    _detect_plot_rows, _make_color_masks, _measure_pixel_counts, _rows_from_pixels, _scale_values_to_total,
    _select_palette_masks, _unmix_bar_colors, price_intervals,
    extract_declaration_info, _read_declaration_tables,
)
from gd_pdf_tool.image_labels import extract_image_label_tables, merge_image_labels
from toolbox.daily_pdf_page import default_output_for_pdf


def sample_chart(heights, *, gray=False, grid_overlay=False):
    """Render known fractional heights by exact pixel-area integration."""
    colors = np.array(((79, 129, 189), (192, 80, 77), (155, 187, 89)), dtype=float)
    if gray:
        colors[2] = (150, 150, 150)
    width = 80 + len(heights) * 32
    array = np.full((720, width, 3), 255.0)
    for y in range(50, 551, 50):
        array[y, 40:width - 40] = 210
    for index, parts in enumerate(heights):
        center = 56 + index * 32
        bottom = 550.0
        weights = np.zeros((720, 3))
        for channel, height in enumerate(parts):
            top = bottom - height
            yy = np.arange(720)
            weights[:, channel] = np.maximum(0, np.minimum(yy + 1, bottom) - np.maximum(yy, top))
            bottom = top
        active = weights.sum(axis=1) > 0
        pixels = 255 + weights @ (colors - 255)
        array[active, center - 5:center + 5] = pixels[active, None]
        # Regular label strokes exercise category fitting without font assets.
        array[575:635, center - 3:center + 3] = 40
    if grid_overlay:
        for y in range(100, 550, 50):
            array[y:y + 2, 40:width - 40] = 210
    return Image.fromarray(np.rint(array).astype(np.uint8))


class DailyPdfParserUnitTests(unittest.TestCase):
    def test_price_intervals_preserve_report_order(self) -> None:
        self.assertEqual(
            ["(-∞,0]", "(0, 25]", "(25, 50]", "(50, 75]"],
            price_intervals(4),
        )

    def test_default_output_uses_report_date_from_filename(self) -> None:
        source = Path("/临时/广东电力现货市场2026年7月运行日报（07.17）.pdf")
        self.assertEqual(
            Path("/临时/现货运行日报报价容量区间20260717.xlsx"),
            default_output_for_pdf(source),
        )

    def test_excel_writer_creates_expected_sheet_values_and_formats(self) -> None:
        rows = [
            QuoteRow("(-∞,0]", 4870.2, 590.0, 36424.2),
            QuoteRow("(0, 25]", 0.0, 0.0, 312.1),
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "result.xlsx"
            write_quote_excel(rows, output)
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["Sheet1"]

            self.assertEqual(HEADERS, tuple(cell.value for cell in sheet[1]))
            self.assertEqual("(-∞,0]", sheet["A2"].value)
            self.assertEqual(4870.2, sheet["B2"].value)
            self.assertEqual("0.0", sheet["B2"].number_format)
            self.assertEqual("A2", str(sheet.freeze_panes))
            self.assertEqual(16, sheet.column_dimensions["A"].width)
            self.assertEqual(12, sheet.column_dimensions["D"].width)
            workbook.close()

    def test_declaration_rows_keep_all_four_fields(self) -> None:
        declaration = DeclarationInfo(
            zero_or_below=DeclarationRow(41981.5148, 4870.25, 590.0, 36424.1798),
            above_600=DeclarationRow(28631.7322, 12219.9, 16309.032, 82.0),
        )
        self.assertAlmostEqual(41981.5148, declaration.zero_or_below.total)
        self.assertAlmostEqual(82.0, declaration.above_600.renewable)


class DailyPdfImageLabelTests(unittest.TestCase):
    def table_and_labels(self, *, offset=0):
        # Different column order and movable rows exercise header/geometry mapping.
        values = [
            ["", "", "", "", ""],
            ["", "405(0.40%,0.0%)", "41897.32(41.36%,-0.0\n%)",
             "38133.1798(37.65%,+\n2.2%)", "3327.5(3.29%,-20.0%)"],
            ["", "16799.152(16.59%,+6.0%\n)", "24194.352(23.89%,+3.\n6%)",
             "82(0.08%,0.0%)", "7303.2(7.21%,-1.4%)"],
        ]
        rows = [SimpleNamespace(cells=[
            (column * 100, offset + row * 50, (column + 1) * 100, offset + (row + 1) * 50)
            for column in range(5)
        ]) for row in range(3)]
        table = SimpleNamespace(rows=rows, extract=lambda: [list(row) for row in values])

        def label(text, column, y):
            return {"text": text, "box": {
                "origin": {"x": (column * 100 + 10) / 500, "y": 1 - (offset + y + 10) / 500},
                "size": {"width": 0.12, "height": 0.02},
            }}

        labels = [label(text, column, 10) for column, text in enumerate(
            ("", "气机容量", "申报容量", "新能源容量", "煤机容量")
        ) if text]
        labels += [label("下", 0, 80), label("零价及以", 0, 60),
                   label("以上", 0, 130), label("报价600", 0, 110)]
        # OCR noise must never overwrite PDF numbers or fill a cell with OCR numbers.
        labels += [label("99999", 0, 10), label("误读容量99999", 2, 60)]
        return table, labels

    def test_image_labels_restore_moved_table_and_reordered_columns(self):
        for offset in (0, 120):
            with self.subTest(offset=offset):
                table, labels = self.table_and_labels(offset=offset)
                merged = merge_image_labels(table, labels, 500, 500)
                self.assertEqual("", merged[0][0])
                self.assertEqual(table.extract()[1][2], merged[1][2])
                result = _read_declaration_tables([merged], require_header=True)
                self.assertEqual(DeclarationRow(41897.32, 3327.5, 405, 38133.1798), result.zero_or_below)
                self.assertEqual(DeclarationRow(24194.352, 7303.2, 16799.152, 82), result.above_600)

    def test_incomplete_image_headers_do_not_guess_columns(self):
        table, labels = self.table_and_labels()
        labels = [label for label in labels if label["text"] != "煤机容量"]
        result = _read_declaration_tables(
            [merge_image_labels(table, labels, 500, 500)], require_header=True,
        )
        self.assertIsNone(result.zero_or_below)
        self.assertIsNone(result.above_600)

    def test_pdf_extraction_falls_back_to_image_labels(self):
        table, labels = self.table_and_labels()
        page = SimpleNamespace(
            images=[{}], width=500, height=500,
            extract_tables=lambda: [table.extract()], find_tables=lambda: [table],
            to_image=lambda **kwargs: SimpleNamespace(original=Image.new("RGB", (500, 500), "white")),
        )
        with mock.patch("gd_pdf_tool.parser.pdfplumber.open") as opened, \
             mock.patch("gd_pdf_tool.image_labels.recognize_image_labels", return_value=labels) as recognize:
            opened.return_value.__enter__.return_value.pages = [page]
            result = extract_declaration_info(Path("image-label-report.pdf"))
        self.assertEqual("image-labels", result.source)
        self.assertEqual(41897.32, result.zero_or_below.total)
        self.assertEqual(16799.152, result.above_600.gas)
        recognize.assert_called_once()
        self.assertFalse(recognize.call_args.args[0].exists())

    def test_native_text_tables_do_not_invoke_image_recognition(self):
        table, labels = self.table_and_labels()
        page = SimpleNamespace(extract_tables=lambda: [merge_image_labels(table, labels, 500, 500)])
        with mock.patch("gd_pdf_tool.parser.pdfplumber.open") as opened, \
             mock.patch("gd_pdf_tool.image_labels.extract_image_label_tables") as fallback:
            opened.return_value.__enter__.return_value.pages = [page]
            result = extract_declaration_info(Path("native-text-report.pdf"))
        self.assertEqual("text", result.source)
        self.assertEqual(3327.5, result.zero_or_below.coal)
        fallback.assert_not_called()

    def test_unrelated_image_table_does_not_trigger_recognition(self):
        table = SimpleNamespace(extract=lambda: [["1.45", "19:00", "0.98", "5:00", "29"]])
        page = SimpleNamespace(images=[{}], find_tables=lambda: [table])
        with mock.patch("gd_pdf_tool.image_labels.recognize_image_labels") as recognize:
            self.assertEqual([], extract_image_label_tables(page))
        recognize.assert_not_called()

    def test_image_recognition_error_is_reported_without_silent_fallback(self):
        page = SimpleNamespace(extract_tables=lambda: [])
        with mock.patch("gd_pdf_tool.parser.pdfplumber.open") as opened, \
             mock.patch("gd_pdf_tool.image_labels.extract_image_label_tables", side_effect=RuntimeError("OCR timeout")):
            opened.return_value.__enter__.return_value.pages = [page]
            with self.assertRaisesRegex(ParseError, "图片标签识别失败.*OCR timeout"):
                extract_declaration_info(Path("image-label-report.pdf"))


class DailyPdfChartPrecisionTests(unittest.TestCase):
    def geometry_chart(self, count, *, labels):
        heights = [(60, 20, 300)] + [(0, 0, 0)] * (count - 1)
        for index, height in ((15, 80), (16, 60), (17, 40), (18, 20), (35, 10), (36, 50)):
            heights[index] = (height, 0, 0)
        heights[-1] = (0, 150, 0)
        array = np.array(sample_chart(heights))
        array[575:635] = 255
        if labels:
            for index in range(count):
                center = 56 + index * 32
                array[575:635, center - 8:center + 8] = 40
                if index % 2 == 0:
                    # A detached thin bracket beside an otherwise intact label.
                    array[575:635, center - 13:center - 11] = 40
        return Image.fromarray(array)

    def test_detached_axis_brackets_do_not_shift_capacity_intervals(self):
        for count in (49, 51, 55):
            with self.subTest(count=count):
                result = analyze_chart_image(self.geometry_chart(count, labels=True), 38000)
                self.assertEqual("axis-labels", result.diagnostics["category_geometry_source"])
                self.assertEqual(count, len(result.rows))
                self.assertEqual("(875, 900]", result.rows[36].interval)
                self.assertAlmostEqual(5000, result.rows[36].coal, delta=60)
                self.assertAlmostEqual(15000, result.rows[-1].gas, delta=60)

    def test_unreadable_labels_use_observed_spacing_without_fixed_bin_count(self):
        for count in (49, 51, 55):
            with self.subTest(count=count):
                result = analyze_chart_image(self.geometry_chart(count, labels=False), 38000)
                self.assertEqual("plot-bounds-fallback", result.diagnostics["category_geometry_source"])
                self.assertEqual(count, len(result.rows))
                self.assertAlmostEqual(5000, result.rows[36].coal, delta=60)
                self.assertAlmostEqual(15000, result.rows[-1].gas, delta=60)

    def measure(self, image, count):
        array = np.array(image)
        masks = _select_palette_masks(_make_color_masks(array), 50, 550)
        return _measure_pixel_counts(array, masks, 50, 550, 56, 32, 11, count)

    def test_fractional_short_bars_are_not_rounded_to_whole_pixels(self):
        heights = [(60, 20, 300)] + [(h, 0, 0) for h in (0.25, 0.5, 0.75, 1, 1.5, 2.25, 3.5)]
        measured = self.measure(sample_chart(heights), len(heights))
        np.testing.assert_allclose(measured, heights, atol=0.04)

    def test_thin_stack_boundaries_are_not_counted_twice(self):
        heights = [(60, 20, 300), (3.25, 0.5, 0), (0, 1.5, 0.25), (1.25, 0.25, 0)]
        measured = self.measure(sample_chart(heights), len(heights))
        np.testing.assert_allclose(measured, heights, atol=0.05)

    def test_gray_renewables_do_not_turn_grid_lines_into_capacity(self):
        heights = [(60, 20, 300), (0, 0, 0), (0, 0, 0.5), (0, 0, 2.25)]
        measured = self.measure(sample_chart(heights, gray=True), len(heights))
        np.testing.assert_allclose(measured, heights, atol=0.06)

    def test_short_grid_occlusions_are_restored_without_losing_tall_stack(self):
        heights = [(60, 20, 300), (5, 0, 0), (0, 0, 0)]
        measured = self.measure(sample_chart(heights, grid_overlay=True), len(heights))
        np.testing.assert_allclose(measured, heights, atol=0.06)

    def test_detached_pixels_are_not_short_bars(self):
        heights = [(60, 20, 300), (0, 0, 0), (0, 0, 0)]
        array = np.array(sample_chart(heights))
        array[530, 85:92] = (79, 129, 189)
        array[549, 120] = (192, 80, 77)
        measured = self.measure(Image.fromarray(array), len(heights))
        np.testing.assert_allclose(measured[1:], 0, atol=0.01)

    def test_complete_chart_to_excel_preserves_small_values_and_order(self):
        heights = [(60, 20, 300), (0.25, 0, 0), (0, 0.5, 0), (0, 0, 0.75),
                   (3.25, 0.5, 0), (0, 0, 0), (2, 1, 0)]
        result = analyze_chart_image(sample_chart(heights), 38000)
        self.assertEqual(len(heights), len(result.rows))
        self.assertEqual(550, result.diagnostics['plot_baseline_y'])
        self.assertTrue(result.diagnostics['low_height_segments'])
        expected = np.array(heights) * 100
        actual = [(r.coal, r.gas, r.renewable) for r in result.rows]
        np.testing.assert_allclose(actual, expected, atol=5)
        with tempfile.TemporaryDirectory() as directory:
            output = write_quote_excel(result.rows, Path(directory) / 'precision.xlsx')
            workbook = load_workbook(output)
            sheet = workbook['Sheet1']
            self.assertEqual(HEADERS, tuple(cell.value for cell in sheet[1]))
            self.assertAlmostEqual(25, sheet['B3'].value, delta=4)
            self.assertEqual(0, sheet['C3'].value)
            self.assertEqual('0.0', sheet['B3'].number_format)
            self.assertEqual('(0, 25]', sheet['A3'].value)
            workbook.close()

    def test_positive_summary_cannot_be_invented_in_last_interval(self):
        with self.assertRaisesRegex(ParseError, '全部未识别'):
            _scale_values_to_total([0, 0, 0], 82)

    def test_tinted_zero_line_is_not_replaced_by_axis_text(self):
        array = np.array(sample_chart([(60, 20, 300)] + [(0, 0, 0)] * 6))
        array[550, 40:-40] = (212, 218, 211)
        array[575, 40:-40] = 210
        self.assertEqual(550, _detect_plot_rows(array)[1])

    def test_unknown_transition_pixel_does_not_discard_rest_of_stack(self):
        heights = [(60, 20, 300), (0, 0, 0), (0, 0, 0)]
        array = np.array(sample_chart(heights))
        array[469, 51:61] = (0, 0, 0)
        measured = self.measure(Image.fromarray(array), len(heights))
        np.testing.assert_allclose(measured[0], (60, 20, 299), atol=0.06)

    def test_summary_rounding_keeps_small_allocations_nonnegative(self):
        values = _scale_values_to_total([1] * 24, 1.3)
        self.assertAlmostEqual(1.3, sum(values))
        self.assertGreaterEqual(min(values), 0)
        self.assertEqual([0, 0], _scale_values_to_total([0, 0], 0))

    def test_summary_calibration_precedes_rounding_tiny_measured_segments(self):
        pixels = [(60, 20, 300)] + [(0, 0, 0)] * 24 + [(0.00001, 0, 0), (0.00002, 0, 0)]
        declaration = DeclarationInfo(None, DeclarationRow(82, 82, 0, 0))
        rows = _rows_from_pixels(pixels, 100, 38000, declaration)
        self.assertAlmostEqual(82, sum(row.coal for row in rows[25:]))
        self.assertGreater(rows[25].coal, 0)

    def test_jpeg_low_bars_keep_lower_error_than_integer_row_counting(self):
        heights = [(60, 20, 300)] + [(h, 0, 0) for h in (0.25, 0.5, 0.75, 1, 1.5, 2.25, 3.5, 4.25, 5.5)]
        buffer = io.BytesIO()
        sample_chart(heights).save(buffer, format='JPEG', quality=90)
        image = Image.open(io.BytesIO(buffer.getvalue())).convert('RGB')
        measured = self.measure(image, len(heights))
        array = np.array(image)
        masks = _make_color_masks(array)
        old = []
        for i in range(1, len(heights)):
            center = 56 + i * 32
            old.append(float((masks['blue'][45:551, center - 11:center + 12].sum(axis=1) >= 4).sum()))
        truth = np.array(heights[1:])[:, 0]
        new_error = np.mean(np.abs(np.array(measured[1:])[:, 0] - truth) / truth)
        old_error = np.mean(np.abs(np.array(old) - truth) / truth)
        self.assertLess(new_error, old_error * 0.3)
        self.assertLess(np.sum(np.abs(np.array(measured[1:]) - heights[1:])), 1.1)

    def test_purple_tinted_gas_does_not_create_coal(self):
        palette = [np.array(color, dtype=float) for color in (
            (78, 130, 187), (190, 80, 78), (155, 187, 87)
        )]
        # JPEG-distorted red observed in real short gas bars, including 10/09.
        pixels = np.full((8, 7, 3), 255, dtype=np.uint8)
        pixels[4] = (190, 80, 78)
        pixels[5] = (151, 92, 101)
        pixels[6] = (158, 93, 106)
        coverage = _unmix_bar_colors(pixels, np.full((8, 3), 255.0), palette)
        self.assertEqual(0, coverage[:, :, 0].sum())
        self.assertEqual(0, coverage[:, :, 2].sum())
        self.assertGreater(np.median(coverage, axis=1)[:, 1].sum(), 2.8)

    def test_pure_gas_jpeg_bars_do_not_create_coal_or_renewables(self):
        heights = [(60, 20, 300), (0, 2, 0), (0, 5, 0), (0, 200, 0)]
        for quality in (95, 90, 80):
            with self.subTest(quality=quality):
                buffer = io.BytesIO()
                sample_chart(heights).save(buffer, format='JPEG', quality=quality)
                image = Image.open(io.BytesIO(buffer.getvalue())).convert('RGB')
                measured = np.array(self.measure(image, len(heights)))
                np.testing.assert_allclose(measured[1:, (0, 2)], 0, atol=0.01)
                np.testing.assert_allclose(measured[1:, 1], (2, 5, 200), atol=0.4)

    def test_pure_blue_jpeg_bars_do_not_create_gas(self):
        heights = [(60, 20, 300), (2, 0, 0), (5, 0, 0), (200, 0, 0)]
        buffer = io.BytesIO()
        sample_chart(heights).save(buffer, format='JPEG', quality=90)
        image = Image.open(io.BytesIO(buffer.getvalue())).convert('RGB')
        measured = np.array(self.measure(image, len(heights)))
        np.testing.assert_allclose(measured[1:, 1:], 0, atol=0.01)
        np.testing.assert_allclose(measured[1:, 0], (2, 5, 200), atol=0.4)

    def test_color_validation_preserves_actual_coal_above_900(self):
        heights = [(60, 20, 300)] + [(0, 0, 0)] * 48
        heights[42] = (2.5, 1.5, 0)
        result = analyze_chart_image(sample_chart(heights), 38000)
        self.assertEqual('(1025, 1050]', result.rows[42].interval)
        self.assertAlmostEqual(250, result.rows[42].coal, delta=5)
        self.assertAlmostEqual(150, result.rows[42].gas, delta=5)

    def test_summary_calibration_does_not_amplify_false_coal_in_gas_bar(self):
        heights = [(60, 20, 300)] + [(0, 0, 0)] * 48
        heights[35], heights[42], heights[48] = (5, 0, 0), (0, 2, 0), (0, 200, 0)
        array = np.array(sample_chart(heights))
        center = 56 + 42 * 32
        array[548, center - 5:center + 5] = (151, 92, 101)
        array[549, center - 5:center + 5] = (158, 93, 106)
        declaration = DeclarationInfo(
            DeclarationRow(38000, 6000, 2000, 30000),
            DeclarationRow(20700, 500, 20200, 0),
        )
        result = analyze_chart_image(Image.fromarray(array), 38000, declaration)
        self.assertEqual(0, sum(row.coal for row in result.rows[37:]))
        self.assertAlmostEqual(500, sum(row.coal for row in result.rows[25:]))
        self.assertAlmostEqual(20200, sum(row.gas for row in result.rows[25:]))


if __name__ == "__main__":
    unittest.main()
