from __future__ import annotations

import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook

import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
PROJECT = ROOT / "广东电价预测"


def load_script(name: str, filename: str):
    path = PROJECT / filename
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


class GuangdongDataUpdateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.update = load_script("test_guangdong_update", "update_gdfire_data.py")

    def test_modeling_frame_drops_actual_and_realtime_leakage(self) -> None:
        frame = pd.DataFrame(
            {
                "日期": pd.to_datetime(["2026-07-11"]),
                "时刻": ["00:00"],
                "全省日前平均电价": [350.0],
                "日前统调负荷(MW)": [100000.0],
                "全省实时平均电价": [360.0],
                "实际统调负荷(MW)": [101000.0],
                "全空字段": [np.nan],
                "来源文件": ["接口"],
            }
        )

        result = self.update.build_modeling_frame(frame)

        self.assertIn("全省日前平均电价", result)
        self.assertIn("日前统调负荷(MW)", result)
        self.assertNotIn("全省实时平均电价", result)
        self.assertNotIn("实际统调负荷(MW)", result)
        self.assertNotIn("全空字段", result)
        self.assertNotIn("来源文件", result)

    def test_missing_dates_finds_internal_gap_and_incomplete_day(self) -> None:
        rows = []
        for day, slots in (("2026-07-09", 96), ("2026-07-11", 95)):
            for slot in range(slots):
                rows.append(
                    {
                        "日期": day,
                        "时刻": f"{slot // 4:02d}:{slot % 4 * 15:02d}",
                        "全省日前平均电价": 300.0,
                    }
                )

        result = self.update.missing_dates(pd.DataFrame(rows), "2026-07-11")

        self.assertEqual(
            [pd.Timestamp("2026-07-10").date(), pd.Timestamp("2026-07-11").date()],
            result,
        )

    def test_config_falls_back_to_shared_account_file(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            missing = Path(folder) / "missing.json"
            original = self.update.SHARED_CONFIG_PATH
            shared = Path(folder) / "shared.json"
            shared.write_text('{"username":"u","password":"p"}', encoding="utf-8")
            self.update.SHARED_CONFIG_PATH = shared
            try:
                result = self.update.load_config(missing)
            finally:
                self.update.SHARED_CONFIG_PATH = original
        self.assertEqual("u", result["username"])

    def test_update_uses_relative_initial_path_and_explicit_master(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder).resolve()
            initial = root / "输入" / "初始.xlsx"
            initial.parent.mkdir()
            frame = pd.DataFrame({
                "日期": pd.Timestamp("2026-07-11"),
                "时刻": pd.date_range("2026-07-11", periods=96, freq="15min").strftime("%H:%M"),
                "全省日前平均电价": 300.0,
            })
            frame.to_excel(initial, index=False, sheet_name="整合数据")
            initial_bytes = initial.read_bytes()
            config = root / "config.json"
            config.write_text('{"initialWorkbook":"输入/初始.xlsx","masterWorkbook":"错误.xlsx"}')
            selected = root / "自选输出" / "总表.xlsx"
            with mock.patch.object(self.update, "SCRIPT_DIR", root), mock.patch.object(
                self.update, "load_config", return_value={"username": "u", "password": "p"}
            ), mock.patch.object(self.update, "GDFireClient") as client:
                result = self.update.update_master_workbook(
                    run_date="2026-07-11", config_path=config, master_workbook=Path("自选输出/总表.xlsx")
                )
            client.assert_not_called()
            self.assertEqual(selected, result["master"])
            self.assertFalse((root / "错误.xlsx").exists())
            self.assertEqual(initial_bytes, initial.read_bytes())
            with pd.ExcelFile(selected) as book:
                self.assertEqual(["整合数据"], book.sheet_names)
                readback = pd.read_excel(book)
            self.assertEqual(list(frame.columns), list(readback.columns))
            self.assertEqual(96, len(readback))
            self.assertEqual(300, readback.iloc[-1]["全省日前平均电价"])


class GuangdongPathTests(unittest.TestCase):
    def test_paths_are_module_relative_and_shared_by_toolbox(self) -> None:
        from toolbox.runtime import ToolPaths

        paths_module = load_script("test_guangdong_paths", "gd_price_forecast/paths.py")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder).resolve()
            module_dir = root / "广东电价预测"
            module_dir.mkdir()
            (module_dir / "config.json").write_text(
                '{"masterWorkbook":"业务数据/总表.xlsx","predictionOutputDir":"业务结果/日前"}'
            )
            resolved = paths_module.resolve_paths(module_dir)
            paths = ToolPaths(root)
            self.assertEqual(module_dir / "业务数据/总表.xlsx", resolved["masterWorkbook"])
            self.assertEqual(resolved["masterWorkbook"], paths.guangdong_price_master_workbook)
            self.assertEqual(resolved["predictionOutputDir"], paths.guangdong_prediction_output_dir)
            self.assertEqual(module_dir / "输入/广东数据初始总表.xlsx", resolved["initialWorkbook"])


class GuangdongPageTests(unittest.TestCase):
    def test_page_uses_config_and_passes_selected_paths_to_tasks(self) -> None:
        import tkinter as tk
        from datetime import date
        from toolbox.pages import GuangdongPricePage
        from toolbox.runtime import TaskRegistry, ToolPaths

        try:
            root = tk.Tk()
            root.withdraw()
        except tk.TclError as exc:
            self.skipTest(str(exc))
        try:
            with tempfile.TemporaryDirectory() as folder:
                workspace = Path(folder).resolve()
                module_dir = workspace / "广东电价预测"
                module_dir.mkdir()
                (module_dir / "config.json").write_text(
                    '{"masterWorkbook":"数据/总表.xlsx","predictionOutputDir":"预测"}'
                )
                page = GuangdongPricePage(root, ToolPaths(workspace), TaskRegistry(), today=date(2026, 10, 10))
                self.assertEqual(module_dir / "数据/总表.xlsx", Path(page.master_workbook.get()))
                self.assertEqual(module_dir / "预测", Path(page.prediction_output_dir.get()))
                page.master_workbook.set("自选/总表.xlsx")
                module = mock.Mock()
                with mock.patch.object(page, "run_in_thread") as run, mock.patch(
                    "toolbox.pages.load_module", return_value=module
                ):
                    page.start_update()
                    run.call_args.args[0]()
                selected = module_dir / "自选/总表.xlsx"
                module.update_master_workbook.assert_called_once_with(run_date="2026-10-10", master_workbook=selected)
                selected.parent.mkdir()
                selected.touch()
                page.prediction_output_dir.set("自选/预测")
                with mock.patch("toolbox.pages._open_directory") as open_directory:
                    page.open_business_directory(page.prediction_output_dir.get())
                open_directory.assert_called_once_with(module_dir / "自选/预测")
                with mock.patch.object(page, "run_process") as run:
                    page.start_prediction()
                command = run.call_args.args[0]
                self.assertEqual(str(selected), command[command.index("--input") + 1])
                self.assertEqual(str(module_dir / "自选/预测"), command[command.index("--output-dir") + 1])
                self.assertEqual("1", command[command.index("--forecast-days") + 1])
                self.assertEqual("0", command[command.index("--training-months") + 1])
                self.assertEqual("2026-10-10", command[command.index("--as-of") + 1])
                page.prediction_output_dir.set("")
                with mock.patch.object(page, "run_process") as run, mock.patch("toolbox.pages.messagebox.showerror") as error:
                    page.start_prediction()
                run.assert_not_called()
                error.assert_called_once()
                page.prediction_output_dir.set("自选/预测")
                page.run_date.set("无效日期")
                with mock.patch.object(page, "run_process") as run, mock.patch("toolbox.pages.messagebox.showerror") as error:
                    page.start_prediction()
                run.assert_not_called()
                error.assert_called_once()
                root.update_idletasks()
        finally:
            root.destroy()


@unittest.skipUnless(
    importlib.util.find_spec("joblib") and importlib.util.find_spec("sklearn"),
    "forecast dependencies are installed in the toolbox virtual environment",
)
class GuangdongForecastTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        sys.path.insert(0, str(PROJECT))
        cls.forecast = load_script("test_guangdong_forecast", "train_d1_forecast.py")

    @classmethod
    def tearDownClass(cls) -> None:
        if sys.path and sys.path[0] == str(PROJECT):
            sys.path.pop(0)

    def test_append_prediction_day_keeps_96_time_slots_and_blank_target(self) -> None:
        times = pd.date_range("2026-07-11", periods=96, freq="15min").strftime("%H:%M")
        frame = pd.DataFrame(
            {
                "日期": pd.Timestamp("2026-07-11"),
                "时刻": times,
                "全省日前平均电价": np.arange(96, dtype=float),
                "日前统调负荷(MW)": np.arange(96, dtype=float) + 100000,
            }
        )

        result = self.forecast.append_prediction_days(frame, 1)
        future = result[result["日期"].eq(pd.Timestamp("2026-07-12"))]

        self.assertEqual(96, len(future))
        self.assertEqual(96, future["时刻"].nunique())
        self.assertTrue(future["全省日前平均电价"].isna().all())

    def test_prediction_cli_reads_configured_paths(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            config = Path(folder) / "config.json"
            config.write_text('{"masterWorkbook":"自选/总表.xlsx","predictionOutputDir":"自选/预测"}')
            from toolbox_shared.storage import data_path
            master = data_path(PROJECT / "自选/总表.xlsx")
            with mock.patch.object(sys, "argv", ["train_d1_forecast.py", "--config", str(config)]), mock.patch.object(
                self.forecast, "ensure_master_workbook", return_value=master
            ) as ensure, mock.patch.object(
                self.forecast, "run_pipeline", return_value={"mae": 0.0}
            ) as pipeline, mock.patch("builtins.print"):
                self.forecast.main()
            self.assertEqual(master, ensure.call_args.args[0]["masterWorkbook"])
            self.assertEqual(
                (master, data_path(PROJECT / "自选/预测")), pipeline.call_args.args
            )

    def test_pipeline_exports_actual_workbook_and_rejects_unsupported_days(self) -> None:
        times = pd.date_range("2026-05-01", periods=50 * 96, freq="15min")
        frame = pd.DataFrame({
            "日期": times.normalize(),
            "时刻": times.strftime("%H:%M"),
            "全省日前平均电价": 300 + np.sin(np.arange(len(times)) / 96) * 20,
        })
        with tempfile.TemporaryDirectory() as folder:
            input_path = Path(folder) / "输入.xlsx"
            output_dir = Path(folder) / "自选结果"
            frame.to_excel(input_path, index=False, sheet_name="整合数据")
            original = input_path.read_bytes()
            summary = self.forecast.run_pipeline(
                input_path, output_dir,
                configs=[{"name": "smoke", "kind": "hist_gradient_boosting", "params": {"max_iter": 5}}],
            )
            self.assertTrue(np.isfinite(summary["mae"]))
            self.assertEqual({"D+1日前电价.xlsx", "MAE.json"}, {p.name for p in output_dir.iterdir()})
            book = load_workbook(output_dir / "D+1日前电价.xlsx")
            try:
                self.assertEqual(["D+1日前电价"], book.sheetnames)
                sheet = book.active
                self.assertEqual((97, 3), (sheet.max_row, sheet.max_column))
                self.assertEqual(["日期", "时刻", "日前电价"], [c.value for c in sheet[1]])
                self.assertEqual(pd.Timestamp("2026-06-20"), sheet["A2"].value)
                self.assertEqual("00:00", sheet["B2"].value)
                self.assertEqual("23:45", sheet["B97"].value)
                self.assertTrue(np.isfinite(sheet["C2"].value))
                self.assertEqual("0.00", sheet["C2"].number_format)
                self.assertNotEqual("General", sheet["A2"].number_format)
            finally:
                book.close()
            self.assertEqual(original, input_path.read_bytes())
            with self.assertRaisesRegex(ValueError, "预测天数必须为 1"):
                self.forecast.run_pipeline(input_path, output_dir, forecast_days=2)

    def test_training_window_keeps_only_recent_three_months(self) -> None:
        dates = pd.date_range("2025-01-01", "2026-07-11", freq="D")
        frame = pd.DataFrame({"日期": dates, "时刻": "00:00", "全省日前平均电价": 1.0})

        result = self.forecast.slice_training_window(frame, 3)

        self.assertEqual(pd.Timestamp("2026-04-11"), result["日期"].min())
        self.assertEqual(pd.Timestamp("2026-07-11"), result["日期"].max())

    def test_sparse_exogenous_field_is_not_used_as_a_feature(self) -> None:
        from gd_price_forecast.features import build_features

        dates = pd.date_range("2026-06-01", periods=10, freq="D")
        rows = []
        for day in dates:
            for slot in range(96):
                rows.append(
                    {
                        "日期": day,
                        "时刻": f"{slot // 4:02d}:{(slot % 4) * 15:02d}",
                        "全省日前平均电价": float(slot),
                        "完整字段": float(slot + 1),
                        "仅一天字段": float(slot) if day == dates[-1] else np.nan,
                    }
                )
        features, _, _ = build_features(pd.DataFrame(rows))

        self.assertIn("完整字段", features)
        self.assertNotIn("仅一天字段", features)

    def test_master_workbook_loader_drops_realtime_and_actual_columns(self) -> None:
        from gd_price_forecast.data import load_model_data

        times = pd.date_range("2026-07-11", periods=96, freq="15min").strftime("%H:%M")
        frame = pd.DataFrame(
            {
                "日期": pd.Timestamp("2026-07-11"),
                "时刻": times,
                "全省日前平均电价": 300.0,
                "日前统调负荷(MW)": 100000.0,
                "实际统调负荷(MW)": 101000.0,
                "全省实时平均电价": 310.0,
            }
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "master.xlsx"
            frame.to_excel(path, index=False, sheet_name="整合数据")
            result = load_model_data(path)

        self.assertIn("日前统调负荷(MW)", result)
        self.assertNotIn("实际统调负荷(MW)", result)
        self.assertNotIn("全省实时平均电价", result)

    def test_d1_output_is_one_sheet_with_96_price_points(self) -> None:
        from gd_price_forecast.reporting import write_d1_price_workbook

        times = pd.date_range("2026-07-13", periods=96, freq="15min")
        prediction = pd.DataFrame(
            {
                "日期": times.normalize(),
                "时刻": times.strftime("%H:%M"),
                "模型预测": np.arange(96),
                "D-1基线": np.arange(96),
                "预测日前电价": np.arange(96, dtype=float) + 300,
            }
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "D+1日前电价.xlsx"
            write_d1_price_workbook(path, prediction)
            with pd.ExcelFile(path) as workbook:
                sheet_names = workbook.sheet_names
            result = pd.read_excel(path, sheet_name="D+1日前电价")

        self.assertEqual(["D+1日前电价"], sheet_names)
        self.assertEqual(["日期", "时刻", "日前电价"], list(result.columns))
        self.assertEqual(96, len(result))
        self.assertEqual(96, result["时刻"].nunique())

    def test_recursive_prediction_can_exclude_historical_missing_dates(self) -> None:
        import unittest.mock as mock

        times = pd.date_range("2026-07-12", periods=96, freq="15min")
        old_missing = pd.DataFrame(
            {
                "日期": pd.Timestamp("2026-07-01"),
                "时刻": times.strftime("%H:%M"),
                "全省日前平均电价": np.nan,
            }
        )
        d1 = old_missing.copy()
        d1["日期"] = pd.Timestamp("2026-07-13")
        frame = pd.concat([old_missing, d1], ignore_index=True)
        identity = frame[["日期", "时刻"]].copy()
        features = pd.DataFrame({"price_lag_1d": 300.0}, index=frame.index)
        target = frame["全省日前平均电价"]

        with mock.patch.object(
            self.forecast, "build_features", return_value=(features, target, identity)
        ), mock.patch.object(
            self.forecast, "predict_model", return_value=np.full(96, 310.0)
        ):
            result = self.forecast._recursive_prediction(
                frame,
                model=object(),
                blend_weight=0.7,
                prediction_dates=[pd.Timestamp("2026-07-13")],
            )

        self.assertEqual([pd.Timestamp("2026-07-13")], list(result["日期"].unique()))


class ForecastMasterBootstrapTests(unittest.TestCase):
    """首次安装不再内置总表，预测入口要在缺失时自动生成。"""

    @classmethod
    def setUpClass(cls) -> None:
        sys.path.insert(0, str(PROJECT))
        cls.train = load_script("test_guangdong_train", "train_d1_forecast.py")

    @classmethod
    def tearDownClass(cls) -> None:
        if sys.path and sys.path[0] == str(PROJECT):
            sys.path.pop(0)

    def test_missing_master_is_generated_from_grid_api(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            master = Path(folder) / "输出/广东电价数据总表.xlsx"
            calls = []

            def fake_update(*, master_workbook=None):
                calls.append(Path(master_workbook))
                Path(master_workbook).parent.mkdir(parents=True, exist_ok=True)
                Path(master_workbook).write_bytes(b"generated")
                return {"master": str(master_workbook)}

            stub = mock.Mock(update_master_workbook=fake_update)
            with mock.patch.dict(sys.modules, {"update_gdfire_data": stub}):
                result = self.train.ensure_master_workbook({"masterWorkbook": master})

            self.assertEqual(master, result)
            self.assertTrue(master.is_file())
            self.assertEqual([master], calls)

    def test_existing_master_is_used_without_generating(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            master = Path(folder) / "总表.xlsx"
            master.write_bytes(b"local")
            stub = mock.Mock()
            with mock.patch.dict(sys.modules, {"update_gdfire_data": stub}):
                result = self.train.ensure_master_workbook({"masterWorkbook": master})
            self.assertEqual(master, result)
            stub.update_master_workbook.assert_not_called()

    def test_explicit_missing_input_is_not_silently_replaced(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            missing = Path(folder) / "指定输入.xlsx"
            stub = mock.Mock()
            with mock.patch.dict(sys.modules, {"update_gdfire_data": stub}):
                with self.assertRaisesRegex(FileNotFoundError, "指定的预测输入工作簿不存在"):
                    self.train.ensure_master_workbook(
                        {"masterWorkbook": Path(folder) / "总表.xlsx"}, missing
                    )
            stub.update_master_workbook.assert_not_called()


if __name__ == "__main__":
    unittest.main()
