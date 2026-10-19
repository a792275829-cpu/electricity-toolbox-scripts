import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook


SCRIPT = Path(__file__).resolve().parents[1] / "generate_electricity_report.py"
SPEC = importlib.util.spec_from_file_location("generate_electricity_report", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


class ReportSummaryTests(unittest.TestCase):
    def test_renewable_values_and_full_price_use_bounded_read(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in MODULE.RENEWABLE_UNITS:
            sheet = workbook.create_sheet(name)
            for hour in range(24):
                sheet.cell(hour + 4, 1, f"{hour}:00")
                for column, value in ((2, 10), (3, 300), (4, 12), (5, 350), (12, 8), (13, 250)):
                    sheet.cell(hour + 4, column, value)
            sheet["AE28"] = 321.5
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "renewables.xlsx"
            workbook.save(path)
            builder = MODULE.ReportBuilder(path)
            builder.wb = load_workbook(path, read_only=True, data_only=True)
            try:
                from openpyxl.worksheet._read_only import ReadOnlyWorksheet
                with mock.patch.object(ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")):
                    builder._parse_renewables()
            finally:
                builder.wb.close()
        self.assertEqual(builder.summary["海风"]["da_mwh"], 240)
        self.assertEqual(builder.summary["海风"]["rt_mwh"], 288)
        self.assertEqual(builder.summary["海风"]["full_price"], 321.5)
        self.assertAlmostEqual(builder.summary["海风"]["profit_wanyuan"], 0.24)

    def test_large_sheets_parse_offer_and_hourly_rows(self):
        workbook = Workbook()
        compensation = workbook.active
        compensation.title = "运行补偿"
        compensation.cell(100, 1, "汕头#1机组")
        compensation.cell(101, 17, "厂用电率")
        compensation.cell(101, 18, 0.05)
        compensation.cell(102, 17, "段号")
        for row, start, end, price in ((103, 0, 100, 300), (104, 100, 200, 300)):
            compensation.cell(row, 18, start)
            compensation.cell(row, 19, end)
            compensation.cell(row, 20, price)
        compensation.cell(1000, 1, "末尾")

        hourly = workbook.create_sheet("汕头")
        hourly.cell(1, 11, "偏差")
        hourly.cell(50, 1, "1号机")
        for hour in range(24):
            row = 53 + hour
            hourly.cell(row, 1, f"{hour}:00")
            hourly.cell(row, 2, 10)
            hourly.cell(row, 3, 200)
            hourly.cell(row, 5, 12)
            hourly.cell(row, 6, 300)
            hourly.cell(row, 7, 11)
            hourly.cell(row, 8, 280)
            hourly.cell(row, 10, 100)
        hourly.cell(1000, 1, "末尾")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "trade.xlsx"
            workbook.save(path)
            builder = MODULE.ReportBuilder(path)
            builder.wb = load_workbook(path, read_only=True, data_only=True)
            try:
                cards = builder._parse_offer_cards()
                blocks = builder._parse_blocks("汕头", "汕头")
            finally:
                builder.wb.close()

        self.assertEqual(cards["汕头#1机组"]["plant_rate"], 0.05)
        self.assertEqual(cards["汕头#1机组"]["segments"], [
            {"start": 0.0, "end": 200.0, "price": 300.0, "source_count": 2}
        ])
        self.assertEqual(len(blocks["汕头#1机组"]), 24)
        self.assertEqual(blocks["汕头#1机组"][0]["profit_ml"], 0.04)

    def test_new_summary_layout_is_parsed(self):
        workbook = Workbook()
        summary = workbook.active
        summary.title = "新"
        brief = workbook.create_sheet("快报")
        summary.cell(8, 1, "火电")
        summary.cell(8, 4, "中长期")
        summary.cell(8, 7, "日前出清")
        summary.cell(10, 3, 321.5)
        summary.cell(10, 4, 100.25)
        summary.cell(10, 7, 120.5)
        summary.cell(10, 8, 456.7)
        summary.cell(10, 9, 20.25)
        summary.cell(10, 10, 500.0)
        summary.cell(10, 11, 12.3)
        brief.cell(7, 3, 576.76)
        brief.cell(7, 11, 0.8624)

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "new_summary.xlsx"
            workbook.save(path)
            builder = MODULE.ReportBuilder(path)
            builder.wb = load_workbook(path, read_only=True, data_only=True)
            try:
                from openpyxl.worksheet._read_only import ReadOnlyWorksheet
                with mock.patch.object(ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")):
                    builder._parse_summary()
            finally:
                builder.wb.close()

        parsed = builder.summary["分公司火电"]
        self.assertEqual(576.76, parsed["capacity_10mw"])
        self.assertEqual(321.5, parsed["cost"])
        self.assertEqual(0.8624, parsed["load_rate"])
        self.assertEqual(1205.0, parsed["da_mwh"])
        self.assertEqual(1002.5, parsed["ml_mwh"])
        self.assertEqual(202.5, parsed["dev_mwh"])
        self.assertEqual(12.3, parsed["profit_wanyuan"])

    def test_unrecognized_workbook_reports_available_sheets(self):
        workbook = Workbook()
        workbook.active.title = "其他数据"
        builder = MODULE.ReportBuilder("unused.xlsx")
        builder.wb = workbook

        with self.assertRaisesRegex(ValueError, "当前工作表：其他数据"):
            builder._parse_summary()

    def test_company_layout_locates_rows_by_labels_and_rejects_missing_units(self):
        workbook = Workbook()
        summary = workbook.active
        summary.title = "新"
        brief = workbook.create_sheet("快报")
        summary["A8"] = "南方分公司"
        summary["D8"] = "中长期"
        summary["G8"] = "日前出清"
        labels = {
            10: ("火电机组合计", None), 11: ("燃煤机组合计", None),
            13: ("汕头", "合计"), 14: (None, "#1机组"),
            15: (None, "#2机组"), 16: (None, "#3机组"),
            18: ("海门", "合计"), 19: (None, "#1机组"),
            20: (None, "#2机组"), 21: (None, "#3机组"),
            22: (None, "#4机组"), 23: ("燃气机组合计", None),
            25: ("东莞", "合计"), 28: ("新能源", None),
            30: ("分公司", None),
        }
        for row, (plant, unit) in labels.items():
            if plant: summary.cell(row, 1, plant)
            if unit: summary.cell(row, 2, unit)
            for column, value in ((3, 300 + row), (4, row), (7, row * 2), (8, 400 + row), (9, row), (10, 500 + row), (11, row / 10)):
                summary.cell(row, column, value)
        brief["C8"] = 120
        brief["K8"] = 0.8
        builder = MODULE.ReportBuilder("unused.xlsx")
        builder.wb = workbook
        builder._parse_summary()
        self.assertEqual(builder.summary["汕头"]["da_mwh"], 260)
        self.assertEqual(builder.summary["汕头"]["capacity_10mw"], 120)
        self.assertEqual(builder.summary["汕头"]["load_rate"], 0.8)
        self.assertEqual(builder.summary["海门#4机组"]["da_price"], 422)
        self.assertEqual(builder.summary["东莞"]["profit_wanyuan"], 2.5)
        self.assertEqual(builder.summary["分公司火电"]["da_mwh"], 200)
        summary["B22"] = None
        with self.assertRaisesRegex(ValueError, "缺少标签.*海门#4机组"):
            builder._parse_summary()


if __name__ == "__main__":
    unittest.main()
