from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from toolbox.runtime import load_module


ROOT = Path(__file__).resolve().parents[2]


class ExcelStreamReadTests(unittest.TestCase):
    def test_sparse_read_preserves_values_formats_and_scans_each_sheet_once(self):
        reader = load_module("excel_stream_reader_test", ROOT / "toolbox_shared/excel_read.py")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "数据"
            sheet["A2"] = datetime(2026, 10, 5)
            sheet["B2"] = 0.125
            sheet["B2"].number_format = "0.00%"
            sheet["C2"] = False
            sheet["D2"] = "=1+1"
            sheet["B20"] = 42
            workbook.create_sheet("补充")["C3"] = "完成"
            workbook.save(path)
            reopened = load_workbook(path, read_only=True, data_only=False)
            requests = [("数据", cell) for cell in ("A2", "B2", "C2", "D2", "B20", "A30")]
            requests += [("补充", "C3")]
            original_iter = ReadOnlyWorksheet.iter_rows
            try:
                with mock.patch.object(ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")), \
                     mock.patch.object(ReadOnlyWorksheet, "iter_rows", autospec=True, side_effect=original_iter) as scanned:
                    result = reader.read_cell_snapshots(reopened, requests)
                self.assertEqual(scanned.call_count, 2)
            finally:
                reopened.close()
            self.assertEqual(result[("数据", "A2")][0], datetime(2026, 10, 5))
            self.assertEqual(result[("数据", "B2")], (0.125, "0.00%"))
            self.assertEqual(result[("数据", "C2")][0], False)
            self.assertEqual(result[("数据", "D2")][0], "=1+1")
            self.assertEqual(result[("数据", "B20")][0], 42)
            self.assertIsNone(result[("数据", "A30")][0])
            self.assertEqual(result[("补充", "C3")][0], "完成")

    def test_spot_plan_inherits_styles_from_blank_date_rows_without_random_reads(self):
        module = load_module("spot_stream_plan_test", ROOT / "现货价格更新/update_spot_prices.py")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prices.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = module.DAY_AHEAD_SHEET
            sheet["A2"] = date(2026, 10, 1)
            sheet["B2"].number_format = "0.00%"
            sheet["C2"].number_format = "0.000"
            sheet["C3"].fill = PatternFill("solid", fgColor="00FF00")
            sheet["D3"].number_format = "0.0000"
            sheet["A4"] = date(2026, 10, 5)
            sheet["A5"] = date(2026, 10, 4)
            sheet["B5"] = "=1+1"
            workbook.save(path)
            expected = tuple(float(index) for index in range(24))
            reopened = load_workbook(path, read_only=True, data_only=False)
            try:
                with mock.patch.object(ReadOnlyWorksheet, "cell", side_effect=AssertionError("random access")):
                    expectations, updated, skipped, _ = module.plan_price_sheet(
                        reopened, module.DAY_AHEAD_SHEET,
                        {date(2026, 10, 5): expected, date(2026, 10, 4): expected}, overwrite=False,
                    )
                self.assertEqual(updated, [date(2026, 10, 5)])
                self.assertEqual(skipped, [date(2026, 10, 4)])
                self.assertEqual(expectations[0].number_formats[:3], ("0.00%", "General", "0.0000"))
            finally:
                reopened.close()

    def test_duplicate_date_after_target_still_blocks_update(self):
        module = load_module("spot_stream_duplicate_test", ROOT / "现货价格更新/update_spot_prices.py")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = module.DAY_AHEAD_SHEET
        sheet["A2"] = date(2026, 10, 5)
        sheet["A20"] = date(2026, 10, 5)
        with self.assertRaisesRegex(RuntimeError, "日期重复.*第 2、20 行"):
            module.plan_price_sheet(workbook, sheet.title, {date(2026, 10, 5): (1.0,) * 24}, overwrite=True)

    def test_spot_readback_accepts_generator_and_rejects_format_or_value_changes(self):
        module = load_module("spot_stream_verify_test", ROOT / "现货价格更新/update_spot_prices.py")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prices.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = module.DAY_AHEAD_SHEET
            for column in range(2, 26):
                sheet.cell(2, column, 300 + column)
                sheet.cell(2, column).number_format = "0.000"
            expected = module.CellExpectation(sheet.title, 2, tuple(300.0 + c for c in range(2, 26)), ("0.000",) * 24)
            workbook.save(path)
            module.verify_workbook(path, (item for item in [expected]))
            sheet["Y2"].number_format = "0.0%"
            workbook.save(path)
            with self.assertRaisesRegex(RuntimeError, "格式回读失败"):
                module.verify_workbook(path, [expected])
            sheet["Y2"].number_format = "0.000"
            sheet["Y2"] = 999
            workbook.save(path)
            with self.assertRaisesRegex(RuntimeError, "保存回读失败"):
                module.verify_workbook(path, [expected])


if __name__ == "__main__":
    unittest.main()
