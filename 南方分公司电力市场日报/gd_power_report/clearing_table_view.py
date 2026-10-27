"""Excel table presentation: styles, display values and merged-cell layout.

No report date inference, business aggregation or file access lives here.
"""
from __future__ import annotations
import re
from colorsys import hls_to_rgb, rgb_to_hls
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import Any
from xml.etree import ElementTree
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, range_boundaries

ERROR_VALUES = {"#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}

def _is_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() in ERROR_VALUES


def _display_value(value: Any, number_format: str = "General") -> str:
    if value is None or isinstance(value, MergedCell):
        return ""
    if _is_excel_error(value):
        return "—"
    if isinstance(value, (datetime, date)):
        title_formats = {
            "分公司日前出清复盘": "分公司日前出清复盘",
            "火电分时明细": "火电分时明细",
            "新能源分时明细": "新能源分时明细",
            "分公司实时出清复盘": "分公司实时出清复盘",
        }
        for marker, title in title_formats.items():
            if marker in number_format:
                return f"{title}（{value.year}/{value.month}/{value.day}）"
        return f"{value.month}月{value.day}日"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if abs(value) >= 1_000_000_000_000:
            return "—"
        if "%" in number_format:
            decimals_match = re.search(r"\.([0#]+)%", number_format)
            decimals = len(decimals_match.group(1)) if decimals_match else 0
            return f"{float(value) * 100:.{decimals}f}%"
        if abs(value) < 1e-10:
            return "0"
        format_section = str(number_format or "General").split(";", 1)[0]
        placeholder = re.search(r"([#0?][#0?,]*)(?:\.([0#?]+))?", format_section)
        if placeholder and str(number_format).lower() != "general":
            decimal_pattern = placeholder.group(2) or ""
            decimals = len(decimal_pattern)
            rounded = Decimal(str(value)).quantize(
                Decimal(1).scaleb(-decimals),
                rounding=ROUND_HALF_UP,
            )
            if rounded == 0:
                rounded = abs(rounded)
            use_grouping = "," in placeholder.group(1)
            text = f"{rounded:,.{decimals}f}" if use_grouping else f"{rounded:.{decimals}f}"
            optional_decimals = len(decimal_pattern) - len(decimal_pattern.rstrip("#"))
            if optional_decimals and "." in text:
                integer, fraction = text.split(".", 1)
                required = decimals - optional_decimals
                fraction = fraction.rstrip("0")
                if len(fraction) < required:
                    fraction = fraction.ljust(required, "0")
                text = f"{integer}.{fraction}" if fraction else integer
            return text
        if float(value).is_integer():
            return f"{int(value)}"
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def _theme_colors(workbook: Any) -> list[str]:
    if not workbook.loaded_theme:
        return []
    try:
        root = ElementTree.fromstring(workbook.loaded_theme)
    except (ElementTree.ParseError, TypeError):
        return []
    scheme = next((item for item in root.iter() if item.tag.endswith("}clrScheme")), None)
    if scheme is None:
        return []
    colors_by_name: dict[str, str] = {}
    for entry in scheme:
        color = next(iter(entry), None)
        name = entry.tag.rsplit("}", 1)[-1]
        raw = "000000" if color is None else color.attrib.get("lastClr") or color.attrib.get("val") or "000000"
        colors_by_name[name] = raw[-6:].upper()
    # Excel's theme index order is bg1, tx1, bg2, tx2, accents, links;
    # clrScheme stores the first four as dk1, lt1, dk2, lt2.
    index_order = (
        "lt1",
        "dk1",
        "lt2",
        "dk2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    )
    return [colors_by_name.get(name, "000000") for name in index_order]


def _apply_tint(rgb: str, tint: float) -> str:
    if not tint:
        return rgb.upper()
    red, green, blue = (int(rgb[index : index + 2], 16) / 255 for index in (0, 2, 4))
    hue, lightness, saturation = rgb_to_hls(red, green, blue)
    if tint < 0:
        lightness *= 1 + tint
    else:
        lightness = lightness * (1 - tint) + tint
    tinted = hls_to_rgb(hue, max(0, min(1, lightness)), saturation)
    return "".join(f"{round(channel * 255):02X}" for channel in tinted)


def _rgb(value: Any, theme_colors: list[str], *, auto: str | None = None) -> str | None:
    if value is None:
        return None
    color_type = getattr(value, "type", None)
    rgb = ""
    if color_type == "rgb":
        rgb = str(getattr(value, "rgb", ""))
    elif color_type == "theme":
        theme = getattr(value, "theme", None)
        if isinstance(theme, int) and 0 <= theme < len(theme_colors):
            rgb = theme_colors[theme]
    elif color_type == "indexed":
        indexed = getattr(value, "indexed", None)
        if isinstance(indexed, int) and 0 <= indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[indexed]
    elif color_type == "auto" or getattr(value, "auto", False):
        return auto
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) == 6:
        tint = float(getattr(value, "tint", 0) or 0)
        return f"#{_apply_tint(rgb, tint)}"
    return None


def _border_rule(side: Any, edge: str, theme_colors: list[str]) -> str | None:
    style = getattr(side, "style", None)
    if not style:
        return None
    width = "1px"
    line_style = "solid"
    if style in {"medium", "mediumDashed", "mediumDashDot", "mediumDashDotDot"}:
        width = "2px"
    elif style == "thick":
        width = "3px"
    elif style == "double":
        width, line_style = "3px", "double"
    elif style in {"dashed", "dashDot", "dashDotDot", "slantDashDot"}:
        line_style = "dashed"
    elif style in {"dotted", "hair"}:
        line_style = "dotted"
    color = _rgb(getattr(side, "color", None), theme_colors, auto="#000000") or "#000000"
    return f"border-{edge}:{width} {line_style} {color}"


def _cell_style(cell: Any, theme_colors: list[str]) -> str:
    rules: list[str] = []
    fill = None
    if getattr(cell.fill, "fill_type", None) == "solid":
        fill = _rgb(getattr(cell.fill, "fgColor", None), theme_colors)
    font_color = _rgb(getattr(cell.font, "color", None), theme_colors, auto="#000000")
    if fill:
        rules.append(f"background-color:{fill}")
    if font_color:
        rules.append(f"color:{font_color}")
    if cell.font.name:
        family = str(cell.font.name).replace("'", "\\'")
        rules.append(f"font-family:'{family}','PingFang SC','Microsoft YaHei',sans-serif")
    if cell.font.sz:
        source_size = float(cell.font.sz)
        screen_size = max(9, min(16, source_size * 2 / 3))
        print_size = max(1.9, min(4.6, source_size * 4.6 / 24))
        rules.append(f"--excel-font-size:{screen_size:.2f}px")
        rules.append(f"--excel-print-font-size:{print_size:.2f}pt")
    if cell.font.bold:
        rules.append("font-weight:700")
    if cell.font.italic:
        rules.append("font-style:italic")
    decorations = []
    if cell.font.underline:
        decorations.append("underline")
    if cell.font.strike:
        decorations.append("line-through")
    if decorations:
        rules.append(f"text-decoration:{' '.join(decorations)}")
    horizontal = cell.alignment.horizontal
    if horizontal in {"left", "center", "right", "justify"}:
        rules.append(f"text-align:{horizontal}")
    if cell.alignment.vertical in {"top", "center", "bottom"}:
        vertical = "middle" if cell.alignment.vertical == "center" else cell.alignment.vertical
        rules.append(f"vertical-align:{vertical}")
    if cell.alignment.wrap_text:
        rules.append("white-space:normal")
    if cell.alignment.indent:
        rules.append(f"padding-left:{3 + int(cell.alignment.indent) * 9}px")
    for edge in ("left", "right", "top", "bottom"):
        rule = _border_rule(getattr(cell.border, edge), edge, theme_colors)
        if rule:
            rules.append(rule)
    if (
        isinstance(cell.value, (int, float))
        and not isinstance(cell.value, bool)
        and cell.value < 0
        and "[red]" in str(cell.number_format).lower()
    ):
        rules.append("color:#FF0000")
    return ";".join(rules)


def _table_model(sheet: Any, cell_range: str, theme_colors: list[str]) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    merge_by_anchor: dict[tuple[int, int], tuple[int, int]] = {}
    merged_children: set[tuple[int, int]] = set()
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row >= min_row
            and merged.max_row <= max_row
            and merged.min_col >= min_col
            and merged.max_col <= max_col
        ):
            anchor = (merged.min_row, merged.min_col)
            merge_by_anchor[anchor] = (
                merged.max_row - merged.min_row + 1,
                merged.max_col - merged.min_col + 1,
            )
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    if (row, column) != anchor:
                        merged_children.add((row, column))

    rows: list[dict[str, Any]] = []
    for row in range(min_row, max_row + 1):
        if sheet.row_dimensions[row].hidden:
            continue
        cells: list[dict[str, Any]] = []
        for column in range(min_col, max_col + 1):
            if (row, column) in merged_children:
                continue
            cell = sheet.cell(row, column)
            rowspan, colspan = merge_by_anchor.get((row, column), (1, 1))
            cells.append(
                {
                    "value": _display_value(cell.value, cell.number_format),
                    "rowspan": rowspan,
                    "colspan": colspan,
                    "style": _cell_style(cell, theme_colors),
                }
            )
        height = sheet.row_dimensions[row].height
        rows.append({"cells": cells, "height": round(float(height), 1) if height else None})

    widths = []
    for column in range(min_col, max_col + 1):
        width = sheet.column_dimensions[get_column_letter(column)].width or 9
        widths.append(max(64, min(210, round(float(width) * 7 + 10))))
    return {
        "range": cell_range,
        "rows": rows,
        "column_widths": widths,
        "min_width": sum(widths),
    }
