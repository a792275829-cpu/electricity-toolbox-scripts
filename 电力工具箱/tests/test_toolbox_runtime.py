from __future__ import annotations
import importlib

import gc
import subprocess
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
import runpy
import os
import shutil
from contextlib import ExitStack
from datetime import date
from datetime import timedelta
from pathlib import Path
from unittest import mock


class RuntimeTests(unittest.TestCase):
    """Workspace paths, module loading and shared runtime helpers."""
    def test_file_dialog_state_remembers_module_and_common_directories(self) -> None:
        from toolbox.file_dialogs import FileDialogState

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "first"
            second = root / "second"
            first.mkdir()
            second.mkdir()
            settings = root / "file_dialog_settings.json"
            state = FileDialogState(settings)

            state.remember_file("report_excel", first / "source.xlsx")
            self.assertEqual(state.initial_directory("report_excel"), first)
            self.assertEqual(state.initial_directory("group_upload_input"), first)

            state.remember_directory("group_upload_input", second)
            reloaded = FileDialogState(settings)
            self.assertEqual(reloaded.initial_directory("report_excel"), first)
            self.assertEqual(reloaded.initial_directory("group_upload_input"), second)
            self.assertEqual(reloaded.initial_directory("new_module"), second)

    def test_file_dialog_state_uses_existing_current_path_then_fallback(self) -> None:
        from toolbox.file_dialogs import FileDialogState

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "workbooks"
            folder.mkdir()
            state = FileDialogState(root / "settings.json")

            self.assertEqual(
                state.initial_directory("market_workbook", folder / "market.xlsx"),
                folder,
            )

    def test_tool_paths_resolve_from_workspace(self) -> None:
        from toolbox.runtime import ToolPaths

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            paths = ToolPaths(workspace)

            self.assertEqual(paths.workspace, workspace.resolve())
            self.assertEqual(
                paths.online_energy,
                workspace.resolve() / "上网电量抓取" / "export_online_energy.py",
            )
            self.assertEqual(
                paths.private_uploader,
                workspace.resolve()
                / "电力工具脚本"
                / "private-data-uploader-tool"
                / "scripts"
                / "upload-private-data.mjs",
            )
            self.assertEqual(
                paths.wps_writer,
                workspace.resolve() / "wps自动" / "wps_excel_to_kdocs_gui.py",
            )
            self.assertEqual(
                paths.market_table_update,
                workspace.resolve() / "市场表更新" / "update_market_table.py",
            )
            self.assertEqual(
                paths.daily_pdf_parser,
                workspace.resolve()
                / "广东运行日报PDF解析"
                / "run_daily_pdf_parser.py",
            )
            self.assertEqual(
                paths.spot_price_update,
                workspace.resolve() / "现货价格更新" / "update_spot_prices.py",
            )
            self.assertEqual(
                paths.transaction_info_update,
                workspace.resolve()
                / "交易系统信息更新"
                / "update_transaction_info.py",
            )

    def test_load_module_rejects_missing_file(self) -> None:
        from toolbox.runtime import load_module

        with tempfile.TemporaryDirectory() as directory:
            missing = Path(directory) / "missing.py"
            with self.assertRaisesRegex(FileNotFoundError, "missing.py"):
                load_module("missing_tool", missing)

    def test_virtualenv_dependency_is_not_treated_as_local_module(self) -> None:
        from toolbox.runtime import _module_is_local

        dependency = mock.Mock(
            __file__=str(
                Path(sys.prefix)
                / "lib"
                / "python3.11"
                / "site-packages"
                / "dependency.py"
            )
        )

        self.assertFalse(
            _module_is_local(dependency, (Path(sys.prefix).resolve().parent,))
        )

    def test_report_export_proxy_reads_environment(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("report_export_online_energy_test", paths.report_scripts_dir / "export_online_energy.py")

        with mock.patch.dict(os.environ, {"HTTPS_PROXY": "http://127.0.0.1:7890"}, clear=True):
            self.assertEqual(module.playwright_proxy(), {"server": "http://127.0.0.1:7890"})

    def test_login_state_paths_are_separate_for_parallel_tools(self) -> None:
        from toolbox_shared.storage import data_path
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        online = load_module("online_energy_state_paths_test", paths.online_energy)
        market = load_module("market_table_state_paths_test", paths.market_table_update)
        group = load_module("group_upload_state_paths_test", paths.group_upload)
        spot = load_module("spot_price_state_paths_test", paths.spot_price_update)
        transaction = load_module(
            "transaction_info_state_paths_test", paths.transaction_info_update
        )
        wind = load_module("wind_forecast_state_paths_test", paths.wind_forecast)

        loaded_auth = {
            "spot": spot.load_auth_module(),
            "transaction": transaction.load_auth_module(),
            "wind": wind.load_auth_module(),
        }
        for name, directory in (
            ("spot", paths.spot_price_update_dir),
            ("transaction", paths.transaction_info_update_dir),
            ("wind", paths.wind_forecast_dir),
        ):
            auth = loaded_auth[name]
            self.assertEqual(data_path(directory / "auth_state.json"), auth.AUTH_STATE_PATH)
            self.assertEqual(data_path(directory / "auth_state.json.lock"), auth.AUTH_LOCK_PATH)
            self.assertEqual(data_path(directory / ".browser-profile"), auth.PROFILE_DIR)
            self.assertEqual(data_path(directory / "config.json"), auth.CONFIG_PATH)

        self.assertEqual(online.AUTH_STATE_PATH, data_path(paths.online_energy_dir / "auth_state.json"))
        self.assertEqual(online.PROFILE_DIR, data_path(paths.online_energy_dir / ".browser-profile"))
        self.assertEqual(market.online_export_module.AUTH_STATE_PATH, data_path(paths.market_table_update_dir / "auth_state.json"))
        self.assertEqual(market.online_export_module.PROFILE_DIR, data_path(paths.market_table_update_dir / ".browser-profile"))
        self.assertEqual(group.AUTH_STATE_PATH, data_path(paths.group_upload_dir / "auth_state.json"))
        self.assertEqual(group.PROFILE_DIR, data_path(paths.group_upload_dir / ".browser-profile"))
        self.assertEqual(spot.AUTH_STATE_PATH, data_path(paths.spot_price_update_dir / "auth_state.json"))
        self.assertEqual(spot.PROFILE_DIR, data_path(paths.spot_price_update_dir / ".browser-profile"))
        self.assertEqual(
            transaction.AUTH_STATE_PATH,
            data_path(paths.transaction_info_update_dir / "auth_state.json"),
        )
        self.assertEqual(
            transaction.PROFILE_DIR,
            data_path(paths.transaction_info_update_dir / ".browser-profile"),
        )

        auth_paths = {
            online.AUTH_STATE_PATH,
            market.online_export_module.AUTH_STATE_PATH,
            group.AUTH_STATE_PATH,
            spot.AUTH_STATE_PATH,
            transaction.AUTH_STATE_PATH,
            wind.AUTH_STATE_PATH,
        }
        profile_paths = {
            online.PROFILE_DIR,
            market.online_export_module.PROFILE_DIR,
            group.PROFILE_DIR,
            spot.PROFILE_DIR,
            transaction.PROFILE_DIR,
            wind.PROFILE_DIR,
        }
        self.assertEqual(len(auth_paths), 6)
        self.assertEqual(len(profile_paths), 6)

    def test_online_energy_process_is_alive_detects_current_process(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("online_energy_process_alive_test", paths.online_energy)

        self.assertTrue(module.process_is_alive(os.getpid()))
        self.assertFalse(module.process_is_alive(-1))

    def test_online_energy_lock_cleanup_does_not_mask_body_exception(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("online_energy_lock_cleanup_test", paths.online_energy)

        with tempfile.TemporaryDirectory() as directory:
            lock_path = Path(directory) / "auth_state.json.lock"
            with mock.patch.object(module, "AUTH_LOCK_PATH", lock_path):
                with mock.patch.object(Path, "unlink", side_effect=PermissionError("locked")):
                    with self.assertRaisesRegex(ValueError, "original failure"):
                        with module.AuthStateLock():
                            raise ValueError("original failure")

    def test_report_summary_reads_shifted_renewable_rows(self) -> None:
        from openpyxl import Workbook
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        sys.path.insert(0, str(paths.report_scripts_dir))
        self.addCleanup(lambda: sys.path.remove(str(paths.report_scripts_dir)))
        module = load_module(
            "report_generate_red_marked_shifted_renewables_test",
            paths.report_scripts_dir / "generate_red_marked_report.py",
        )

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "出清情况（分公司2026.7.2).xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "快报  (日前)"
            for cell, value in {
                "F7": 10116.2898,
                "G7": 422.0653,
                "F8": 1317.6129,
                "G8": 379.0106,
                "F13": 8042.6026,
                "G13": 419.1546,
                "F19": 756.0743,
                "G19": 528.0599,
            }.items():
                sheet[cell] = value
            sheet["A28"] = "新能源"
            sheet["G29"] = "日前出清"
            sheet["G30"] = "电量"
            sheet["H30"] = "电价"
            for row, name, energy, price in [
                (31, "分公司", 247.5663, 332.4490),
                (32, "汕头海风", 189.3145, 285.4164),
                (33, "鮀莲光伏", 24.4543, 512.7376),
                (34, "归湖光伏", 33.7975, 465.4510),
            ]:
                sheet.cell(row, 1).value = name
                sheet.cell(row, 7).value = energy
                sheet.cell(row, 8).value = price
            workbook.save(workbook_path)

            clearing = module.load_day_ahead_clearing_summary("2026-07-02", workbook_path)
            prices = module.load_online_price_summary("2026-07-02", workbook_path)

        self.assertEqual(clearing["companies"]["海上风电"]["energy"], 189.3145)
        self.assertEqual(clearing["companies"]["鮀莲"]["price"], 512.7376)
        self.assertEqual(clearing["companies"]["归湖"]["energy"], 33.7975)
        self.assertEqual(prices["companies"]["海上风电"], 285.4164)

    def test_report_energy_text_includes_qingyuan_descriptions(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        sys.path.insert(0, str(paths.report_scripts_dir))
        self.addCleanup(lambda: sys.path.remove(str(paths.report_scripts_dir)))
        module = load_module(
            "report_generate_red_marked_qingyuan_text_test",
            paths.report_scripts_dir / "generate_red_marked_report.py",
        )

        online_text = module.ensure_qingyuan_online_energy_text(
            "7月1日分公司上网电量9003万千瓦时，其中，汕头1526万、海门6120万、"
            "东莞1232万、海上风电58万、光伏67万千瓦时；分公司日前均价404厘/千瓦时，"
            "其中汕头387厘、海门406厘、东莞411厘、海上风电347厘、鮀莲361厘、归湖358厘/千瓦时。"
        )
        self.assertIn("光伏67万千瓦时、清远0万千瓦时；", online_text)

        clearing_text = module.build_day_ahead_clearing_text(
            "2026-07-02",
            {
                "totalEnergy": 10116.2898,
                "averagePrice": 422.0653,
                "companies": {
                    "汕头": {"energy": 1317.6129, "price": 379.0106},
                    "海门": {"energy": 8042.6026, "price": 419.1546},
                    "东莞": {"energy": 756.0743, "price": 528.0599},
                    "海上风电": {"energy": 189.3145, "price": 285.4164},
                    "鮀莲": {"energy": 24.4543, "price": 512.7376},
                    "归湖": {"energy": 33.7975, "price": 465.4510},
                },
            },
        )
        self.assertIn("归湖34万千瓦时、清远（计划）0万千瓦时；", clearing_text)

    def test_utf8_environment_sets_python_encoding(self) -> None:
        from toolbox.runtime import utf8_environment

        environment = utf8_environment()

        self.assertEqual(environment["PYTHONUTF8"], "1")
        self.assertEqual(environment["PYTHONIOENCODING"], "utf-8")

    def test_python_executable_uses_current_interpreter_on_macos(self) -> None:
        if sys.platform != "darwin":
            self.skipTest("macOS-specific runtime expectation")

        from toolbox.runtime import python_executable, pythonw_executable

        self.assertEqual(Path(python_executable()).resolve(), Path(sys.executable).resolve())
        self.assertEqual(Path(pythonw_executable()).resolve(), Path(sys.executable).resolve())

    def test_node_executable_finds_local_node_on_macos(self) -> None:
        from toolbox.runtime import node_executable

        with tempfile.TemporaryDirectory() as folder:
            local_node = Path(folder) / ".local" / "bin" / "node"
            local_node.parent.mkdir(parents=True)
            local_node.touch(mode=0o755)
            with (
                mock.patch("toolbox.runtime.sys.platform", "darwin"),
                mock.patch("toolbox.runtime.shutil.which", return_value=None),
                mock.patch.dict(os.environ, {"HOME": folder}),
            ):
                self.assertEqual(node_executable(), str(local_node))

    def test_task_registry_tracks_and_terminates_owned_process(self) -> None:
        from toolbox.runtime import TaskRegistry

        registry = TaskRegistry()
        process = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(30)"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            text=True,
        )
        self.addCleanup(lambda: process.poll() is None and process.kill())

        registry.register_process(process)
        self.assertTrue(registry.has_running_tasks())

        registry.terminate_all()
        deadline = time.time() + 5
        while process.poll() is None and time.time() < deadline:
            time.sleep(0.05)

        self.assertIsNotNone(process.poll())
        self.assertFalse(registry.has_running_tasks())


if __name__ == "__main__":
    unittest.main()
