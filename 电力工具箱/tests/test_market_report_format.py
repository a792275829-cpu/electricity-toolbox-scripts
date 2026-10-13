from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "南方分公司电力市场日报"))
from gd_power_report import pipeline
from gd_power_report.collector import market
from gd_power_report.company_clearing_xlsx import _market_cards, _strategy_lines
from gd_power_report.config import AppConfig
from gd_power_report.legacy_concise_report import _split_strategy_profit


def thermal_payload(coal, gas, *, realtime=False):
    period = "REAL_TIME" if realtime else "DAY_AHEAD"
    return {
        "data": [
            {"dataType": "2" if realtime else "1",
             "dataItemGroup": f"{period}_CLEARANCE_{fuel}", "unitCount": values}
            for fuel, values in (("COAL", coal), ("GAS", gas))
        ]
    }


class MarketReportFormatTests(unittest.TestCase):
    def test_regional_strategy_keeps_paragraphs_and_each_plants_profit(self):
        for title in ("粤东策略", "粤东区域运行策略", "粤东区域运行策略研究"):
            with self.subTest(title=title):
                text = (
                    f"2.{title}：策略总述。\n负荷推演：需求变化。\n"
                    "东莞厂也受需求变化影响。\n策略调整及效果：调整结果。\n"
                    "汕头厂各机组日前收益为：15万元、26万元与317万元。\n"
                    "海门厂各机组日前收益分别为：70万元、60万元、416万元与26万元。\n"
                    "3.东莞厂独立策略。"
                )
                rows = _strategy_lines(text)
                self.assertEqual([row["plant"] for row in rows], [title, "东莞厂"])
                paragraphs = _split_strategy_profit(rows[0])["paragraphs"]
                self.assertEqual(len(paragraphs), 6)
                self.assertIn("东莞厂也受需求变化影响。", paragraphs)
                self.assertIn("汕头厂各机组日前偏差收益分别为：15万元", paragraphs[-2])
                self.assertIn("海门厂各机组日前偏差收益分别为：70万元", paragraphs[-1])

    def test_coal_uses_first_mode_and_gas_uses_1900_for_both_dates(self):
        current = thermal_payload([115] * 12 + [113] * 12, [90] + [30] * 18 + [55] + [40] * 3 + [80])
        previous = thermal_payload([114] * 24, [100] + [31] * 18 + [28] + [41] * 3 + [70])
        card = market.parse_day_ahead_thermal_unit_count_card(current, previous)
        counts = card["thermal_breakdown"]
        self.assertEqual((counts["coal"], counts["gas"], counts["coal_delta"], counts["gas_delta"]), (115, 55, 1, 27))
        self.assertEqual((card["value"], card["delta"]), (170, 28))

    def test_missing_1900_does_not_use_another_hour(self):
        gas = [50] * 24
        gas[19] = None
        card = market.parse_realtime_thermal_unit_count_card(
            thermal_payload([100] * 24, gas, realtime=True), thermal_payload([100] * 24, [40] * 24)
        )
        self.assertTrue(card["pending"])
        self.assertIsNone(card["thermal_breakdown"]["gas"])

    def test_numeric_excel_text_preserves_market_units_and_delta(self):
        current, previous = Workbook(), Workbook()
        try:
            basic = current.create_sheet("基础信息")
            daily = current.create_sheet("日报")
            old_basic = previous.create_sheet("基础信息")
            for cell in ("G5", "B5", "C5"):
                basic[cell] = " 12340 "
            for cell in ("M5", "N5"):
                daily[cell] = "12340"
            for cell in ("B5", "C5"):
                old_basic[cell] = "12000"
            cards = {card["label"]: card for card in _market_cards(current, previous)}
            self.assertEqual(cards["检修容量"]["value"], 1234)
            self.assertEqual(cards["统调负荷峰值"]["delta"], 34)
        finally:
            current.close()
            previous.close()

    def test_incomplete_thermal_comparison_stops_before_export(self):
        report = {
            "report_date": "2026-10-07", "report_date_label": "10月7日",
            "day_ahead": {"cards": []}, "operations": {"rows": []},
            "market": {"day_ahead": [], "realtime": []},
        }
        previous = {
            "report_date": "2026-10-06", "report_date_label": "10月6日",
            "cards": [], "operations": {"rows": []}, "source_path": "previous.xlsx",
        }
        raw = {
            "market_day_ahead_cards": [{"label": "市场开机", "thermal_breakdown": {"coal_delta": 1, "gas_delta": None}}],
            "market_realtime_cards": [{"label": "在线机组数", "thermal_breakdown": {"coal_delta": 0, "gas_delta": 1}}],
        }
        detail = {"report_date": "2026-10-05", "realtime": {"cards": [
            {"label": "实时出清电量", "value": 10}
        ]}}
        with tempfile.TemporaryDirectory() as directory, \
             mock.patch.object(pipeline, "parse_daily_report_workbook", side_effect=[copy.deepcopy(report), detail]), \
             mock.patch.object(pipeline, "parse_day_ahead_comparison_workbook", return_value=previous), \
             mock.patch.object(pipeline, "parse_realtime_unit_workbook", return_value={"source_path": "d2.xlsx"}), \
             mock.patch.object(pipeline, "summarize_realtime_units", return_value=[{"realtime": 10}]), \
             mock.patch.object(pipeline, "read_full_day_clearing_table", return_value={}), \
             mock.patch.object(pipeline, "collect_realtime_market_comparison", return_value=raw), \
             mock.patch.object(pipeline, "_render_concise_outputs") as export:
            config = AppConfig.from_workspace(directory)
            with self.assertRaises(RuntimeError) as caught:
                pipeline.run_pipeline(config, date(2026, 10, 7), "day.xlsx",
                                      comparison_clearing_xlsx_path="previous.xlsx", realtime_detail_xlsx_path="d2.xlsx")
            self.assertIn("煤机、气机台数", str(caught.exception.__cause__))
            export.assert_not_called()
            self.assertFalse(config.outputs_dir.exists())


if __name__ == "__main__":
    unittest.main()
