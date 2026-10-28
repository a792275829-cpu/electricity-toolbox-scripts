from __future__ import annotations

import argparse
import json
import math
import re
import sys
import uuid
import warnings
import zipfile
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel, to_excel
from playwright.sync_api import sync_playwright
from xml.etree import ElementTree


SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SCRIPT_DIR.parent
if str(WORKSPACE_DIR) not in sys.path:
    sys.path.insert(0, str(WORKSPACE_DIR))
from toolbox_shared.auth_loader import load_login_module, standalone_login
from toolbox_shared.dates import parse_date_input
from toolbox_shared.credentials import apply_credentials
from toolbox_shared.excel_read import read_cell_snapshots
from toolbox_shared.workbook_commit import (
    FileFingerprint, assert_workbook_unchanged, commit_workbooks, workbook_fingerprint,
)

BASE_URL = "https://xhxt.chng.com.cn"
APP_URL = f"{BASE_URL}/gdfire/PrivateDataManageCE/TradeResult"
PRICE_API_URL = f"{BASE_URL}/gdfire/api/data/spot/trade/unit"
from toolbox_shared.storage import data_path

PROFILE_DIR = data_path(SCRIPT_DIR / ".browser-profile")
AUTH_STATE_PATH = data_path(SCRIPT_DIR / "auth_state.json")
AUTH_LOCK_PATH = Path(str(AUTH_STATE_PATH) + ".lock")
CONFIG_PATH = data_path(SCRIPT_DIR / "config.json")
FALLBACK_CONFIG_PATH = data_path(WORKSPACE_DIR / "上网电量抓取" / "config.json")

DEFAULT_WORKBOOK_DIR = Path.home() / "Desktop" / "日常工作"
DEFAULT_GUIHU_WORKBOOK = DEFAULT_WORKBOOK_DIR / "2026年现货价格分析（2026.8.14更新）-归湖.xlsx"
DEFAULT_TUOLIAN_WORKBOOK = DEFAULT_WORKBOOK_DIR / "2026年现货价格分析（2026.8.14更新）-鮀莲.xlsx"
DEFAULT_HAIFENG_WORKBOOK = DEFAULT_WORKBOOK_DIR / "2026年现货价格分析-海风.xlsx"
STATION_KEYS = ("归湖", "鮀莲", "海风")
DAY_AHEAD_SHEET = "日前价格"
REAL_TIME_SHEET = "实时价格"
EXPECTED_POINT_COUNT = 24


@dataclass(frozen=True, slots=True)
class Config:
    username: str = ""
    password: str = ""
    headless: bool = True


@dataclass(frozen=True, slots=True)
class TargetSpec:
    key: str
    workbook: Path
    tenant_keywords: tuple[str, ...]
    unit_keywords: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class TargetPrices:
    target: TargetSpec
    tenant_name: str
    unit_id: str
    unit_name: str
    day_ahead: dict[date, tuple[float, ...]]
    real_time: dict[date, tuple[float, ...]]


@dataclass(frozen=True, slots=True)
class CellExpectation:
    sheet_name: str
    row: int
    values: tuple[float, ...]
    number_formats: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class DateExpectation:
    sheet_name: str
    row: int
    value: date
    number_format: str


@dataclass(frozen=True, slots=True)
class PreparedWorkbook:
    original: Path
    temporary: Path
    source_fingerprint: FileFingerprint
    expectations: tuple[CellExpectation, ...]
    updated_dates: tuple[date, ...]
    skipped_dates: tuple[date, ...]
    date_expectations: tuple[DateExpectation, ...] = ()


def log(message: str) -> None:
    print(message, flush=True)


def _load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload if isinstance(payload, dict) else {}


def load_config() -> Config:
    raw = apply_credentials({**_load_json(FALLBACK_CONFIG_PATH), **_load_json(CONFIG_PATH)})
    return Config(
        username=str(raw.get("username") or ""),
        password=str(raw.get("password") or ""),
        headless=bool(raw.get("headless", True)),
    )


def load_auth_module():
    return load_login_module(
        WORKSPACE_DIR,
        module_name="spot_price_auth",
        state_dir=SCRIPT_DIR,
        config_path=CONFIG_PATH,
        app_url=APP_URL,
    )


def validate_date(value: str) -> date:
    return parse_date_input(value)


def date_range(start_date: date, end_date: date) -> tuple[date, ...]:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期。")
    return tuple(
        start_date + timedelta(days=offset)
        for offset in range((end_date - start_date).days + 1)
    )


def default_targets(
    guihu_workbook: Path = DEFAULT_GUIHU_WORKBOOK,
    tuolian_workbook: Path = DEFAULT_TUOLIAN_WORKBOOK,
    haifeng_workbook: Path = DEFAULT_HAIFENG_WORKBOOK,
    *,
    selected_stations: Iterable[str] | None = None,
) -> tuple[TargetSpec, ...]:
    targets = (
        TargetSpec(
            key="归湖",
            workbook=Path(guihu_workbook).expanduser(),
            tenant_keywords=("潮州潮安",),
            unit_keywords=("华归", "归湖"),
        ),
        TargetSpec(
            key="鮀莲",
            workbook=Path(tuolian_workbook).expanduser(),
            tenant_keywords=("汕头金平",),
            unit_keywords=("鮀莲",),
        ),
        TargetSpec(
            key="海风",
            workbook=Path(haifeng_workbook).expanduser(),
            tenant_keywords=("华能广东汕头海上风电有限责任公司",),
            unit_keywords=("广澳海风",),
        ),
    )
    selected = set(STATION_KEYS if selected_stations is None else selected_stations)
    if not selected:
        raise ValueError("请至少选择一个更新场站。")
    unknown = selected - set(STATION_KEYS)
    if unknown:
        raise ValueError(f"未知场站：{'、'.join(sorted(unknown))}")
    return tuple(target for target in targets if target.key in selected)


def normalize_name(value: Any) -> str:
    text = str(value or "").strip().lower()
    for token in (" ", "\t", "\r", "\n", "（", "）", "(", ")", "-", "_", "·"):
        text = text.replace(token, "")
    return text


def resolve_tenant(companies: Iterable[dict[str, Any]], target: TargetSpec) -> dict[str, Any]:
    rows = [row for row in companies if row.get("tenantId") and row.get("name")]
    matches = [
        row
        for row in rows
        if all(normalize_name(keyword) in normalize_name(row.get("name")) for keyword in target.tenant_keywords)
    ]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        names = "、".join(str(row.get("name")) for row in matches)
        raise RuntimeError(f"{target.key} 匹配到多个公司：{names}")
    available = "、".join(str(row.get("name")) for row in rows) or "无"
    raise RuntimeError(
        f"未找到 {target.key} 对应公司（关键词：{'/'.join(target.tenant_keywords)}）。"
        f"当前账号可见公司：{available}"
    )


def resolve_unit(units: Iterable[dict[str, Any]], target: TargetSpec, auth) -> dict[str, Any]:
    rows = [row for row in units if row.get("id") and auth.unit_display_name(row)]
    preferred = [
        row
        for row in rows
        if any(
            normalize_name(keyword) in normalize_name(auth.unit_display_name(row))
            for keyword in target.unit_keywords
        )
    ]
    if len(preferred) == 1:
        return preferred[0]
    if len(preferred) > 1:
        names = "、".join(auth.unit_display_name(row) for row in preferred)
        raise RuntimeError(f"{target.key} 匹配到多个场站：{names}")

    renewable = [row for row in rows if auth.is_renewable_unit(row)]
    if len(renewable) == 1:
        log(
            f"{target.key} 未按名称匹配到场站，当前公司只有一个新能源场站，"
            f"将使用：{auth.unit_display_name(renewable[0])}"
        )
        return renewable[0]
    available = "、".join(auth.unit_display_name(row) for row in rows) or "无"
    raise RuntimeError(
        f"未找到 {target.key} 对应场站（关键词：{'/'.join(target.unit_keywords)}）。"
        f"当前公司可见场站：{available}"
    )


def _number(value: Any, *, label: str) -> float:
    if value in (None, "", "-") or isinstance(value, bool):
        raise ValueError(f"{label} 为空。")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} 不是有效数值：{value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label} 不是有限数值：{value!r}")
    return number


def parse_price_payload(
    payload: dict[str, Any],
    *,
    unit_id: str,
    expected_dates: tuple[date, ...],
    label: str,
    skip_incomplete: bool = False,
) -> dict[date, tuple[float, ...]]:
    data = payload.get("data")
    if isinstance(data, dict):
        rows = data.get("unitDateList") or []
    elif isinstance(data, list):
        rows = data
    else:
        rows = []
    if not isinstance(rows, list):
        raise RuntimeError(f"{label} 返回格式异常：unitDateList 不是数组。")

    expected = set(expected_dates)
    result: dict[date, tuple[float, ...]] = {}
    for row in rows:
        if not isinstance(row, dict) or str(row.get("unitId") or "") != str(unit_id):
            continue
        raw_date = str(row.get("date") or "")[:10]
        try:
            row_date = datetime.strptime(raw_date, "%Y-%m-%d").date()
        except ValueError:
            continue
        if row_date not in expected:
            continue
        raw_prices = row.get("acPrice")
        if not isinstance(raw_prices, list) or len(raw_prices) != EXPECTED_POINT_COUNT:
            count = len(raw_prices) if isinstance(raw_prices, list) else 0
            message = (
                f"{label} {row_date.isoformat()} 应有 24 个电价点，实际为 {count} 个。"
            )
            if skip_incomplete:
                log(f"跳过：{message}")
                continue
            raise RuntimeError(message)
        try:
            result[row_date] = tuple(
                _number(value, label=f"{label} {row_date.isoformat()} {hour:02d}:00")
                for hour, value in enumerate(raw_prices)
            )
        except ValueError as exc:
            if skip_incomplete:
                log(f"跳过：{exc}")
                continue
            raise

    missing = [item.isoformat() for item in expected_dates if item not in result]
    if missing:
        message = f"{label} 缺少日期：{'、'.join(missing)}"
        if skip_incomplete:
            log(f"跳过：{message}")
        else:
            raise RuntimeError(message)
    return result


def fetch_price_series(
    context,
    auth,
    *,
    unit_id: str,
    expected_dates: tuple[date, ...],
    data_type: int,
    label: str,
) -> dict[date, tuple[float, ...]]:
    response = context.request.get(
        PRICE_API_URL,
        params={
            "dataType": data_type,
            "startDate": expected_dates[0].isoformat(),
            "endDate": expected_dates[-1].isoformat(),
            "unitIds": unit_id,
        },
        timeout=60_000,
    )
    payload = auth.parse_json_response(response, f"读取{label}")
    return parse_price_payload(
        payload,
        unit_id=unit_id,
        expected_dates=expected_dates,
        label=label,
        skip_incomplete=True,
    )


def collect_target_prices(
    start_date: date,
    end_date: date,
    targets: tuple[TargetSpec, ...],
    *,
    config: Config | None = None,
) -> tuple[TargetPrices, ...]:
    expected_dates = date_range(start_date, end_date)
    config = config or load_config()
    auth = load_auth_module()
    collected: list[TargetPrices] = []
    with auth.AuthStateLock():
        with sync_playwright() as playwright:
            context = auth.launch_context(playwright, headless=config.headless)
            try:
                auth.ensure_login(context, config)
                companies = auth.get_companies(context)
                for target in targets:
                    tenant = resolve_tenant(companies, target)
                    tenant_name = str(tenant["name"])
                    log(f"{target.key}：切换公司 {tenant_name}")
                    auth.switch_company(context, str(tenant["tenantId"]))
                    units = auth.get_units(context)
                    unit = resolve_unit(units, target, auth)
                    unit_id = str(unit["id"])
                    unit_name = auth.unit_display_name(unit)
                    log(f"{target.key}：使用场站 {unit_name}")
                    day_ahead = fetch_price_series(
                        context,
                        auth,
                        unit_id=unit_id,
                        expected_dates=expected_dates,
                        data_type=1,
                        label=f"{target.key} 日前出清电价",
                    )
                    real_time = fetch_price_series(
                        context,
                        auth,
                        unit_id=unit_id,
                        expected_dates=expected_dates,
                        data_type=2,
                        label=f"{target.key} 实时出清电价",
                    )
                    collected.append(
                        TargetPrices(
                            target=target,
                            tenant_name=tenant_name,
                            unit_id=unit_id,
                            unit_name=unit_name,
                            day_ahead=day_ahead,
                            real_time=real_time,
                        )
                    )
                    day_missing = len(expected_dates) - len(day_ahead)
                    real_missing = len(expected_dates) - len(real_time)
                    log(
                        f"{target.key}：日前完整 {len(day_ahead)} 天、跳过 {day_missing} 天；"
                        f"实时完整 {len(real_time)} 天、跳过 {real_missing} 天。"
                    )
                auth.save_auth_state(context)
            finally:
                context.close()
    return tuple(collected)


def excel_cell_date(value: Any, *, epoch: datetime) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch=epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    return None


def _iter_dated_rows(sheet, *, epoch: datetime, max_column: int = 1):
    result: dict[date, int] = {}
    for row, cells in enumerate(sheet.iter_rows(
        min_row=2, max_row=sheet.max_row, min_col=1, max_col=max_column,
    ), start=2):
        parsed = excel_cell_date(cells[0].value, epoch=epoch)
        if parsed is not None:
            if parsed in result:
                raise RuntimeError(
                    f"{sheet.title} 日期重复：{parsed.isoformat()}（第 {result[parsed]}、{row} 行）"
                )
            result[parsed] = row
        yield row, cells, parsed


def date_row_index(sheet, *, epoch: datetime) -> dict[date, int]:
    return {parsed: row for row, _, parsed in _iter_dated_rows(sheet, epoch=epoch)
            if parsed is not None}


def values_match(existing: Iterable[Any], expected: tuple[float, ...]) -> bool:
    rows = list(existing)
    if len(rows) != len(expected):
        return False
    for old, new in zip(rows, expected):
        if isinstance(old, bool) or not isinstance(old, (int, float)):
            return False
        if not math.isclose(float(old), new, rel_tol=1e-9, abs_tol=1e-9):
            return False
    return True


def effective_number_format(sheet, row: int, column: int) -> str:
    return effective_number_formats(sheet, (row,), column)[row]


def effective_number_formats(sheet, rows: Iterable[int], column: int) -> dict[int, str]:
    """Resolve inherited formats in one forward scan, including styled blanks."""
    requested = set(rows)
    result = {row: "General" for row in requested}
    if not requested or max(requested) < 2:
        return result
    inherited = "General"
    last_row = 1
    for row, (cell,) in enumerate(sheet.iter_rows(
        min_row=2, max_row=max(requested), min_col=column, max_col=column,
    ), start=2):
        last_row = row
        if getattr(cell, "has_style", False):
            inherited = str(getattr(cell, "number_format", None) or "General")
        if row in requested:
            result[row] = inherited
    for row in requested:
        if row > last_row:
            result[row] = inherited
    return result


def plan_date_extensions(sheet, prices: dict[date, tuple[float, ...]], *, epoch) -> dict[int, date]:
    if not prices:
        return {}
    dated_rows = {}
    content_rows = set()
    for row, cells, parsed in _iter_dated_rows(
        sheet, epoch=epoch, max_column=max(sheet.max_column or 1, 25)
    ):
        if parsed is not None:
            dated_rows[parsed] = row
        if any(cell.value not in (None, "") for cell in cells):
            content_rows.add(row)
    if not dated_rows:
        raise RuntimeError(f"{sheet.title} 没有可接续的日期行。")
    last_date = max(dated_rows)
    if max(prices) <= last_date:
        return {}
    last_row = dated_rows[last_date]
    if last_row != max(dated_rows.values()) or any(row > last_row for row in content_rows):
        raise RuntimeError(f"{sheet.title} 日期表尾后存在其他内容，不能自动接续日期行。")
    missing_history = sorted(day for day in prices if day <= last_date and day not in dated_rows)
    if missing_history:
        raise RuntimeError(f"{sheet.title} 中间缺少日期 {missing_history[0].isoformat()}，不会猜测行位置。")
    count = (max(prices) - last_date).days
    if last_row + count > 1_048_576:
        raise RuntimeError(f"{sheet.title} 接续日期将超出 Excel 行数限制。")
    return {last_row + offset: last_date + timedelta(days=offset) for offset in range(1, count + 1)}


def plan_price_sheet(
    workbook,
    sheet_name: str,
    prices: dict[date, tuple[float, ...]],
    *,
    overwrite: bool,
    date_extensions: dict[int, date] | None = None,
) -> tuple[
    list[CellExpectation],
    list[date],
    list[date],
    dict[int, tuple[float, ...]],
]:
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Excel 中未找到工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    if not prices:
        date_row_index(sheet, epoch=workbook.epoch)
        return [], [], [], {}
    snapshots = {}
    inherited_formats = ["General"] * 24
    date_extensions = date_extensions or {}
    for row, cells, parsed in _iter_dated_rows(sheet, epoch=workbook.epoch, max_column=25):
        for index, cell in enumerate(cells[1:]):
            if getattr(cell, "has_style", False):
                inherited_formats[index] = str(cell.number_format or "General")
        parsed = date_extensions.get(row, parsed)
        if parsed in prices:
            snapshots[parsed] = (row, [cell.value for cell in cells[1:]], tuple(inherited_formats))
    for row, day in date_extensions.items():
        if day in prices and day not in snapshots:
            snapshots[day] = (row, [None] * 24, tuple(inherited_formats))
    expectations: list[CellExpectation] = []
    updated: list[date] = []
    skipped: list[date] = []
    row_values: dict[int, tuple[float, ...]] = {}
    for target_date, expected in sorted(prices.items()):
        snapshot = snapshots.get(target_date)
        if snapshot is None:
            raise RuntimeError(
                f"{sheet_name} 未找到日期 {target_date.isoformat()}，不会猜测或新增行。"
            )
        row, existing, formats = snapshot
        if values_match(existing, expected):
            skipped.append(target_date)
            continue
        if not overwrite and any(value not in (None, "") for value in existing):
            skipped.append(target_date)
            log(f"{sheet_name} {target_date.isoformat()} 已有数据，按设置跳过。")
            continue
        expectations.append(
            CellExpectation(
                sheet_name=sheet_name,
                row=row,
                values=expected,
                number_formats=formats,
            )
        )
        updated.append(target_date)
        row_values[row] = expected
    return expectations, updated, skipped, row_values


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def xlsx_sheet_paths(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    result: dict[str, str] = {}
    sheets = workbook_root.find(f"{{{MAIN_NS}}}sheets")
    if sheets is None:
        return result
    for sheet in sheets.findall(f"{{{MAIN_NS}}}sheet"):
        rel_id = sheet.attrib.get(f"{{{REL_NS}}}id")
        target = relationships.get(str(rel_id), "")
        if not target:
            continue
        target = target.lstrip("/")
        result[str(sheet.attrib.get("name") or "")] = (
            target if target.startswith("xl/") else f"xl/{target}"
        )
    return result


def _column_name(number: int) -> str:
    if not 1 <= number <= 26:
        raise ValueError(f"不支持的列号：{number}")
    return chr(ord("A") + number - 1)


def _cell_pattern(cell_ref: str) -> re.Pattern[str]:
    escaped = re.escape(cell_ref)
    return re.compile(
        rf'<c\b(?=[^>]*\br="{escaped}")[^>]*?(?:/>|>.*?</c>)',
        flags=re.DOTALL,
    )


def _style_from_cell_xml(cell_xml: str) -> str | None:
    match = re.search(r'\bs="([^"]+)"', cell_xml)
    return match.group(1) if match else None


def _nearest_prior_style(xml_text: str, column: str, row_number: int) -> str | None:
    pattern = re.compile(
        rf'<c\b(?=[^>]*\br="{re.escape(column)}(\d+)")[^>]*?(?:/>|>.*?</c>)',
        flags=re.DOTALL,
    )
    best_row = -1
    best_style: str | None = None
    for match in pattern.finditer(xml_text):
        candidate_row = int(match.group(1))
        if candidate_row >= row_number or candidate_row <= best_row:
            continue
        style = _style_from_cell_xml(match.group(0))
        if style is not None:
            best_row = candidate_row
            best_style = style
    return best_style


def _numeric_cell_xml(cell_ref: str, value: float, style: str | None) -> str:
    style_attribute = f' s="{style}"' if style is not None else ""
    return f'<c r="{cell_ref}"{style_attribute}><v>{format(value, ".15g")}</v></c>'


def extend_sheet_dates(xml_bytes: bytes, dates: dict[int, date], *, epoch) -> bytes:
    if not dates:
        return xml_bytes
    xml_text = xml_bytes.decode("utf-8")
    first_row = min(dates)
    for merged in re.finditer(r'<mergeCell\b[^>]*\bref="([^"]+)"', xml_text):
        min_col, min_row, _, max_row = range_boundaries(merged.group(1))
        if min_col <= 25 and any(min_row <= row <= max_row for row in dates):
            raise RuntimeError(f"接续日期区域包含合并单元格 {merged.group(1)}，停止写入。")
    template = re.search(
        rf'<row\b(?=[^>]*\br="{first_row - 1}")[^>]*?(?:/>|>.*?</row>)', xml_text, re.DOTALL
    )
    if template is None:
        raise RuntimeError(f"工作表 XML 未找到接续模板第 {first_row - 1} 行。")
    # Reuse the tail row's height and other row formatting, without copying prices.
    template_opening = template.group(0).split(">", 1)[0].rstrip("/") + ">"
    for row, day in sorted(dates.items()):
        pattern = re.compile(rf'<row\b(?=[^>]*\br="{row}")[^>]*?(?:/>|>.*?</row>)', re.DOTALL)
        match = pattern.search(xml_text)
        if match:
            row_xml = match.group(0)
            if row_xml.endswith("/>"):
                row_xml = row_xml[:-2] + "></row>"
        else:
            opening = re.sub(r'\br="\d+"', f'r="{row}"', template_opening, count=1)
            row_xml = opening + "</row>"
        cell_ref = f"A{row}"
        cell = _cell_pattern(cell_ref).search(row_xml)
        style = (_style_from_cell_xml(cell.group(0)) if cell else None) or _nearest_prior_style(xml_text, "A", row)
        replacement = _numeric_cell_xml(cell_ref, to_excel(day, epoch), style)
        if cell:
            row_xml = row_xml[:cell.start()] + replacement + row_xml[cell.end():]
        else:
            end_opening = row_xml.index(">") + 1
            row_xml = row_xml[:end_opening] + replacement + row_xml[end_opening:]
        for column in range(2, 26):
            reference = f"{_column_name(column)}{row}"
            if not _cell_pattern(reference).search(row_xml):
                style = _nearest_prior_style(xml_text, _column_name(column), row)
                if style is not None:
                    row_xml = row_xml.replace("</row>", f'<c r="{reference}" s="{style}"/></row>', 1)
        if match:
            xml_text = xml_text[:match.start()] + row_xml + xml_text[match.end():]
        else:
            next_row = next((item for item in re.finditer(r'<row\b[^>]*\br="(\d+)"', xml_text)
                             if int(item.group(1)) > row), None)
            position = next_row.start() if next_row else xml_text.find("</sheetData>")
            if position < 0:
                raise RuntimeError("工作表 XML 缺少 sheetData 结束标记。")
            xml_text = xml_text[:position] + row_xml + xml_text[position:]
    dimension = re.search(r'<dimension\s+ref="([^"]+)"', xml_text)
    if dimension:
        min_col, min_row, max_col, max_row = range_boundaries(dimension.group(1))
        reference = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max(max_col, 25))}{max(max_row, max(dates))}"
        xml_text = xml_text[:dimension.start(1)] + reference + xml_text[dimension.end(1):]
    return xml_text.encode("utf-8")


def patch_sheet_xml(
    xml_bytes: bytes,
    row_values: dict[int, tuple[float, ...]],
) -> bytes:
    xml_text = xml_bytes.decode("utf-8")
    for row_number, values in sorted(row_values.items()):
        row_pattern = re.compile(
            rf'(<row\b(?=[^>]*\br="{row_number}")[^>]*>)(.*?)(</row>)',
            flags=re.DOTALL,
        )
        row_match = row_pattern.search(xml_text)
        if row_match is None:
            raise RuntimeError(f"工作表 XML 未找到第 {row_number} 行。")
        opening, content, closing = row_match.groups()
        if 'spans="' in opening:
            opening = re.sub(r'spans="[^"]*"', 'spans="1:25"', opening, count=1)
        else:
            opening = opening[:-1] + ' spans="1:25">'

        for column_number, value in enumerate(values, start=2):
            column = _column_name(column_number)
            cell_ref = f"{column}{row_number}"
            pattern = _cell_pattern(cell_ref)
            existing = pattern.search(content)
            if existing is not None:
                style = _style_from_cell_xml(existing.group(0))
                replacement = _numeric_cell_xml(cell_ref, value, style)
                content = content[: existing.start()] + replacement + content[existing.end() :]
            else:
                style = _nearest_prior_style(xml_text, column, row_number)
                content += _numeric_cell_xml(cell_ref, value, style)

        replacement_row = opening + content + closing
        xml_text = xml_text[: row_match.start()] + replacement_row + xml_text[row_match.end() :]
    return xml_text.encode("utf-8")


def verify_package_preservation(
    original: Path,
    patched: Path,
    changed_parts: set[str],
) -> None:
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(patched) as result:
        if source.namelist() != result.namelist():
            raise RuntimeError(f"XLSX 包结构发生变化：{original.name}")
        for name in source.namelist():
            if name in changed_parts:
                continue
            if source.read(name) != result.read(name):
                raise RuntimeError(f"XLSX 非目标部件被改动：{original.name} / {name}")


def patch_xlsx_values(
    original: Path,
    temporary: Path,
    changes: dict[str, dict[int, tuple[float, ...]]],
    *,
    date_extensions: dict[str, dict[int, date]] | None = None,
    epoch=None,
) -> None:
    sheet_paths = xlsx_sheet_paths(original)
    part_changes: dict[str, dict[int, tuple[float, ...]]] = {}
    part_dates = {}
    for sheet_name, rows in changes.items():
        part = sheet_paths.get(sheet_name)
        if not part:
            raise RuntimeError(f"XLSX 包内未找到工作表：{sheet_name}")
        part_changes[part] = rows
        part_dates[part] = (date_extensions or {}).get(sheet_name, {})

    with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(temporary, "w") as output:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in part_changes:
                data = extend_sheet_dates(data, part_dates[info.filename], epoch=epoch)
                data = patch_sheet_xml(data, part_changes[info.filename])
            output.writestr(info, data)
    verify_package_preservation(original, temporary, set(part_changes))


def verify_workbook(
    path: Path, expectations: Iterable[CellExpectation],
    date_expectations: Iterable[DateExpectation] = (),
) -> None:
    expectations = tuple(expectations)
    date_expectations = tuple(date_expectations)
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
        )
        workbook = openpyxl.load_workbook(
            path, data_only=False, read_only=True, keep_links=True
        )
    try:
        references = [
            (item.sheet_name, f"{_column_name(column)}{item.row}")
            for item in expectations for column in range(2, 26)
        ]
        references.extend((item.sheet_name, f"A{item.row}") for item in date_expectations)
        snapshots = read_cell_snapshots(workbook, references)
        for item in date_expectations:
            value, number_format = snapshots[(item.sheet_name, f"A{item.row}")]
            if excel_cell_date(value, epoch=workbook.epoch) != item.value or number_format != item.number_format:
                raise RuntimeError(f"日期回读失败：{path.name} / {item.sheet_name} / A{item.row}")
        for expectation in expectations:
            cells = [snapshots[(expectation.sheet_name, f"{_column_name(column)}{expectation.row}")]
                     for column in range(2, 26)]
            actual = [value for value, _ in cells]
            if not values_match(actual, expectation.values):
                raise RuntimeError(
                    f"保存回读失败：{path.name} / {expectation.sheet_name} / 第 {expectation.row} 行"
                )
            formats = tuple(number_format for _, number_format in cells)
            if formats != expectation.number_formats:
                raise RuntimeError(
                    f"格式回读失败：{path.name} / {expectation.sheet_name} / 第 {expectation.row} 行"
                )
    finally:
        workbook.close()


def prepare_workbook(prices: TargetPrices, *, overwrite: bool) -> PreparedWorkbook:
    path = prices.target.workbook.resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"仅支持 .xlsx 工作簿：{path}")

    source_fingerprint = workbook_fingerprint(path)
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
        )
        workbook = openpyxl.load_workbook(
            path, data_only=False, read_only=True, keep_links=True
        )
    try:
        epoch = workbook.epoch
        extensions = {
            name: plan_date_extensions(workbook[name], values, epoch=epoch)
            for name, values in ((DAY_AHEAD_SHEET, prices.day_ahead), (REAL_TIME_SHEET, prices.real_time))
            if name in workbook.sheetnames
        }
        date_formats = {name: effective_number_formats(workbook[name], dates, 1)
                        for name, dates in extensions.items() if dates}
        date_expectations = tuple(
            DateExpectation(name, row, day, date_formats[name][row])
            for name, dates in extensions.items() for row, day in dates.items()
        )
        day_expectations, day_updated, day_skipped, day_rows = plan_price_sheet(
            workbook, DAY_AHEAD_SHEET, prices.day_ahead, overwrite=overwrite,
            date_extensions=extensions.get(DAY_AHEAD_SHEET),
        )
        real_expectations, real_updated, real_skipped, real_rows = plan_price_sheet(
            workbook, REAL_TIME_SHEET, prices.real_time, overwrite=overwrite,
            date_extensions=extensions.get(REAL_TIME_SHEET),
        )
        expectations = tuple(day_expectations + real_expectations)
    finally:
        workbook.close()
    try:
        patch_xlsx_values(
            path,
            temporary,
            {
                DAY_AHEAD_SHEET: day_rows,
                REAL_TIME_SHEET: real_rows,
            },
            date_extensions=extensions,
            epoch=epoch,
        )
        verify_workbook(temporary, expectations, date_expectations)
        assert_workbook_unchanged(path, source_fingerprint)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    for name, dates in extensions.items():
        if dates:
            log(f"{name}：已准备接续 {min(dates.values())} 至 {max(dates.values())}，共 {len(dates)} 行。")
    return PreparedWorkbook(
        original=path,
        temporary=temporary,
        source_fingerprint=source_fingerprint,
        expectations=expectations,
        updated_dates=tuple(sorted(set(day_updated + real_updated))),
        skipped_dates=tuple(sorted(set(day_skipped + real_skipped))),
        date_expectations=date_expectations,
    )


def commit_prepared_workbooks(
    prepared: tuple[PreparedWorkbook, ...], *, create_backup: bool
) -> dict[Path, Path | None]:
    return commit_workbooks(
        prepared, lambda path, item: verify_workbook(path, item.expectations, item.date_expectations),
        create_backup=create_backup,
    )


def update_spot_price_workbooks(
    start_date: date,
    end_date: date,
    targets: tuple[TargetSpec, ...],
    *,
    config: Config | None = None,
    overwrite: bool = True,
    create_backup: bool = True,
) -> tuple[PreparedWorkbook, ...]:
    if not targets:
        raise ValueError("请至少选择一个更新场站。")
    if len({target.workbook.resolve() for target in targets}) != len(targets):
        raise ValueError("已选场站不能指向同一个工作簿。")
    collected = collect_target_prices(
        start_date, end_date, targets, config=config
    )
    prepared: list[PreparedWorkbook] = []
    try:
        for prices in collected:
            if not prices.day_ahead and not prices.real_time:
                log(f"{prices.target.key}：日前和实时均无完整数据，不改动工作簿。")
                continue
            item = prepare_workbook(prices, overwrite=overwrite)
            prepared.append(item)
        backups = commit_prepared_workbooks(tuple(prepared), create_backup=create_backup)
    except Exception:
        for item in prepared:
            item.temporary.unlink(missing_ok=True)
        raise

    for item in prepared:
        log(
            f"已更新 {item.original}：{len(item.updated_dates)} 个日期；"
            f"未改动 {len(item.skipped_dates)} 个日期。"
        )
        backup = backups.get(item.original)
        if backup:
            log(f"备份：{backup}")
    return tuple(prepared)


def interactive_login(config: Config | None = None) -> None:
    standalone_login(
        load_auth_module(),
        config or load_config(),
        app_url=APP_URL,
        log=log,
        label="现货价格模块独立",
    )


def main() -> int:
    yesterday = date.today() - timedelta(days=1)
    parser = argparse.ArgumentParser(description="抓取所选场站的日前/实时出清电价并更新分析工作簿")
    parser.add_argument("--start-date", default=yesterday.isoformat(), help="开始日期 YYYY-MM-DD")
    parser.add_argument("--end-date", help="结束日期 YYYY-MM-DD，默认等于开始日期")
    parser.add_argument("--guihu-workbook", type=Path, default=DEFAULT_GUIHU_WORKBOOK)
    parser.add_argument("--tuolian-workbook", type=Path, default=DEFAULT_TUOLIAN_WORKBOOK)
    parser.add_argument("--haifeng-workbook", type=Path, default=DEFAULT_HAIFENG_WORKBOOK)
    parser.add_argument(
        "--stations", nargs="+", choices=STATION_KEYS,
        help="选择更新场站，可选一个或多个，默认全部：归湖 鮀莲 海风",
    )
    parser.add_argument("--skip-existing", action="store_true", help="已有非空价格时跳过，不覆盖")
    parser.add_argument("--no-backup", action="store_true", help="不保留时间戳备份")
    parser.add_argument("--headed", action="store_true", help="显示抓取浏览器")
    parser.add_argument("--dry-run", action="store_true", help="只抓取并验证数据，不写入 Excel")
    parser.add_argument("--login", action="store_true", help="登录并保存本模块独立会话")
    args = parser.parse_args()

    try:
        config = replace(load_config(), headless=not args.headed)
        if args.login:
            interactive_login(config)
            return 0
        start_date = validate_date(args.start_date)
        end_date = validate_date(args.end_date or args.start_date)
        targets = default_targets(
            args.guihu_workbook, args.tuolian_workbook, args.haifeng_workbook,
            selected_stations=args.stations,
        )
        if args.dry_run:
            collected = collect_target_prices(start_date, end_date, targets, config=config)
            for item in collected:
                log(
                    f"校验通过：{item.target.key} / {item.tenant_name} / {item.unit_name} / "
                    f"{len(item.day_ahead)} 天日前 + {len(item.real_time)} 天实时。"
                )
            return 0
        update_spot_price_workbooks(
            start_date,
            end_date,
            targets,
            config=config,
            overwrite=not args.skip_existing,
            create_backup=not args.no_backup,
        )
        return 0
    except KeyboardInterrupt:
        print("已取消。", file=sys.stderr)
        return 130
    except Exception as exc:
        from toolbox_shared.cli_login import emit_login_error
        emit_login_error(exc)
        print(f"失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
