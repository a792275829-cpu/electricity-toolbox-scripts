"""Sparse formula/value views for inspecting XLSX without editing its package."""
from array import array
from dataclasses import dataclass
from xml.etree.ElementTree import iterparse
from pathlib import Path

from openpyxl.reader.excel import ExcelReader
from openpyxl.styles import Border
from openpyxl.styles.cell_style import StyleArray
from openpyxl.styles.numbers import BUILTIN_FORMATS, BUILTIN_FORMATS_MAX_SIZE
from openpyxl.styles.stylesheet import apply_stylesheet
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.worksheet._reader import WorkSheetParser


class UnsupportedInspection(ValueError):
    """The full workbook loader must handle this template's merge semantics."""


@dataclass(frozen=True, slots=True)
class InspectionCell:
    value: object
    coordinate: str
    number_format: str
    _style: object
    _styles: object

    @property
    def style_id(self):
        return self._styles.add(self._style)


class InspectionSheet:
    def __init__(self, title, cells, max_row, formats, styles):
        self.title = title
        self._cells = cells
        self.max_row = max_row or 1
        self._formats = formats
        self._styles = styles

    def cell(self, row, column):
        if row < 1 or column < 1:
            raise ValueError("Row or column values must be at least 1")
        coordinate = f"{get_column_letter(column)}{row}"
        item = self._cells.get((row, column))
        if item is None:
            # Synthesized openpyxl Cells start with an empty StyleArray, even
            # when the file's existing style 0 has a custom number format.
            return InspectionCell(None, coordinate, "General", StyleArray(), self._styles)
        value, style = item
        return InspectionCell(value, coordinate, self._formats[style], self._styles[style], self._styles)

    def __getitem__(self, coordinate):
        return self.cell(*coordinate_to_tuple(coordinate))

    def iter_rows(self, min_row=1, max_row=None, min_col=1, max_col=None, values_only=False):
        for row in range(min_row, (self.max_row if max_row is None else max_row) + 1):
            cells = tuple(self.cell(row, col) for col in range(min_col, (max_col or 14) + 1))
            yield tuple(cell.value for cell in cells) if values_only else cells


class InspectionWorkbook:
    def __init__(self, epoch, sheets):
        self.epoch = epoch
        self._sheets = sheets

    def __iter__(self):
        return iter(self._sheets.values())

    def __getitem__(self, title):
        return self._sheets[title]

    def close(self):
        pass  # The ZIP is closed before either sparse view is returned.


def load_inspection_pair(path, *, columns, target_columns, sheet_filter, requests=None, capture_cache=True):
    """One streaming parse keeps selected formula cells and column-A caches.

    Existing openpyxl handles date epochs, shared formulas, array formulas,
    strings and errors. Merges in target columns below the two header rows
    require the full loader's border/style propagation and are rejected here.
    """
    reader = ExcelReader(Path(path), read_only=True, data_only=False, keep_links=False)
    try:
        reader.read_manifest()
        reader.read_strings()
        reader.read_workbook()
        book = reader.wb
        apply_stylesheet(reader.archive, book)
        formats = []
        for style in book._cell_styles:
            fid = style.numFmtId
            formats.append(BUILTIN_FORMATS.get(fid, "General") if fid < BUILTIN_FORMATS_MAX_SIZE
                           else book._number_formats[fid - BUILTIN_FORMATS_MAX_SIZE])
        formula_sheets, cache_sheets = {}, {}
        columns = frozenset(columns)
        for descriptor, relation in reader.parser.find_sheets():
            if "chartsheet" in relation.Type:
                continue
            if relation.target not in reader.valid_files:
                raise UnsupportedInspection("工作表部件缺失")
            if not sheet_filter(descriptor.name):
                # Merges in an otherwise ignored sheet can mutate shared style
                # arrays when the full loader propagates borders. Be conservative.
                with reader.archive.open(relation.target) as source:
                    for _, element in iterparse(source):
                        if element.tag.rsplit("}", 1)[-1] == "mergeCell":
                            raise UnsupportedInspection("非月份工作表含合并样式")
                        element.clear()
                continue
            cells, cached, row_styles = {}, {}, {}
            options = dict(epoch=book.epoch, date_formats=book._date_formats,
                           timedelta_formats=book._timedelta_formats)
            cache_parser = WorkSheetParser(None, reader.shared_strings, data_only=True, **options)

            class InspectionParser(WorkSheetParser):
                def parse_cell(self, element):
                    # Register even unselected shared-formula masters by using
                    # the original parser for every cell; retain only columns.
                    position = element.get("r")
                    cache_parser.row_counter = self.row_counter
                    cache_parser.col_counter = self.col_counter
                    if capture_cache and ((position and coordinate_to_tuple(position)[1] == 1) or (not position and self.col_counter == 0)):
                        item = cache_parser.parse_cell(element)
                        cached[(item["row"], item["column"])] = (item["value"], item["style_id"])
                    return super().parse_cell(element)

            max_row = 0
            with reader.archive.open(relation.target) as source:
                parser = InspectionParser(source, reader.shared_strings, data_only=False, **options)
                for _, row in parser.parse():
                    if row:
                        row_styles[row[0]["row"]] = array("I", (value for item in row for value in (item["column"], item["style_id"])))
                    for item in row:
                        max_row = max(max_row, item["row"])
                        key = (item["row"], item["column"])
                        if item["column"] in columns and (requests is None or key in requests.get(descriptor.name, ())):
                            cells[key] = (item["value"], item["style_id"])
            merges = parser.merged_cells.mergeCell if parser.merged_cells else ()
            def style_at(row, column):
                packed = row_styles.get(row, ())
                return next((packed[i + 1] for i in range(0, len(packed), 2) if packed[i] == column), None)

            for merge in merges:
                left, top, right, bottom = range_boundaries(str(merge.ref))
                end_style = style_at(bottom, right)
                if end_style is not None:
                    start_style = style_at(top, left) or 0
                    start_border = book._borders[book._cell_styles[start_style].borderId]
                    end_border = book._borders[book._cell_styles[end_style].borderId]
                    if start_border + Border(right=end_border.right, bottom=end_border.bottom) != start_border:
                        raise UnsupportedInspection("合并区域需要传播锚点边框")
                if bottom > 2 and any(left <= column <= right for column in target_columns) and (left != right or top != bottom):
                    raise UnsupportedInspection("数据区合并单元格需要原样式处理")
                max_row = max(max_row, bottom)
                # Full worksheets discard all non-anchor merged cell values.
                for column in columns:
                    if left <= column <= right:
                        for row in range(top, bottom + 1):
                            if (row, column) != (top, left):
                                cells.pop((row, column), None)
                                cached.pop((row, column), None)
            formula_sheets[descriptor.name] = InspectionSheet(descriptor.name, cells, max_row, formats, book._cell_styles)
            cache_sheets[descriptor.name] = InspectionSheet(descriptor.name, cached, max_row, formats, book._cell_styles)
        return InspectionWorkbook(book.epoch, formula_sheets), InspectionWorkbook(book.epoch, cache_sheets)
    finally:
        reader.archive.close()
