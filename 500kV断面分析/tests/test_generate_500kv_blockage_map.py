import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook

from src.generate_500kv_blockage_map import generate_blockage_map


class Generate500kVBlockageMapTest(unittest.TestCase):
    def test_in_memory_pipeline_matches_standalone_workbook_and_html_bytes(self):
        from src import analyze_500kv_section_gaps as analysis
        from src.build_500kv_map_html import build_map_html
        from src.deduplicate_500kv_prices import process_workbook
        from src.generate_500kv_blockage_map import _run_pipeline
        with TemporaryDirectory() as directory:
            root = Path(directory)
            topology, raw, cleaned = root / "topology.md", root / "raw.xlsx", root / "cleaned.xlsx"
            topology.write_text("## 1. 节点分区\n### 片区\n甲、乙、丙\n## 2. 接线\n### 片区\n甲-乙、乙-丙\n## 3. 约定\n")
            book = Workbook()
            sheet = book.active
            times = [f"{i // 4:02d}:{i % 4 * 15:02d}" for i in range(12)]
            sheet.append(["节点名称", "数据项", *times])
            for index, name in enumerate(("甲", "乙", "丙")):
                sheet.append([f"其他{name}站500kV", "电价(元/MWh)", *[
                    (100.1234567890123 + i * 99.9876543210123) * (-1 if index == 1 else 1)
                    for i in range(12)
                ]])
            book.save(raw)
            book.close()
            process_workbook(raw, cleaned, topology)
            result = analysis.analyze_section_gaps(topology, cleaned, root)
            html = root / "map.html"
            build_map_html(topology, raw, cleaned, result.output_path, html)
            reference = html.read_bytes()
            result.output_path.unlink()
            with patch.object(analysis, "write_analysis_workbook", side_effect=AssertionError("intermediate export")), patch.object(
                analysis, "read_price_workbook", wraps=analysis.read_price_workbook
            ) as read:
                optimized = _run_pipeline(topology, raw, root, 80.0, html, cleaned)
            self.assertEqual(read.call_count, 1)
            self.assertEqual(html.read_bytes(), reference)
            self.assertEqual(optimized.flagged_edges, result.flagged_edges)
            self.assertEqual(optimized.flagged_intervals, result.flagged_intervals)
            self.assertFalse(result.output_path.exists())

    def test_generate_blockage_map_runs_full_pipeline_from_raw_prices(self):
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            raw_prices = tmp_path / "实时节点电价查询.xlsx"
            topology = tmp_path / "topology.md"
            output_dir = tmp_path / "out"

            topology.write_text(
                """# test

## 1. 按 PDF 重划片区

### 佛山片
凤城

### 东莞片
水乡

## 2. 跨片区主通道边

### 佛山片 - 东莞片
凤城-水乡

## 3. next
""",
                encoding="utf-8",
            )

            wb = Workbook()
            ws = wb.active
            ws.title = "实时节点电价查询"
            ws.append(["节点名称", "数据项", "00:00", "00:15", "00:30", "00:45", "01:00", "01:15", "01:30", "01:45", "02:00"])
            ws.append(["其他凤城站500kV", "电价(元/MWh)", 500, 600, 700, 800, 900, 1000, 1100, 1200, 1300])
            ws.append(["其他水乡站500kV", "电价(元/MWh)", 500, -100, -200, -300, -400, -500, -600, -700, -800])
            ws.append(["其他一总降站220kV", "电价(元/MWh)", 1, 1, 1, 1, 1, 1, 1, 1, 1])
            wb.save(raw_prices)

            result = generate_blockage_map(
                topology_path=topology,
                raw_price_path=raw_prices,
                output_dir=output_dir,
                threshold=500,
            )

            self.assertTrue(result.html_path.exists())
            self.assertEqual(
                sorted(path.name for path in output_dir.iterdir()),
                ["实时节点电价查询_500kV分析.html"],
            )

            html = result.html_path.read_text(encoding="utf-8")
            self.assertIn("广东省电网 500kV 节点断面地图", html)
            self.assertIn('"edge": "凤城-水乡"', html)
            self.assertIn('"is_section": true', html)

    def test_generate_blockage_map_removes_legacy_excel_outputs(self):
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            raw_prices = tmp_path / "实时节点电价查询.xlsx"
            topology = tmp_path / "topology.md"
            output_dir = tmp_path / "out"
            topology.write_text(
                """# test

## 2. 跨片区主通道边

### 佛山片 - 东莞片
凤城-水乡

## 3. next
""",
                encoding="utf-8",
            )

            wb = Workbook()
            ws = wb.active
            ws.title = "实时节点电价查询"
            ws.append(["节点名称", "数据项", "00:00"])
            ws.append(["其他凤城站500kV", "电价(元/MWh)", 100])
            ws.append(["其他水乡站500kV", "电价(元/MWh)", 700])
            wb.save(raw_prices)

            output_dir.mkdir()
            (output_dir / "500kv_section_price_gap_analysis.xlsx").write_bytes(b"legacy")
            (output_dir / "实时节点电价查询_500kV去重.xlsx").write_bytes(b"legacy")
            (output_dir / "guangdong_500kv_section_map.html").write_text("legacy", encoding="utf-8")

            result = generate_blockage_map(
                topology_path=topology,
                raw_price_path=raw_prices,
                output_dir=output_dir,
                threshold=500,
            )

            self.assertTrue(result.html_path.exists())
            self.assertEqual(
                sorted(path.name for path in output_dir.iterdir()),
                ["实时节点电价查询_500kV分析.html"],
            )


if __name__ == "__main__":
    unittest.main()
