from datetime import datetime
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal, assert_series_equal
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "500kV断面分析"), str(ROOT / "广东电价预测"), str(ROOT / "广东运行日报PDF解析")]
from toolbox.runtime import load_module
from toolbox_shared.native_excel import native_binary, read_model_frame
from toolbox_shared import native_compute as compute
from toolbox_shared import native_write as writer
from src.deduplicate_500kv_prices import _indexed_topology_matcher, match_topology_node
from gd_price_forecast.features import build_features, FORECAST_COLUMNS
from gd_pdf_tool.parser import _unmix_bar_colors_python


class IndexedAndFeatureTests(unittest.TestCase):
    def test_index_preserves_exact_longest_alias_and_containment_priority(self):
        nodes = ["宝安", "宝安站", "宝安换流站", "凤城", "凤城站", "徐闻开关站", "柘林电厂"]
        names = [None, "", "其他宝安换流站500kVACF1母", "凤城站500kV2母", "前缀柘林电厂#1M", "徐闻开关站徐闻开关站任一个", "未知站"]
        names += [f"其他{node}500kV#{branch}" for node in nodes for branch in range(10)]
        match = _indexed_topology_matcher(nodes)
        for name in names * 2:
            self.assertEqual(match(name), match_topology_node(name, nodes))

    def test_feature_values_dtypes_order_and_missing_dates_match(self):
        rng = np.random.default_rng(42)
        dates = pd.date_range("2026-09-01", periods=24).delete(10)
        frame = pd.DataFrame({"日期": np.repeat(dates,96), "时刻": [f"{h:02d}:{m:02d}" for h in range(24) for m in (0,15,30,45)] * len(dates)})
        frame["全省日前平均电价"] = rng.normal(300,100,len(frame))
        frame.loc[5:10,"全省日前平均电价"] = np.nan
        for column in FORECAST_COLUMNS:
            frame[column] = rng.normal(10000,100,len(frame))
        frame["温度"] = rng.normal(25,1,len(frame))
        for exogenous in ("legacy","none","published","lagged"):
            results = []
            for mode in ("python","auto"):
                with patch.dict(os.environ, TOOLBOX_FEATURE_BACKEND=mode):
                    results.append(build_features(frame,exogenous_mode=exogenous,holidays=True,supply=True,weather=True))
            assert_frame_equal(results[0][0],results[1][0],check_exact=True)
            assert_series_equal(results[0][1],results[1][1],check_exact=True)
            assert_frame_equal(results[0][2],results[1][2],check_exact=True)


@unittest.skipUnless(native_binary(), "build native/price-xlsx first")
class RemainingNativeTests(unittest.TestCase):
    def test_old_binary_capability_failure_falls_back_without_waiting_for_stream(self):
        with patch.object(compute,"_supports_stream",return_value=False), patch.dict(os.environ,TOOLBOX_COMPUTE_BACKEND="auto"):
            with self.assertWarns(RuntimeWarning):
                with compute.native_session():
                    self.assertIsNone(compute.unmix(np.zeros((128,32,3),dtype=np.uint8),np.zeros((128,3)),[None,None,None]))
            self.assertIsNone(compute._session.get())
        with patch.object(compute,"_supports_stream",return_value=False), patch.dict(os.environ,TOOLBOX_COMPUTE_BACKEND="rust"):
            with self.assertRaisesRegex(RuntimeError,"重新构建"):
                with compute.native_session():pass

    def test_frame_reader_keeps_pandas_header_na_dtype_and_date_behavior(self):
        with tempfile.TemporaryDirectory() as tmp:
            for use_mac in (False,True):
                path = Path(tmp) / "frame.xlsx"
                book = Workbook()
                if use_mac:
                    book.epoch = CALENDAR_MAC_1904
                sheet = book.active
                sheet.title = "第一张"
                sheet.append(["日期","编号","编号","缺失","公式","错误","真假"])
                sheet.append([datetime(2026,10,20),"001",3.5,"NA","=1+1","#DIV/0!",False])
                sheet.append([datetime(2026,10,21),"002",2,None,None,None,True])
                book.save(path)
                for preferred in ("不存在","第一张"):
                    frames=[]
                    for mode in ("python","rust"):
                        with patch.dict(os.environ,TOOLBOX_XLSX_BACKEND=mode):
                            frames.append(read_model_frame(path,preferred))
                    assert_frame_equal(frames[0],frames[1],check_exact=True)

    def test_float_and_uint8_unmix_profiles_match_numpy(self):
        rng=np.random.default_rng(123)
        palette=[np.array([20.,50.,220.]),np.array([220.,30.,40.]),np.array([30.,160.,50.])]
        for width in (9,10):
            for dtype in (np.uint8,np.float64):
                pixels=rng.uniform(0,255,(128,width,3)).astype(dtype)
                background=np.full((128,3),255.)
                with patch.dict(os.environ,TOOLBOX_COMPUTE_BACKEND="python"):
                    expected=_unmix_bar_colors_python(pixels,background,palette,filter_sources=False)
                with patch.dict(os.environ,TOOLBOX_COMPUTE_BACKEND="rust"), compute.native_session():
                    actual=compute.unmix(pixels,background,palette)
                    profile=compute.unmix(pixels,background,palette,profile=True)
                np.testing.assert_allclose(actual,expected,rtol=1e-12,atol=1e-12)
                np.testing.assert_allclose(profile,np.median(expected,axis=1,keepdims=True),rtol=1e-12,atol=1e-12)

    def test_compute_session_terminates_after_exception(self):
        import subprocess
        processes=[]
        real=compute.subprocess.Popen
        def record(*args,**kwargs):
            process=real(*args,**kwargs);processes.append(process);return process
        with patch.dict(os.environ,TOOLBOX_COMPUTE_BACKEND="rust"), patch.object(compute.subprocess,"Popen",side_effect=record):
            with self.assertRaisesRegex(ValueError,"intentional"):
                with compute.native_session():
                    raise ValueError("intentional")
        self.assertIsNone(compute._session.get())
        self.assertTrue(all(p.poll() is not None for p in processes))

    def test_native_sheet_xml_preserves_escapes_formulas_styles_and_blanks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);book=Workbook();sheet=book.active
            sheet.append(["  中文<&>\r\n文本  ",1/3,False,"=1+1","#N/A",""])
            sheet["B1"].number_format="0.0000"
            sheet["A1"].font=Font(bold=True,color="112233")
            sheet["A1"].fill=PatternFill("solid",fgColor="ABCDEF")
            sheet.merge_cells("A2:C2");sheet["A2"]="合并"
            sheet.freeze_panes="B3";sheet.column_dimensions["A"].width=25
            sheet.row_dimensions[1].height=32
            old=root/"old.xlsx";new=root/"new.xlsx"
            book.save(old)
            with patch.dict(os.environ,TOOLBOX_XLSX_BACKEND="rust"):
                writer.save_workbook(book,new)
            a=load_workbook(old);b=load_workbook(new)
            try:
                for ra,rb in zip(a.active,b.active):
                    for ca,cb in zip(ra,rb):
                        self.assertEqual((ca.value,ca.data_type,ca.number_format,ca._style),(cb.value,cb.data_type,cb.number_format,cb._style))
                self.assertEqual(a.active.freeze_panes,b.active.freeze_panes)
                self.assertEqual(str(a.active.merged_cells),str(b.active.merged_cells))
            finally:a.close();b.close()

    def test_failed_native_summary_write_preserves_existing_output(self):
        module=load_module("native_failure_summary",ROOT/"电力工具脚本/电量汇总工具/summarize_511_excel.py")
        with tempfile.TemporaryDirectory() as tmp:
            output=Path(tmp)/"summary.xlsx";output.write_bytes(b"old")
            book=Workbook();book.active["A1"]="new"
            with patch.dict(os.environ,TOOLBOX_XLSX_BACKEND="rust"), patch.object(writer,"call",return_value={"written":False}):
                with self.assertRaises(ValueError):module.save_workbook_atomically(book,output)
            self.assertEqual(output.read_bytes(),b"old")
            self.assertEqual(list(output.parent.iterdir()),[output])
