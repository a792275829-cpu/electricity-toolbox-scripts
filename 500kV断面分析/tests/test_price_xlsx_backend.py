import datetime
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src import price_xlsx_backend as backend
from src.deduplicate_500kv_prices import process_workbook


def snapshot(path):
    book = load_workbook(path, data_only=False)
    try:
        return [(s.title, [(tuple((c.value, c.data_type, c.number_format) for c in row))
                           for row in s.iter_rows()]) for s in book]
    finally:
        book.close()


class BackendFallbackTests(unittest.TestCase):
    def test_unavailable_native_auto_falls_back_and_strict_fails(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(backend, "native_binary", return_value=None):
            path = Path(tmp) / "input.xlsx"
            Workbook().save(path)
            with patch.dict(os.environ, TOOLBOX_500KV_BACKEND="auto"):
                source, output = backend.open_price_workbooks(path)
                source.close()
                output.close()
            with patch.dict(os.environ, TOOLBOX_500KV_BACKEND="rust"):
                with self.assertRaisesRegex(RuntimeError, "未安装"):
                    backend.open_price_workbooks(path)

    def test_invalid_backend_is_rejected(self):
        with patch.dict(os.environ, TOOLBOX_500KV_BACKEND="invalid"):
            with self.assertRaises(ValueError):
                backend.backend_mode()


@unittest.skipUnless(backend.native_binary(), "build native/price-xlsx first")
class RustParityTests(unittest.TestCase):
    def test_both_modes_preserve_values_types_formats_and_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "input.xlsx"
            topo = root / "topology.md"
            topo.write_text("## 2. 拓扑\n### 测试\n木棉站-凤城站\n## 3. 结束\n")
            book = Workbook()
            sheet = book.active
            sheet.title = "中文表"
            sheet.append([None, "节点名称", "数据项", "00:00", "00:15", "00:30"])
            sheet.append([None, "其他木棉站500kV1母", "电价", "100.00", None, -1.25])
            sheet.append([None, "其他木棉站500kV2母", "电价", 200, "", "-1.250"])
            sheet.append([None, "其他凤城站500kV1母", "电价", 100, None, -1.25])
            sheet.append([None, "其他凤城站500kV2母", "电价", 100, None, -1.25])
            sheet.append([None, "其他非拓扑220kV", "电价", 10, 20, 30])
            sheet.append([None, None, None])
            other = book.create_sheet("第二张")
            other.append(["节点名称", "数据项", "00:00", "00:15", "00:30"])
            other.append(["其他木棉站500kV3母", "电价", 300, None, -1.25])
            book.save(path)
            for topology in (None, topo):
                results = []
                for mode in ("python", "rust"):
                    with patch.dict(os.environ, TOOLBOX_500KV_BACKEND=mode):
                        output = root / f"{mode}.xlsx"
                        result = process_workbook(path, output, topology_path=topology)
                        results.append(((result.kept_rows, result.removed_rows, result.processed_sheets), snapshot(output)))
                self.assertEqual(results[0], results[1])

    def test_native_failure_does_not_replace_existing_output(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, TOOLBOX_500KV_BACKEND="rust"):
            output = Path(tmp) / "existing.xlsx"
            output.write_bytes(b"original")
            book = backend.PriceWorkbook(backend.native_binary())
            book.create_sheet("sheet").append(["value"])
            with patch.object(backend, "_call", side_effect=ValueError("failure")):
                with self.assertRaises(ValueError):
                    book.save(output)
            self.assertEqual(output.read_bytes(), b"original")
            self.assertEqual(list(Path(tmp).iterdir()), [output])

    def test_fractional_averages_match_openpyxl_sixteen_digit_rounding(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "input.xlsx"
            topology = root / "topology.md"
            topology.write_text("## 2. 拓扑\n### 测试\n木棉站-凤城站\n## 3. 结束\n")
            book = Workbook()
            book.active.append(["节点名称", "数据项", "00:00", "00:15"])
            for branch, value in enumerate((100.12, 300.0, 585.46)):
                book.active.append([f"木棉站500kV{branch}母", "电价", value, "-0.00125"])
            book.save(path)
            snapshots = []
            for mode in ("python", "rust"):
                with patch.dict(os.environ, TOOLBOX_500KV_BACKEND=mode):
                    output = root / f"{mode}.xlsx"
                    process_workbook(path, output, topology_path=topology)
                    snapshots.append(snapshot(output))
            self.assertEqual(snapshots[0], snapshots[1])

    def test_dates_trigger_compatibility_reader(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            book = Workbook()
            book.active.append([datetime.datetime(2026, 10, 20)])
            book.save(path)
            with patch.dict(os.environ, TOOLBOX_500KV_BACKEND="auto"):
                with self.assertWarns(RuntimeWarning):
                    source, output = backend.open_price_workbooks(path)
                try:
                    self.assertEqual(source.active["A1"].value, datetime.datetime(2026, 10, 20))
                finally:
                    source.close()
                    output.close()

    def test_header_below_row_one_is_not_promoted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            book = Workbook()
            book.active.append([])
            book.active.append(["节点名称", "数据项", "00:00"])
            book.active.append(["木棉站500kV", "电价", 100])
            book.save(path)
            for mode in ("rust", "python"):
                with patch.dict(os.environ, TOOLBOX_500KV_BACKEND=mode):
                    with self.assertRaisesRegex(ValueError, "没有找到"):
                        process_workbook(path, Path(tmp) / f"{mode}.xlsx")
