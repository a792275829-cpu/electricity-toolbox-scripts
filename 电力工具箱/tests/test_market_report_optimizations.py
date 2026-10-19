from __future__ import annotations

import copy
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import ExitStack
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

import fcntl

from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

ROOT = Path(__file__).resolve().parents[2]
REPORT_ROOT = ROOT / "南方分公司电力市场日报"
sys.path.insert(0, str(REPORT_ROOT))
import run_market_report as runner
from gd_power_report import output_commit, pipeline, workbook_reader
from gd_power_report.collector import market
from gd_power_report.config import AppConfig


def write_pdf(path, pages=1, version="new"):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": version})
    with path.open("wb") as stream:
        writer.write(stream)


class ReportInputAndErrorTests(unittest.TestCase):
    def test_cli_validation_date_inference_and_pipeline_open_each_workbook_once(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / f"{index}.xlsx" for index in range(3)]
            for index, path in enumerate(paths):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "分公司" if index == 1 else "新快报"
                sheet["A2" if index == 1 else "A6"] = date(2026, 10, 7) - timedelta(days=index)
                workbook.save(path)
                workbook.close()
            argv = ["run_market_report", "2026-10-07", "--day-report", str(paths[0]),
                    "--previous-report", str(paths[1]), "--d2-report", str(paths[2])]
            opened = []
            original_loader = workbook_reader._load_workbook

            def load(*args, **kwargs):
                workbook = original_loader(*args, **kwargs)
                workbook.close = mock.Mock(wraps=workbook.close)
                opened.append(workbook)
                return workbook

            @workbook_reader.reuse_workbook_inputs
            def generate(*args, **kwargs):
                # The pipeline inherits the entry point's live cache.
                for path, workbook in zip(paths, opened):
                    self.assertIs(workbook_reader.load_workbook(path, data_only=True, read_only=False), workbook)
                return pipeline.RunResult(Path(directory), Path("phone.pdf"), Path("print.pdf"))

            with mock.patch.object(sys, "argv", argv), \
                 mock.patch.object(workbook_reader, "_load_workbook", side_effect=load) as loader, \
                 mock.patch.object(runner, "run_pipeline", side_effect=generate), \
                 mock.patch("builtins.print"):
                self.assertEqual(runner.main(), 0)
                self.assertEqual(loader.call_count, 3)
            self.assertIsNone(workbook_reader._workbooks.get())
            for workbook in opened:
                workbook.close.assert_called_once()

    def inputs(self, stack, *, rows=None, total=10, detail_date="2026-10-05"):
        report = {
            "report_date": "2026-10-07", "report_date_label": "10月7日",
            "day_ahead": {"cards": []}, "realtime": {"cards": []},
            "operations": {"rows": []}, "appendices": {},
            "market": {"day_ahead": [], "realtime": [{"label": "全省实时均价", "value": 111}]},
        }
        previous = {
            "report_date": "2026-10-06", "report_date_label": "10月6日",
            "cards": [], "operations": {}, "source_path": "previous.xlsx",
        }
        detail = {"report_date": detail_date, "realtime": {"cards": [
            {"label": "实时出清电量", "value": total}
        ]}, "market": {"day_ahead": [{"label": "全省日前均价", "value": 250}],
                        "realtime": [{"label": "全省实时均价", "value": 300}]}}
        stack.enter_context(mock.patch.object(pipeline, "parse_daily_report_workbook", side_effect=[copy.deepcopy(report), detail]))
        stack.enter_context(mock.patch.object(pipeline, "parse_day_ahead_comparison_workbook", return_value=previous))
        units = stack.enter_context(mock.patch.object(pipeline, "parse_realtime_unit_workbook", return_value={"source_path": "d2.xlsx"}))
        stack.enter_context(mock.patch.object(pipeline, "summarize_realtime_units", return_value=rows if rows is not None else [{"realtime": 10}]))
        full_day = stack.enter_context(mock.patch.object(pipeline, "read_full_day_clearing_table", return_value={}))
        collect = stack.enter_context(mock.patch.object(pipeline, "collect_realtime_market_comparison"))
        render = stack.enter_context(mock.patch.object(pipeline, "_render_concise_outputs"))
        return units, full_day, collect, render

    def run_report(self, directory, progress=lambda _: None):
        return pipeline.run_pipeline(AppConfig.from_workspace(directory), date(2026, 10, 7), "day.xlsx", progress,
                                     comparison_clearing_xlsx_path="previous.xlsx", realtime_detail_xlsx_path="d2.xlsx")

    def test_invalid_d2_fails_before_any_website_request(self):
        for options, message in (({"rows": [{"realtime": None}]}, "缺少实时出清电量"),
                                 ({"total": 20}, "核对不一致"),
                                 ({"detail_date": "2026-10-04"}, "应为 2026-10-05")):
            with self.subTest(options=options), tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
                _, _, collect, render = self.inputs(stack, **options)
                with self.assertRaisesRegex(ValueError, message):
                    self.run_report(directory)
                collect.assert_not_called()
                render.assert_not_called()
                self.assertFalse(AppConfig.from_workspace(directory).outputs_dir.exists())

    def test_unit_date_or_hour_error_fails_before_website_request(self):
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            units, _, collect, render = self.inputs(stack)
            units.side_effect = ValueError("汕头1号机缺少小时19:00")
            with self.assertRaisesRegex(ValueError, "缺少小时19:00"):
                self.run_report(directory)
            collect.assert_not_called()
            render.assert_not_called()

    def test_network_error_keeps_dates_and_original_cause(self):
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            _, full_day, collect, render = self.inputs(stack)
            original = RuntimeError("实时负荷接口 HTTP 503")

            def fetch(*args, **kwargs):
                full_day.assert_called_once_with("d2.xlsx", date(2026, 10, 5))
                raise original

            collect.side_effect = fetch
            messages = []
            with self.assertRaises(RuntimeError) as caught:
                self.run_report(directory, messages.append)
            text = str(caught.exception)
            for phrase in ("日前 2026-10-07", "实时 2026-10-05", "实时负荷接口 HTTP 503"):
                self.assertIn(phrase, text)
            self.assertIs(caught.exception.__cause__, original)
            self.assertEqual(messages[-1], text)
            render.assert_not_called()

    def test_missing_market_payload_keeps_collector_diagnostic(self):
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            _, _, collect, render = self.inputs(stack)
            collect.return_value = {"errors": {"market_day_ahead": "2026-10-07 日前负荷接口 HTTP 502"}}
            with self.assertRaisesRegex(RuntimeError, "日前负荷接口 HTTP 502"):
                self.run_report(directory)
            render.assert_not_called()

    def test_complete_pipeline_publishes_pair_and_keeps_d2_price_basis(self):
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            units, full_day, collect, render = self.inputs(stack)
            collect.return_value = {
                f"market_{period}_cards": [{"label": label, "value": 100,
                                           "thermal_breakdown": {"coal_delta": 1, "gas_delta": 2}}]
                for period, label in (("day_ahead", "市场开机"), ("realtime", "在线机组数"))
            }

            def export(run_dir, report, progress):
                full_day.assert_called_once()
                collect.assert_called_once()
                units.assert_called_once_with("d2.xlsx", date(2026, 10, 5))
                self.assertEqual(report["realtime_date"], "2026-10-05")
                self.assertEqual(report["realtime"]["cards"][0]["value"], 10)
                price = next(card for card in report["market"]["realtime"] if card["label"] == "全省实时均价")
                self.assertEqual((price["value"], price["delta"], price["delta_prefix"]), (300, 50, "较日前"))
                phone, printing = run_dir / "phone.pdf", run_dir / "print.pdf"
                write_pdf(phone)
                write_pdf(printing, 2)
                return run_dir / "phone.html", phone, printing

            render.side_effect = export
            result = self.run_report(directory)
            self.assertEqual(result.run_dir.name, "2026-10-07")
            self.assertEqual(len(PdfReader(result.preview_pdf_path).pages), 1)
            self.assertEqual(len(PdfReader(result.print_pdf_path).pages), 2)
            self.assertEqual(set(result.run_dir.iterdir()), {result.preview_pdf_path, result.print_pdf_path})

    def test_missing_thermal_comparison_identifies_date_and_fuel(self):
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            _, _, collect, render = self.inputs(stack)
            collect.return_value = {
                "market_day_ahead_cards": [{"label": "市场开机", "thermal_breakdown": {"coal_delta": 1, "gas_delta": None}}],
                "errors": {"thermal_unit_count": "2026-10-06 前一日日前机组台数：HTTP 503"},
            }
            with self.assertRaises(RuntimeError) as caught:
                self.run_report(directory)
            for phrase in ("2026-10-07 日前", "缺少气机", "2026-10-06 前一日日前", "HTTP 503"):
                self.assertIn(phrase, str(caught.exception))
            render.assert_not_called()

    def test_thermal_retries_preserve_first_exception_and_dispose_responses(self):
        context = mock.Mock()
        responses = [mock.Mock(ok=False, status=code) for code in (503, 502, 500)]
        context.request.post.side_effect = responses
        with mock.patch.object(market, "_retry_pause"):
            with self.assertRaises(RuntimeError) as caught:
                market._fetch_thermal_payload(context, "https://example.invalid", "2026-10-05",
                                              ("group",), "2", "实时机组台数", lambda _: None)
        for phrase in ("2026-10-05", "分时分类型出清接口", "首次异常", "HTTP 503", "HTTP 500"):
            self.assertIn(phrase, str(caught.exception))
        self.assertIn("HTTP 503", str(caught.exception.__cause__))
        for response in responses:
            response.dispose.assert_called_once()


class PdfPairPublicationTests(unittest.TestCase):
    def sources(self, directory):
        sources = (directory / "phone.pdf", directory / "print.pdf")
        for index, path in enumerate(sources):
            write_pdf(path, pages=index + 1)
        return sources

    def test_success_publishes_both_verified_files_and_cleans_staging(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sources = self.sources(root)
            targets = output_commit.publish_pdf_pair(root / "output", sources)
            for index, (source, target) in enumerate(zip(sources, targets)):
                self.assertEqual(target.read_bytes(), source.read_bytes())
                self.assertEqual(len(PdfReader(target).pages), index + 1)
            self.assertEqual(set((root / "output").iterdir()), set(targets))

    def test_second_replace_failure_restores_existing_or_absent_pair(self):
        for existing in ((True, True), (False, False), (True, False), (False, True)):
            with self.subTest(existing=existing), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                sources = self.sources(root)
                output = root / "output"
                output.mkdir()
                originals = {}
                for source, present in zip(sources, existing):
                    if present:
                        target = output / source.name
                        write_pdf(target, version="previous")
                        originals[target] = target.read_bytes()
                replace = os.replace

                def fail_second(source, target):
                    if Path(source).name == "1.pdf":
                        raise OSError("second replacement failed")
                    replace(source, target)

                with mock.patch.object(output_commit.os, "replace", side_effect=fail_second):
                    with self.assertRaisesRegex(OSError, "second replacement failed"):
                        output_commit.publish_pdf_pair(output, sources)
                self.assertEqual(set(output.iterdir()), set(originals))
                for path, content in originals.items():
                    self.assertEqual(path.read_bytes(), content)
                # A failed publication releases its lock, allowing an immediate retry.
                output_commit.publish_pdf_pair(output, sources)

    def test_invalid_second_pdf_leaves_previous_pair_untouched(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sources = self.sources(root)
            output = root / "output"
            targets = output_commit.publish_pdf_pair(output, sources)
            originals = [target.read_bytes() for target in targets]
            sources[1].write_bytes(b"%PDF-1.7\ninvalid")
            with self.assertLogs("pypdf", level="WARNING"), self.assertRaises(PdfReadError):
                output_commit.publish_pdf_pair(output, sources)
            self.assertEqual([target.read_bytes() for target in targets], originals)
            self.assertEqual(set(output.iterdir()), set(targets))

    def test_failed_rollback_preserves_recovery_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sources = self.sources(root)
            output = root / "output"
            targets = output_commit.publish_pdf_pair(output, sources)
            original = targets[0].read_bytes()
            write_pdf(sources[0], version="changed")
            replace = os.replace

            def fail_second_and_restore(source, target):
                if Path(source).name in {"1.pdf", "0.previous.pdf"}:
                    raise OSError("replacement blocked")
                replace(source, target)

            with mock.patch.object(output_commit.os, "replace", side_effect=fail_second_and_restore):
                with self.assertRaises(OSError) as caught:
                    output_commit.publish_pdf_pair(output, sources)
            recovery = list(output.glob(".report-staging-*/0.previous.pdf"))
            self.assertEqual(len(recovery), 1)
            self.assertEqual(recovery[0].read_bytes(), original)
            self.assertIn(str(recovery[0].parent), caught.exception.__notes__[0])

    def test_another_process_cannot_publish_while_date_directory_is_locked(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sources = self.sources(root)
            output = root / "output"
            output.mkdir()
            descriptor = os.open(output, os.O_RDONLY)
            try:
                fcntl.flock(descriptor, fcntl.LOCK_EX | fcntl.LOCK_NB)
                command = [sys.executable, "-c",
                           "from pathlib import Path; import sys; from gd_power_report.output_commit import publish_pdf_pair; "
                           "publish_pdf_pair(Path(sys.argv[1]), (Path(sys.argv[2]), Path(sys.argv[3])))",
                           str(output), *map(str, sources)]
                result = subprocess.run(command, env={**os.environ, "PYTHONPATH": str(REPORT_ROOT)},
                                        capture_output=True, text=True, timeout=15)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("同日期日报正在保存", result.stderr)
                self.assertEqual(list(output.iterdir()), [])
                # Independent report dates remain available.
                output_commit.publish_pdf_pair(root / "other-date", sources)
            finally:
                os.close(descriptor)
            output_commit.publish_pdf_pair(output, sources)
            self.assertFalse(list(root.rglob("*.lock")))


if __name__ == "__main__":
    unittest.main()
