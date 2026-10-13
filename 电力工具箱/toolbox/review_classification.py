"""Organize clearing exports in place; filenames give dates/markets, sheets give stations."""
from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path
from typing import Callable

from .page_support import clearing_file_date
from .tasks import CancellationToken


STATION_ALIASES = {
    "东莞": ("东莞", "谢岗"),
    "归湖": ("归湖", "华归", "潮安", "潮州"),
    "汕头": ("汕头",),
    "海门": ("海门",),
    "海风": ("海风", "海上风电", "广澳"),
    "鮀莲": ("鮀莲", "金平"),
}


def default_template_folder() -> Path:
    return Path.home() / "Desktop" / "日常工作" / "复盘" / "模板"


def _station_matches(text: str) -> set[str]:
    matches = {station for station, aliases in STATION_ALIASES.items()
               if any(alias in text for alias in aliases)}
    # The renewable company's own name also includes 汕头.
    if matches & {"海风", "鮀莲"}:
        matches.discard("汕头")
    return matches


def identify_station(path: Path) -> str:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("暂不支持自动识别 .xls，请先另存为 .xlsx")
    # Page construction only needs folder helpers, not the Excel parser.
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        matches: set[str] = set()
        for title in workbook.sheetnames:
            matches.update(_station_matches(title))
    finally:
        workbook.close()
    if len(matches) != 1:
        detail = "工作表涉及多个单元：" + "、".join(sorted(matches)) if matches else "工作表名中无明确单元标识"
        raise ValueError(detail)
    return matches.pop()


def _excel_files(folder: Path, excluded: set[Path]) -> list[Path]:
    return sorted(
        (path for path in folder.iterdir()
         if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}
         and not path.name.startswith("~$") and "汇总" not in path.stem
         and path.resolve() not in excluded),
        key=lambda path: path.name,
    )


def missing_review_stations(source: Path) -> tuple[str, ...]:
    present: set[str] = set()
    # Preserve the existing single-station summary input.
    own_station = _station_matches(source.name)
    if len(own_station) == 1 and _excel_files(source, set()):
        present.update(own_station)
    for folder in source.iterdir():
        if folder.is_dir() and folder.name in STATION_ALIASES and _excel_files(folder, set()):
            present.add(folder.name)
    return tuple(station for station in STATION_ALIASES if station not in present)


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _move_without_overwrite(source: Path, destination: Path) -> None:
    """Verify the copy before removing its source; never overwrite another file."""
    expected = _digest(source)
    created = False
    try:
        with source.open("rb") as reader, destination.open("xb") as writer:
            created = True
            shutil.copyfileobj(reader, writer)
        shutil.copystat(source, destination)
        if _digest(destination) != expected or _digest(source) != expected:
            raise OSError(f"移动校验失败，源文件可能已变化：{source.name}")
        source.unlink()
    except BaseException:
        if created:
            destination.unlink(missing_ok=True)
        raise


def prepare_review_folder(
    source: Path,
    *,
    logger: Callable[[str], None],
    token: CancellationToken,
    template: Path | None = None,
    excluded: tuple[Path, ...] = (),
) -> Path:
    source = source.resolve()
    template = template or default_template_folder()
    if source == template.resolve():
        raise ValueError("模板目录只用于参考，请选择存放出清文件的文件夹。")
    if template.is_dir():
        stations = {path.name for path in template.iterdir() if path.is_dir() and not path.name.startswith(".")}
        if stations != set(STATION_ALIASES):
            raise ValueError("模板目录应包含且仅包含：东莞、归湖、汕头、海门、海风、鮀莲。")

    assignments: list[tuple[Path, str, str]] = []
    errors = []
    excluded_paths = {path.resolve() for path in excluded}
    candidates = _excel_files(source, excluded_paths)
    # Normalize exports that were already sorted into station folders, too.
    for folder in sorted(source.iterdir()):
        if folder.is_dir() and not folder.is_symlink() and len(_station_matches(folder.name)) == 1:
            candidates.extend(_excel_files(folder, excluded_paths))
    for path in candidates:
        token.raise_if_cancelled()
        markets = [market for market in ("日前", "实时") if market in path.stem]
        if not markets:
            logger(f"跳过非日前/实时文件：{path.name}")
            continue
        run_date = clearing_file_date(path)
        if len(markets) != 1 or run_date is None:
            errors.append(f"{path.name}：文件名需包含一个有效日期和日前或实时类型")
            continue
        try:
            station = identify_station(path)
        except Exception as exc:
            errors.append(f"{path.name}：{exc}")
            continue
        base_name = f"{markets[0]}交易结果查询({run_date.isoformat()}).xlsx"
        # Keep an existing valid trailing counter so repeated runs are stable.
        valid_name = re.fullmatch(
            re.escape(base_name[:-5]) + r"(?:\s*\(\d+\)|\s*\d+)?\.xlsx", path.name,
        )
        normalized_name = path.name if valid_name else base_name
        assignments.append((path, station, normalized_name))
        logger(f"识别：{path.name} → {station}（{run_date}，{markets[0]}）")
    if errors:
        for error in errors:
            logger(f"无法归类：{error}")
        raise ValueError("自动归类已停止，请处理以下文件后重试：\n" + "\n".join(errors))

    single_station = bool(assignments) and _station_matches(source.name) == {assignments[0][1]} and all(
        station == assignments[0][1] for _, station, _ in assignments
    )
    if assignments:
        # Plan every destination first. Preserve all different same-name exports.
        planned: list[tuple[Path, Path]] = []
        reserved: set[Path] = set()
        for path, station, normalized_name in assignments:
            folder = source if single_station else source / station
            if folder.exists() and (not folder.is_dir() or folder.is_symlink()):
                raise ValueError(f"单元目录不可用：{folder}")
            destination = folder / normalized_name
            if destination == path:
                reserved.add(destination)
                continue
            index = 2
            while destination.exists() or destination in reserved:
                stem = re.sub(r"(?:\s*\(\d+\)|\s*\d+)$", "", Path(normalized_name).stem)
                destination = folder / f"{stem}({index}).xlsx"
                index += 1
            reserved.add(destination)
            planned.append((path, destination))
        if not single_station:
            for station in STATION_ALIASES:
                folder = source / station
                if folder.exists() and (not folder.is_dir() or folder.is_symlink()):
                    raise ValueError(f"单元目录不可用：{folder}")
            for station in STATION_ALIASES:
                (source / station).mkdir(exist_ok=True)
        moved: list[tuple[Path, Path]] = []
        try:
            for path, destination in planned:
                token.raise_if_cancelled()
                _move_without_overwrite(path, destination)
                moved.append((path, destination))
                logger(f"整理：{path.name} → {destination.parent.name}/{destination.name}")
        except BaseException:
            # Return already moved files to their original locations on failure.
            for original, destination in reversed(moved):
                try:
                    _move_without_overwrite(destination, original)
                except OSError as exc:
                    logger(f"归类回退未完成，文件保留在 {destination}：{exc}")
            raise
        if moved:
            logger(f"归类及命名整理完成：{source}（{len(moved)} 个文件；文件内容和日期保持不变）")
    missing = missing_review_stations(source)
    if missing:
        logger("提示：缺少单元数据：" + "、".join(missing) + "；继续处理已有数据。")
    return source
