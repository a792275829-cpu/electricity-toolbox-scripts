"""Compare identical input using both backends; all business output is temporary."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import statistics
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "500kV断面分析"))
from src.deduplicate_500kv_prices import process_workbook
from src.generate_500kv_blockage_map import _run_pipeline
from openpyxl import load_workbook


def snapshot(path):
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        return [(s.title, [[(c.value, c.data_type, c.number_format) for c in row]
                           for row in s.iter_rows()]) for s in book]
    finally:
        book.close()


def benchmark(input_path, topology, repeats):
    report = {"input": input_path.name, "input_bytes": input_path.stat().st_size,
              "repeats": repeats, "backends": {}}
    previous = os.environ.get("TOOLBOX_500KV_BACKEND")
    reference = None
    try:
        with tempfile.TemporaryDirectory(prefix="500kv-benchmark-") as tmp:
            root = Path(tmp)
            # Identical paths also allow an exact comparison of HTML metadata.
            for mode in ("python", "rust"):
                os.environ["TOOLBOX_500KV_BACKEND"] = mode
                durations, pipeline_durations = [], []
                for _ in range(repeats):
                    cleaned = root / "cleaned.xlsx"
                    start = time.perf_counter()
                    result = process_workbook(input_path, cleaned, topology_path=topology)
                    durations.append(time.perf_counter() - start)
                    cells = snapshot(cleaned)
                    start = time.perf_counter()
                    pipeline = _run_pipeline(topology, input_path, root, 80.0,
                                             root / "map.html", cleaned)
                    pipeline_durations.append(time.perf_counter() - start)
                    html = (root / "map.html").read_bytes()
                    counts = dict(kept_rows=result.kept_rows, removed_rows=result.removed_rows,
                                  processed_sheets=result.processed_sheets,
                                  flagged_edges=pipeline.flagged_edges,
                                  flagged_intervals=pipeline.flagged_intervals)
                    current = (cells, html, counts)
                    if reference is None:
                        reference = current
                    elif current != reference:
                        raise AssertionError(f"{mode}: workbook/HTML/count parity failed")
                report["backends"][mode] = dict(
                    dedupe_seconds=durations, pipeline_seconds=pipeline_durations,
                    dedupe_median=statistics.median(durations),
                    pipeline_median=statistics.median(pipeline_durations), **counts)
            report["parity"] = "all cell values/types/formats, counts and HTML bytes identical"
            report["html_sha256"] = hashlib.sha256(reference[1]).hexdigest()
            for stage in ("dedupe", "pipeline"):
                report[stage + "_speedup"] = (report["backends"]["python"][stage + "_median"] /
                                             report["backends"]["rust"][stage + "_median"])
    finally:
        if previous is None:
            os.environ.pop("TOOLBOX_500KV_BACKEND", None)
        else:
            os.environ["TOOLBOX_500KV_BACKEND"] = previous
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--topology", type=Path, default=ROOT / "500kV断面分析/500kV节点拓扑.md")
    parser.add_argument("--repeats", type=int, default=3)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    report = benchmark(args.input.resolve(), args.topology.resolve(), args.repeats)
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
