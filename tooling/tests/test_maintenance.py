from dataclasses import FrozenInstanceError
from datetime import date, datetime
import importlib.util
import json
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))


class MaintenanceTests(unittest.TestCase):
    def test_every_source_test_folder_is_registered(self):
        # Include untracked source tests during development, but never caches,
        # private data or bundled copies excluded by Git.
        spec = importlib.util.spec_from_file_location('suite_inventory', ROOT/'tooling/run_tests.py')
        runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runner)
        names = subprocess.check_output(
            ['git', 'ls-files', '--cached', '--others', '--exclude-standard', '-z'], cwd=ROOT).decode().split('\0')
        unregistered = [name for name in names if Path(name).match('test_*.py')
                        and str(Path(name).parent) not in runner.SUITES.values()]
        self.assertEqual(unregistered, [], 'Register new tool suites in tooling/run_tests.py')

    def test_market_contract_snapshot_and_leap_year_dates(self):
        from toolbox_market_contracts import MARKET_SERIES
        from openpyxl.utils import column_index_from_string
        snapshot = json.loads((ROOT / "tooling/fixtures/market_contract.json").read_text())
        actual = {k: dict(source_days=v.source_days, sheet=v.sheet,
                         columns=[v.first_column, v.last_column],
                         fields=json.loads(json.dumps(v.fields)),
                         percent_fields=sorted(v.percent_fields)) for k, v in MARKET_SERIES.items()}
        self.assertEqual(actual, snapshot)
        self.assertEqual(MARKET_SERIES['actual_load'].source_date(date(2024, 3, 2)), date(2024, 2, 29))
        self.assertEqual(MARKET_SERIES['generation_settlement'].source_date(date(2026, 1, 3)), date(2025, 12, 28))
        for item in MARKET_SERIES.values():
            if item.fields:
                self.assertEqual(len(item.fields), column_index_from_string(item.last_column)
                                 - column_index_from_string(item.first_column) + 1)
        with self.assertRaises(FrozenInstanceError):
            MARKET_SERIES['actual_load'].source_days = 0
        with self.assertRaises(TypeError):
            MARKET_SERIES['new'] = None

    def test_wps_values_import_without_gui_and_contracts(self):
        from toolbox_wps_values import (normalize_cell_value, compare_matrix,
                                       range_shape, rects_overlap, target_range_from_start)
        self.assertEqual(normalize_cell_value(datetime(2026, 1, 1, 12)), '2026-01-01 12:00:00')
        self.assertEqual(compare_matrix([[0, None, .125]], [['0', '', '.125']], 1, 3), [])
        self.assertTrue(compare_matrix([[1]], [[2]], 1, 1))
        self.assertEqual(range_shape('Z2', 'AA3'), (2, 2))
        self.assertEqual(target_range_from_start('Z2', 2, 2), 'Z2:AA3')
        self.assertFalse(rects_overlap((1, 1, 2, 2), (3, 1, 4, 2)))

    def test_batch_preflight_fail_fast_and_partial_completion(self):
        from toolbox_trade_batch import iter_batch_reports
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            output = root / 'output'
            calls = []

            def render(source, *, out_path):
                calls.append(source.name)
                if source.stem == 'bad':
                    raise ValueError('fixture failure')
                out_path.write_text(source.name)
                return out_path

            with self.assertRaisesRegex(ValueError, '冲突'):
                list(iter_batch_reports([root/'a/day.xlsx', root/'b/DAY.xlsx'], output, render))
            self.assertFalse(output.exists())
            self.assertEqual(calls, [])
            with self.assertRaisesRegex(ValueError, 'fixture failure'):
                list(iter_batch_reports([root/'bad.xlsx', root/'good.xlsx'], output, render))
            self.assertEqual(calls, ['bad.xlsx'])
            results = list(iter_batch_reports([root/'bad.xlsx', root/'good.xlsx'], output,
                                              render, continue_on_error=True))
            self.assertIsInstance(results[0][3], ValueError)
            self.assertEqual(results[1][2].read_text(), 'good.xlsx')

    def test_private_openpyxl_adapter_matches_locked_version(self):
        import openpyxl
        from openpyxl.cell._writer import _set_attributes
        from openpyxl.worksheet._reader import WorkSheetParser
        from openpyxl.worksheet._writer import WorksheetWriter
        lock = (ROOT / 'packaging/tauri/requirements-macos.lock').read_text()
        self.assertIn(f'openpyxl=={openpyxl.__version__} ', lock)
        self.assertTrue(callable(_set_attributes))
        self.assertTrue(callable(WorksheetWriter.write_rows))
        self.assertTrue(callable(WorkSheetParser.parse))

    def test_packaged_runtime_includes_extracted_modules(self):
        sys.path.insert(0, str(ROOT/'packaging/tauri'))
        from prepare_payload import runtime_source
        names = ('toolbox_market_contracts.py', 'toolbox_wps_values.py',
                     'toolbox_trade_report.py', 'toolbox_province_excel.py',
                     'toolbox_trade_batch.py', '500kV断面分析/src/map_html_view.py',
                     '南方分公司电力市场日报/gd_power_report/clearing_table_view.py')
        for name in names:
            self.assertTrue(runtime_source(Path(name)), name)
        self.assertFalse(runtime_source(Path('tooling/fixtures/market_contract.json')))
        import os
        from build_support import export_source
        with tempfile.TemporaryDirectory() as temp:
            exported = Path(temp)
            export_source(exported, predicate=runtime_source)
            for name in names:
                self.assertTrue((exported/name).is_file(), name)
            check = subprocess.run([sys.executable, '-B', '-c',
                "import sys; from pathlib import Path; "
                "sys.path[:0]=['.', '电力工具箱', '500kV断面分析', '南方分公司电力市场日报']; "
                "import toolbox_trade_batch, toolbox_market_contracts, toolbox_wps_values; "
                "import toolbox_trade_report, toolbox_province_excel; "
                "from src.build_500kv_map_html import _render_html; "
                "from gd_power_report.company_clearing_xlsx import _display_value; "
                "assert _display_value(.125, '0.0%') == '12.5%'; "
                "assert '<!doctype html>' in _render_html({'metadata':{}}); "
                "assert Path(toolbox_trade_batch.__file__).is_relative_to(Path.cwd())"],
                cwd=exported, env=dict(os.environ, PYTHONPATH=str(exported)),
                capture_output=True, text=True)
            self.assertEqual(check.returncode, 0, check.stdout + check.stderr)

    def test_clearing_view_roundtrip_values_formats_and_merged_cells(self):
        sys.path.insert(0, str(ROOT/'南方分公司电力市场日报'))
        from gd_power_report import company_clearing_xlsx as legacy
        from gd_power_report import clearing_table_view as view
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font
        from openpyxl.writer.theme import theme_xml
        self.assertIs(legacy._table_model, view._table_model)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'presentation.xlsx'
            book = Workbook()
            book.loaded_theme = theme_xml
            sheet = book.active
            sheet['A1'] = .125
            sheet['A1'].number_format = '0.0%'
            sheet['B1'] = datetime(2026, 10, 27)
            sheet['B1'].number_format = 'yyyy-mm-dd'
            sheet['A2'] = 1234.565
            sheet['A2'].number_format = '#,##0.00'
            sheet['B2'] = '#DIV/0!'
            sheet['A3'] = '合并值'
            sheet.merge_cells('A3:B3')
            sheet['A3'].font = Font(bold=True, color='FF112233')
            before = view._table_model(sheet, 'A1:B3', view._theme_colors(book))
            book.save(path)
            book.close()
            reopened = load_workbook(path)
            try:
                result = view._table_model(reopened.active, 'A1:B3', view._theme_colors(reopened))
                self.assertEqual(result, before)
                self.assertEqual([cell['value'] for row in result['rows'] for cell in row['cells']],
                                 ['12.5%', '10月27日', '1,234.57', '—', '合并值'])
                self.assertEqual(result['rows'][2]['cells'][0]['colspan'], 2)
                self.assertEqual(reopened.active['A1'].value, .125)
                self.assertEqual(reopened.active['A1'].number_format, '0.0%')
            finally:
                reopened.close()

    def test_runner_isolates_modules_and_reports_failure_timeout_and_empty(self):
        spec = importlib.util.spec_from_file_location('maintenance_runner', ROOT/'tooling/run_tests.py')
        runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runner)
        with tempfile.TemporaryDirectory() as temp, patch('builtins.print'):
            folder = Path(temp)
            (folder/'test_a.py').write_text('import unittest, os\nos.environ["ISOLATION_FIXTURE"]="set"\nclass Test(unittest.TestCase):\n def test_ok(self): self.assertTrue(True)\n')
            (folder/'test_b.py').write_text('import unittest, os\nclass Test(unittest.TestCase):\n def test_isolated(self): self.assertNotIn("ISOLATION_FIXTURE",os.environ)\n')
            report = folder/'report.json'
            with patch.dict(runner.SUITES, fixture=str(folder)):
                self.assertEqual(runner.run(['fixture'], report=report), 0)
                self.assertEqual(len(json.loads(report.read_text())), 2)
                (folder/'test_a.py').write_text('raise RuntimeError("first meaningful error")\n')
                self.assertEqual(runner.run(['fixture'], pattern='test_a.py', report=report), 1)
                self.assertIn('first meaningful error', json.loads(report.read_text())[0]['output'])
                (folder/'test_a.py').write_text('import time\ntime.sleep(10)\n')
                self.assertEqual(runner.run(['fixture'], pattern='test_a.py', timeout=.2, report=report), 1)
                self.assertEqual(json.loads(report.read_text())[0]['returncode'], 124)
                (folder/'test_a.py').write_text('')
                self.assertEqual(runner.run(['fixture'], pattern='test_a.py'), 1)


if __name__ == '__main__':
    unittest.main()
