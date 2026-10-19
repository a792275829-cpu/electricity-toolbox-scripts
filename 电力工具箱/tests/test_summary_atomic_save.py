from __future__ import annotations

import unittest
from pathlib import Path
from types import SimpleNamespace
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.styles.numbers import NumberFormat


class SummaryAtomicSaveTests(unittest.TestCase):
    def load_summary_module(self):
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        return load_module("summary_atomic_save_test", paths.summary)

    def test_complete_workbook_replaces_existing_output(self) -> None:
        module = self.load_summary_module()
        workbook = module.Workbook()
        workbook.active.title = "汕头"
        workbook.active["A1"] = "verified"

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "summary.xlsx"
            output.write_bytes(b"old incomplete output")

            module.save_workbook_atomically(workbook, output)

            reopened = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(reopened.sheetnames, ["汕头"])
                self.assertEqual(reopened["汕头"]["A1"].value, "verified")
            finally:
                reopened.close()
            self.assertEqual(list(output.parent.glob(".*.tmp.xlsx")), [])

    def test_failed_save_preserves_existing_output(self) -> None:
        module = self.load_summary_module()

        class InterruptedWorkbook:
            @staticmethod
            def save(path) -> None:
                Path(path).write_bytes(b"PK\x03\x04partial")
                raise RuntimeError("interrupted")

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "summary.xlsx"
            output.write_bytes(b"previous valid output")

            with self.assertRaisesRegex(RuntimeError, "interrupted"):
                module.save_workbook_atomically(InterruptedWorkbook(), output)

            self.assertEqual(output.read_bytes(), b"previous valid output")
            self.assertEqual(list(output.parent.glob(".*.tmp.xlsx")), [])

    def test_number_format_object_is_normalized_before_save(self) -> None:
        module = self.load_summary_module()
        malformed_format = NumberFormat(numFmtId=176, formatCode="0.000")
        source_cell = SimpleNamespace(value=12.345, number_format=malformed_format)

        parsed_cell = module.read_numeric_cell(source_cell)
        self.assertEqual(parsed_cell.number_format, "0.000")

        values = [module.NumericCell(12.345, malformed_format)] + [
            module.NumericCell(0.0) for _ in module.TIME_SLOTS[1:]
        ]
        unit = module.UnitData(
            station_key="汕头",
            station_name="汕头",
            unit_name="1号机",
            values={"dayahead": {"energy": values, "price": values}},
        )

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "summary.xlsx"
            module.write_workbook([unit], output, lambda _message: None)

            reopened = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(reopened["汕头"]["B4"].value, 12.345)
                self.assertEqual(reopened["汕头"]["B4"].number_format, "0.000")
            finally:
                reopened.close()

    def test_corrupt_workbook_number_format_table_is_repaired_before_save(self) -> None:
        module = self.load_summary_module()
        workbook = module.Workbook()
        cell = workbook.active["A1"]
        cell.value = 12.345
        cell.number_format = "0.000"
        workbook._number_formats[0] = NumberFormat(
            numFmtId=176,
            formatCode="0.000",
        )

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "summary.xlsx"
            module.save_workbook_atomically(workbook, output)

            reopened = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(reopened.active["A1"].value, 12.345)
                self.assertEqual(reopened.active["A1"].number_format, "0.000")
            finally:
                reopened.close()


if __name__ == "__main__":
    unittest.main()
