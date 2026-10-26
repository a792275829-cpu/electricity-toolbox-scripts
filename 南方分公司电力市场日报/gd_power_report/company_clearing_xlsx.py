from __future__ import annotations

import re
from colorsys import hls_to_rgb, rgb_to_hls
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from math import isfinite
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

from gd_power_report.workbook_reader import load_workbook, reuse_workbook_inputs
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, range_boundaries


SHEET_NAME = "新快报"
ENERGY_METRICS_SHEET_NAME = "新"
REPORT_SHEET_ALIASES = ("新",)
FALLBACK_SHEET_PREFIX = "快报"
DAY_AHEAD_RANGE = "A6:Q66"
REALTIME_RANGE = "A68:Q84"
ERROR_VALUES = {"#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
WORKBOOK_DATE_RE = re.compile(r"分公司(?P<year>\d{4})\.(?P<month>\d{1,2})\.(?P<day>\d{1,2})")


def _looks_like_report_sheet(sheet: Any) -> bool:
    day_text = str(sheet["A4"].value or "")
    if "日前市场简析" not in day_text or int(sheet.max_column or 0) < 17:
        return False
    return any(
        isinstance(sheet.cell(row, 1).value, str)
        and "实时市场简析" in sheet.cell(row, 1).value
        for row in range(1, min(int(sheet.max_row or 1), 200) + 1)
    )


def _select_report_sheet(workbook: Any, preferred: str | None = None) -> Any:
    if preferred:
        if preferred not in workbook.sheetnames:
            raise ValueError(f"出清情况 Excel 缺少工作表：{preferred}")
        return workbook[preferred]
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]
    for name in REPORT_SHEET_ALIASES:
        if name in workbook.sheetnames and _looks_like_report_sheet(workbook[name]):
            return workbook[name]
    candidates = [
        workbook[name]
        for name in workbook.sheetnames
        if name.strip().startswith(FALLBACK_SHEET_PREFIX)
        and _looks_like_report_sheet(workbook[name])
    ]
    if not candidates:
        raise ValueError(
            f"出清情况 Excel 缺少可识别的日报工作表：{SHEET_NAME}或最新版快报"
        )

    def version_key(sheet: Any) -> tuple[int, int]:
        match = re.search(r"\((\d+)\)\s*$", sheet.title)
        version = int(match.group(1)) if match else 0
        return version, workbook.sheetnames.index(sheet.title)

    return max(candidates, key=version_key)


def _is_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() in ERROR_VALUES


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if _is_excel_error(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _meaningful_number(value: Any) -> float | None:
    number = _number(value)
    if number is None or abs(number) < 1e-9:
        return None
    return number


def _difference(current: float | None, comparison: float | None) -> float | None:
    if current is None or comparison is None:
        return None
    return current - comparison


def _day_ahead_energy_metrics(workbook: Any) -> tuple[float | None, float | None]:
    """Read day-ahead energy price and fee from the day recap company row."""

    sheet = next(
        (workbook[name] for name in (SHEET_NAME, ENERGY_METRICS_SHEET_NAME, "分公司")
         if name in workbook.sheetnames and _find_company_summary_row(workbook[name]) is not None),
        None,
    )
    if sheet is None:
        return None, None
    metrics = _day_ahead_summary_metrics(sheet)
    return metrics["long_term_energy_price"], metrics["electricity_fee"]


def _energy_comparison_metrics(
    workbook: Any,
    realtime_day_ahead_workbook: Any | None = None,
) -> dict[str, float | None]:
    """Read energy metrics, using a separate same-date workbook for realtime deltas."""

    day_ahead_price, day_ahead_fee = _day_ahead_energy_metrics(workbook)
    realtime_price = None
    realtime_fee = None
    if ENERGY_METRICS_SHEET_NAME in workbook.sheetnames:
        sheet = workbook[ENERGY_METRICS_SHEET_NAME]
        realtime_energy_parts = [_number(sheet[cell].value) for cell in ("F72", "F81")]
        realtime_price_parts = [_number(sheet[cell].value) for cell in ("P72", "P81")]
        if all(energy is not None for energy in realtime_energy_parts):
            realtime_energy = 0.0
            realtime_amount = 0.0
            for energy, price in zip(realtime_energy_parts, realtime_price_parts):
                if abs(energy) < 1e-9:
                    continue
                if price is None:
                    break
                realtime_energy += energy
                realtime_amount += energy * price
            else:
                if abs(realtime_energy) >= 1e-9:
                    realtime_price = realtime_amount / realtime_energy
                    realtime_fee = realtime_amount / 1000
    realtime_day_ahead_price = None
    realtime_day_ahead_fee = None
    if realtime_day_ahead_workbook is not None:
        realtime_day_ahead_price, realtime_day_ahead_fee = (
            _day_ahead_energy_metrics(realtime_day_ahead_workbook)
        )

    return {
        "day_ahead_price": day_ahead_price,
        "day_ahead_fee": day_ahead_fee,
        "realtime_price": realtime_price,
        "realtime_fee": realtime_fee,
        "realtime_day_ahead_price": realtime_day_ahead_price,
        "realtime_day_ahead_fee": realtime_day_ahead_fee,
    }


def _workbook_energy_date(workbook: Any, source: Path) -> date | None:
    if ENERGY_METRICS_SHEET_NAME in workbook.sheetnames:
        detected = _as_date(workbook[ENERGY_METRICS_SHEET_NAME]["A6"].value)
        if detected is not None:
            return detected
    if "分公司" in workbook.sheetnames:
        detected = _as_date(workbook["分公司"]["A2"].value)
        if detected is not None:
            return detected
    match = WORKBOOK_DATE_RE.search(source.name)
    if not match:
        return None
    return date(int(match["year"]), int(match["month"]), int(match["day"]))


def _realtime_day_ahead_workbook(
    current_source: Path,
    realtime_date: date | None,
    preferred_path: str | Path | None,
) -> tuple[Any | None, Path | None]:
    """Find the same-date workbook for independent price and market metrics."""

    if realtime_date is None:
        return None, None
    candidates: list[Path] = []
    if preferred_path and str(preferred_path).strip():
        candidates.append(Path(preferred_path).expanduser().resolve())
    detected = find_company_clearing_workbook(realtime_date, current_source.parent)
    if detected is not None:
        candidates.append(detected.resolve())

    seen: set[Path] = set()
    for candidate in candidates:
        if candidate in seen or candidate == current_source or not candidate.is_file():
            continue
        seen.add(candidate)
        try:
            workbook = load_workbook(candidate, data_only=True, read_only=False)
        except Exception:
            continue
        if _workbook_energy_date(workbook, candidate) != realtime_date:
            continue
        # Missing price inputs must not discard independently valid market
        # data, such as this date's must-run capacity. Each metric validates
        # its own inputs after the dated workbook has been selected.
        return workbook, candidate
    return None, None


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                pass
    return None


def _date_label(value: date) -> str:
    weekdays = "一二三四五六日"
    return f"{value.month}月{value.day}日（周{weekdays[value.weekday()]}）"


def _section_text(sheet: Any, marker: str, preferred_coordinate: str) -> str:
    preferred = str(sheet[preferred_coordinate].value or "")
    if marker in preferred:
        return preferred
    for row in range(1, min(int(sheet.max_row or 1), 200) + 1):
        value = sheet.cell(row, 1).value
        if isinstance(value, str) and marker in value:
            return value
    return preferred


def _dated_day_narrative(workbook: Any, report_date: date) -> str | None:
    candidates: list[tuple[int, int, str]] = []
    for sheet_index, sheet in enumerate(workbook.worksheets):
        if not sheet.title.strip().startswith(FALLBACK_SHEET_PREFIX):
            continue
        for row in range(1, min(int(sheet.max_row or 1), 200) + 1):
            value = sheet.cell(row, 1).value
            if not isinstance(value, str) or (
                "日前简析" not in value and "日前市场简析" not in value
            ):
                continue
            nearby_dates = [
                _as_date(sheet.cell(candidate_row, 1).value)
                for candidate_row in range(max(1, row - 2), min(int(sheet.max_row or row), row + 2) + 1)
                if candidate_row != row
            ]
            if report_date in nearby_dates:
                candidates.append((sheet_index, row, value))
    if candidates:
        return max(candidates, key=lambda item: (item[0], item[1]))[2]
    return None


def _date_matched_day_narrative(workbook: Any, report_date: date, fallback_sheet: Any) -> str:
    return _dated_day_narrative(workbook, report_date) or _section_text(
        fallback_sheet, "日前市场简析", "A4"
    )


def _date_matched_realtime_narrative(
    workbook: Any,
    realtime_date: date,
    fallback_sheet: Any,
) -> str:
    candidates: list[tuple[int, int, int, str]] = []
    for sheet_index, sheet in enumerate(workbook.worksheets):
        if not sheet.title.strip().startswith(FALLBACK_SHEET_PREFIX):
            continue
        for row in range(1, min(int(sheet.max_row or 1), 200) + 1):
            value = sheet.cell(row, 1).value
            if not isinstance(value, str) or (
                "实时简析" not in value and "实时市场简析" not in value
            ):
                continue
            nearby_dates = [
                _as_date(sheet.cell(candidate_row, 1).value)
                for candidate_row in range(
                    max(1, row - 3),
                    min(int(sheet.max_row or row), row + 3) + 1,
                )
                if candidate_row != row
            ]
            if realtime_date not in nearby_dates:
                continue
            fresh_price_marker = int(
                "全天实时均价" in value or "全省实时均价" in value
            )
            candidates.append((fresh_price_marker, sheet_index, row, value))
    if candidates:
        return max(candidates, key=lambda item: item[:3])[3]
    return _section_text(fallback_sheet, "实时市场简析", "A68")


def _date_matched_company_summary(workbook: Any, report_date: date, fallback: str) -> str:
    if "分公司" not in workbook.sheetnames:
        return fallback
    sheet = workbook["分公司"]
    if _as_date(sheet["A2"].value) != report_date:
        return fallback
    for row in range(1, min(int(sheet.max_row or 1), 80) + 1):
        label = _normalized_label(sheet.cell(row, 1).value)
        if "整体出清情况" not in label:
            continue
        text = "\n".join(
            str(sheet.cell(row, column).value or "")
            for column in range(1, min(int(sheet.max_column or 1), 24) + 1)
        ).strip()
        if "日前：" in text:
            return text
    return fallback


def _display_value(value: Any, number_format: str = "General") -> str:
    if value is None or isinstance(value, MergedCell):
        return ""
    if _is_excel_error(value):
        return "—"
    if isinstance(value, (datetime, date)):
        title_formats = {
            "分公司日前出清复盘": "分公司日前出清复盘",
            "火电分时明细": "火电分时明细",
            "新能源分时明细": "新能源分时明细",
            "分公司实时出清复盘": "分公司实时出清复盘",
        }
        for marker, title in title_formats.items():
            if marker in number_format:
                return f"{title}（{value.year}/{value.month}/{value.day}）"
        return f"{value.month}月{value.day}日"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if abs(value) >= 1_000_000_000_000:
            return "—"
        if "%" in number_format:
            decimals_match = re.search(r"\.([0#]+)%", number_format)
            decimals = len(decimals_match.group(1)) if decimals_match else 0
            return f"{float(value) * 100:.{decimals}f}%"
        if abs(value) < 1e-10:
            return "0"
        format_section = str(number_format or "General").split(";", 1)[0]
        placeholder = re.search(r"([#0?][#0?,]*)(?:\.([0#?]+))?", format_section)
        if placeholder and str(number_format).lower() != "general":
            decimal_pattern = placeholder.group(2) or ""
            decimals = len(decimal_pattern)
            rounded = Decimal(str(value)).quantize(
                Decimal(1).scaleb(-decimals),
                rounding=ROUND_HALF_UP,
            )
            if rounded == 0:
                rounded = abs(rounded)
            use_grouping = "," in placeholder.group(1)
            text = f"{rounded:,.{decimals}f}" if use_grouping else f"{rounded:.{decimals}f}"
            optional_decimals = len(decimal_pattern) - len(decimal_pattern.rstrip("#"))
            if optional_decimals and "." in text:
                integer, fraction = text.split(".", 1)
                required = decimals - optional_decimals
                fraction = fraction.rstrip("0")
                if len(fraction) < required:
                    fraction = fraction.ljust(required, "0")
                text = f"{integer}.{fraction}" if fraction else integer
            return text
        if float(value).is_integer():
            return f"{int(value)}"
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def _theme_colors(workbook: Any) -> list[str]:
    if not workbook.loaded_theme:
        return []
    try:
        root = ElementTree.fromstring(workbook.loaded_theme)
    except (ElementTree.ParseError, TypeError):
        return []
    scheme = next((item for item in root.iter() if item.tag.endswith("}clrScheme")), None)
    if scheme is None:
        return []
    colors_by_name: dict[str, str] = {}
    for entry in scheme:
        color = next(iter(entry), None)
        name = entry.tag.rsplit("}", 1)[-1]
        raw = "000000" if color is None else color.attrib.get("lastClr") or color.attrib.get("val") or "000000"
        colors_by_name[name] = raw[-6:].upper()
    # Excel's theme index order is bg1, tx1, bg2, tx2, accents, links;
    # clrScheme stores the first four as dk1, lt1, dk2, lt2.
    index_order = (
        "lt1",
        "dk1",
        "lt2",
        "dk2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    )
    return [colors_by_name.get(name, "000000") for name in index_order]


def _apply_tint(rgb: str, tint: float) -> str:
    if not tint:
        return rgb.upper()
    red, green, blue = (int(rgb[index : index + 2], 16) / 255 for index in (0, 2, 4))
    hue, lightness, saturation = rgb_to_hls(red, green, blue)
    if tint < 0:
        lightness *= 1 + tint
    else:
        lightness = lightness * (1 - tint) + tint
    tinted = hls_to_rgb(hue, max(0, min(1, lightness)), saturation)
    return "".join(f"{round(channel * 255):02X}" for channel in tinted)


def _rgb(value: Any, theme_colors: list[str], *, auto: str | None = None) -> str | None:
    if value is None:
        return None
    color_type = getattr(value, "type", None)
    rgb = ""
    if color_type == "rgb":
        rgb = str(getattr(value, "rgb", ""))
    elif color_type == "theme":
        theme = getattr(value, "theme", None)
        if isinstance(theme, int) and 0 <= theme < len(theme_colors):
            rgb = theme_colors[theme]
    elif color_type == "indexed":
        indexed = getattr(value, "indexed", None)
        if isinstance(indexed, int) and 0 <= indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[indexed]
    elif color_type == "auto" or getattr(value, "auto", False):
        return auto
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) == 6:
        tint = float(getattr(value, "tint", 0) or 0)
        return f"#{_apply_tint(rgb, tint)}"
    return None


def _border_rule(side: Any, edge: str, theme_colors: list[str]) -> str | None:
    style = getattr(side, "style", None)
    if not style:
        return None
    width = "1px"
    line_style = "solid"
    if style in {"medium", "mediumDashed", "mediumDashDot", "mediumDashDotDot"}:
        width = "2px"
    elif style == "thick":
        width = "3px"
    elif style == "double":
        width, line_style = "3px", "double"
    elif style in {"dashed", "dashDot", "dashDotDot", "slantDashDot"}:
        line_style = "dashed"
    elif style in {"dotted", "hair"}:
        line_style = "dotted"
    color = _rgb(getattr(side, "color", None), theme_colors, auto="#000000") or "#000000"
    return f"border-{edge}:{width} {line_style} {color}"


def _cell_style(cell: Any, theme_colors: list[str]) -> str:
    rules: list[str] = []
    fill = None
    if getattr(cell.fill, "fill_type", None) == "solid":
        fill = _rgb(getattr(cell.fill, "fgColor", None), theme_colors)
    font_color = _rgb(getattr(cell.font, "color", None), theme_colors, auto="#000000")
    if fill:
        rules.append(f"background-color:{fill}")
    if font_color:
        rules.append(f"color:{font_color}")
    if cell.font.name:
        family = str(cell.font.name).replace("'", "\\'")
        rules.append(f"font-family:'{family}','PingFang SC','Microsoft YaHei',sans-serif")
    if cell.font.sz:
        source_size = float(cell.font.sz)
        screen_size = max(9, min(16, source_size * 2 / 3))
        print_size = max(1.9, min(4.6, source_size * 4.6 / 24))
        rules.append(f"--excel-font-size:{screen_size:.2f}px")
        rules.append(f"--excel-print-font-size:{print_size:.2f}pt")
    if cell.font.bold:
        rules.append("font-weight:700")
    if cell.font.italic:
        rules.append("font-style:italic")
    decorations = []
    if cell.font.underline:
        decorations.append("underline")
    if cell.font.strike:
        decorations.append("line-through")
    if decorations:
        rules.append(f"text-decoration:{' '.join(decorations)}")
    horizontal = cell.alignment.horizontal
    if horizontal in {"left", "center", "right", "justify"}:
        rules.append(f"text-align:{horizontal}")
    if cell.alignment.vertical in {"top", "center", "bottom"}:
        vertical = "middle" if cell.alignment.vertical == "center" else cell.alignment.vertical
        rules.append(f"vertical-align:{vertical}")
    if cell.alignment.wrap_text:
        rules.append("white-space:normal")
    if cell.alignment.indent:
        rules.append(f"padding-left:{3 + int(cell.alignment.indent) * 9}px")
    for edge in ("left", "right", "top", "bottom"):
        rule = _border_rule(getattr(cell.border, edge), edge, theme_colors)
        if rule:
            rules.append(rule)
    if (
        isinstance(cell.value, (int, float))
        and not isinstance(cell.value, bool)
        and cell.value < 0
        and "[red]" in str(cell.number_format).lower()
    ):
        rules.append("color:#FF0000")
    return ";".join(rules)


def _table_model(sheet: Any, cell_range: str, theme_colors: list[str]) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    merge_by_anchor: dict[tuple[int, int], tuple[int, int]] = {}
    merged_children: set[tuple[int, int]] = set()
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row >= min_row
            and merged.max_row <= max_row
            and merged.min_col >= min_col
            and merged.max_col <= max_col
        ):
            anchor = (merged.min_row, merged.min_col)
            merge_by_anchor[anchor] = (
                merged.max_row - merged.min_row + 1,
                merged.max_col - merged.min_col + 1,
            )
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    if (row, column) != anchor:
                        merged_children.add((row, column))

    rows: list[dict[str, Any]] = []
    for row in range(min_row, max_row + 1):
        if sheet.row_dimensions[row].hidden:
            continue
        cells: list[dict[str, Any]] = []
        for column in range(min_col, max_col + 1):
            if (row, column) in merged_children:
                continue
            cell = sheet.cell(row, column)
            rowspan, colspan = merge_by_anchor.get((row, column), (1, 1))
            cells.append(
                {
                    "value": _display_value(cell.value, cell.number_format),
                    "rowspan": rowspan,
                    "colspan": colspan,
                    "style": _cell_style(cell, theme_colors),
                }
            )
        height = sheet.row_dimensions[row].height
        rows.append({"cells": cells, "height": round(float(height), 1) if height else None})

    widths = []
    for column in range(min_col, max_col + 1):
        width = sheet.column_dimensions[get_column_letter(column)].width or 9
        widths.append(max(64, min(210, round(float(width) * 7 + 10))))
    return {
        "range": cell_range,
        "rows": rows,
        "column_widths": widths,
        "min_width": sum(widths),
    }


def _workbook_source(path: str | Path, label: str = "出清情况 Excel") -> Path:
    source = Path(path).expanduser().resolve()
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError(f"{label} 不存在或格式不正确：{source}")
    return source


@reuse_workbook_inputs
def validate_report_workbook(path: str | Path) -> Path:
    source = _workbook_source(path)
    workbook = load_workbook(source, data_only=True, read_only=False)
    _select_report_sheet(workbook)
    return source


@reuse_workbook_inputs
def infer_report_date(path: str | Path) -> date:
    source = _workbook_source(path)
    workbook = load_workbook(source, data_only=True, read_only=False)
    sheet = _select_report_sheet(workbook)
    for coordinate in ("A6", "A2", "A7"):
        detected = _as_date(sheet[coordinate].value)
        if detected:
            return detected
    match = WORKBOOK_DATE_RE.search(source.name)
    if match:
        return date(int(match["year"]), int(match["month"]), int(match["day"]))
    raise ValueError("无法从“新快报”工作表识别报告日期。")


@reuse_workbook_inputs
def validate_day_ahead_comparison_workbook(path: str | Path) -> Path:
    """Validate the D workbook used for company day-ahead comparisons.

    Older clearing workbooks do not always contain the narrative quick-report
    sheet, so the stable ``分公司`` worksheet is the contract for this input.
    """

    source = _workbook_source(path, "D 日前对比表")
    workbook = load_workbook(source, data_only=True, read_only=False)
    if "分公司" not in workbook.sheetnames:
        raise ValueError("D 日前对比表缺少“分公司”工作表。")
    return source


@reuse_workbook_inputs
def infer_day_ahead_comparison_date(path: str | Path) -> date:
    source = _workbook_source(path, "D 日前对比表")
    workbook = load_workbook(source, data_only=True, read_only=False)
    if "分公司" not in workbook.sheetnames:
        raise ValueError("D 日前对比表缺少“分公司”工作表。")
    candidates = [
        ("分公司", "A2"),
        ("日报", "A1"),
        ("日前", "A1"),
        ("快报 ", "G2"),
        ("快报  (日前)", "G2"),
    ]
    for sheet_name, coordinate in candidates:
        if sheet_name in workbook.sheetnames:
            detected = _as_date(workbook[sheet_name][coordinate].value)
            if detected:
                return detected
    match = WORKBOOK_DATE_RE.search(source.name)
    if match:
        return date(int(match["year"]), int(match["month"]), int(match["day"]))
    raise ValueError("无法从 D 日前对比表识别日期。")


def _normalized_label(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _find_header_column(sheet: Any, marker: str, *, max_row: int = 12, max_column: int = 24) -> int | None:
    target = _normalized_label(marker)
    for row in range(1, min(int(sheet.max_row or max_row), max_row) + 1):
        for column in range(1, min(int(sheet.max_column or max_column), max_column) + 1):
            if target in _normalized_label(sheet.cell(row, column).value):
                return column
    return None


def _find_company_summary_row(sheet: Any) -> int | None:
    for row in range(1, min(int(sheet.max_row or 1), 80) + 1):
        label = _normalized_label(sheet.cell(row, 1).value)
        if label == "新能源":
            return None
        if label in {"分公司火电", "分公司", "火电机组合计"}:
            return row
    return None


def _merged_header_value(sheet: Any, row: int, column: int) -> Any:
    """Return a header value, including the anchor of a merged header cell."""

    cell = sheet.cell(row, column)
    if not isinstance(cell, MergedCell):
        return cell.value
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= column <= merged.max_col
        ):
            return sheet.cell(merged.min_row, merged.min_col).value
    return cell.value


def _find_header_column_by_hierarchy(
    sheet: Any,
    parent: str,
    children: tuple[str, ...],
    *,
    before_row: int,
    after_row: int = 2,
    exact_parent: bool = False,
) -> int | None:
    """Find a metric column from its merged parent and child header labels."""

    parent_label = _normalized_label(parent)
    child_labels = {_normalized_label(child) for child in children}
    max_column = min(int(sheet.max_column or 1), 80)
    max_row = min(max(before_row - 1, 1), int(sheet.max_row or 1))
    for row in range(max(2, after_row), max_row + 1):
        for column in range(1, max_column + 1):
            child = _normalized_label(sheet.cell(row, column).value)
            if child not in child_labels:
                continue
            for parent_row in range(row - 1, max(0, row - 4), -1):
                header = _normalized_label(_merged_header_value(sheet, parent_row, column))
                matches_parent = header == parent_label if exact_parent else parent_label in header
                if matches_parent:
                    return column
    return None


def _find_standalone_header_column(
    sheet: Any,
    label: str,
    *,
    before_row: int,
    lookback: int = 3,
    standalone_only: bool = False,
) -> int | None:
    """Find a non-hierarchical header immediately above a recap row."""

    target = _normalized_label(label)
    max_column = min(int(sheet.max_column or 1), 80)
    for row in range(max(1, before_row - lookback), before_row):
        for column in range(1, max_column + 1):
            if _normalized_label(_merged_header_value(sheet, row, column)) == target:
                if standalone_only and row > 1:
                    above = _merged_header_value(sheet, row - 1, column)
                    if (
                        above not in (None, "")
                        and not _normalized_label(above).startswith("单位")
                        and _as_date(above) is None
                    ):
                        continue
                return column
    return None


def _day_ahead_summary_metrics(sheet: Any) -> dict[str, Any]:
    """Read company day-ahead KPIs from labelled recap headers, not offsets."""

    summary_row = _find_company_summary_row(sheet)
    if summary_row is None:
        return {
            "energy": None,
            "price": None,
            "long_term_energy_price": None,
            "capacity_energy_price": None,
            "electricity_fee": None,
            "capacity_fee": None,
            "deviation": None,
            "margin": None,
            "long_term_energy": None,
            "long_term_margin": None,
            "day_total_energy_price": None,
            "day_total_capacity_price": None,
        }

    def value(
        parent: str, *children: str, meaningful: bool = False, exact_parent: bool = False
    ) -> float | None:
        column = _find_header_column_by_hierarchy(
            sheet,
            parent,
            children,
            before_row=summary_row,
            exact_parent=exact_parent,
        )
        if column is None:
            return None
        parser = _meaningful_number if meaningful else _number
        number = parser(sheet.cell(summary_row, column).value)
        if exact_parent and number is not None and not isfinite(number):
            return None
        return number

    energy = value("日前出清", "电量", meaningful=True)
    energy_price = value("日前出清", "电价", meaningful=True)
    deviation_energy = value("日前偏差", "电量")
    # All day-ahead price entry points use these fees and this recap row's
    # clearing energy. Explicit price caches and legacy quantity/price models
    # cannot substitute for a missing input.
    price_energy = value("日前出清", "电量", exact_parent=True)
    if price_energy == 0:
        price_energy = None
    electricity_fee = value("当日合计", "电能电费", "电能量电费", exact_parent=True)
    capacity_total_column = _find_header_column_by_hierarchy(
        sheet, "当日合计", ("容量电费",), before_row=summary_row, exact_parent=True
    )
    capacity_fee = value("当日合计", "容量电费", exact_parent=True)
    if capacity_total_column is None:
        standalone_capacity_column = _find_standalone_header_column(
            sheet, "容量电费", before_row=summary_row,
            lookback=summary_row - 1, standalone_only=True,
        )
        if standalone_capacity_column is not None:
            capacity_fee = _number(sheet.cell(summary_row, standalone_capacity_column).value)
            if capacity_fee is not None and not isfinite(capacity_fee):
                capacity_fee = None
    long_term_energy_price = (
        electricity_fee / price_energy * 1000
        if price_energy is not None and electricity_fee is not None else None
    )
    capacity_energy_price = (
        (electricity_fee + capacity_fee) / price_energy * 1000
        if price_energy is not None and electricity_fee is not None and capacity_fee is not None
        else None
    )
    if long_term_energy_price is not None and not isfinite(long_term_energy_price):
        long_term_energy_price = None
    if capacity_energy_price is not None and not isfinite(capacity_energy_price):
        capacity_energy_price = None

    return {
        "energy": energy,
        "price": energy_price,
        "long_term_energy_price": long_term_energy_price,
        "capacity_energy_price": capacity_energy_price,
        "electricity_fee": electricity_fee,
        "capacity_fee": capacity_fee,
        "deviation": deviation_energy,
        "margin": value("日前偏差", "边际贡献"),
        "long_term_energy": value("中长期", "电量", exact_parent=True),
        "long_term_margin": value("中长期", "边际贡献", "边际收益", exact_parent=True),
        "day_total_energy_price": long_term_energy_price,
        "day_total_capacity_price": capacity_energy_price,
    }


def _day_ahead_price_card(
    label: str, value: float | None, metrics: dict[str, Any], comparison_key: str
) -> dict[str, Any]:
    card = _card(label, value, "元/兆瓦时")
    card["price_basis"] = "day_fee"
    comparison_value = metrics[comparison_key]
    card["day_total_value"] = (
        comparison_value
        if comparison_value is not None and isfinite(comparison_value)
        else None
    )
    return card


def _company_day_ahead_text(sheet: Any) -> str:
    for row in range(1, min(int(sheet.max_row or 1), 80) + 1):
        if "整体出清情况" not in _normalized_label(sheet.cell(row, 1).value):
            continue
        text = "\n".join(
            str(sheet.cell(row, column).value or "")
            for column in range(1, min(int(sheet.max_column or 1), 24) + 1)
        )
        return text.split("实时：", 1)[0]
    return ""


def _legacy_quick_report_profit(workbook: Any) -> float | None:
    """Read thermal plus renewable company profit from legacy quick reports."""

    for sheet in reversed(workbook.worksheets):
        if not sheet.title.strip().startswith(FALLBACK_SHEET_PREFIX):
            continue
        profit_columns: list[int] = []
        company_rows: list[int] = []
        for row in range(1, min(int(sheet.max_row or 1), 30) + 1):
            if _normalized_label(sheet.cell(row, 1).value) == "分公司":
                company_rows.append(row)
            for column in range(1, min(int(sheet.max_column or 1), 60) + 1):
                if _normalized_label(sheet.cell(row, column).value) == "盈亏":
                    profit_columns.append(column)
        for column in reversed(sorted(set(profit_columns))):
            values = [
                value
                for row in company_rows
                if (value := _number(sheet.cell(row, column).value)) is not None
            ]
            if values:
                return sum(values)
    return None


def _day_ahead_company_profit(sheet: Any) -> float | None:
    """Return thermal plus renewable profit from the day-ahead recap."""

    day_range, _ = _appendix_ranges(sheet)
    _, start, _, end = range_boundaries(day_range)
    renewable_marker = next(
        (
            row
            for row in range(start, end + 1)
            if _normalized_label(sheet.cell(row, 1).value) == "新能源"
        ),
        None,
    )
    thermal_end = renewable_marker - 1 if renewable_marker is not None else end
    thermal_row = next(
        (
            row
            for row in range(start, thermal_end + 1)
            if _normalized_label(sheet.cell(row, 1).value) == "分公司"
        ),
        None,
    )
    thermal = _number(sheet.cell(thermal_row, 17).value) if thermal_row else None
    if renewable_marker is None:
        return thermal
    renewable_row = next(
        (
            row
            for row in range(renewable_marker + 1, end + 1)
            if _normalized_label(sheet.cell(row, 1).value) == "分公司"
        ),
        None,
    )
    renewable = _number(sheet.cell(renewable_row, 17).value) if renewable_row else None
    if thermal is None or renewable is None:
        return None
    return thermal + renewable


def _realtime_company_profit(sheet: Any, thermal_row: int) -> float | None:
    """Return thermal plus renewable profit from the realtime recap."""

    _, realtime_range = _appendix_ranges(sheet)
    _, start, _, end = range_boundaries(realtime_range)
    thermal = _number(sheet.cell(thermal_row, 15).value)
    renewable_row = next(
        (
            row
            for row in range(start, end + 1)
            if _normalized_label(sheet.cell(row, 1).value) == "分公司新能源"
        ),
        None,
    )
    if renewable_row is None:
        return thermal
    renewable = _number(sheet.cell(renewable_row, 15).value)
    if thermal is None or renewable is None:
        return None
    return thermal + renewable


def _comparison_day_ahead_profit(workbook: Any) -> float | None:
    for name in (SHEET_NAME, *REPORT_SHEET_ALIASES):
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        day_range, _ = _appendix_ranges(sheet)
        _, start, _, end = range_boundaries(day_range)
        has_renewable = any(
            _normalized_label(sheet.cell(row, 1).value) == "新能源"
            for row in range(start, end + 1)
        )
        value = _day_ahead_company_profit(sheet)
        if has_renewable or value is not None:
            return value
    return _legacy_quick_report_profit(workbook)


@reuse_workbook_inputs
def parse_day_ahead_comparison_workbook(path: str | Path) -> dict[str, Any]:
    """Read D company day-ahead summary metrics for D+1 comparisons."""

    source = _workbook_source(path, "D 日前对比表")
    workbook = load_workbook(source, data_only=True, read_only=False)
    if "分公司" not in workbook.sheetnames:
        raise ValueError("D 日前对比表缺少“分公司”工作表。")
    sheet = workbook["分公司"]
    summary_row = _find_company_summary_row(sheet)
    if summary_row is None:
        raise ValueError("D 日前对比表的“分公司”工作表缺少分公司汇总行。")

    recap_sheet = next(
        (
            workbook[name]
            for name in (SHEET_NAME, *REPORT_SHEET_ALIASES)
            if name in workbook.sheetnames
            and _find_company_summary_row(workbook[name]) is not None
        ),
        sheet,
    )
    comparison_metrics = _day_ahead_summary_metrics(sheet)
    recap_metrics = _day_ahead_summary_metrics(recap_sheet)
    energy = comparison_metrics["energy"] or recap_metrics["energy"]
    capacity_energy_price = recap_metrics["capacity_energy_price"]
    report_date = infer_day_ahead_comparison_date(source)
    return {
        "report_date": report_date.isoformat(),
        "report_date_label": _date_label(report_date),
        "source_path": str(source),
        "operations": _operations(workbook),
        "cards": [
            _card("日前出清电量", energy, "万千瓦时"),
            _card(
                "日前加权均价",
                comparison_metrics["price"] or recap_metrics["price"],
                "元/兆瓦时",
            ),
            _day_ahead_price_card(
                "含中长期电能电价",
                recap_metrics["long_term_energy_price"],
                recap_metrics,
                "day_total_energy_price",
            ),
            _day_ahead_price_card(
                "含容量综合电价", capacity_energy_price, recap_metrics, "day_total_capacity_price"
            ),
            _card(
                "较中长期偏差",
                comparison_metrics["deviation"] or recap_metrics["deviation"],
                "万千瓦时",
            ),
            _card("日前边际收益", recap_metrics["margin"], "万元"),
            _card("中长期电量", recap_metrics["long_term_energy"], "万千瓦时"),
            _card("中长期边际收益", recap_metrics["long_term_margin"], "万元"),
        ],
    }


def find_company_clearing_workbook(
    analysis_date: date,
    search_dir: str | Path | None = None,
) -> Path | None:
    directory = Path(search_dir) if search_dir else Path.home() / "Downloads"
    if not directory.is_dir():
        return None
    pattern = f"出清情况（分公司{analysis_date.year}.{analysis_date.month}.{analysis_date.day}*.xlsx"
    candidates = [
        path
        for path in directory.glob(pattern)
        if path.is_file() and not path.name.startswith("~$")
    ]
    return max(candidates, key=lambda path: (path.stat().st_mtime, path.name)) if candidates else None


def find_previous_workbook(path: str | Path, report_date: date | None = None) -> Path | None:
    source = Path(path).resolve()
    target_date = report_date or infer_report_date(source)
    candidates: list[tuple[date, Path]] = []
    for candidate in source.parent.glob("*.xlsx"):
        if candidate == source or candidate.name.startswith("~$"):
            continue
        match = WORKBOOK_DATE_RE.search(candidate.name)
        if not match:
            continue
        candidate_date = date(int(match["year"]), int(match["month"]), int(match["day"]))
        if candidate_date < target_date:
            candidates.append((candidate_date, candidate))
    if not candidates:
        return None
    latest_date = max(item[0] for item in candidates)
    same_day = [item[1] for item in candidates if item[0] == latest_date]
    return max(same_day, key=lambda item: (item.stat().st_mtime, item.name))


def _operations(workbook: Any) -> dict[str, str]:
    if "基础信息" not in workbook.sheetnames:
        return {}
    sheet = workbook["基础信息"]
    plant_names = {
        "汕头": "汕头厂",
        "汕头厂": "汕头厂",
        "海门": "海门厂",
        "海门厂": "海门厂",
        "东莞": "东莞厂",
        "东莞厂": "东莞厂",
        "清远": "清远厂",
        "清远厂": "清远厂",
    }
    groups: dict[str, list[tuple[str, str]]] = {plant: [] for plant in set(plant_names.values())}
    current_plant: str | None = None
    set_numbers = "一二三四五六七八九十"
    for row in range(1, min(int(sheet.max_row or 1), 150) + 1):
        plant_label = str(sheet.cell(row, 1).value or "").strip()
        if plant_label in plant_names:
            current_plant = plant_names[plant_label]
            continue
        if plant_label:
            current_plant = None
        if current_plant is None:
            continue
        unit_label = str(sheet.cell(row, 2).value or "").strip()
        status = str(sheet.cell(row, 3).value or "").strip()
        if not unit_label or unit_label in {"合计", "约束情况"} or not status:
            continue
        if current_plant in {"东莞厂", "清远厂"}:
            index = len(groups[current_plant])
            unit = set_numbers[index] if index < len(set_numbers) else str(index + 1)
        else:
            unit_match = re.search(r"#\s*\d+", unit_label)
            unit = unit_match.group(0).replace(" ", "") if unit_match else unit_label.replace("机组", "")
        groups[current_plant].append((unit, status))

    legacy_groups = {
        "汕头厂": [(15, "#1"), (16, "#2"), (17, "#3")],
        "海门厂": [(20, "#1"), (21, "#2"), (22, "#3"), (23, "#4")],
        "东莞厂": [(26, "一"), (27, "二"), (28, "三"), (29, "四")],
        "清远厂": [(32, "一"), (33, "二")],
    }
    for plant, units in legacy_groups.items():
        if groups[plant]:
            continue
        for row, unit in units:
            status = str(sheet.cell(row, 3).value or "").strip()
            if status:
                groups[plant].append((unit, status))

    result: dict[str, str] = {}
    for plant, units in groups.items():
        units_by_status: dict[str, list[str]] = {}
        for unit, status in units:
            units_by_status.setdefault(status, []).append(unit)
        if units_by_status:
            unit_suffix = "套机组" if plant in {"东莞厂", "清远厂"} else "机组"
            pieces = [
                f"{'、'.join(status_units)}{unit_suffix}{status}"
                for status, status_units in units_by_status.items()
            ]
            result[plant] = "；".join(pieces) + "。"
    return result


def _operation_text_section(value: Any) -> str:
    text = str(value or "")
    if "分公司机组运行方式" not in text:
        return ""
    return text.split("分公司机组运行方式", 1)[1].split("（三）", 1)[0]


def _parse_operation_lines(section: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for match in re.finditer(
        r"(?m)^\s*(汕头|海门|东莞|清远)(?:厂)?\s*[:：]\s*(.+?)\s*$",
        section,
    ):
        content = " ".join(match.group(2).strip().split())
        if content:
            result[f"{match.group(1)}厂"] = content.rstrip("。") + "。"
    return result


def _operations_from_text(value: Any) -> dict[str, str]:
    return _parse_operation_lines(_operation_text_section(value))


def _dated_operations_from_text(value: Any, reference_date: date) -> dict[date, dict[str, str]]:
    section = _operation_text_section(value)
    date_matches = list(
        re.finditer(
            r"(?m)^\s*(?:(\d{4})年)?(\d{1,2})月(\d{1,2})日\s*$",
            section,
        )
    )
    result: dict[date, dict[str, str]] = {}
    for index, match in enumerate(date_matches):
        year = int(match.group(1)) if match.group(1) else reference_date.year
        detected = date(year, int(match.group(2)), int(match.group(3)))
        if not match.group(1):
            if detected < reference_date - timedelta(days=180):
                detected = date(year + 1, detected.month, detected.day)
            elif detected > reference_date + timedelta(days=180):
                detected = date(year - 1, detected.month, detected.day)
        block_end = date_matches[index + 1].start() if index + 1 < len(date_matches) else len(section)
        operations = _parse_operation_lines(section[match.end() : block_end])
        if operations:
            result[detected] = operations
    return result


def _same_operation(left: str, right: str) -> bool:
    def normalized(value: str) -> str:
        return re.sub(r"\s+", "", value).replace("停备", "备用").rstrip("。")

    return normalized(left) == normalized(right)


def _operation_table(
    current_workbook: Any,
    current_date: date,
    previous_workbook: Any | None,
    previous_date: date | None,
    quick_sheet: Any | None = None,
) -> dict[str, Any]:
    quick = quick_sheet or _select_report_sheet(current_workbook)
    quick_text = _section_text(quick, "日前市场简析", "A4")
    dated_operations = _dated_operations_from_text(quick_text, current_date)
    workbook_operations = _operations(current_workbook)
    undated_operations = _operations_from_text(quick_text)
    current = (
        dated_operations.get(current_date)
        or workbook_operations
        or undated_operations
    )
    comparison_date = current_date - timedelta(days=1)
    comparison = dated_operations.get(comparison_date)
    if comparison is None and previous_workbook is not None and previous_date == comparison_date:
        comparison = _operations(previous_workbook)
    if comparison is None:
        comparison = {}
    dates = [_date_label(current_date), _date_label(comparison_date)]
    rows = []
    for plant in ("汕头厂", "海门厂", "东莞厂", "清远厂"):
        current_text = current.get(plant, workbook_operations.get(plant, "—"))
        comparison_text = comparison.get(plant)
        if comparison_text and current_text != "—":
            comparison_display = (
                "与前一日运行方式一致。"
                if _same_operation(current_text, comparison_text)
                else comparison_text
            )
        else:
            comparison_display = "无对比数据"
        statuses = [current_text, comparison_display]
        rows.append({"unit": plant, "statuses": statuses})
    return {"dates": dates, "rows": rows}


def _strategy_lines(value: Any) -> list[dict[str, str]]:
    text = str(value or "")
    if "各厂报价策略与收益" in text:
        text = text.split("各厂报价策略与收益", 1)[1]
    # Regional headings can be numbered or explicit colon headings. Plant names
    # in ordinary body sentences are not section boundaries.
    regional_title = r"粤东(?:区域(?:运行策略(?:研究)?)?|策略(?:研究)?)"
    title_end = r"(?=[ \t]*(?:[:：]|\r?$))"
    if re.search(rf"(?m)^[ \t]*(?:\d+[ \t]*[.、][ \t]*)?{regional_title}{title_end}", text):
        sections = list(re.finditer(
            rf"(?m)^[ \t]*(?P<numbered>\d+[ \t]*[.、][ \t]*)?"
            rf"(?P<plant>{regional_title}{title_end}|东莞(?:电厂|厂)?)[ \t]*"
            r"(?(numbered)[:：]?|(?:[:：]|(?=\r?$)))[ \t]*", text,
        ))
        result = []
        for index, section in enumerate(sections):
            end = sections[index + 1].start() if index + 1 < len(sections) else len(text)
            body = "\n".join(line.strip() for line in text[section.end():end].splitlines() if line.strip())
            if body:
                plant = section['plant'] if section['plant'].startswith('粤东') else '东莞厂'
                result.append({"plant": plant, "text": body})
        return result
    matches = list(
        re.finditer(
            r"(?m)^\s*(?:\d+\s*[.、]\s*)?(?P<lead>结合)?\s*"
            r"(?P<plant>汕头|海门|东莞)(?:电厂|厂)?\s*[:：]?\s*",
            text,
        )
    )
    strategies: list[dict[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        content = " ".join(text[match.end() : end].strip().split())
        if match.group("lead"):
            content = "结合" + content
        if content:
            strategies.append({"plant": f"{match.group('plant')}厂", "text": content})
    return strategies



def _sync_strategy_profits(
    strategies: list[dict[str, str]], sheet: Any, report_date: date
) -> list[dict[str, str]]:
    """Sync complete, canonically ordered profit lists from the dated recap.

    The caller must first verify the narrative date. Only the existing thermal
    fleet's complete unit sets are supported; partial or ambiguous data leaves
    the original sentence intact instead of assigning values by row position.
    """
    day_range, _ = _appendix_ranges(sheet)
    _, start, _, end = range_boundaries(day_range)
    if _as_date(sheet.cell(start, 1).value) != report_date:
        return strategies
    summary_row = _find_company_summary_row(sheet)
    if summary_row is None or summary_row <= start:
        return strategies
    if not any(
        "万元" in str(sheet.cell(row, 1).value or "")
        for row in range(start, summary_row)
    ):
        return strategies

    columns = set()
    for row in range(start, summary_row):
        for column in range(1, min(int(sheet.max_column or 1), 80) + 1):
            if _normalized_label(sheet.cell(row, column).value) != "边际贡献":
                continue
            if any(
                _normalized_label(_merged_header_value(sheet, parent_row, column))
                .replace("(", "（").replace(")", "）")
                in {"日前偏差", "日前偏差（较中长期）"}
                for parent_row in range(max(start, row - 3), row)
            ):
                columns.add(column)
    if len(columns) != 1:
        return strategies
    profit_column = columns.pop()
    unit_orders = {
        "汕头厂": ("#1机组", "#2机组", "#3机组"),
        "海门厂": ("#1机组", "#2机组", "#3机组", "#4机组"),
        "东莞厂": ("一套机组", "二套机组", "三套机组", "四套机组"),
    }
    amount_pattern = r"[+-]?\d+(?:\.\d+)?\s*万元"
    sentence_pattern = re.compile(
        r"(?:^|(?<=[。；;]))\s*各机组(?:日前(?:偏差)?)?收益(?:分别)?为[：:]?\s*"
        rf"(?P<amounts>{amount_pattern}(?:\s*[、，,与和及]\s*{amount_pattern})+)"
        r"(?=[。；;]|$)"
    )
    result = []
    for strategy in strategies:
        if strategy["plant"].startswith("粤东"):
            def sync_regional_line(match):
                plant = match.group("plant") + "厂"
                synced = _sync_strategy_profits(
                    [{"plant": plant, "text": match.group("profit")}], sheet, report_date
                )[0]["text"]
                return match.group("prefix") + synced
            regional_text = re.sub(
                r"(?m)^(?P<prefix>[ \t]*(?P<plant>汕头|海门)(?:电厂|厂)?)(?P<profit>各机组[^\n]+)",
                sync_regional_line, strategy["text"],
            )
            result.append({**strategy, "text": regional_text})
            continue
        expected_units = unit_orders.get(strategy["plant"])
        matches = list(sentence_pattern.finditer(strategy["text"]))
        if expected_units is None or len(matches) != 1:
            result.append(strategy)
            continue
        match = matches[0]
        amounts = list(re.finditer(r"[+-]?\d+(?:\.\d+)?", match["amounts"]))
        if len(amounts) != len(expected_units):
            result.append(strategy)
            continue
        values = {}
        valid = True
        for row in range(summary_row + 1, end + 1):
            plant = _normalized_label(_merged_header_value(sheet, row, 1))
            if plant not in {strategy["plant"], strategy["plant"].removesuffix("厂")}:
                continue
            unit = _normalized_label(sheet.cell(row, 2).value)
            if unit == "合计":
                continue
            value = sheet.cell(row, profit_column).value
            if (
                unit not in expected_units or unit in values
                or isinstance(value, bool) or not isinstance(value, (int, float))
                or not isfinite(value) or abs(value) >= 1_000_000_000_000
            ):
                valid = False
                break
            values[unit] = value
        if not valid or set(values) != set(expected_units):
            result.append(strategy)
            continue
        replacement = match["amounts"]
        for amount, unit in reversed(list(zip(amounts, expected_units))):
            replacement = (
                replacement[:amount.start()] + _display_value(values[unit], "0")
                + replacement[amount.end():]
            )
        result.append({
            **strategy,
            "text": strategy["text"][:match.start("amounts")] + replacement
            + strategy["text"][match.end("amounts"):],
        })
    return result


def _time_rows(sheet: Any) -> list[int]:
    rows_by_hour: dict[int, int] = {}
    for row in range(1, min(int(sheet.max_row or 1), 300) + 1):
        value = str(sheet.cell(row, 1).value or "").strip()
        match = re.fullmatch(r"([01]?\d|2[0-3]):[0-5]\d", value)
        if match:
            rows_by_hour.setdefault(int(match.group(1)), row)
    return [rows_by_hour[hour] for hour in sorted(rows_by_hour)]


def _series(sheet: Any, energy_column: int, price_column: int) -> dict[str, list[Any] | bool]:
    labels = [f"{hour}:00" for hour in range(24)]
    energy: list[float | None] = []
    price: list[float | None] = []
    row_by_hour = {
        int(str(sheet.cell(row, 1).value).split(":", 1)[0]): row
        for row in _time_rows(sheet)
    }
    for hour in range(24):
        row = row_by_hour.get(hour)
        energy.append(_number(sheet.cell(row, energy_column).value) if row else None)
        price.append(_number(sheet.cell(row, price_column).value) if row else None)
    ready = any(value is not None and abs(value) > 1e-9 for value in energy)
    if not ready:
        energy = [None] * len(labels)
        price = [None] * len(labels)
    return {"labels": labels, "energy": energy, "price": price, "ready": ready}


def _valid_values(values: Iterable[Any]) -> list[float]:
    return [number for value in values if (number := _meaningful_number(value)) is not None]


def _card(label: str, value: Any, unit: str, *, delta: Any = None) -> dict[str, Any]:
    number = _number(value)
    return {
        "label": label,
        "value": number,
        "unit": unit,
        "delta": _number(delta),
        "pending": number is None,
    }


def _must_run_capacity(workbook: Any | None) -> float | None:
    """Read the must-run capacity from the labelled basic-information row.

    The workbook stores capacity in MW while the report displays 万千瓦.
    Looking up the label instead of a fixed cell keeps the parser stable when
    the row moves in later daily workbooks.
    """

    if workbook is None or "基础信息" not in workbook.sheetnames:
        return None
    sheet = workbook["基础信息"]
    labels = {"必开总容量", "必开机组容量"}
    max_row = min(int(sheet.max_row or 1), 200)
    max_column = min(int(sheet.max_column or 1), 50)
    for row in range(1, max_row + 1):
        for column in range(1, max_column + 1):
            if _normalized_label(sheet.cell(row, column).value) not in labels:
                continue
            for value_column in range(column + 1, min(column + 3, max_column) + 1):
                capacity_mw = _number(sheet.cell(row, value_column).value)
                if capacity_mw is not None and capacity_mw >= 0:
                    return capacity_mw / 10
            return None
    return None


def _market_cards(
    current: Any,
    previous: Any | None,
    realtime: bool = False,
    must_run_workbook: Any | None = None,
) -> list[dict[str, Any]]:
    basic = current["基础信息"] if "基础信息" in current.sheetnames else None
    old_basic = previous["基础信息"] if previous and "基础信息" in previous.sheetnames else None
    daily = current["日报"] if "日报" in current.sheetnames else None
    maintenance = _number(basic["G5"].value) / 10 if basic and _number(basic["G5"].value) is not None else None
    must_run_capacity = _must_run_capacity(
        must_run_workbook if realtime else current
    )
    if realtime:
        peak = _number(daily["M5"].value) / 10 if daily and _number(daily["M5"].value) is not None else None
        valley = _number(daily["N5"].value) / 10 if daily and _number(daily["N5"].value) is not None else None
        starts = daily["D9"].value if daily else None
        return [
            _card("统调实际负荷峰值", peak, "万千瓦"),
            _card("统调实际负荷谷值", valley, "万千瓦"),
            _card("检修容量", maintenance, "万千瓦"),
            _card("市场开机", starts, "台"),
            _card("必开机组容量", must_run_capacity, "万千瓦"),
        ]
    peak = _number(basic["B5"].value) / 10 if basic and _number(basic["B5"].value) is not None else None
    valley = _number(basic["C5"].value) / 10 if basic and _number(basic["C5"].value) is not None else None
    old_peak = _number(old_basic["B5"].value) / 10 if old_basic and _number(old_basic["B5"].value) is not None else None
    old_valley = _number(old_basic["C5"].value) / 10 if old_basic and _number(old_basic["C5"].value) is not None else None
    starts = daily["D9"].value if daily else None
    return [
        _card("统调负荷峰值", peak, "万千瓦", delta=peak - old_peak if peak is not None and old_peak is not None else None),
        _card("统调负荷谷值", valley, "万千瓦", delta=valley - old_valley if valley is not None and old_valley is not None else None),
        _card("检修容量", maintenance, "万千瓦"),
        _card("市场开机", starts, "台"),
        _card("必开机组容量", must_run_capacity, "万千瓦"),
    ]


def _text_number(text: str, pattern: str) -> float | None:
    match = re.search(pattern, text)
    return float(match.group(1).replace(",", "")) if match else None


def _text_delta(text: str, pattern: str) -> float | None:
    match = re.search(pattern, text)
    if not match:
        return None
    number = float(match.group(2).replace(",", ""))
    return -number if match.group(1) in {"下降", "减少"} else number


def _market_average_price_card(text: str, *, realtime: bool = False) -> dict[str, Any]:
    if realtime:
        marker = r"(?:全省实时(?:发电侧)?(?:出清)?均价|全省日前发电侧出清均价|全天实时均价)"
        label = "全省实时均价"
    else:
        marker = r"(?:全省日前(?:发电侧)?(?:出清)?均价|全省日前均价)"
        label = "全省日前均价"
    unit = r"元/(?:兆瓦时|MWh)"
    value = _text_number(text, rf"{marker}(?:为)?\s*([\d,.]+)\s*{unit}")
    delta = None
    delta_prefix = None
    for prefix in (("较日前", "环比") if realtime else ("环比",)):
        delta = _text_delta(
            text,
            rf"{marker}(?:为)?\s*[\d,.]+\s*{unit}[^。；]*?{prefix}(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*{unit}",
        )
        if delta is not None:
            delta_prefix = prefix
            break
    card = _card(label, value, "元/兆瓦时", delta=delta)
    if delta_prefix is not None:
        card["delta_prefix"] = delta_prefix
    return card


def _market_cards_from_quick(
    workbook: Any,
    previous: Any | None,
    sheet: Any,
    realtime: bool = False,
    text_override: str | None = None,
    must_run_workbook: Any | None = None,
) -> list[dict[str, Any]]:
    fallback = _market_cards(
        workbook,
        previous,
        realtime=realtime,
        must_run_workbook=must_run_workbook,
    )
    fallback_by_label = {card["label"]: card for card in fallback}
    day_text = text_override or _section_text(sheet, "日前市场简析", "A4")
    if realtime:
        text = text_override or _section_text(sheet, "实时市场简析", "A68")
        peak = _text_number(text, r"全天(?:最大|峰值)(?:值)?(?:为|达到)?\s*([\d,.]+)\s*万千瓦")
        valley = _text_number(text, r"(?:全天)?(?:最小|谷值)(?:值)?(?:为|达到)?\s*([\d,.]+)\s*万千瓦")
        maintenance = _text_number(text, r"检修容量(?:为)?\s*([\d,.]+)\s*万千瓦")
        starts = _text_number(text, r"(?:全市场|市场)开机(?:台数)?(?:为)?\s*([\d,.]+)\s*台")
        peak_delta = _text_delta(
            text,
            r"全天(?:最大|峰值)(?:值)?(?:为|达到)?\s*[\d,.]+\s*万千瓦[^，。]*[，,]\s*较日前(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*万千瓦",
        )
        valley_delta = _text_delta(
            text,
            r"(?:全天)?(?:最小|谷值)(?:值)?(?:为|达到)?\s*[\d,.]+\s*万千瓦[^，。]*[，,]\s*较日前(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*万千瓦",
        )
        starts_delta = _text_delta(
            text,
            r"(?:全市场|市场)开机(?:台数)?(?:为)?\s*[\d,.]+\s*台[^，。]*[，,]\s*较日前(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*台",
        )
        values = {
            "统调实际负荷峰值": peak,
            "统调实际负荷谷值": valley,
            "检修容量": maintenance,
            "市场开机": starts,
        }
        deltas = {
            "统调实际负荷峰值": peak_delta,
            "统调实际负荷谷值": valley_delta,
            "市场开机": starts_delta,
        }
        cards = []
        for fallback_card in fallback:
            card = _card(
                fallback_card["label"],
                values.get(fallback_card["label"])
                if values.get(fallback_card["label"]) is not None
                else fallback_card["value"],
                fallback_card["unit"],
                delta=deltas.get(fallback_card["label"]),
            )
            if card["delta"] is not None:
                card["delta_prefix"] = "较日前"
            cards.append(card)
        cards.append(_market_average_price_card(text, realtime=True))
        return cards

    peak = _text_number(day_text, r"全天(?:最大|峰值)(?:值)?(?:为|达到)?\s*([\d,.]+)\s*万千瓦")
    valley = _text_number(day_text, r"(?:全天)?(?:最小|谷值)(?:值)?(?:为|达到)?\s*([\d,.]+)\s*万千瓦")
    maintenance = _text_number(day_text, r"检修容量(?:为)?\s*([\d,.]+)\s*万千瓦")
    starts = _text_number(day_text, r"(?:全市场|市场)开机(?:台数)?(?:为)?\s*([\d,.]+)\s*台")
    peak_delta = _text_delta(
        day_text,
        r"全天(?:最大|峰值)(?:值)?(?:为|达到)?\s*[\d,.]+\s*万千瓦[^，。]*[，,]\s*环比(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*万千瓦",
    )
    valley_delta = _text_delta(
        day_text,
        r"(?:全天)?(?:最小|谷值)(?:值)?(?:为|达到)?\s*[\d,.]+\s*万千瓦[^，。]*[，,]\s*环比(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*万千瓦",
    )
    starts_delta = _text_delta(
        day_text,
        r"(?:全市场|市场)开机(?:台数)?(?:为)?\s*[\d,.]+\s*台[^，。]*[，,]\s*环比(下降|上涨|上升|减少|增加)\s*([\d,.]+)\s*台",
    )
    specs = [
        ("统调负荷峰值", peak, "万千瓦", peak_delta),
        ("统调负荷谷值", valley, "万千瓦", valley_delta),
        ("检修容量", maintenance, "万千瓦", None),
        ("市场开机", starts, "台", starts_delta),
        (
            "必开机组容量",
            fallback_by_label["必开机组容量"]["value"],
            "万千瓦",
            None,
        ),
    ]
    cards = [
        _card(
            label,
            value if value is not None else fallback_by_label[label]["value"],
            unit,
            delta=delta if delta is not None else fallback_by_label[label].get("delta"),
        )
        for label, value, unit, delta in specs
    ]
    cards.append(_market_average_price_card(day_text))
    return cards


def _find_summary_row(sheet: Any, labels: set[str], start: int, end: int, value_column: int) -> int:
    for row in range(start, end + 1):
        label = str(sheet.cell(row, 1).value or "").strip()
        if label in labels and _number(sheet.cell(row, value_column).value) is not None:
            return row
    for row in range(start, end + 1):
        if _number(sheet.cell(row, value_column).value) is not None:
            return row
    return start


def _summary_cards(
    sheet: Any,
    series: dict[str, Any],
    realtime: bool = False,
    text_override: str | None = None,
    energy_metrics: dict[str, float | None] | None = None,
) -> list[dict[str, Any]]:
    energy_metrics = energy_metrics or {}
    if realtime:
        _, realtime_range = _appendix_ranges(sheet)
        _, start, _, end = range_boundaries(realtime_range)
        row = next(
            (
                candidate
                for candidate in range(start, end + 1)
                if _normalized_label(sheet.cell(candidate, 1).value)
                in {"分公司火电", "分公司"}
            ),
            None,
        )

        def value(parent: str, child: str, *, meaningful: bool = False) -> float | None:
            if row is None:
                return None
            column = _find_header_column_by_hierarchy(
                sheet,
                parent,
                (child,),
                before_row=row,
                after_row=start,
            )
            if column is None:
                return None
            return (_meaningful_number if meaningful else _number)(sheet.cell(row, column).value)

        def standalone(label: str) -> float | None:
            if row is None:
                return None
            column = _find_standalone_header_column(sheet, label, before_row=row)
            return _number(sheet.cell(row, column).value) if column is not None else None

        def standalone_any(*labels: str) -> float | None:
            for label in labels:
                candidate = standalone(label)
                if candidate is not None:
                    return candidate
            return None

        energy = value("实时", "电量", meaningful=True)
        price = value("实时", "电价", meaningful=True)
        day_ahead_price = value("日前", "电价", meaningful=True)
        price_delta = (
            price - day_ahead_price
            if price is not None and day_ahead_price is not None
            else None
        )
        return [
            _card("实时出清电量", energy, "万千瓦时"),
            _card("实时加权均价", price, "元/兆瓦时", delta=price_delta),
            _card(
                "电能电价",
                standalone_any(
                    "电能电价",
                    "电能均价",
                    "电能量均价（不含容量/机制）",
                ),
                "元/兆瓦时",
            ),
            _card("含容量综合电价", standalone("综合电价"), "元/兆瓦时"),
            _card("实时较日前偏差电量", value("实时偏差", "电量"), "万千瓦时"),
            _card("实时偏差收益", value("实时偏差", "收益"), "万元"),
            _card("实际上网偏差电量", value("实际上网偏差", "电量"), "万千瓦时"),
            _card("实际上网偏差收益", value("实际上网偏差", "收益"), "万元"),
        ]

    metrics = _day_ahead_summary_metrics(sheet)
    return [
        _card("日前出清电量", metrics["energy"], "万千瓦时"),
        _card("日前加权均价", metrics["price"], "元/兆瓦时"),
        _day_ahead_price_card(
            "含中长期电能电价", metrics["long_term_energy_price"], metrics, "day_total_energy_price"
        ),
        _day_ahead_price_card(
            "含容量综合电价", metrics["capacity_energy_price"], metrics, "day_total_capacity_price"
        ),
        _card("较中长期偏差", metrics["deviation"], "万千瓦时"),
        _card("日前边际收益", metrics["margin"], "万元"),
        _card("中长期电量", metrics["long_term_energy"], "万千瓦时"),
        _card("中长期边际收益", metrics["long_term_margin"], "万元"),
    ]


def _appendix_ranges(sheet: Any) -> tuple[str, str]:
    time_rows = _time_rows(sheet)
    first_time_row = min(time_rows) if time_rows else 68

    def title_row(marker: str, start: int, end: int) -> int | None:
        date_fallback = None
        for row in range(max(1, start), min(end, int(sheet.max_row or end)) + 1):
            cell = sheet.cell(row, 1)
            if cell.value not in (None, "") and marker in str(cell.number_format or ""):
                return row
            if date_fallback is None and _as_date(cell.value):
                date_fallback = row
        return date_fallback

    day_start = title_row("分公司日前出清复盘", 5, first_time_row - 1)
    if day_start is None:
        _, day_start, _, _ = range_boundaries(DAY_AHEAD_RANGE)
    detected_day_end = max(time_rows) if time_rows else day_start + 60
    if str(sheet.cell(detected_day_end + 1, 1).value or "").strip() == "合计":
        detected_day_end += 1
    day_end = max(day_start + 60, detected_day_end)

    realtime_text_row = None
    for row in range(day_end + 1, min(int(sheet.max_row or day_end + 60), day_end + 100) + 1):
        value = sheet.cell(row, 1).value
        if isinstance(value, str) and "实时市场简析" in value:
            realtime_text_row = row
            break
    realtime_start = title_row(
        "分公司实时出清复盘",
        (realtime_text_row + 1) if realtime_text_row else day_end + 1,
        (realtime_text_row + 12) if realtime_text_row else day_end + 20,
    )
    if realtime_start is None:
        _, realtime_start, _, _ = range_boundaries(REALTIME_RANGE)

    realtime_end = realtime_start + 16
    scan_limit = min(int(sheet.max_row or realtime_start + 40), realtime_start + 60)
    for row in range(realtime_start, scan_limit + 1):
        has_value = any(sheet.cell(row, column).value not in (None, "") for column in range(1, 18))
        if has_value:
            realtime_end = max(realtime_end, row)
        elif row >= realtime_start + 5:
            break
    return f"A{day_start}:Q{day_end}", f"A{realtime_start}:Q{realtime_end}"


def _date_in_range(sheet: Any, cell_range: str) -> date | None:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            detected = _as_date(sheet.cell(row, column).value)
            if detected is not None:
                return detected
    return None


@reuse_workbook_inputs
def parse_realtime_report_workbook_for_date(
    path: str | Path,
    target_date: date,
) -> dict[str, Any]:
    """Parse the report sheet whose realtime appendix matches ``target_date``.

    A source workbook can contain several numbered quick-report revisions.  The
    latest sheet is not always the one that owns the requested D-1 appendix, so
    matching the embedded appendix date is more reliable than matching a sheet
    name or file modification time.
    """

    source = Path(path).expanduser().resolve()
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError(f"D-1 实时输入表不存在或格式不正确：{source}")
    workbook = load_workbook(source, data_only=True, read_only=False)
    candidates = [
        workbook[name]
        for name in workbook.sheetnames
        if name == SHEET_NAME
        or (
            name.strip().startswith(FALLBACK_SHEET_PREFIX)
            and _looks_like_report_sheet(workbook[name])
        )
    ]
    matches: list[Any] = []
    detected_dates: list[date] = []
    for sheet in candidates:
        _, realtime_range = _appendix_ranges(sheet)
        detected = _date_in_range(sheet, realtime_range)
        if detected is not None:
            detected_dates.append(detected)
        if detected == target_date:
            matches.append(sheet)

    if not matches:
        labels = "、".join(sorted({value.isoformat() for value in detected_dates})) or "未识别"
        raise ValueError(
            f"D-1 实时输入表中未找到 {target_date.isoformat()} 的实时附表"
            f"（已识别：{labels}）。"
        )

    def version_key(sheet: Any) -> tuple[int, int]:
        match = re.search(r"\((\d+)\)\s*$", sheet.title)
        version = int(match.group(1)) if match else 0
        return version, workbook.sheetnames.index(sheet.title)

    selected = max(matches, key=version_key)
    return parse_daily_report_workbook(source, sheet_name=selected.title)


@reuse_workbook_inputs
def parse_daily_report_workbook(
    path: str | Path,
    previous_path: str | Path | None = None,
    sheet_name: str | None = None,
    realtime_day_ahead_path: str | Path | None = None,
) -> dict[str, Any]:
    source = _workbook_source(path)
    workbook = load_workbook(source, data_only=True, read_only=False)
    sheet = _select_report_sheet(workbook, sheet_name)
    theme_colors = _theme_colors(workbook)
    report_date = _as_date(sheet["A6"].value) or infer_report_date(source)

    previous_source = Path(previous_path).resolve() if previous_path else find_previous_workbook(source, report_date)
    previous_workbook = None
    previous_date = None
    if previous_source and previous_source.is_file():
        try:
            previous_workbook = load_workbook(previous_source, data_only=True, read_only=False)
            previous_sheet = _select_report_sheet(previous_workbook)
            previous_date = _as_date(previous_sheet["A6"].value)
            if previous_date is None:
                match = WORKBOOK_DATE_RE.search(previous_source.name)
                if match:
                    previous_date = date(
                        int(match["year"]), int(match["month"]), int(match["day"])
                    )
            if "基础信息" not in previous_workbook.sheetnames:
                previous_workbook = None
                previous_date = None
        except Exception:
            previous_workbook = None
            previous_date = None

    day_series = _series(sheet, 2, 3)
    realtime_series = {
        "labels": [f"{hour}:00" for hour in range(24)],
        "energy": [None] * 24,
        "price": [None] * 24,
        "ready": False,
    }
    market_previous = (
        previous_workbook
        if previous_date == report_date - timedelta(days=1)
        else None
    )
    day_ahead_appendix_range, realtime_appendix_range = _appendix_ranges(sheet)
    realtime_date = _date_in_range(sheet, realtime_appendix_range)
    realtime_comparison_workbook, realtime_comparison_source = (
        _realtime_day_ahead_workbook(
            source,
            realtime_date,
            realtime_day_ahead_path,
        )
    )
    energy_metrics = _energy_comparison_metrics(
        workbook,
        realtime_comparison_workbook,
    )
    dated_day_narrative = _dated_day_narrative(workbook, report_date)
    day_narrative_text = dated_day_narrative or _section_text(sheet, "日前市场简析", "A4")
    strategies = _strategy_lines(day_narrative_text)
    if dated_day_narrative is not None:
        strategies = _sync_strategy_profits(strategies, sheet, report_date)
    realtime_narrative_text = _date_matched_realtime_narrative(
        workbook,
        realtime_date or report_date,
        sheet,
    )
    day_summary_text = _date_matched_company_summary(
        workbook,
        report_date,
        day_narrative_text,
    )
    return {
        "report_date": report_date.isoformat(),
        "report_date_label": _date_label(report_date),
        "source_path": str(source),
        "source_sheet": sheet.title,
        "realtime_date": realtime_date.isoformat() if realtime_date else None,
        "realtime_energy_comparison_source_path": (
            str(realtime_comparison_source) if realtime_comparison_source else None
        ),
        "realtime_energy_comparison_date": (
            realtime_date.isoformat()
            if realtime_comparison_source is not None and realtime_date is not None
            else None
        ),
        "previous_path": str(previous_source) if previous_source else None,
        "market": {
            "day_ahead": _market_cards_from_quick(
                workbook,
                market_previous,
                sheet,
                text_override=day_narrative_text,
            ),
            "realtime": _market_cards_from_quick(
                workbook,
                market_previous,
                sheet,
                realtime=True,
                text_override=realtime_narrative_text,
                must_run_workbook=realtime_comparison_workbook,
            ),
        },
        "operations": _operation_table(
            workbook,
            report_date,
            previous_workbook,
            previous_date,
            quick_sheet=sheet,
        ),
        "day_ahead": {
            "cards": _summary_cards(
                sheet,
                day_series,
                text_override=day_summary_text,
                energy_metrics=energy_metrics,
            ),
            "series": day_series,
        },
        "realtime": {
            "cards": _summary_cards(
                sheet,
                realtime_series,
                realtime=True,
                energy_metrics=energy_metrics,
            ),
            "series": realtime_series,
        },
        "strategies": strategies,
        "appendices": {
            "day_ahead": _table_model(sheet, day_ahead_appendix_range, theme_colors),
            "realtime": _table_model(sheet, realtime_appendix_range, theme_colors),
        },
    }


def parse_company_clearing_workbook(path: str | Path) -> list[dict[str, Any]]:
    """Compatibility helper returning the company-fire 24-point day-ahead curve."""
    data = parse_daily_report_workbook(path)
    series = data["day_ahead"]["series"]
    return [
        {
            "time": label,
            "hour": index,
            "unit": "分公司火电",
            "energy_10k_kwh": series["energy"][index],
            "price_yuan_mwh": series["price"][index],
        }
        for index, label in enumerate(series["labels"])
    ]
