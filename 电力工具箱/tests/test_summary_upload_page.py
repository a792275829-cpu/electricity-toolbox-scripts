from __future__ import annotations

import gc
import tempfile
import threading
import tkinter as tk
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from toolbox.runtime import TaskRegistry, ToolPaths
from toolbox.summary_upload_page import SummaryUploadPage
from toolbox.tasks import CancellationToken
from toolbox.tasks.process import ProcessExecutionError


class SummaryUploadPageTests(unittest.TestCase):
    def setUp(self) -> None:
        try:
            self.root = tk.Tk()
        except tk.TclError as exc:
            self.skipTest(f"Tk unavailable: {exc}")
        self.root.withdraw()
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name) / "10-10"
        self.folder.mkdir()
        self.page = SummaryUploadPage(
            self.root, ToolPaths(Path(__file__).resolve().parents[2]), TaskRegistry(),
            today=date(2026, 10, 9), review_root=self.folder.parent,
        )
        self.page.file_dialogs.settings_path = self.folder.parent / "dialogs.json"

    def tearDown(self) -> None:
        self.page._defaults_thread.join(timeout=2)
        self.page.destroy()
        self.root.destroy()
        gc.collect()

    def test_slow_directory_scan_does_not_block_page_or_replace_user_input(self) -> None:
        entered = threading.Event()
        release = threading.Event()

        def slow_scan(*_args):
            self.assertIsNot(threading.current_thread(), threading.main_thread())
            entered.set()
            release.wait(3)
            return self.folder

        with (
            mock.patch("toolbox.summary_upload_page.default_review_folder_for_date", side_effect=slow_scan),
            mock.patch("toolbox.summary_upload_page.FileDialogState") as dialogs,
        ):
            page = SummaryUploadPage(
                self.root, self.page.paths, TaskRegistry(),
                today=date(2026, 10, 9), review_root=self.folder.parent,
            )
            try:
                self.assertTrue(entered.wait(1))
                self.assertTrue(page._defaults_thread.is_alive())
                page.pack()
                self.root.update_idletasks()
                self.assertEqual(page.winfo_manager(), "pack")
                self.assertEqual(page.run_button.cget("text"), "开始汇总")
                dialogs.assert_not_called()
                selected = self.folder.parent / "手动选择"
                page.input_dir.set(str(selected))
                page.output_file.set(str(selected / "自定义.xlsx"))
                release.set()
                page._defaults_thread.join(2)
                page._apply_default_input()
                self.assertEqual(Path(page.input_dir.get()), selected)
                self.assertEqual(Path(page.output_file.get()), selected / "自定义.xlsx")
            finally:
                release.set()
                page._defaults_thread.join(2)
                page.destroy()

    def test_background_default_keeps_latest_folder_fallback_and_updates_output(self) -> None:
        page = SummaryUploadPage(
            self.root, self.page.paths, TaskRegistry(),
            today=date(2026, 10, 10), review_root=self.folder.parent,
        )
        try:
            page._defaults_thread.join(2)
            page._apply_default_input()
            self.assertEqual(Path(page.input_dir.get()), self.folder)
            self.assertEqual(Path(page.output_file.get()), self.folder / "10-10_汇总.xlsx")
        finally:
            page.destroy()

    def test_hidden_default_scan_keeps_polling_until_result_is_applied(self) -> None:
        release = threading.Event()

        def slow_scan(*_args):
            release.wait(3)
            return self.folder

        with mock.patch("toolbox.summary_upload_page.default_review_folder_for_date", side_effect=slow_scan):
            page = SummaryUploadPage(
                self.root, self.page.paths, TaskRegistry(),
                today=date(2026, 10, 10), review_root=self.folder.parent,
            )
            try:
                self.root.after(100, self.root.quit)
                self.root.mainloop()
                self.assertIsNotNone(page._poll_after_id)
                release.set()
                page._defaults_thread.join(2)
                self.root.after(100, self.root.quit)
                self.root.mainloop()
                self.assertEqual(Path(page.input_dir.get()), self.folder)
                self.assertEqual(Path(page.output_file.get()), self.folder / "10-10_汇总.xlsx")
                self.assertIsNone(page._poll_after_id)
            finally:
                release.set()
                page._defaults_thread.join(2)
                page.destroy()

    def test_directory_scan_failure_leaves_page_usable(self) -> None:
        with mock.patch("toolbox.summary_upload_page.default_review_folder_for_date", side_effect=OSError("目录不可读")):
            page = SummaryUploadPage(
                self.root, self.page.paths, TaskRegistry(),
                today=date(2026, 10, 9), review_root=self.folder.parent,
            )
            try:
                page._defaults_thread.join(2)
                page._apply_default_input()
                self.assertEqual(Path(page.input_dir.get()), self.folder)
                page._drain_log_queue()
                self.assertIn("目录不可读", page.log_text.get("1.0", "end"))
                self.assertFalse(page.busy)
            finally:
                page.destroy()

    def prepare_summary_inputs(self) -> None:
        station = self.folder / "汕头"
        station.mkdir()
        for market, day, energy, price in (
            ("日前", "2026-10-10", 1.125, 300.25),
            ("实时", "2026-10-08", 2.125, 400.25),
        ):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "汕头1号机"
            sheet.append(["时刻", "电量", "结算电价"])
            for hour in range(24):
                sheet.append([f"{hour:02d}:00", energy + hour, price])
                sheet.cell(hour + 2, 2).number_format = "0.000"
                sheet.cell(hour + 2, 3).number_format = "0.00"
            workbook.save(station / f"1号机{market}交易结果查询({day}).xlsx")
            workbook.close()

    def start_worker(self, mode: str):
        self.page.mode.set(mode)
        with (
            mock.patch.object(self.page, "run_in_thread") as run,
            mock.patch("toolbox.summary_upload_page.messagebox.askokcancel", return_value=True),
            mock.patch("toolbox.summary_upload_page.node_executable", return_value="node"),
        ):
            self.page.start_selected()
        run.assert_called_once()
        return run.call_args.args[0], run.call_args.kwargs

    def test_modes_show_only_relevant_controls_and_keep_shared_input(self) -> None:
        for mode, summary, upload in (("summary", True, False), ("upload", False, True), ("both", True, True)):
            self.page.mode.set(mode)
            self.assertFalse(hasattr(self.page, "summary_form"))
            self.assertEqual(bool(self.page.preview_button.winfo_manager()), upload)
            self.assertEqual(bool(self.page.open_button.winfo_manager()), summary)
            self.assertEqual(Path(self.page.input_dir.get()), self.folder)

    def test_changing_input_updates_default_output_but_keeps_custom_output(self) -> None:
        new_folder = self.folder.parent / "10-11"
        self.page.input_dir.set(str(new_folder))
        self.assertEqual(Path(self.page.output_file.get()), new_folder / "10-11_汇总.xlsx")
        self.page.output_file.set(str(self.folder.parent / "custom.xlsx"))
        self.page.input_dir.set(str(self.folder))
        self.assertEqual(Path(self.page.output_file.get()), self.folder.parent / "custom.xlsx")

    def test_summary_runs_without_node_and_output_preserves_values_and_formats(self) -> None:
        self.prepare_summary_inputs()
        originals = {p: p.read_bytes() for p in self.folder.rglob("*.xlsx")}
        with (
            mock.patch.object(self.page, "run_in_thread") as run,
            mock.patch("toolbox.summary_upload_page.node_executable", side_effect=AssertionError("Node required")),
            mock.patch("toolbox.summary_upload_page.messagebox.askokcancel") as confirm,
        ):
            self.page.start_selected()
            output = run.call_args.args[0](CancellationToken())
        confirm.assert_not_called()
        self.assertTrue(run.call_args.kwargs["cancellable"])
        self.assertEqual(output, self.folder / "10-10_汇总.xlsx")
        workbook = load_workbook(output, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["汕头"])
            sheet = workbook["汕头"]
            for cell, value, number_format in (
                ("B4", 1.125, "0.000"), ("C4", 300.25, "0.00"),
                ("E4", 2.125, "0.000"), ("F4", 400.25, "0.00"),
                ("B27", 24.125, "0.000"), ("E27", 25.125, "0.000"),
            ):
                self.assertEqual(sheet[cell].value, value)
                self.assertEqual(sheet[cell].number_format, number_format)
            self.assertEqual(sheet["B3"].value, "日前电量\n（MWh）")
            self.assertEqual(sheet["E3"].value, "实时电量\n（MWh）")
        finally:
            workbook.close()
        self.assertEqual(
            sorted(path.read_bytes() for path in (self.folder / "汕头").glob("*.xlsx")),
            sorted(originals.values()),
        )

    def test_upload_does_not_load_summary_or_require_output(self) -> None:
        self.page.output_file.set("invalid-output")
        worker, options = self.start_worker("upload")
        with (
            mock.patch("toolbox.summary_upload_page.load_module", side_effect=AssertionError("summary loaded")),
            mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner,
        ):
            self.assertIsNone(worker(CancellationToken()))
        command = runner.return_value.run.call_args.args[0]
        self.assertEqual(command[-3:], ["--execute", "--source", str(self.folder.resolve())])
        self.assertFalse(options["cancellable"])

    def test_both_completes_summary_before_uploading_original_directory(self) -> None:
        self.prepare_summary_inputs()
        worker, options = self.start_worker("both")

        def upload(command, **_kwargs):
            self.assertTrue((self.folder / "10-10_汇总.xlsx").is_file())
            self.assertEqual(command[-1], str(self.folder.resolve()))

        with mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner:
            runner.return_value.run.side_effect = upload
            output = worker(CancellationToken())
        runner.return_value.run.assert_called_once()
        self.assertEqual(output, self.folder / "10-10_汇总.xlsx")
        self.assertFalse(options["cancellable"])

    def test_failed_summary_prevents_upload(self) -> None:
        worker, _ = self.start_worker("both")
        with mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner:
            with self.assertRaisesRegex(ValueError, "没有找到可汇总"):
                worker(CancellationToken())
        runner.assert_not_called()

    def test_failed_upload_keeps_summary_and_propagates_exit_code(self) -> None:
        self.prepare_summary_inputs()
        worker, _ = self.start_worker("both")
        with mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner:
            runner.return_value.run.side_effect = ProcessExecutionError(7)
            with self.assertRaises(ProcessExecutionError) as error:
                worker(CancellationToken())
        self.assertEqual(error.exception.returncode, 7)
        self.assertTrue((self.folder / "10-10_汇总.xlsx").is_file())

    def test_declining_upload_starts_neither_operation(self) -> None:
        self.page.mode.set("both")
        with (
            mock.patch("toolbox.summary_upload_page.messagebox.askokcancel", return_value=False),
            mock.patch.object(self.page, "run_in_thread") as run,
        ):
            self.page.start_selected()
        run.assert_not_called()

    def test_preview_uses_plan_without_starting_summary(self) -> None:
        with (
            mock.patch.object(self.page, "run_in_thread") as run,
            mock.patch("toolbox.summary_upload_page.node_executable", return_value="node"),
            mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner,
        ):
            self.page.preview_upload()
            run.call_args.args[0](CancellationToken())
        self.assertEqual(runner.return_value.run.call_args.args[0][-3:], ["--plan", "--source", str(self.folder.resolve())])
        self.assertEqual(runner.return_value.run.call_args.kwargs["cwd"], self.page.paths.private_uploader_dir)

    def test_flat_workbooks_are_classified_before_summary_and_upload(self) -> None:
        self.prepare_summary_inputs()
        station = self.folder / "汕头"
        for path in station.iterdir():
            path.rename(self.folder / path.name)
        station.rmdir()
        worker, _ = self.start_worker("both")
        with mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner:
            output = worker(CancellationToken())
        self.assertTrue(output.is_file())
        prepared = self.folder
        self.assertEqual(runner.return_value.run.call_args.args[0][-1], str(prepared.resolve()))
        self.assertEqual(len(list(prepared.glob("汕头/*.xlsx"))), 2)
        self.assertEqual(len(list(self.folder.glob("*交易结果查询*.xlsx"))), 0)
        self.assertFalse((self.folder / "已归类").exists())
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["汕头"])
            self.assertEqual(workbook["汕头"]["B4"].value, 1.125)
        finally:
            workbook.close()

    def test_unrecognized_flat_file_prevents_summary_and_upload(self) -> None:
        workbook = Workbook()
        workbook.save(self.folder / "日前交易结果查询(2026-10-10).xlsx")
        workbook.close()
        worker, _ = self.start_worker("both")
        with mock.patch("toolbox.summary_upload_page.ProcessRunner") as runner:
            with self.assertRaisesRegex(ValueError, "自动归类已停止"):
                worker(CancellationToken())
        runner.assert_not_called()
        self.assertFalse((self.folder / "10-10_汇总.xlsx").exists())

    def test_missing_stations_appear_in_completion_message(self) -> None:
        self.prepare_summary_inputs()
        worker, options = self.start_worker("summary")
        output = worker(CancellationToken())
        with mock.patch("toolbox.summary_upload_page.messagebox.showinfo") as info:
            options["on_success"](output)
        self.assertIn("缺少单元数据", info.call_args.args[1])
        self.assertIn("海门", info.call_args.args[1])
        self.assertNotIn("缺少单元数据：汕头", info.call_args.args[1])

    def test_invalid_input_does_not_start_or_confirm_upload(self) -> None:
        self.page.input_dir.set("")
        self.page.mode.set("both")
        with (
            mock.patch("toolbox.summary_upload_page.messagebox.showerror") as error,
            mock.patch("toolbox.summary_upload_page.messagebox.askokcancel") as confirm,
            mock.patch.object(self.page, "run_in_thread") as run,
        ):
            self.page.start_selected()
        error.assert_called_once()
        confirm.assert_not_called()
        run.assert_not_called()

    def test_both_rejects_summary_inside_upload_station_folder(self) -> None:
        self.page.output_file.set(str(self.folder / "汕头" / "汇总.xlsx"))
        self.page.mode.set("both")
        with (
            mock.patch("toolbox.summary_upload_page.messagebox.showerror") as error,
            mock.patch.object(self.page, "run_in_thread") as run,
        ):
            self.page.start_selected()
        error.assert_called_once()
        run.assert_not_called()

    def test_busy_page_does_not_start_second_operation(self) -> None:
        self.page.set_busy(True)
        with mock.patch.object(self.page, "run_in_thread") as run, mock.patch.object(self.page, "run_process") as process:
            self.page.start_selected()
            self.page.preview_upload()
        run.assert_not_called()
        process.assert_not_called()
        for widget in self.page._busy_widgets:
            self.assertIn("disabled", widget.state())


if __name__ == "__main__":
    unittest.main()
