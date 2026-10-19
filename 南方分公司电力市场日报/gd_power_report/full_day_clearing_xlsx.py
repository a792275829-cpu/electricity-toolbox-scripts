"""Read the explicitly dated full-day DA/RT table from the selected realtime workbook."""
from __future__ import annotations

from datetime import date, time
from hashlib import sha256
from math import isfinite
from pathlib import Path
import re
from typing import Any

from gd_power_report.workbook_reader import load_workbook, reuse_workbook_inputs
from openpyxl.utils import get_column_letter

from gd_power_report.company_clearing_xlsx import (
    _as_date, _merged_header_value, _normalized_label, _number,
    _select_report_sheet, _table_model, _theme_colors,
)

TITLE_MARKERS = ('分公司全天出清明细', '分公司出清全天出清明细')


def _title_date(cell: Any) -> date | None:
    text = _normalized_label(cell.value)
    if not any(marker in text or marker in str(cell.number_format) for marker in TITLE_MARKERS):
        return None
    direct = _as_date(cell.value)
    if direct is not None:
        return direct
    match = re.search(r'(\d{4})[年/.-](\d{1,2})[月/.-](\d{1,2})', text)
    if match:
        try:
            return date(*(int(value) for value in match.groups()))
        except ValueError:
            return None
    return None


def _hour(cell: Any) -> int | None:
    value = cell.value
    if isinstance(value, time):
        return value.hour if value.minute == value.second == value.microsecond == 0 else None
    text = str(value or '').strip().replace('：', ':')
    match = re.fullmatch(r'([01]?\d|2[0-3]):00(?::00)?', text)
    return int(match.group(1)) if match else None


@reuse_workbook_inputs
def read_full_day_clearing_table(path: str | Path, target_date: date) -> dict[str, Any]:
    """Never substitute a current-report-day table for the requested realtime day."""
    source = Path(path).expanduser().resolve()
    result: dict[str, Any] = {
        'date': target_date.isoformat(),
        'title': f'分公司全天出清明细（{target_date.year}/{target_date.month}/{target_date.day}）',
        'unit': '单位：万千瓦时、元/MWh',
        'source_path': str(source), 'source_sheet': None,
        'pending': True, 'pending_message': '所选实时日期工作簿未找到同日全天出清明细。',
        'table': None, 'source_warnings': [],
    }
    workbook = load_workbook(source, data_only=True, read_only=False)
    try:
        sheet = workbook['新'] if '新' in workbook.sheetnames else _select_report_sheet(workbook)
    except ValueError:
        result['pending_message'] = '工作簿缺少可读取的日报工作表。'
        return result
    result['source_sheet'] = sheet.title
    workbook_date = _as_date(sheet['A6'].value)
    if workbook_date is not None and workbook_date != target_date:
        result['pending_message'] = '所选工作簿日期与实时市场日期不一致。'
        return result
    candidates = [row for row in range(1, min(sheet.max_row, 1000) + 1)
                  if _title_date(sheet.cell(row, 1)) == target_date]
    if len(candidates) != 1:
        if len(candidates) > 1:
            result['pending_message'] = '工作簿中存在多个同日全天出清明细，无法唯一定位。'
        return result
    title_row = candidates[0]
    first_hour = next((row for row in range(title_row + 2, min(sheet.max_row, title_row + 12) + 1)
                       if _hour(sheet.cell(row, 1)) == 0), None)
    header_start = next((row for row in range(title_row + 1, (first_hour or title_row + 12))
                         if _normalized_label(sheet.cell(row, 1).value) in {'时刻', '时间'}), None)
    if first_hour is None or header_start is None:
        result['pending_message'] = '全天出清明细缺少完整表头或0:00记录。'
        return result
    end_row = first_hour + 24
    if (end_row > sheet.max_row or
            [_hour(sheet.cell(row, 1)) for row in range(first_hour, end_row)] != list(range(24)) or
            _normalized_label(sheet.cell(end_row, 1).value) not in {'合计', '总计', '全天合计'} or
            any(sheet.row_dimensions[row].hidden for row in range(header_start, end_row + 1))):
        result['pending_message'] = '全天出清明细的24小时、合计行或可见表头不完整。'
        return result
    metric_row = first_hour - 1
    columns = [col for col in range(2, min(sheet.max_column, 81) + 1)
               if sheet.cell(metric_row, col).value not in (None, '')]
    end_column = max(columns, default=1)
    if end_column < 5 or (end_column - 1) % 4:
        result['pending_message'] = '全天出清明细的日前/实时量价列不完整。'
        return result
    groups = []
    for column in range(2, end_column + 1):
        offset = (column - 2) % 4
        metric = '电量' if offset % 2 == 0 else '电价'
        period = '日前出清' if offset < 2 else '实时出清'
        if (_normalized_label(sheet.cell(metric_row, column).value) != metric or
                _normalized_label(_merged_header_value(sheet, metric_row - 1, column)) != period):
            result['pending_message'] = '全天出清明细表头与日前/实时电量电价结构不符。'
            return result
        if offset == 0:
            group = _normalized_label(_merged_header_value(sheet, header_start, column))
            if not group or any(_normalized_label(_merged_header_value(sheet, header_start, c)) != group
                                for c in range(column, column + 4)):
                result['pending_message'] = '全天出清明细缺少完整的单位列组。'
                return result
            groups.append(group)
    cell_range = f'A{header_start}:{get_column_letter(end_column)}{end_row}'
    table = _table_model(sheet, cell_range, _theme_colors(workbook))
    header_rows = first_hour - header_start
    raw_values = []
    for offset, row in enumerate(range(first_hour, end_row + 1)):
        cells = table['rows'][header_rows + offset]['cells']
        if len(cells) != end_column or any(c['rowspan'] != 1 or c['colspan'] != 1 for c in cells):
            result['pending_message'] = '全天出清明细数据区存在跨列或跨行合并，无法可靠对应量价。'
            return result
        cells[0]['value'] = f'{offset}:00' if offset < 24 else '合计'
        values = []
        for column in range(2, end_column + 1):
            source_cell = sheet.cell(row, column)
            value = _number(source_cell.value)
            if value is None or not isfinite(value):
                result['source_warnings'].append({'cell': source_cell.coordinate, 'reason': 'missing_or_invalid_cached_number'})
                cells[column - 1]['value'] = '—'
                value = None
            values.append(value)
        raw_values.append(values)
    result.update(pending=False, pending_message='', table=table, header_rows=header_rows,
                  title_cell=f'A{title_row}', source_range=cell_range, groups=groups,
                  source_sha256=sha256(source.read_bytes()).hexdigest(), raw_values=raw_values)
    return result

