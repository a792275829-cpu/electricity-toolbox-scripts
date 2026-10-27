"""Native sheet-data serialization with openpyxl's unchanged style/package XML."""
import os
import warnings
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.cell._writer import _set_attributes
from openpyxl.compat.strings import safe_string
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

from .native_excel import backend_mode, native_binary, call


class SkeletonWriter(WorksheetWriter):
    def write_rows(self):
        xf = self.xf.send(True)
        with xf.element("sheetData"):
            pass
        self.xf.send(None)


def save_workbook(workbook, path):
    mode = backend_mode()
    binary = native_binary() if mode != "python" else None
    if not isinstance(workbook, Workbook) or mode == "python" or (binary is None and mode == "auto"):
        return workbook.save(path)
    if binary is None:
        raise RuntimeError("未安装 Rust 写表后端")
    prepared = {}
    writers = []
    try:
        if workbook.write_only:
            raise ValueError("Rust 写表不支持 write_only 工作簿")
        for ws in workbook.worksheets:
            if ws._charts or ws._images or ws._tables or ws._pivots or ws.legacy_drawing:
                raise ValueError("复杂工作簿使用原 Python 写入")
            writer = SkeletonWriter(ws)
            writers.append(writer)
            writer.write_top()
            rows = []
            for row_number, cells in writer.rows():
                attrs = {"r": str(row_number)}
                attrs.update(dict(ws.row_dimensions.get(row_number, {})))
                row = []
                for cell in cells:
                    if cell._comment or cell.hyperlink or (cell.value is not None and not isinstance(cell.value, (str,int,float,bool))):
                        raise ValueError("特殊单元格使用 Python 写入")
                    if cell._value is None and not cell.has_style:
                        continue
                    value, cell_attrs = _set_attributes(cell, cell.has_style)
                    if value is not None:
                        value = str(value)[1:] if cell.data_type == "f" else safe_string(value)
                    row.append(dict(attrs=cell_attrs, value=value, kind=cell.data_type))
                rows.append(dict(attrs=attrs, cells=row))
            writer.write_rows()
            writer.write_tail()
            writer.close()
            prepared[id(ws)] = (writer, dict(path=writer.out, rows=rows))
        response = call(binary, "sheet_data", "", payload=[item[1] for item in prepared.values()])
        if response.get("written") is not True:
            raise ValueError("Rust 写表未确认完成")
        class NativeExcelWriter(ExcelWriter):
            def write_worksheet(self, ws):
                ws._drawing = SpreadsheetDrawing()
                writer = prepared[id(ws)][0]
                ws._rels = writer._rels
                self._archive.write(writer.out, ws.path[1:])
                self.manifest.append(ws)
        with ZipFile(path, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            NativeExcelWriter(workbook, archive).write_data()
    except Exception as exc:
        if mode == "rust":
            raise
        warnings.warn(f"Rust 写表不可用，改用 Python: {exc}", RuntimeWarning)
        workbook.save(path)
    finally:
        for writer in writers:
            try:
                writer.close()
            except Exception:
                pass
            if os.path.exists(writer.out):
                writer.cleanup()


def read_table_layout(path, preferred="整合数据"):
    """Read only two rows plus column metadata, never instantiate historical cells."""
    from copy import copy
    from xml.etree.ElementTree import iterparse
    from openpyxl.cell.cell import Cell
    from openpyxl.reader.excel import ExcelReader
    from openpyxl.styles.stylesheet import apply_stylesheet
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    class HeaderSheet(ReadOnlyWorksheet):
        def _get_size(self):
            # openpyxl's dimension probe otherwise reads all of sheetData when
            # the optional dimension element is absent. Only two rows are needed.
            pass

    reader = ExcelReader(path, read_only=True, data_only=False, keep_links=False)
    try:
        reader.read_manifest()
        reader.read_strings()
        reader.read_workbook()
        workbook = reader.wb
        apply_stylesheet(reader.archive, workbook)
        sheets = [(descriptor, relation) for descriptor, relation in reader.parser.find_sheets()
                  if "chartsheet" not in relation.Type and relation.target in reader.valid_files]
        descriptor, relation = next((item for item in sheets if item[0].name == preferred), sheets[0])
        sheet = HeaderSheet(workbook, descriptor.name, relation.target, reader.shared_strings)
        widths = {}
        with workbook._archive.open(sheet._worksheet_path) as source:
            for event, element in iterparse(source, events=("start",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "sheetData":
                    break
                if tag == "col":
                    # Match ColumnDimension's keyed lookup, including grouped columns.
                    widths[int(element.attrib["min"])] = float(element.attrib.get("width", 13))
        rows = list(sheet.iter_rows(min_row=1, max_row=2))
        layout = {}
        for index, cell in enumerate(rows[0], 1):
            if cell.value is None:
                continue
            def style(source):
                result = {name: copy(getattr(source, name)) for name in (
                    "font", "fill", "border", "alignment", "protection", "number_format")}
                array = getattr(source, "style_array", getattr(source, "_style", None))
                for name in ("quotePrefix", "pivotButton"):
                    result[name] = bool(getattr(array, name, False))
                return result
            data_cell = rows[1][index - 1] if index <= len(rows[1]) else None
            if not hasattr(data_cell, "style_array"):
                data_cell = Cell(sheet, row=2, column=index)
            layout[str(cell.value)] = {
                "width": widths.get(index, 13),
                "header_style": style(cell), "data_style": style(data_cell),
            }
        return layout
    finally:
        reader.archive.close()


def apply_table_style(cell, style):
    for name, value in style.items():
        setattr(cell, name, value)


def save_frame(frame, path, sheet_name, layout):
    """Write a fresh dataframe table through a bounded-memory native XML stream.

    Returns False for Python mode or optional-helper failure; the caller retains
    its original streaming writer and atomic replacement policy.
    """
    import json
    import tempfile
    from datetime import date, datetime, time, timedelta
    from pathlib import Path
    import pandas as pd
    from openpyxl.cell.cell import Cell, ERROR_CODES, ILLEGAL_CHARACTERS_RE
    from openpyxl.utils.datetime import to_excel, time_to_days, timedelta_to_days

    mode = backend_mode()
    if mode == "python":
        return False
    binary = native_binary()
    if binary is None:
        if mode == "rust":
            raise RuntimeError("未安装 Rust 写表后端")
        return False
    writers = []
    try:
        if "write_frame" not in call(binary, "capabilities", "").get("operations", []):
            raise ValueError("Rust 后端不支持流式表格写入")
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet.freeze_panes = "A2"
        # Register one style per column/type, rather than one Cell per data value.
        samples = {}
        def encode(value, column, header=False):
            if pd.isna(value):
                value = None
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
            kind = ("date" if isinstance(value, (datetime, date, time, timedelta)) else
                    "bool" if isinstance(value, bool) else
                    "str" if isinstance(value, str) else "number")
            key = (column, header, kind)
            if key not in samples:
                cell = Cell(sheet, row=1 if header else 2, column=column + 1, value=value)
                column_layout = layout.get(str(frame.columns[column]))
                if column_layout:
                    apply_table_style(cell, column_layout["header_style" if header else "data_style"])
                samples[key] = cell.style_id
            style_id = samples[key]
            if value is None:
                return ["n", "", style_id] if style_id else None
            if kind == "date":
                if isinstance(value, time):
                    value = time_to_days(value)
                elif isinstance(value, timedelta):
                    value = timedelta_to_days(value)
                else:
                    value = to_excel(value, book.epoch)
            if isinstance(value, str):
                if value.startswith("=") and len(value) > 1:
                    return ["n", value[1:], style_id, True]
                return ["e" if value in ERROR_CODES else "inlineStr", value, style_id]
            return ["b" if kind == "bool" else "n", safe_string(value), style_id]
        with tempfile.TemporaryDirectory(prefix="toolbox-table-") as directory:
            rows_path = Path(directory) / "rows.jsonl"
            with rows_path.open("w", encoding="utf-8") as rows_file:
                rows_file.write(json.dumps([encode(column, i, True) for i, column in enumerate(frame.columns)], ensure_ascii=False) + "\n")
                for row in frame.itertuples(index=False, name=None):
                    rows_file.write(json.dumps([encode(value, i) for i, value in enumerate(row)], ensure_ascii=False) + "\n")
            for i, column in enumerate(frame.columns, 1):
                if str(column) in layout:
                    from openpyxl.utils import get_column_letter
                    sheet.column_dimensions[get_column_letter(i)].width = layout[str(column)]["width"]
            writer = SkeletonWriter(sheet)
            writers.append(writer)
            writer.write_top()
            writer.write_rows()
            writer.write_tail()
            writer.close()
            # Match openpyxl write_only: omit dimensions instead of advertising
            # the empty skeleton's A1:A1 or truncating future streaming reads.
            xml = Path(writer.out).read_text()
            xml = xml.replace('<dimension ref="A1:A1"/>', "", 1)
            Path(writer.out).write_text(xml)
            response = call(binary, "write_frame", "", payload=dict(path=writer.out, rows_path=str(rows_path)))
            if response.get("written") is not True or response.get("rows") != len(frame) + 1:
                raise ValueError("Rust 流式写入行数不匹配")
            class TableWriter(ExcelWriter):
                def write_worksheet(self, ws):
                    ws._drawing = SpreadsheetDrawing()
                    ws._rels = writer._rels
                    self._archive.write(writer.out, ws.path[1:])
                    self.manifest.append(ws)
            with ZipFile(path, "w", ZIP_DEFLATED, allowZip64=True) as archive:
                TableWriter(book, archive).write_data()
        return True
    except Exception as exc:
        if mode == "rust":
            raise
        warnings.warn(f"Rust 流式写表不可用，改用 Python: {exc}", RuntimeWarning)
        return False
    finally:
        for writer in writers:
            try:
                writer.close()
            except Exception:
                pass
            if os.path.exists(writer.out):
                writer.cleanup()
