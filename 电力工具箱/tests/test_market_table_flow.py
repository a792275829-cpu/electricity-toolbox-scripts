from __future__ import annotations
import importlib

import gc
import subprocess
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
import runpy
import os
import shutil
from contextlib import ExitStack
from datetime import date
from datetime import timedelta
from pathlib import Path
from unittest import mock

def destroy_tk(root: tk.Misc) -> None:
    try:
        root.update_idletasks()
    except tk.TclError:
        pass
    root.destroy()


class MarketTableFlowTests(unittest.TestCase):
    """Market table update: date offsets, field order and cell writes."""
    def setUp(self) -> None:
        try:
            self.root = tk.Tk()
            self.root.withdraw()
        except tk.TclError as exc:
            self.skipTest(f"Tk unavailable: {exc}")
        from toolbox.runtime import TaskRegistry, ToolPaths

        self.paths = ToolPaths(Path(__file__).resolve().parents[2])
        self.registry = TaskRegistry()

    def tearDown(self) -> None:
        if hasattr(self, "root"):
            # Finish asynchronous page imports before destroying their Tk owner.
            for child in self.root.winfo_children():
                loader = getattr(child, "_writer_loader", None)
                if loader is not None and loader.ident is not None:
                    loader.join(5)
                    self.assertFalse(loader.is_alive(), "WPS page import did not finish")
            destroy_tk(self.root)
            gc.collect()  # Tk variables must be finalized on the main thread.

    def test_market_table_update_writes_actual_load_values_by_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_test", self.paths.market_table_update)
        values = [1000 + index for index in range(20)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["AH1"] = "实际运行"
            ws["A5"] = date(2026, 7, 1)
            ws["A6"] = date(2026, 7, 2)
            ws["A7"] = date(2026, 7, 3)
            wb.save(workbook_path)

            summary = module.write_actual_load_values(
                workbook_path,
                date(2026, 7, 2),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "AH6:BA6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(34, 54)],
            values,
        )

    def test_market_table_update_reports_permission_error_before_fetching(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_permission_error_test", self.paths.market_table_update)

        with mock.patch.object(module.openpyxl, "load_workbook", side_effect=PermissionError("Operation not permitted")):
            with self.assertRaisesRegex(RuntimeError, "无法打开 Excel 文件"):
                module.ensure_workbook_accessible(Path("/Users/auren/Desktop/日常工作/market.xlsx"))

    def test_market_table_update_skips_non_empty_market_cells(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_skip_non_empty_market_test", self.paths.market_table_update)
        values = [7000 + index for index in range(4)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["BB1"] = "实时市场"
            ws["A6"] = date(2026, 7, 2)
            ws["BB6"] = "保留"
            ws["BD6"] = 999
            wb.save(workbook_path)

            summary = module.write_realtime_market_values(
                workbook_path,
                date(2026, 7, 2),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["range"], "BB6:BE6")
        self.assertEqual(summary["count"], 2)
        self.assertEqual(summary["skipped"], 2)
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(54, 58)],
            ["保留", 7001, 999, 7003],
        )

    def test_market_table_update_writes_day_ahead_load_values_by_base_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_day_ahead_test", self.paths.market_table_update)
        values = [2000 + index for index in range(20)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["B1"] = "日前"
            ws["A5"] = date(2026, 7, 1)
            ws["A6"] = date(2026, 7, 2)
            ws["A7"] = date(2026, 7, 3)
            wb.save(workbook_path)

            summary = module.write_day_ahead_load_values(
                workbook_path,
                date(2026, 7, 2),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "B6:U6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(2, 22)],
            values,
        )

    def test_market_table_update_writes_day_ahead_market_values_by_source_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_day_ahead_market_test", self.paths.market_table_update)
        values = [3000 + index for index in range(12)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["V1"] = "发电侧申报均价"
            ws["A5"] = date(2026, 7, 4)
            ws["A6"] = date(2026, 7, 5)
            ws["A7"] = date(2026, 7, 6)
            wb.save(workbook_path)

            summary = module.write_day_ahead_market_values(
                workbook_path,
                date(2026, 7, 5),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "V6:AG6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(22, 34)],
            values,
        )

    def test_market_table_update_writes_realtime_market_values_by_source_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_realtime_market_write_test", self.paths.market_table_update)
        values = [4000 + index for index in range(4)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["BB1"] = "实时市场"
            ws["A5"] = date(2026, 7, 4)
            ws["A6"] = date(2026, 7, 5)
            ws["A7"] = date(2026, 7, 6)
            wb.save(workbook_path)

            summary = module.write_realtime_market_values(
                workbook_path,
                date(2026, 7, 5),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "BB6:BE6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(54, 58)],
            values,
        )

    def test_market_table_update_writes_generation_settlement_values_by_source_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_generation_settlement_write_test", self.paths.market_table_update)
        values = [5000 + index for index in range(12)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["BF1"] = "发电侧结算"
            ws["A5"] = date(2026, 6, 30)
            ws["A6"] = date(2026, 7, 1)
            ws["A7"] = date(2026, 7, 2)
            wb.save(workbook_path)

            summary = module.write_generation_settlement_values(
                workbook_path,
                date(2026, 7, 1),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "BF6:BQ6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(58, 70)],
            values,
        )

    def test_market_table_update_overwrites_generation_settlement_percent_cells(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_generation_settlement_overwrite_percent_test", self.paths.market_table_update)
        values = [
            17.25,
            -0.51,
            -0.029,
            14.41,
            0.835,
            2.02,
            0.117,
            0.485,
            0.368,
            7.28,
            421.6,
            372.6,
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["BF1"] = "发电侧结算"
            ws["A6"] = date(2026, 7, 1)
            ws["BH6"] = -2.9
            ws["BH6"].number_format = "0.00%"
            ws["BJ6"] = 83.5
            ws["BJ6"].number_format = "0.00%"
            ws["BL6"] = 11.7
            ws["BL6"].number_format = "0.00%"
            wb.save(workbook_path)

            summary = module.write_generation_settlement_values(
                workbook_path,
                date(2026, 7, 1),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["count"], 12)
        self.assertEqual(summary["skipped"], 0)
        self.assertEqual(updated_ws["BH6"].value, -0.029)
        self.assertEqual(updated_ws["BJ6"].value, 0.835)
        self.assertEqual(updated_ws["BL6"].value, 0.117)

    def test_market_table_update_writes_user_settlement_values_by_source_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_user_settlement_write_test", self.paths.market_table_update)
        values = [6000 + index for index in range(7)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年市场情况"
            ws["A1"] = "日期"
            ws["BR1"] = "用电侧结算"
            ws["A5"] = date(2026, 6, 30)
            ws["A6"] = date(2026, 7, 1)
            ws["A7"] = date(2026, 7, 2)
            wb.save(workbook_path)

            summary = module.write_user_settlement_values(
                workbook_path,
                date(2026, 7, 1),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年市场情况"]

        self.assertEqual(summary["sheet"], "2026年市场情况")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "BR6:BX6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(70, 77)],
            values,
        )

    def test_market_table_update_extracts_market_board_online_unit_values(self) -> None:
        from toolbox.runtime import load_module

        module = load_module(
            "market_table_update_market_board_extract_test",
            self.paths.market_table_update,
        )
        payload = {
            "dataApplyConfigDOList": [
                {
                    "date": "2026-07-16 00:00:00",
                    "dataType": 2,
                    "onlineUnit": 999,
                    "overHaulCapacity": 8888.8,
                },
                {
                    "date": "2026-07-16 00:00:00",
                    "dataType": 1,
                    "onlineUnit": 473.0,
                    "overHaulCapacity": 2985.6,
                },
            ]
        }

        self.assertEqual(
            module.extract_market_board_online_unit_values(
                payload,
                date(2026, 7, 16),
            ),
            [473, 2985.6],
        )

    def test_market_table_update_writes_market_board_values_to_aa_and_ae(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module(
            "market_table_update_market_board_write_test",
            self.paths.market_table_update,
        )

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A4"] = date(2026, 7, 15)
            ws["A5"] = "=A4+1"
            ws["AA5"] = "已有在线机组数"
            wb.save(workbook_path)

            summary = module.write_market_board_online_unit_values(
                workbook_path,
                date(2026, 7, 16),
                [473, 2985.6],
            )

            updated = load_workbook(workbook_path, data_only=False)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(summary["sheet"], "2026年运行方式及成本")
        self.assertEqual(summary["row"], 5)
        self.assertEqual(summary["range"], "AA5,AE5")
        self.assertEqual(summary["count"], 1)
        self.assertEqual(summary["skipped"], 1)
        self.assertEqual(updated_ws["AA5"].value, "已有在线机组数")
        self.assertEqual(updated_ws["AE5"].value, 2985.6)

    def test_market_table_update_extracts_day_ahead_coal_and_gas_unit_counts(self) -> None:
        from toolbox.runtime import load_module

        module = load_module(
            "market_table_update_spot_clearing_extract_test",
            self.paths.market_table_update,
        )
        coal_counts = [124, 122, 122, 122, 122] + [121] * 19
        gas_counts = [42, 25, 23, 22, 22, 24, 27, 29, 41, 43, 45, 45,
                      45, 45, 45, 44, 45, 45, 45, 46, 44, 41, 40, 38]
        payload = [
            {
                "date": "2026-07-16 00:00:00",
                "dataType": "2",
                "dataItemGroup": "REAL_TIME_CLEARANCE_COAL",
                "unitCount": [999] * 24,
            },
            {
                "date": "2026-07-16 00:00:00",
                "dataType": "1",
                "dataItemGroup": "DAY_AHEAD_CLEARANCE_COAL",
                "unitCount": coal_counts,
            },
            {
                "date": "2026-07-16 00:00:00",
                "dataType": "1",
                "dataItemGroup": "DAY_AHEAD_CLEARANCE_GAS",
                "unitCount": gas_counts,
            },
        ]

        self.assertEqual(
            module.extract_spot_clearing_unit_count_values(
                payload,
                date(2026, 7, 16),
            ),
            [121, 46],
        )

    def test_market_table_update_writes_spot_clearing_values_to_ab_and_ac(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module(
            "market_table_update_spot_clearing_write_test",
            self.paths.market_table_update,
        )

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A4"] = date(2026, 7, 15)
            ws["A5"] = "=A4+1"
            ws["AB5"] = "已有燃煤台数"
            wb.save(workbook_path)

            summary = module.write_spot_clearing_unit_count_values(
                workbook_path,
                date(2026, 7, 16),
                [121, 46],
            )

            updated = load_workbook(workbook_path, data_only=False)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(summary["sheet"], "2026年运行方式及成本")
        self.assertEqual(summary["row"], 5)
        self.assertEqual(summary["range"], "AB5:AC5")
        self.assertEqual(summary["count"], 1)
        self.assertEqual(summary["skipped"], 1)
        self.assertEqual(updated_ws["AB5"].value, "已有燃煤台数")
        self.assertEqual(updated_ws["AC5"].value, 46)

    def test_market_table_update_writes_unit_cost_prices_by_formula_date(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_cost_write_test", self.paths.market_table_update)
        values = [300 + index for index in range(11)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A1"] = "日期"
            ws["N1"] = "火电成本"
            ws["A4"] = date(2026, 1, 1)
            ws["A5"] = "=A4+1"
            ws["A6"] = "=A5+1"
            wb.save(workbook_path)

            summary = module.write_unit_cost_price_values(
                workbook_path,
                date(2026, 1, 3),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(summary["sheet"], "2026年运行方式及成本")
        self.assertEqual(summary["row"], 6)
        self.assertEqual(summary["range"], "N6:X6")
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(14, 25)],
            values,
        )
        self.assertEqual(updated_ws["N6"].number_format, updated_ws["N5"].number_format)

    def test_market_table_update_skips_non_empty_unit_cost_cells(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import load_module

        module = load_module("market_table_update_skip_non_empty_cost_test", self.paths.market_table_update)
        values = [800 + index for index in range(11)]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A1"] = "日期"
            ws["N1"] = "火电成本"
            ws["A4"] = date(2026, 1, 1)
            ws["A5"] = "=A4+1"
            ws["A6"] = "=A5+1"
            ws["N6"] = "保留"
            ws["P6"] = 999
            wb.save(workbook_path)

            summary = module.write_unit_cost_price_values(
                workbook_path,
                date(2026, 1, 3),
                values,
            )

            updated = load_workbook(workbook_path, data_only=True)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(summary["range"], "N6:X6")
        self.assertEqual(summary["count"], 9)
        self.assertEqual(summary["skipped"], 2)
        expected = values[:]
        expected[0] = "保留"
        expected[2] = 999
        self.assertEqual(
            [updated_ws.cell(6, column).value for column in range(14, 25)],
            expected,
        )

    def test_market_table_update_preserves_previous_row_column_formats(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        from toolbox.runtime import load_module

        module = load_module("market_table_update_cost_format_test", self.paths.market_table_update)

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A4"] = date(2026, 1, 1)
            ws["A5"] = "=A4+1"
            for column in range(14, 25):
                source = ws.cell(4, column)
                source.number_format = "0.0000"
                source.font = Font(name="Arial", bold=True, color="FF0000")
                source.fill = PatternFill("solid", fgColor="FFF2CC")
            wb.save(workbook_path)

            module.write_unit_cost_price_values(
                workbook_path,
                date(2026, 1, 2),
                [400 + index for index in range(11)],
            )

            updated = load_workbook(workbook_path)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(updated_ws["N5"].number_format, "0.0000")
        self.assertEqual(updated_ws["N5"].font.name, "Arial")
        self.assertTrue(updated_ws["N5"].font.bold)
        self.assertEqual(updated_ws["N5"].font.color.rgb, "00FF0000")
        self.assertEqual(updated_ws["N5"].fill.fgColor.rgb, "00FFF2CC")

    def test_market_table_update_copies_previous_day_unit_operation_mode(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        from toolbox.runtime import load_module

        module = load_module("market_table_update_operation_mode_test", self.paths.market_table_update)
        source_values = [
            "运行", "运行", "备用", "二分列",
            "运行", "运行", "运行", "运行",
            "备用", "调峰", "调峰", "备用",
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "market.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "2026年运行方式及成本"
            ws["A4"] = date(2026, 7, 3)
            ws["A5"] = "=A4+1"
            ws["B5"] = "已有运行"
            ws["D5"] = "已有备用"
            ws["N5"] = 999
            for offset, value in enumerate(source_values, start=2):
                cell = ws.cell(4, offset)
                cell.value = value
                cell.font = Font(name="Arial", bold=True, color="FF0000")
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            wb.save(workbook_path)

            summary = module.copy_previous_day_unit_operation_mode(
                workbook_path,
                date(2026, 7, 4),
            )

            updated = load_workbook(workbook_path)
            updated_ws = updated["2026年运行方式及成本"]

        self.assertEqual(summary["sheet"], "2026年运行方式及成本")
        self.assertEqual(summary["sourceRow"], 4)
        self.assertEqual(summary["row"], 5)
        self.assertEqual(summary["range"], "B5:M5")
        expected_values = source_values[:]
        expected_values[0] = "已有运行"
        expected_values[2] = "已有备用"
        self.assertEqual([updated_ws.cell(5, column).value for column in range(2, 14)], expected_values)
        self.assertEqual(summary["count"], 10)
        self.assertEqual(summary["skipped"], 2)
        self.assertIsNotNone(updated_ws["A5"].value)
        self.assertEqual(updated_ws["N5"].value, 999)
        self.assertEqual(updated_ws["C5"].font.name, "Arial")
        self.assertTrue(updated_ws["C5"].font.bold)
        self.assertEqual(updated_ws["C5"].fill.fgColor.rgb, "00FFF2CC")

    def test_market_table_update_extracts_page_actual_load_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_order_test", self.paths.market_table_update)
        dataset = {
            "data": {
                "dataNetLoadDTOList": [
                    {"loadType": "2", "actualPeriodList": [{"max": 41, "min": 42, "avg": 43, "sum": 44}]},
                    {"loadType": "1", "actualPeriodList": [{"max": 11, "min": 12, "avg": 13, "sum": 14}]},
                    {"loadType": "3", "actualPeriodList": [{"max": 51, "min": 52, "avg": 53, "sum": 54}]},
                    {"loadType": "6", "actualPeriodList": [{"max": 21, "min": 22, "avg": 23, "sum": 24}]},
                    {"loadType": "4", "actualPeriodList": [{"max": 31, "min": 32, "avg": 33, "sum": 34}]},
                ]
            }
        }

        self.assertEqual(
            module.extract_actual_load_values(dataset, date(2026, 7, 2)),
            [
                11, 12, 13, 14,
                21, 22, 23, 24,
                31, 32, 33, 34,
                41, 42, 43, 44,
                51, 52, 53, 54,
            ],
        )

    def test_market_table_update_treats_none_actual_load_metrics_as_unavailable(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_none_actual_load_test", self.paths.market_table_update)
        dataset = {
            "data": {
                "dataNetLoadDTOList": [
                    {
                        "loadType": "1",
                        "actualPeriodList": [{"max": None, "min": None, "avg": None, "sum": None}],
                    }
                ]
            }
        }

        with self.assertRaisesRegex(RuntimeError, "没有返回实际运行数据"):
            module.extract_actual_load_values(dataset, date(2026, 7, 5))

    def test_market_table_update_extracts_day_ahead_market_field_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_market_order_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-07-05 00:00:00",
                "info": {
                    "generateSideDeclareAvgPrice": 376.6,
                    "totalGenerateSideDealEle": 15.64,
                    "totalCoalDealEle": 9.26,
                    "totalGasDealEle": 1.33,
                    "totalNuclearDealEle": 2.42,
                    "totalNewEnergyDealEle": 2.58,
                    "totalPowerSideDealEle": 25.27,
                    "generateSideAvgPrice": 295.63,
                    "coalDealMaxPrice": 1157.53,
                    "coalDealAvgPrice": 318.36,
                    "gasDealMaxPrice": 1376.96,
                    "gasDealAvgPrice": 460.69,
                },
            }
        ]

        self.assertEqual(
            module.extract_day_ahead_market_values(payload, date(2026, 7, 5)),
            [
                376.6,
                15.64,
                9.26,
                1.33,
                2.42,
                2.58,
                25.27,
                295.63,
                1157.53,
                318.36,
                1376.96,
                460.69,
            ],
        )

    def test_market_table_update_extracts_realtime_market_field_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_realtime_market_order_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-07-05 00:00:00",
                "info": {
                    "generateSideDeclareAvgPrice": 376.6,
                    "totalGenerateSideDealEle": 15.64,
                    "totalCoalDealEle": 9.26,
                    "totalGasDealEle": 1.33,
                    "totalNuclearDealEle": 2.42,
                    "totalNewEnergyDealEle": 2.58,
                    "totalPowerSideDealEle": 25.27,
                    "generateSideAvgPrice": 295.63,
                    "coalDealMaxPrice": 1157.53,
                    "coalDealAvgPrice": 318.36,
                    "gasDealMaxPrice": 1376.96,
                    "gasDealAvgPrice": 460.69,
                },
            }
        ]

        self.assertEqual(
            module.extract_realtime_market_values(payload, date(2026, 7, 5)),
            [
                15.64,
                295.63,
                318.36,
                460.69,
            ],
        )

    def test_market_table_update_extracts_realtime_market_live_deal_energy_field(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_realtime_market_live_field_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-06-30 00:00:00",
                "info": {
                    "generateSideDealEle": 17.01,
                    "generateSideAvgPrice": 341.2,
                    "coalDealAvgPrice": 320.5,
                    "gasDealAvgPrice": 442.8,
                },
            }
        ]

        self.assertEqual(
            module.extract_realtime_market_values(payload, date(2026, 6, 30)),
            [
                17.01,
                341.2,
                320.5,
                442.8,
            ],
        )

    def test_market_table_update_extracts_generation_settlement_field_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_generation_settlement_order_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-06-30 00:00:00",
                "info": {
                    "marketUnitMltEleRadio": 83.5,
                    "spotNegDeviationRadio": 36.8,
                    "marketUnitBaseEle": -0.51,
                    "marketUnitDeviationEle": 2.02,
                    "spotPosDeviationRadio": 48.5,
                    "generateAvgPrice": 372.6,
                    "marketUnitFee": 7.28,
                    "marketUnitDeviationEleRadio": 11.7,
                    "marketUnitMltEle": 14.41,
                    "generateAvgPriceWithCompensate": 421.6,
                    "marketUnitBaseEleRadio": -2.9,
                    "marketUnitOnlineEle": 17.25,
                },
            }
        ]

        self.assertEqual(
            module.extract_generation_settlement_values(payload, date(2026, 6, 30)),
            [
                17.25,
                -0.51,
                -0.029,
                14.41,
                0.835,
                2.02,
                0.117,
                0.485,
                0.368,
                7.28,
                421.6,
                372.6,
            ],
        )

    def test_market_table_update_extracts_user_settlement_field_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_user_settlement_order_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-06-30 00:00:00",
                "info": {
                    "marketUserConsumeEle": 25.77,
                    "userAvePrice": 418.7,
                    "mltEle": 21.1,
                    "mltEleRadio": 81.9,
                    "spotDeviationEle": 4.67,
                    "spotDeviationRadio": 18.1,
                    "getMarketUserConsumeFee": 10.79,
                },
            }
        ]

        self.assertEqual(
            module.extract_user_settlement_values(payload, date(2026, 6, 30)),
            [
                25.77,
                21.1,
                0.819,
                4.67,
                0.181,
                10.79,
                418.7,
            ],
        )

    def test_market_table_update_converts_generation_settlement_percent_strings(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_generation_settlement_percent_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-06-30 00:00:00",
                "info": {
                    "marketUnitOnlineEle": "17.25",
                    "marketUnitBaseEle": "-0.51",
                    "marketUnitBaseEleRadio": "-2.9%",
                    "marketUnitMltEle": "14.41",
                    "marketUnitMltEleRadio": 83.5,
                    "marketUnitDeviationEle": "2.02",
                    "marketUnitDeviationEleRadio": "11.7%",
                    "spotPosDeviationRadio": 48.5,
                    "spotNegDeviationRadio": "36.8%",
                    "marketUnitFee": "7.28",
                    "generateAvgPriceWithCompensate": "421.6",
                    "generateAvgPrice": "372.6",
                },
            }
        ]

        values = module.extract_generation_settlement_values(payload, date(2026, 6, 30))

        self.assertEqual(
            values,
            [
                17.25,
                -0.51,
                -0.029,
                14.41,
                0.835,
                2.02,
                0.117,
                0.485,
                0.368,
                7.28,
                421.6,
                372.6,
            ],
        )

    def test_market_table_update_converts_user_settlement_percent_strings(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_user_settlement_percent_test", self.paths.market_table_update)
        payload = [
            {
                "date": "2026-06-30 00:00:00",
                "info": {
                    "marketUserConsumeEle": "25.77",
                    "userAvePrice": "418.7",
                    "mltEle": "21.1",
                    "mltEleRadio": 81.9,
                    "spotDeviationEle": "4.67",
                    "spotDeviationRadio": "18.1%",
                    "getMarketUserConsumeFee": "10.79",
                },
            }
        ]

        values = module.extract_user_settlement_values(payload, date(2026, 6, 30))

        self.assertEqual(
            values,
            [
                25.77,
                21.1,
                0.819,
                4.67,
                0.181,
                10.79,
                418.7,
            ],
        )

    def test_market_table_update_skips_mapping_when_values_are_unavailable(self) -> None:
        from types import SimpleNamespace
        from toolbox.runtime import load_module

        module = load_module("market_table_update_skip_test", self.paths.market_table_update)

        class DummyLock:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return None

        class DummyPlaywright:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return None

        context = SimpleNamespace(close=mock.Mock())
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(module, "load_config", return_value=SimpleNamespace(headless=True)))
            stack.enter_context(mock.patch.object(module, "ensure_workbook_accessible"))
            stack.enter_context(mock.patch.object(module, "AuthStateLock", return_value=DummyLock()))
            stack.enter_context(mock.patch.object(module, "sync_playwright", return_value=DummyPlaywright()))
            stack.enter_context(mock.patch.object(module, "launch_context", return_value=context))
            ensure_login = stack.enter_context(mock.patch.object(module, "ensure_login"))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_load_values", side_effect=RuntimeError("没有返回日前数据")))
            write_day_ahead_load = stack.enter_context(mock.patch.object(module, "write_day_ahead_load_values"))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_market_values", return_value=[1] * 12))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_market_values", return_value={"range": "V6:AG6", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_realtime_market_values", return_value=[4] * 4))
            stack.enter_context(mock.patch.object(module, "write_realtime_market_values", return_value={"range": "BB5:BE5", "sheet": "2026年市场情况", "count": 4}))
            stack.enter_context(mock.patch.object(module, "fetch_generation_settlement_values", return_value=[5] * 12))
            stack.enter_context(mock.patch.object(module, "write_generation_settlement_values", return_value={"range": "BF1:BQ1", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_user_settlement_values", return_value=[6] * 7))
            stack.enter_context(mock.patch.object(module, "write_user_settlement_values", return_value={"range": "BR1:BX1", "sheet": "2026年市场情况", "count": 7}))
            stack.enter_context(mock.patch.object(module, "fetch_actual_load_values", return_value=[2] * 20))
            stack.enter_context(mock.patch.object(module, "write_actual_load_values", return_value={"range": "AH5:BA5", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_unit_cost_price_values", return_value=[3] * 11))
            stack.enter_context(mock.patch.object(module, "write_unit_cost_price_values", return_value={"range": "N7:X7", "sheet": "2026年运行方式及成本", "count": 11}))
            fetch_market_board = stack.enter_context(mock.patch.object(module, "fetch_market_board_online_unit_values", return_value=[473, 2985.6]))
            write_market_board = stack.enter_context(mock.patch.object(module, "write_market_board_online_unit_values", return_value={"range": "AA7,AE7", "sheet": "2026年运行方式及成本", "count": 2}))
            fetch_spot_clearing = stack.enter_context(mock.patch.object(module, "fetch_spot_clearing_unit_count_values", return_value=[121, 46]))
            write_spot_clearing = stack.enter_context(mock.patch.object(module, "write_spot_clearing_unit_count_values", return_value={"range": "AB7:AC7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "copy_previous_day_unit_operation_mode", return_value={"range": "B7:M7", "sheet": "2026年运行方式及成本", "sourceRow": 6}))
            result = module.run_update(Path("market.xlsx"), date(2026, 7, 6), headless=True)

        self.assertIsNone(result["dayAheadLoad"])
        self.assertEqual(result["dayAheadMarket"]["range"], "V6:AG6")
        self.assertEqual(result["realtimeMarket"]["range"], "BB5:BE5")
        self.assertEqual(result["generationSettlement"]["range"], "BF1:BQ1")
        self.assertEqual(result["userSettlement"]["range"], "BR1:BX1")
        self.assertEqual(result["actualLoad"]["range"], "AH5:BA5")
        self.assertEqual(result["unitCostPrices"]["range"], "N7:X7")
        self.assertEqual(result["marketBoardOnlineUnit"]["range"], "AA7,AE7")
        self.assertEqual(result["spotClearingUnitCounts"]["range"], "AB7:AC7")
        self.assertEqual(result["unitOperationMode"]["range"], "B7:M7")
        fetch_market_board.assert_called_once_with(context, date(2026, 7, 6))
        write_market_board.assert_called_once_with(
            Path("market.xlsx"),
            date(2026, 7, 6),
            [473, 2985.6],
        )
        fetch_spot_clearing.assert_called_once_with(context, date(2026, 7, 6))
        write_spot_clearing.assert_called_once_with(
            Path("market.xlsx"),
            date(2026, 7, 6),
            [121, 46],
        )
        write_day_ahead_load.assert_not_called()
        ensure_login.assert_called_once()
        context.close.assert_called_once()

    def test_market_table_update_retries_transient_playwright_fetch_errors(self) -> None:
        from playwright.sync_api import Error as PlaywrightError
        from toolbox.runtime import load_module

        module = load_module("market_table_update_retry_fetch_test", self.paths.market_table_update)
        fetcher = mock.Mock(side_effect=[PlaywrightError("socket disconnected"), [1, 2, 3]])

        self.assertEqual(
            module.fetch_values_or_skip("现货运行日报日前市场", fetcher),
            [1, 2, 3],
        )
        self.assertEqual(fetcher.call_count, 2)

    def test_market_table_update_skips_after_repeated_playwright_fetch_errors(self) -> None:
        from playwright.sync_api import Error as PlaywrightError
        from toolbox.runtime import load_module

        module = load_module("market_table_update_skip_fetch_error_test", self.paths.market_table_update)
        fetcher = mock.Mock(side_effect=PlaywrightError("socket disconnected"))

        self.assertIsNone(
            module.fetch_values_or_skip("现货运行日报日前市场", fetcher)
        )
        self.assertEqual(fetcher.call_count, 3)

    def test_market_table_update_uses_base_date_minus_two_for_realtime_market(self) -> None:
        from types import SimpleNamespace
        from toolbox.runtime import load_module

        module = load_module("market_table_update_realtime_market_date_test", self.paths.market_table_update)

        class DummyLock:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return None

        class DummyPlaywright:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return None

        context = SimpleNamespace(close=mock.Mock())
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(module, "load_config", return_value=SimpleNamespace(headless=True)))
            stack.enter_context(mock.patch.object(module, "ensure_workbook_accessible"))
            stack.enter_context(mock.patch.object(module, "AuthStateLock", return_value=DummyLock()))
            stack.enter_context(mock.patch.object(module, "sync_playwright", return_value=DummyPlaywright()))
            stack.enter_context(mock.patch.object(module, "launch_context", return_value=context))
            stack.enter_context(mock.patch.object(module, "ensure_login"))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_load_values", return_value=[1] * 20))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_load_values", return_value={"range": "B7:U7", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_market_values", return_value=[2] * 12))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_market_values", return_value={"range": "V6:AG6", "sheet": "2026年市场情况", "count": 12}))
            fetch_realtime_market = stack.enter_context(mock.patch.object(module, "fetch_realtime_market_values", return_value=[3] * 4))
            write_realtime_market = stack.enter_context(mock.patch.object(module, "write_realtime_market_values", return_value={"range": "BB5:BE5", "sheet": "2026年市场情况", "count": 4}))
            stack.enter_context(mock.patch.object(module, "fetch_generation_settlement_values", return_value=[6] * 12))
            stack.enter_context(mock.patch.object(module, "write_generation_settlement_values", return_value={"range": "BF1:BQ1", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_user_settlement_values", return_value=[7] * 7))
            stack.enter_context(mock.patch.object(module, "write_user_settlement_values", return_value={"range": "BR1:BX1", "sheet": "2026年市场情况", "count": 7}))
            stack.enter_context(mock.patch.object(module, "fetch_actual_load_values", return_value=[4] * 20))
            stack.enter_context(mock.patch.object(module, "write_actual_load_values", return_value={"range": "AH5:BA5", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_unit_cost_price_values", return_value=[5] * 11))
            stack.enter_context(mock.patch.object(module, "write_unit_cost_price_values", return_value={"range": "N7:X7", "sheet": "2026年运行方式及成本", "count": 11}))
            stack.enter_context(mock.patch.object(module, "fetch_market_board_online_unit_values", return_value=[473, 2985.6]))
            stack.enter_context(mock.patch.object(module, "write_market_board_online_unit_values", return_value={"range": "AA7,AE7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "fetch_spot_clearing_unit_count_values", return_value=[121, 46]))
            stack.enter_context(mock.patch.object(module, "write_spot_clearing_unit_count_values", return_value={"range": "AB7:AC7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "copy_previous_day_unit_operation_mode", return_value={"range": "B7:M7", "sheet": "2026年运行方式及成本", "sourceRow": 6}))
            result = module.run_update(Path("market.xlsx"), date(2026, 7, 6), headless=True)

        self.assertEqual(result["realtimeMarket"]["range"], "BB5:BE5")
        fetch_realtime_market.assert_called_once_with(context, date(2026, 7, 4))
        write_realtime_market.assert_called_once_with(Path("market.xlsx"), date(2026, 7, 4), [3] * 4)

    def test_market_table_update_uses_base_date_minus_six_for_generation_settlement(self) -> None:
        from types import SimpleNamespace
        from toolbox.runtime import load_module

        module = load_module("market_table_update_generation_settlement_date_test", self.paths.market_table_update)

        class DummyLock:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return None

        class DummyPlaywright:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return None

        context = SimpleNamespace(close=mock.Mock())
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(module, "load_config", return_value=SimpleNamespace(headless=True)))
            stack.enter_context(mock.patch.object(module, "ensure_workbook_accessible"))
            stack.enter_context(mock.patch.object(module, "AuthStateLock", return_value=DummyLock()))
            stack.enter_context(mock.patch.object(module, "sync_playwright", return_value=DummyPlaywright()))
            stack.enter_context(mock.patch.object(module, "launch_context", return_value=context))
            stack.enter_context(mock.patch.object(module, "ensure_login"))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_load_values", return_value=[1] * 20))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_load_values", return_value={"range": "B7:U7", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_market_values", return_value=[2] * 12))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_market_values", return_value={"range": "V6:AG6", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_realtime_market_values", return_value=[3] * 4))
            stack.enter_context(mock.patch.object(module, "write_realtime_market_values", return_value={"range": "BB5:BE5", "sheet": "2026年市场情况", "count": 4}))
            fetch_generation_settlement = stack.enter_context(mock.patch.object(module, "fetch_generation_settlement_values", return_value=[4] * 12))
            write_generation_settlement = stack.enter_context(mock.patch.object(module, "write_generation_settlement_values", return_value={"range": "BF1:BQ1", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_user_settlement_values", return_value=[7] * 7))
            stack.enter_context(mock.patch.object(module, "write_user_settlement_values", return_value={"range": "BR1:BX1", "sheet": "2026年市场情况", "count": 7}))
            stack.enter_context(mock.patch.object(module, "fetch_actual_load_values", return_value=[5] * 20))
            stack.enter_context(mock.patch.object(module, "write_actual_load_values", return_value={"range": "AH5:BA5", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_unit_cost_price_values", return_value=[6] * 11))
            stack.enter_context(mock.patch.object(module, "write_unit_cost_price_values", return_value={"range": "N7:X7", "sheet": "2026年运行方式及成本", "count": 11}))
            stack.enter_context(mock.patch.object(module, "fetch_market_board_online_unit_values", return_value=[473, 2985.6]))
            stack.enter_context(mock.patch.object(module, "write_market_board_online_unit_values", return_value={"range": "AA7,AE7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "fetch_spot_clearing_unit_count_values", return_value=[121, 46]))
            stack.enter_context(mock.patch.object(module, "write_spot_clearing_unit_count_values", return_value={"range": "AB7:AC7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "copy_previous_day_unit_operation_mode", return_value={"range": "B7:M7", "sheet": "2026年运行方式及成本", "sourceRow": 6}))
            result = module.run_update(Path("market.xlsx"), date(2026, 7, 6), headless=True)

        self.assertEqual(result["generationSettlement"]["range"], "BF1:BQ1")
        fetch_generation_settlement.assert_called_once_with(context, date(2026, 6, 30))
        write_generation_settlement.assert_called_once_with(Path("market.xlsx"), date(2026, 6, 30), [4] * 12)

    def test_market_table_update_uses_base_date_minus_six_for_user_settlement(self) -> None:
        from types import SimpleNamespace
        from toolbox.runtime import load_module

        module = load_module("market_table_update_user_settlement_date_test", self.paths.market_table_update)

        class DummyLock:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return None

        class DummyPlaywright:
            def __enter__(self):
                return object()

            def __exit__(self, exc_type, exc, tb):
                return None

        context = SimpleNamespace(close=mock.Mock())
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(module, "load_config", return_value=SimpleNamespace(headless=True)))
            stack.enter_context(mock.patch.object(module, "ensure_workbook_accessible"))
            stack.enter_context(mock.patch.object(module, "AuthStateLock", return_value=DummyLock()))
            stack.enter_context(mock.patch.object(module, "sync_playwright", return_value=DummyPlaywright()))
            stack.enter_context(mock.patch.object(module, "launch_context", return_value=context))
            stack.enter_context(mock.patch.object(module, "ensure_login"))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_load_values", return_value=[1] * 20))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_load_values", return_value={"range": "B7:U7", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_day_ahead_market_values", return_value=[2] * 12))
            stack.enter_context(mock.patch.object(module, "write_day_ahead_market_values", return_value={"range": "V6:AG6", "sheet": "2026年市场情况", "count": 12}))
            stack.enter_context(mock.patch.object(module, "fetch_realtime_market_values", return_value=[3] * 4))
            stack.enter_context(mock.patch.object(module, "write_realtime_market_values", return_value={"range": "BB5:BE5", "sheet": "2026年市场情况", "count": 4}))
            stack.enter_context(mock.patch.object(module, "fetch_generation_settlement_values", return_value=[4] * 12))
            stack.enter_context(mock.patch.object(module, "write_generation_settlement_values", return_value={"range": "BF1:BQ1", "sheet": "2026年市场情况", "count": 12}))
            fetch_user_settlement = stack.enter_context(mock.patch.object(module, "fetch_user_settlement_values", return_value=[5] * 7))
            write_user_settlement = stack.enter_context(mock.patch.object(module, "write_user_settlement_values", return_value={"range": "BR1:BX1", "sheet": "2026年市场情况", "count": 7}))
            stack.enter_context(mock.patch.object(module, "fetch_actual_load_values", return_value=[6] * 20))
            stack.enter_context(mock.patch.object(module, "write_actual_load_values", return_value={"range": "AH5:BA5", "sheet": "2026年市场情况", "count": 20}))
            stack.enter_context(mock.patch.object(module, "fetch_unit_cost_price_values", return_value=[7] * 11))
            stack.enter_context(mock.patch.object(module, "write_unit_cost_price_values", return_value={"range": "N7:X7", "sheet": "2026年运行方式及成本", "count": 11}))
            stack.enter_context(mock.patch.object(module, "fetch_market_board_online_unit_values", return_value=[473, 2985.6]))
            stack.enter_context(mock.patch.object(module, "write_market_board_online_unit_values", return_value={"range": "AA7,AE7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "fetch_spot_clearing_unit_count_values", return_value=[121, 46]))
            stack.enter_context(mock.patch.object(module, "write_spot_clearing_unit_count_values", return_value={"range": "AB7:AC7", "sheet": "2026年运行方式及成本", "count": 2}))
            stack.enter_context(mock.patch.object(module, "copy_previous_day_unit_operation_mode", return_value={"range": "B7:M7", "sheet": "2026年运行方式及成本", "sourceRow": 6}))
            result = module.run_update(Path("market.xlsx"), date(2026, 7, 6), headless=True)

        self.assertEqual(result["userSettlement"]["range"], "BR1:BX1")
        fetch_user_settlement.assert_called_once_with(context, date(2026, 6, 30))
        write_user_settlement.assert_called_once_with(Path("market.xlsx"), date(2026, 6, 30), [5] * 7)

    def test_market_table_update_uses_base_date_minus_two_for_actual_load(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_actual_load_date_test", self.paths.market_table_update)

        self.assertEqual(
            module.actual_load_source_date(date(2026, 7, 4)),
            date(2026, 7, 2),
        )

    def test_market_table_update_extracts_unit_cost_price_order(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("market_table_update_cost_order_test", self.paths.market_table_update)
        payload = [
            {
                "unitCostInfoDTOList": [
                    {"unitId": "hm2", "priceOfDyCost": 344.8668888},
                    {"unitId": "hm1", "priceOfDyCost": 341.4059744},
                    {"unitId": "hm4", "priceOfDyCost": 342.5596126},
                    {"unitId": "hm3", "priceOfDyCost": 343.7132507},
                ]
            },
            {
                "unitCostInfoDTOList": [
                    {"unitId": "st3", "priceOfDyCost": 424.7904072},
                    {"unitId": "st1", "priceOfDyCost": 394.3277667},
                    {"unitId": "st2", "priceOfDyCost": 395.4994067},
                ]
            },
            {
                "unitCostInfoDTOList": [
                    {"unitId": "xg56", "priceOfDyCost": 620.36619563},
                    {"unitId": "xg78", "priceOfDyCost": 620.36619563},
                    {"unitId": "xg12", "priceOfDyCost": 620.36619563},
                    {"unitId": "xg34", "priceOfDyCost": 620.36619563},
                ]
            },
        ]

        with mock.patch.object(
            module,
            "UNIT_COST_PRICE_UNIT_IDS",
            ("st1", "st2", "st3", "hm1", "hm2", "hm3", "hm4", "xg12", "xg34", "xg56", "xg78"),
        ):
            self.assertEqual(
                module.extract_unit_cost_price_values(payload),
                [
                    394.3278,
                    395.4994,
                    424.7904,
                    341.406,
                    344.8669,
                    343.7133,
                    342.5596,
                    620.3662,
                    620.3662,
                    620.3662,
                    620.3662,
                ],
            )


if __name__ == "__main__":
    unittest.main()
