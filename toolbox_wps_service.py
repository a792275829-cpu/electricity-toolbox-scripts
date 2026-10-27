"""WPS mappings, workbook reads and verified online writes; no Tk dependency."""
from __future__ import annotations
import json
import os
import signal
import subprocess
import sys
import threading
import time as time_module
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
except ImportError:  # pragma: no cover - surfaced in GUI
    load_workbook = None
    get_column_letter = None
    range_boundaries = None
    coordinate_from_string = None
    column_index_from_string = None

try:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright
except ImportError:  # pragma: no cover - surfaced in GUI
    PlaywrightTimeoutError = Exception
    sync_playwright = None

from toolbox_wps_paths import (
    BASE_DIR,
    DATA_DIR,
    CONFIG_PATH,
    DEFAULT_CONFIG_PATH,
    PROFILE_DIR,
    WPS_CHROME_PID_PATH,
    WPS_CHROME_STATE_PATH,
    WPS_CHROME_LAUNCH_VERSION,
    WPS_CHROME_START_LOCK,
    LOG_DIR,
    MAX_RECENT_URLS,
    DOCUMENT_CONFIGS_MIN_HEIGHT,
    CDP_CONNECT_TIMEOUT_MS,
    WPS_LOGIN_STATUS_URL,
    LOGIN_STATUS_POLL_MS,
)
from toolbox_wps_config import (CONFIG_EXPORT_SCHEMA, CONFIG_EXPORT_VERSION, _document_config_from_dict, _region_config_from_dict, _string_list, build_config_export, load_config_export, normalize_runtime_config, runtime_config_from_export_payload, save_config_export)
from toolbox_wps_values import (
    normalize_cell_value,
    normalize_for_compare,
    values_equivalent,
    parse_start_cell,
    normalize_cell_address,
    range_from_cells,
    range_shape,
    split_range,
    target_range_from_start,
    target_rect,
    rect_from_cells,
    rects_overlap,
    coerce_wps_matrix,
    compare_matrix,
)

from toolbox_wps_browser import (
    browser_has_page,
    ensure_background_cdp_target,
    find_or_open_page,
    cdp_candidates,
    resolve_cdp_endpoint,
    find_macos_chrome,
    managed_cdp_host_port,
    build_wps_chrome_command,
    try_resolve_cdp_endpoint,
    cdp_page_targets,
    ensure_wps_cdp_page,
    managed_chrome_state_version,
    managed_chrome_command,
    is_managed_chrome_command,
    discover_managed_chrome_pid,
    remember_managed_chrome,
    clear_managed_chrome_state,
    managed_chrome_pid,
    managed_chrome_launch_version,
    managed_chrome_is_running,
    stop_managed_chrome,
    managed_chrome_is_headless,
    stop_outdated_managed_chrome,
    _ensure_wps_dedicated_chrome,
    ensure_wps_dedicated_chrome,
    open_browser_context,
)

WPS_APPLICATION_GUARD = """
  const api = window.WPSOpenApi;
  if (typeof api?.EtApplication !== "function") {
    throw new Error("WPS_API_NOT_READY: 操作前表格连接已变化");
  }
  const app = api.EtApplication();
  if (!app || typeof app.Worksheets !== "function") {
    throw new Error("WPS_API_NOT_READY: 表格尚未就绪");
  }
"""


@dataclass
class MappingRow:
    kdocs_url: str
    local_file: str
    source_sheet: str
    source_start: str
    source_end: str
    target_sheet: str
    target_start: str
    target_end: str
    source_type: str = "excel"
    source_url: str = ""



@dataclass
class RegionRow:
    source_sheet: str
    source_start: str
    source_end: str
    target_sheet: str
    target_start: str
    target_end: str



@dataclass
class DocumentConfig:
    name: str
    kdocs_url: str
    local_file: str
    regions: list[RegionRow]
    source_type: str = "excel"
    source_url: str = ""



@dataclass
class PreparedMapping:
    index: int
    kdocs_url: str
    workbook_path: str
    source_sheet: str
    source_range: str
    target_sheet: str
    target_start: str
    target_range: str
    rows: int
    cols: int
    values: list[list[Any]]



@dataclass(frozen=True)
class WpsLoginStatus:
    state: str
    message: str



def read_wps_login_status(context) -> WpsLoginStatus:
    """Confirm the current browser session with WPS, without exposing account data."""
    response = context.request.get(WPS_LOGIN_STATUS_URL, timeout=10000)
    data = response.json()
    if not isinstance(data, dict):
        return WpsLoginStatus("unknown", "暂时无法确认，请稍后刷新登录状态。")
    if response.status == 200 and data.get("result") == "ok" and data.get("userid"):
        return WpsLoginStatus("logged_in", "已登录")
    if data.get("result") == "userNotLogin" or response.status == 401:
        return WpsLoginStatus("logged_out", "未登录或登录已失效，请点击“登录 WPS”。")
    return WpsLoginStatus("unknown", "暂时无法确认，请检查网络后刷新登录状态。")



def check_wps_login_status(browser_mode: str, cdp_url: str, log) -> WpsLoginStatus:
    # A copied browser profile is not proof of login: cookies may be encrypted
    # for the previous Mac. Existing profiles must be verified online.
    if browser_mode != "cdp" and not PROFILE_DIR.is_dir():
        return WpsLoginStatus("logged_out", "本机尚未登录，请点击“登录 WPS”。")
    if sync_playwright is None:
        raise RuntimeError("缺少浏览器自动化组件，请先运行工具箱初始化程序。")
    with sync_playwright() as playwright:
        _browser, context, close_context = open_browser_context(
            playwright, browser_mode, cdp_url, log, for_login_check=True
        )
        try:
            return read_wps_login_status(context)
        finally:
            if close_context:
                context.close()



def load_config() -> dict[str, Any]:
    path = CONFIG_PATH if CONFIG_PATH.exists() else DEFAULT_CONFIG_PATH
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}



def save_config(data: dict[str, Any]) -> None:
    CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")



def update_recent_urls(urls: list[str], url: str) -> list[str]:
    normalized = url.strip()
    if not normalized:
        return urls[:MAX_RECENT_URLS]
    result = [normalized]
    for item in urls:
        item = item.strip()
        if item and item != normalized and item not in result:
            result.append(item)
        if len(result) >= MAX_RECENT_URLS:
            break
    return result



def load_excel_workbook(workbook_path: Path, **kwargs):
    """Open an OOXML workbook with a useful error for partial/corrupt files."""
    if workbook_path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        try:
            is_zip = zipfile.is_zipfile(workbook_path)
        except OSError as exc:
            raise OSError(f"无法读取表格文件：{workbook_path}") from exc
        if not is_zip:
            try:
                with workbook_path.open("rb") as source:
                    starts_as_zip = source.read(4).startswith(b"PK")
            except OSError:
                starts_as_zip = False
            if starts_as_zip:
                reason = (
                    "表格文件尚未生成完整。"
                    "请重新生成汇总表后再选择。"
                )
            else:
                reason = (
                    "所选文件的内容与表格格式不符。"
                    "请用表格软件打开，另存为 .xlsx 文件后重试。"
                )
            raise ValueError(f"{reason}\n\n文件：{workbook_path}")
    try:
        return load_workbook(workbook_path, **kwargs)
    except zipfile.BadZipFile as exc:
        raise ValueError(
            "表格文件不完整或已损坏，请重新生成后再试。"
            f"\n\n文件：{workbook_path}"
        ) from exc



def read_excel_range(workbook_path: Path, sheet_name: str, source_range: str) -> tuple[list[list[Any]], int, int]:
    return read_excel_ranges(workbook_path, [(sheet_name, source_range)])[0]



def read_excel_ranges(
    workbook_path: Path, requests: list[tuple[str, str]]
) -> list[tuple[list[list[Any]], int, int]]:
    """Read several regions of a workbook with one scan per source sheet."""
    if load_workbook is None or range_boundaries is None:
        raise RuntimeError("缺少表格读取组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到表格文件：{workbook_path}")
    if not requests:
        return []

    wb = load_excel_workbook(workbook_path, data_only=True, read_only=True)
    try:
        grouped: dict[str, list[tuple[int, tuple[int, int, int, int]]]] = {}
        for index, (sheet_name, source_range) in enumerate(requests):
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"找不到来源工作表：{sheet_name}")
            grouped.setdefault(sheet_name, []).append((index, range_boundaries(source_range)))

        results: list[tuple[list[list[Any]], int, int] | None] = [None] * len(requests)
        for sheet_name, regions in grouped.items():
            first_row = min(bounds[1] for _, bounds in regions)
            last_row = max(bounds[3] for _, bounds in regions)
            first_col = min(bounds[0] for _, bounds in regions)
            last_col = max(bounds[2] for _, bounds in regions)
            collected: dict[int, list[list[Any]]] = {index: [] for index, _ in regions}
            for row_number, row in enumerate(
                wb[sheet_name].iter_rows(
                    min_row=first_row,
                    max_row=last_row,
                    min_col=first_col,
                    max_col=last_col,
                    values_only=True,
                ),
                start=first_row,
            ):
                for index, (min_col, min_row, max_col, max_row) in regions:
                    if min_row <= row_number <= max_row:
                        cells = row[min_col - first_col : max_col - first_col + 1]
                        collected[index].append([normalize_cell_value(cell) for cell in cells])
            for index, (min_col, min_row, max_col, max_row) in regions:
                results[index] = (
                    collected[index],
                    max_row - min_row + 1,
                    max_col - min_col + 1,
                )
        if any(result is None for result in results):
            raise RuntimeError("未能完整读取来源区域，请检查表格文件后重试。")
        return [result for result in results if result is not None]
    finally:
        wb.close()



def read_excel_sheet_names(workbook_path: Path) -> list[str]:
    if load_workbook is None:
        raise RuntimeError("缺少表格读取组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到表格文件：{workbook_path}")
    wb = load_excel_workbook(workbook_path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()



def prepare_mappings(
    mappings: list[MappingRow],
    browser_mode: str = "persistent",
    cdp_url: str = "http://127.0.0.1:9222",
    log=None,
) -> list[PreparedMapping]:
    if not mappings:
        raise ValueError("尚未配置写入区域。")

    source_specs: list[tuple[int, MappingRow, str, Path | None, str, str, int, int]] = []
    excel_groups: dict[Path, list[tuple[int, str, str]]] = {}
    target_rects: dict[tuple[str, str], list[tuple[int, tuple[int, int, int, int], str]]] = {}
    log = log or (lambda _message: None)
    for idx, mapping in enumerate(mappings, start=1):
        source_type = mapping.source_type or "excel"
        workbook_path = None
        if source_type == "excel":
            workbook_path_text = mapping.local_file.strip()
            if not workbook_path_text:
                raise ValueError(f"第 {idx} 个写入区域未选择本地文件。")
            workbook_path = Path(workbook_path_text).expanduser()
        elif source_type == "kdocs":
            if not mapping.source_url.strip():
                raise ValueError(f"第 {idx} 个写入区域未填写来源表格链接。")
        else:
            raise ValueError(f"第 {idx} 个写入区域的数据来源类型无法识别：{source_type}")
        if not all(
            [
                mapping.kdocs_url,
                mapping.source_sheet,
                mapping.source_start,
                mapping.source_end,
                mapping.target_sheet,
                mapping.target_start,
                mapping.target_end,
            ]
        ):
            raise ValueError(f"第 {idx} 个写入区域的信息未填完整。")
        source_range = range_from_cells(mapping.source_start, mapping.source_end)
        target_range = range_from_cells(mapping.target_start, mapping.target_end)
        source_rows, source_cols = range_shape(mapping.source_start, mapping.source_end)
        target_rows, target_cols = range_shape(mapping.target_start, mapping.target_end)
        if (source_rows, source_cols) != (target_rows, target_cols):
            raise ValueError(
                f"第 {idx} 个写入区域大小不一致：来源为 {source_rows} 行 {source_cols} 列，"
                f"目标为 {target_rows} 行 {target_cols} 列。"
            )
        rect = rect_from_cells(mapping.target_start, mapping.target_end)
        target_key = (mapping.kdocs_url, mapping.target_sheet)
        for other_idx, other_rect, other_range in target_rects.get(target_key, []):
            if rects_overlap(rect, other_rect):
                raise ValueError(
                    f"目标表格 {mapping.kdocs_url} 的工作表“{mapping.target_sheet}”存在重叠区域："
                    f"第 {other_idx} 个区域 {other_range} 与第 {idx} 个区域 {target_range}。"
                )
        target_rects.setdefault(target_key, []).append((idx, rect, target_range))
        source_specs.append(
            (idx, mapping, source_type, workbook_path, source_range, target_range, source_rows, source_cols)
        )
        if workbook_path is not None:
            excel_groups.setdefault(workbook_path, []).append((idx, mapping.source_sheet, source_range))

    excel_values: dict[int, tuple[list[list[Any]], int, int]] = {}
    for workbook_path, requests in excel_groups.items():
        results = read_excel_ranges(
            workbook_path, [(sheet_name, source_range) for _, sheet_name, source_range in requests]
        )
        for (idx, _, _), result in zip(requests, results):
            excel_values[idx] = result

    online_specs = [spec for spec in source_specs if spec[2] != "excel"]
    online_values = {}
    if online_specs:
        results = read_online_ranges(
            [(mapping.source_url, mapping.source_sheet, source_range, rows, cols)
             for _, mapping, _, _, source_range, _, rows, cols in online_specs],
            browser_mode, cdp_url, log,
        )
        online_values = {spec[0]: values for spec, values in zip(online_specs, results)}

    prepared: list[PreparedMapping] = []
    for idx, mapping, source_type, workbook_path, source_range, target_range, source_rows, source_cols in source_specs:
        if source_type == "excel":
            values, rows, cols = excel_values[idx]
            source_label = str(workbook_path)
        else:
            values = online_values[idx]
            rows, cols = source_rows, source_cols
            source_label = mapping.source_url
        prepared.append(
            PreparedMapping(
                index=idx,
                kdocs_url=mapping.kdocs_url,
                workbook_path=source_label,
                source_sheet=mapping.source_sheet,
                source_range=source_range,
                target_sheet=mapping.target_sheet,
                target_start=mapping.target_start,
                target_range=target_range,
                rows=rows,
                cols=cols,
                values=values,
            )
        )
    return prepared



def wait_for_wps_ready(page, log, timeout_seconds: int) -> dict[str, Any]:
    deadline = time_module.monotonic() + timeout_seconds
    last_state: dict[str, Any] = {"ready": False, "url": "", "title": ""}
    last_log_second = 0
    readiness_script = """
    async () => {
      const api = window.WPSOpenApi;
      if (typeof api?.EtApplication !== "function") {
        return {
          ready: false,
          url: location.href,
          title: document.title,
          hasWPSOpenApi: !!api
        };
      }

      const bounded = (value, milliseconds = 750) => Promise.race([
        Promise.resolve(value),
        new Promise(resolve => setTimeout(() => resolve(false), milliseconds))
      ]);
      let app;
      let applicationReady = false;
      let documentReady = false;
      try {
        app = api.EtApplication();
        applicationReady = Boolean(await bounded(app.Ready));
      } catch (_error) {}
      if (!applicationReady && api.documentReadyPromise) {
        try {
          documentReady = Boolean(await bounded(
            Promise.resolve(api.documentReadyPromise).then(() => true)
          ));
        } catch (_error) {}
      }
      return {
        ready: !!app && typeof app.Worksheets === "function"
          && (applicationReady || documentReady) && window.WPSOpenApi === api,
        url: location.href,
        title: document.title,
        hasSave: typeof api.save === "function"
      };
    }
    """
    while time_module.monotonic() < deadline:
        try:
            state = page.evaluate(readiness_script)
            last_state = state
            if state.get("ready"):
                return state
        except Exception as exc:
            last_state = {"ready": False, "url": page.url, "title": "", "error": str(exc)}

        remaining = int(deadline - time_module.monotonic())
        elapsed_bucket = timeout_seconds - remaining
        if elapsed_bucket // 15 > last_log_second // 15:
            last_log_second = elapsed_bucket
            log(f"仍在等待表格加载，当前页面：{page.url}")
        page.wait_for_timeout(1000)
    return last_state



def evaluate_wps_operation(page, script: str, payload: Any, log, timeout_seconds: int):
    """Recheck readiness after navigation, including the gap before evaluate."""
    first_error: Exception | None = None
    for attempt in range(1, 4):
        ready = wait_for_wps_ready(
            page, log, timeout_seconds=timeout_seconds if attempt == 1 else 60
        )
        if not ready.get("ready"):
            message = (
                "表格加载超时。"
                "请点击“登录 WPS”，确认已登录且能访问该表格后重试。"
                f"当前页面：{ready.get('title')} {ready.get('url')}"
            )
            if first_error is not None:
                message += f" 首次操作失败原因：{first_error}"
            raise RuntimeError(message) from first_error
        try:
            return page.evaluate(script, payload)
        except Exception as exc:
            # The guard runs before any cell assignment. Do not retry arbitrary
            # TypeErrors or API errors that may have occurred during a write.
            if not any(marker in str(exc) for marker in (
                "WPS_API_NOT_READY:", "Execution context was destroyed"
            )) and "navigation" not in str(exc).lower():
                raise
            if first_error is None:
                first_error = exc
            if attempt == 3:
                raise RuntimeError(
                    f"重试 3 次后仍无法操作表格。首次失败原因：{first_error}"
                ) from exc
            log(f"表格连接已变化，等待恢复后重试（第 {attempt + 1}/3 次）。")
            page.wait_for_timeout(1500)



def write_to_wps(
    url: str,
    prepared: list[PreparedMapping],
    log,
    browser_mode: str = "persistent",
    cdp_url: str = "http://127.0.0.1:9222",
) -> dict[str, Any]:
    if sync_playwright is None:
        raise RuntimeError("缺少浏览器自动化组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")

    payload = [
        {
            "index": item.index,
            "targetSheet": item.target_sheet,
            "targetRange": item.target_range,
            "rows": item.rows,
            "cols": item.cols,
            "values": item.values,
        }
        for item in prepared
    ]

    PROFILE_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        try:
            browser, context, close_context = open_browser_context(
                p, browser_mode, cdp_url, log, initial_url=url
            )
        except Exception as exc:
            if browser_mode == "cdp":
                raise RuntimeError(
                    "无法连接已打开的浏览器，请使用独立的用户资料目录和远程调试端口启动，"
                    "并在该窗口登录金山文档。"
                ) from exc
            raise
        try:
            log(f"正在打开目标表格：{url}")
            page = find_or_open_page(context, url)

            log("正在等待表格加载；如出现登录页面，请先在专用浏览器中完成登录。")
            writer_script = "async (payload) => {" + WPS_APPLICATION_GUARD + """
              if (typeof api.save !== "function") {
                throw new Error("WPS_API_NOT_READY: 表格保存功能尚未就绪");
              }
              const results = [];
              for (const item of payload) {
                const sheet = app.Worksheets(item.targetSheet);
                const range = sheet.Range(item.targetRange);
                const assignment = item.rows === 1 && item.cols === 1 ? item.values[0][0] : item.values;
                const mergeCells = await range.MergeCells;
                range.Value = assignment;
                results.push({
                  index: item.index,
                  targetSheet: item.targetSheet,
                  targetRange: item.targetRange,
                  rows: item.rows,
                  cols: item.cols,
                  mergeCells
                });
              }
              await api.save();
              for (const result of results) {
                const sheet = app.Worksheets(result.targetSheet);
                const range = sheet.Range(result.targetRange);
                result.value = await range.Value;
                result.value2 = await range.Value2;
                result.text = await range.Text;
              }
              return { saved: true, results };
            }
            """
            log(f"正在写入 {len(payload)} 个区域…")
            return evaluate_wps_operation(page, writer_script, payload, log, 180)
        finally:
            if close_context:
                context.close()



def read_online_document_info(url: str, browser_mode: str, cdp_url: str, log) -> dict[str, Any]:
    if sync_playwright is None:
        raise RuntimeError("缺少浏览器自动化组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")

    started_at = time_module.monotonic()
    log("正在读取表格名称和工作表列表…")
    with sync_playwright() as p:
        browser, context, close_context = open_browser_context(
            p, browser_mode, cdp_url, log, initial_url=url
        )
        try:
            page = find_or_open_page(context, url)
            log(
                f"已打开表格，用时 {time_module.monotonic() - started_at:.1f} 秒；"
                "正在等待表格就绪…"
            )
            script = "async () => {" + WPS_APPLICATION_GUARD + """
              const count = await app.Worksheets.Count;
              const sheets = [];
              for (let i = 1; i <= Number(count); i++) {
                const ws = app.Worksheets(i);
                sheets.push(await ws.Name);
              }
              return {title: document.title, sheets};
            }
            """
            info = evaluate_wps_operation(page, script, None, log, 45)
            sheets = info["sheets"]
            title = info["title"] or page.title()
            log(
                f"表格信息读取完成，用时 {time_module.monotonic() - started_at:.1f} 秒，"
                f"共 {len(sheets)} 个工作表。"
            )
            return {"title": title, "sheets": sheets}
        finally:
            if close_context:
                context.close()



def read_online_sheet_names(url: str, browser_mode: str, cdp_url: str, log) -> list[str]:
    return read_online_document_info(url, browser_mode, cdp_url, log)["sheets"]



def read_online_range_values(
    url: str,
    sheet_name: str,
    range_text: str,
    rows: int,
    cols: int,
    browser_mode: str,
    cdp_url: str,
    log,
) -> list[list[Any]]:
    return read_online_ranges(
        [(url, sheet_name, range_text, rows, cols)], browser_mode, cdp_url, log
    )[0]



def read_online_ranges(
    requests: list[tuple[str, str, str, int, int]],
    browser_mode: str,
    cdp_url: str,
    log,
) -> list[list[list[Any]]]:
    """Read source regions in order using one connection and browser context."""
    if not requests:
        return []
    if sync_playwright is None:
        raise RuntimeError("缺少浏览器自动化组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")

    with sync_playwright() as p:
        browser, context, close_context = open_browser_context(
            p, browser_mode, cdp_url, log, initial_url=requests[0][0]
        )
        try:
            script = "async (payload) => {" + WPS_APPLICATION_GUARD + """
              const {sheetName, rangeText} = payload;
              const sheet = await app.Worksheets(sheetName);
              const range = await sheet.Range(rangeText);
              return await range.Value;
            }
            """
            results = []
            for url, sheet_name, range_text, rows, cols in requests:
                page = find_or_open_page(context, url)
                raw = evaluate_wps_operation(
                    page, script, {"sheetName": sheet_name, "rangeText": range_text}, log, 120
                )
                matrix = coerce_wps_matrix(raw, rows, cols)
                results.append([[normalize_cell_value(cell) for cell in row] for row in matrix])
            return results
        finally:
            if close_context:
                context.close()
