"""Build a self-contained resource payload on the native release runner."""

from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
import os
import platform
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location(
    "desktop_build_support", Path(__file__).with_name("build_support.py")
)
build_support = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build_support)


def run(args, **kwargs):
    return subprocess.run([str(x) for x in args], check=True, **kwargs)


MODULE_DIRECTORIES = {
    "电力工具箱", "上网电量抓取", "风力预测截图", "市场表更新", "现货价格更新",
    "交易系统信息更新", "功率调整收益统计表更新", "集团每日上传", "广东电价预测",
    "广东运行日报PDF解析", "南方分公司电力市场日报", "每日生产经营情况汇报自动生成工具",
    "500kV断面分析", "wps自动", "电力工具脚本",
}

# Fixed runtime inputs, including templates used indirectly by renderers.
REQUIRED_RUNTIME_ASSETS = tuple(map(Path, (
    "500kV断面分析/500kV节点拓扑.md",
    "500kV断面分析/data/500kv_node_positions.json",
    "500kV断面分析/data/500kv_node_position_candidates.json",
    "500kV断面分析/接线图.png",
    "500kV断面分析/tools/anchor_calibrator/index.html",
    "500kV断面分析/tools/anchor_calibrator/app.js",
    "500kV断面分析/tools/anchor_calibrator/styles.css",
    "风力预测截图/fonts/NotoSerifSC.ttf",
    "每日生产经营情况汇报自动生成工具/模板/每日生产经营情况汇报（自动）.docx",
    "每日生产经营情况汇报自动生成工具/scripts/load_compare_subject_map.json",
    "wps自动/wps_excel_to_kdocs_config.defaults.json",
    "广东运行日报PDF解析/gd_pdf_tool/recognize_image_labels.js",
    "南方分公司电力市场日报/gd_power_report/templates/legacy_print_report.css",
    "南方分公司电力市场日报/gd_power_report/templates/legacy_concise_report.css",
    "南方分公司电力市场日报/gd_power_report/templates/legacy_print_report.html.j2",
    "南方分公司电力市场日报/gd_power_report/templates/legacy_concise_report.html.j2",
    "南方分公司电力市场日报/gd_power_report/templates/report_logo.html.j2",
    "南方分公司电力市场日报/gd_power_report/templates/metric_card_layout.js.j2",
    "desktop/assets/app-icon-256.png",
)))
FORECAST_SEED = Path("广东电价预测/输出/广东电价数据总表.xlsx")


def validate_runtime_assets(source: Path) -> None:
    for relative in REQUIRED_RUNTIME_ASSETS:
        path = source / relative
        if not path.is_file() or not path.stat().st_size or (
            path.suffix == ".md" and not path.read_text(encoding="utf-8").strip()
        ):
            raise FileNotFoundError(f"打包缺少必需功能资源或文件为空：{relative}")


def prepare_forecast_seed(defaults: Path, workbook: Path | None = None) -> Path:
    """Ship a byte-preserving initial workbook; never manage users' live data."""
    from migrate_forecast_data import copy_atomic, resolve_paths
    from openpyxl import load_workbook

    if workbook is None:
        configured = os.environ.get("FORECAST_MASTER_WORKBOOK")
        url = os.environ.get("FORECAST_MASTER_URL")
        if configured:
            workbook = Path(configured).expanduser()
        elif url:
            expected = os.environ.get("FORECAST_MASTER_SHA256", "")
            if len(expected) != 64 or any(c not in "0123456789abcdef" for c in expected):
                raise ValueError("下载预测总表必须设置有效的 FORECAST_MASTER_SHA256")
            workbook = build_support.download(
                url, ROOT / "dist/tauri-build/forecast" / expected / "master.xlsx", expected,
            )
        else:
            workbook = resolve_paths(ROOT / "广东电价预测")["masterWorkbook"]
    workbook = Path(workbook)
    if not workbook.is_file():
        raise FileNotFoundError("打包缺少电价预测总数据表，请提供 --forecast-workbook 或 FORECAST_MASTER_WORKBOOK")
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = book["整合数据"] if "整合数据" in book.sheetnames else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        if not {"日期", "时刻", "全省日前平均电价"}.issubset(headers) or next(rows, None) is None:
            raise ValueError("预测总表缺少必要字段或数据行")
    finally:
        book.close()
    target = defaults / FORECAST_SEED
    copy_atomic(workbook, target)
    return target


def runtime_source(relative: Path) -> bool:
    if not build_support.exportable(relative):
        return False
    if relative.as_posix() in {"desktop/python/native_bridge.py", "desktop/python/service_bridge.py", "desktop/python/task_worker.py", "desktop/python/remote_updater.py", "desktop/python/github_update.py", "desktop/python/delta_update.py", "desktop/python/installer_image.py", "desktop/python/update_ui.py", "desktop/assets/app-icon-256.png"}:
        return True
    if len(relative.parts) == 1:
        return relative.name.startswith("toolbox_") and relative.suffix == ".py"
    return (relative.parts[0] in MODULE_DIRECTORIES
            and not {"tests", "docs", "scripts-dev"}.intersection(relative.parts)
            and not relative.name.startswith("test_")
            and relative.suffix.lower() != ".command")


def prepare_defaults(source: Path, defaults: Path) -> None:
    """Build first-install settings separately from mutable user configuration."""
    shutil.rmtree(defaults, ignore_errors=True)
    defaults.mkdir(parents=True)
    for example in source.rglob("config.example.json"):
        raw = json.loads(example.read_text(encoding="utf-8"))
        for key in ("username", "password"):
            if key in raw:
                raw[key] = ""
        example.write_text(
            json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        path = defaults / example.relative_to(source).with_name("config.json")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
    wps_source = source / "wps自动/wps_excel_to_kdocs_config.defaults.json"
    wps_target = defaults / "wps自动/wps_excel_to_kdocs_config.json"
    wps_target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(wps_source, wps_target)
    config = defaults / "电力工具脚本/private-data-uploader-tool/upload.config.json"
    config.parent.mkdir(parents=True, exist_ok=True)
    config.write_text(
        json.dumps(
            {"baseUrl": "https://xhxt.chng.com.cn", "sourceDir": ".", "headless": True}
        ),
        encoding="utf-8",
    )


def prepare_price_xlsx(source: Path, target: str) -> Path:
    """Always build the helper from this checkout, never ship a stale binary."""
    binary = source / "500kV断面分析/bin/toolbox-price-xlsx"
    run([sys.executable, ROOT / "native/price-xlsx/build.py", "--target", target,
         "--output", binary])
    if not binary.is_file() or not binary.stat().st_size:
        raise FileNotFoundError("500kV Rust Excel 后端构建未生成可执行文件")
    return binary


def prepare(target: str, runtime_from: Path | None = None, forecast_workbook: Path | None = None):
    native = (
        "aarch64-apple-darwin"
        if sys.platform == "darwin" and platform.machine() == "arm64"
        else None
    )
    if target != "aarch64-apple-darwin":
        raise ValueError("仅支持 Apple 芯片 Mac 构建目标")
    if target != native and not runtime_from:
        raise RuntimeError(
            "运行环境必须在对应原生系统上构建，或提供已验证的 --runtime-from"
        )
    payload = ROOT / "desktop/src-tauri/payload"
    cache = ROOT / "dist/tauri-build"
    cache.mkdir(parents=True, exist_ok=True)
    source = payload / "app"
    if source.exists():
        shutil.rmtree(source)
    source.mkdir(parents=True)
    build_support.export_source(source, predicate=runtime_source)
    prepare_price_xlsx(source, target)
    validate_runtime_assets(source)
    defaults = payload / "defaults"
    prepare_defaults(source, defaults)
    prepare_forecast_seed(defaults, forecast_workbook)
    runtime = payload / "runtime"
    marker = runtime / "target.txt"
    if not marker.is_file() or marker.read_text() != target:
        if runtime.exists():
            shutil.rmtree(runtime)
        runtime.mkdir()
        if runtime_from:
            for name in ("python", "node", "ms-playwright"):
                origin = runtime_from / name
                shutil.copytree(origin, runtime / name, symlinks=True)
        else:
            uv = shutil.which("uv")
            if not uv:
                raise RuntimeError("构建机需要 uv")
            identity = "cpython-3.11.15-macos-aarch64-none"
            downloads = cache / "python-download"
            run(
                [
                    uv,
                    "python",
                    "install",
                    identity,
                    "--install-dir",
                    downloads,
                    "--no-bin",
                ]
            )
            shutil.copytree(downloads / identity, runtime / "python", symlinks=True)
            python = runtime / "python/bin/python3"
            run(
                [
                    python,
                    "-m",
                    "pip",
                    "install",
                    "--break-system-packages",
                    "--disable-pip-version-check",
                    "--no-compile",
                    "--no-deps",
                    "--require-hashes",
                    "--only-binary=:all:",
                    "-r",
                    ROOT / "packaging/tauri/requirements-macos.lock",
                ]
            )
            node_name = f"node-v{build_support.NODE_VERSION}-darwin-arm64"
            archive_name = node_name + ".tar.gz"
            base = f"https://nodejs.org/dist/v{build_support.NODE_VERSION}/"
            checksums = build_support.download(
                base + "SHASUMS256.txt",
                cache / f"node-{build_support.NODE_VERSION}-SHASUMS256.txt",
            )
            expected = next(
                line.split()[0]
                for line in checksums.read_text().splitlines()
                if line.split()[-1] == archive_name
            )
            archive = build_support.download(
                base + archive_name, cache / archive_name, expected
            )
            unpack = cache / "node-unpack"
            shutil.rmtree(unpack, ignore_errors=True)
            unpack.mkdir()
            build_support.extract(archive, unpack)
            shutil.move(unpack / node_name, runtime / "node")
            env = {
                **os.environ,
                "PLAYWRIGHT_BROWSERS_PATH": str(runtime / "ms-playwright"),
            }
            run([python, "-m", "playwright", "install", "chromium"], env=env)
        marker.write_text(target)
    from prepare_tk import prepare as prepare_tk

    prepare_tk(runtime / "python", cache)
    uploader = source / "电力工具脚本/private-data-uploader-tool"
    package = json.loads((uploader / "package.json").read_text())
    package["dependencies"]["playwright"] = build_support.PLAYWRIGHT_VERSION
    (uploader / "package.json").write_text(
        json.dumps(package, indent=2), encoding="utf-8"
    )
    (uploader / "package-lock.json").unlink(missing_ok=True)
    npm = shutil.which("npm")
    run(
        [
            npm,
            "install",
            "--ignore-scripts",
            "--omit=optional",
            "--no-audit",
            "--no-fund",
        ],
        cwd=uploader,
    )
    python = runtime / "python/bin/python3"
    registry = (
        runtime
        / "python"
        / "lib/python3.11/site-packages"
        / "playwright/driver/package/browsers.json"
    )
    if json.loads(registry.read_text()) != json.loads(
        (uploader / "node_modules/playwright-core/browsers.json").read_text()
    ):
        raise RuntimeError("Python / Node Playwright 浏览器版本不一致")
    bootstrap = payload / "bootstrap"
    bootstrap.mkdir(exist_ok=True)
    for name in ("prepare_workspace.py", "workspace_layout.py"):
        shutil.copy2(Path(__file__).with_name(name), bootstrap / name)
    # Remove only the former generated layout when reusing a local staging folder.
    shutil.rmtree(payload / "workspace", ignore_errors=True)
    (payload / "prepare_workspace.py").unlink(missing_ok=True)
    for path in list(payload.rglob("__pycache__")):
        shutil.rmtree(path)
    version = json.loads((ROOT / "desktop/package.json").read_text())["version"]
    files = {
        p.relative_to(source).as_posix(): build_support.sha256(p)
        for p in sorted(source.rglob("*"))
        if p.is_file()
    }
    default_hashes = {
        p.relative_to(defaults).as_posix(): build_support.sha256(p)
        for p in sorted(defaults.rglob("*")) if p.is_file()
    }
    build = hashlib.sha256(json.dumps(
        {"files": files, "default_hashes": default_hashes}, sort_keys=True,
    ).encode()).hexdigest()
    manifest = {
        "layout": 2,
        "version": version,
        "target": target,
        "build": build,
        "files": files,
        "defaults": list(default_hashes),
        "default_hashes": default_hashes,
    }
    (payload / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    build_support.normalize_permissions(payload)
    if target == native:
        run(
            [
                python,
                "-B",
                "-c",
                'import openpyxl,playwright,pandas,sklearn,tkinter,PIL,pdfplumber,pypdfium2,docx; print("Bundled Python dependencies OK")',
            ]
        )
    print(
        f"Payload ready: {payload} ({len(files)} managed files, {target})", flush=True
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--target",
        required=True,
        choices=["aarch64-apple-darwin"],
    )
    parser.add_argument("--runtime-from", type=Path)
    parser.add_argument("--forecast-workbook", type=Path, help="随包首次安装的预测总表（不写入 Git）")
    args = parser.parse_args()
    prepare(args.target, args.runtime_from, args.forecast_workbook)
