from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from toolbox.runtime import ToolPaths, load_module


class SummaryStationFolderTests(unittest.TestCase):
    def test_grouped_filenames_from_station_or_parent_folder(self):
        module = load_module(
            "summary_station_folder_test",
            ToolPaths(Path(__file__).resolve().parents[2]).summary,
        )
        with TemporaryDirectory() as directory:
            parent = Path(directory) / "9-28"
            station = parent / "海门"
            station.mkdir(parents=True)
            for market, label, date, price in (
                ("dayahead", "日前", "2026-09-28", "711.47"),
                ("realtime", "实时", "2026-09-26", "625.33"),
            ):
                for pair in ((1, 2), (3, 4)):
                    workbook = Workbook()
                    workbook.remove(workbook.active)
                    for number in pair:
                        sheet = workbook.create_sheet(f"海门#{number}机组")
                        sheet.append(["时刻", "电量(MWh)", "标签", "出清电价", "结算电价"])
                        for hour in range(24):
                            sheet.append([f"{hour:02d}:00", f"{number + hour}.125", "", "1.00", price])
                    workbook.save(station / f"{pair[0]},{pair[1]}号机{label}交易结果查询({date}).xlsx")
            (station / "~$1,2号机日前交易结果查询(2026-09-28).xlsx").write_bytes(b"lock")
            for input_dir in (station, parent):
                with self.subTest(input_dir=input_dir.name):
                    units = module.collect_units(input_dir, lambda _: None)
                    self.assertEqual([unit.unit_name for unit in units], [f"{n}号机" for n in range(1, 5)])
                    self.assertTrue(all(unit.station_name == "海门" for unit in units))
                    self.assertTrue(all(set(unit.values) == {"dayahead", "realtime"} for unit in units))
                    output = module.run_summary(input_dir, input_dir / "汇总.xlsx")
                    reopened = load_workbook(output)
                    try:
                        self.assertEqual(reopened.sheetnames, ["海门"])
                        sheet = reopened["海门"]
                        self.assertEqual(sheet["B3"].value, "日前电量\n（MWh）")
                        self.assertEqual(sheet["K3"].value, "实时电量\n（MWh）")
                        for number in range(1, 5):
                            for base, price in ((2, 711.47), (11, 625.33)):
                                col = base + (number - 1) * 2
                                self.assertEqual(sheet.cell(2, col).value, f"{number}号机")
                                self.assertEqual(sheet.cell(4, col).value, number + .125)
                                self.assertEqual(sheet.cell(27, col).value, number + 23.125)
                                self.assertEqual(sheet.cell(4, col).number_format, "0.000")
                                self.assertEqual(sheet.cell(4, col + 1).value, price)
                                self.assertEqual(sheet.cell(4, col + 1).number_format, "0.00")
                    finally:
                        reopened.close()
            for separator, space in (("、", ""), ("，", ""), ("、", " ")):
                for path in station.glob("*.xlsx"):
                    if path.name.startswith("~$") or "交易结果查询" not in path.name:
                        continue
                    # Rename the same exports to verify punctuation and spacing
                    # do not change which internal unit worksheets are read.
                    suffix = path.name[path.name.index("号机") + 2:]
                    pair = (1, 2) if path.name.startswith("1") else (3, 4)
                    path.rename(station / f"{pair[0]}{separator}{pair[1]}{space}号机{suffix}")
                with self.subTest(separator=separator, space=space):
                    units = module.collect_units(parent, lambda _: None)
                    self.assertEqual([unit.unit_name for unit in units], [f"{n}号机" for n in range(1, 5)])
                    for unit in units:
                        self.assertEqual(set(unit.values), {"dayahead", "realtime"})
                        self.assertEqual(unit.values["dayahead"]["price"][0].value, 711.47)
                        self.assertEqual(unit.values["realtime"]["price"][0].value, 625.33)


if __name__ == "__main__":
    unittest.main()
