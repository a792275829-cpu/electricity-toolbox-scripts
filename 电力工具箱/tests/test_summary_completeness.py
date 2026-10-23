from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from toolbox.runtime import ToolPaths, load_module


class SummaryCompletenessTests(unittest.TestCase):
    def setUp(self):
        self.module = load_module("summary_completeness_test", ToolPaths(Path(__file__).resolve().parents[2]).summary)
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.station = self.root / "归湖"
        self.station.mkdir()
        self.source = self.station / "日前交易结果查询.xlsx"
        self.output = self.root / "汇总.xlsx"

    def write_source(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = "归湖#1"
        sheet.append(["时刻", "电量", "出清电价"])
        for row in rows:
            sheet.append(row)
        book.save(self.source)
        book.close()

    def complete_rows(self):
        return [[f"{hour:02d}:00", 0.0, -12.345] for hour in range(24)]

    def test_incomplete_input_preserves_existing_output(self):
        self.write_source(self.complete_rows()[:1])
        self.output.write_bytes(b"previous report")
        with self.assertRaisesRegex(ValueError, "缺少时段.*01:00"):
            self.module.run_summary(self.root, self.output)
        self.assertEqual(self.output.read_bytes(), b"previous report")

    def test_missing_and_nonfinite_numeric_values_are_rejected(self):
        for value in (None, "", " ", "nan", "inf"):
            with self.subTest(value=value):
                rows = self.complete_rows()
                rows[3][2] = value
                self.write_source(rows)
                with self.assertRaisesRegex(ValueError, "03:00.*缺失或无效"):
                    self.module.run_summary(self.root, self.output)
        self.assertFalse(self.output.exists())

    def test_valid_zero_negative_and_missing_market_readback(self):
        self.write_source(self.complete_rows())
        self.module.run_summary(self.root, self.output)
        book = load_workbook(self.output)
        try:
            sheet = book["潮安"]
            self.assertEqual(sheet["B4"].value, 0)
            self.assertEqual(sheet["C4"].value, -12.345)
            self.assertEqual(sheet["C4"].number_format, "General")
            self.assertIsNone(sheet["E4"].value)
            self.assertIsNone(sheet["F27"].value)
            self.assertEqual(sheet["A27"].value, "23:00")
        finally:
            book.close()
