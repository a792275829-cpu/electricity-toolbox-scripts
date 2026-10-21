from datetime import datetime, time, timedelta
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

import toolbox_native_excel as native


def snapshot(book):
    return [(s.title, [[(c.value, c.number_format, c.data_type) for c in row]
                       for row in s.iter_rows(max_row=10, max_col=3)]) for s in book.worksheets]


class NativeFallbackTests(unittest.TestCase):
    def test_unavailable_and_invalid_backend(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            Workbook().save(path)
            with patch.object(native, "native_binary", return_value=None):
                with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="auto"):
                    book = native.load_value_workbook(path)
                    book.close()
                with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
                    with self.assertRaisesRegex(RuntimeError, "未安装"):
                        native.load_value_workbook(path)
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="invalid"):
                with self.assertRaises(ValueError):
                    native.load_value_workbook(path)


@unittest.skipUnless(native.native_binary(), "build native/price-xlsx first")
class NativeReadParityTests(unittest.TestCase):
    def test_cells_formats_dates_epochs_errors_and_cached_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            for mac_epoch in (False, True):
                path = Path(tmp) / f"{mac_epoch}.xlsx"
                book = Workbook()
                if mac_epoch:
                    book.epoch = CALENDAR_MAC_1904
                sheet = book.active
                sheet.title = "中文与格式"
                sheet.append(["字段", "数值", "备注"])
                sheet.append(["价格", 711.47, "001.2300"])
                sheet["B2"].number_format = '0.00"元"'
                sheet.append(["百分比", .125, False])
                sheet["B3"].number_format = "0.00%"
                sheet.append(["日期", datetime(2026, 10, 20, 1, 2, 3), ""])
                sheet.append(["时间", time(3, 4, 5), "#DIV/0!"])
                sheet.append(["时长", timedelta(hours=49, seconds=4), None])
                sheet["B6"].number_format = "[h]:mm:ss"
                sheet.append(["公式", "=1+1", None])
                sheet["C8"].number_format = "0.0000"
                sheet.merge_cells("A9:C9")
                sheet["A9"] = "合并标题"
                book.save(path)
                # Inject cached formula and an integer above binary64's exact
                # integer range; an ordinary Excel writer rounds the latter.
                with zipfile.ZipFile(path) as source:
                    files = {name: source.read(name) for name in source.namelist()}
                xml = files["xl/worksheets/sheet1.xml"].decode()
                xml = xml.replace("<f>1+1</f><v></v>", "<f>1+1</f><v>2</v>")
                # Add a non-date numeric cell while preserving the range.
                xml = xml.replace('</sheetData>', '<row r="10"><c r="A10" t="n"><v>9007199254740993</v></c></row></sheetData>')
                xml = xml.replace('ref="A1:C9"', 'ref="A1:C10"')
                files["xl/worksheets/sheet1.xml"] = xml.encode()
                with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
                    for name, data in files.items():
                        target.writestr(name, data)
                for read_only in (True, False):
                    with self.subTest(mac_epoch=mac_epoch, read_only=read_only):
                        reference = load_workbook(path, data_only=True, read_only=read_only)
                        with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
                            actual = native.load_value_workbook(path, read_only=read_only)
                        try:
                            if read_only:
                                reference.active.reset_dimensions()
                                actual.active.reset_dimensions()
                            self.assertEqual(snapshot(reference), snapshot(actual))
                        finally:
                            reference.close()
                            actual.close()

    def test_batch_uses_one_process_and_scope_does_not_leak(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
            paths = [Path(tmp) / f"{i}.xlsx" for i in range(3)]
            for i, path in enumerate(paths):
                book = Workbook()
                book.active.append(["时刻", i])
                book.save(path)
            with patch.object(native, "call", wraps=native.call) as called:
                with native.prefetch_workbooks(paths):
                    for i, path in enumerate(paths):
                        book = native.load_value_workbook(path)
                        self.assertEqual(book.active["B1"].value, i)
                        book.close()
                self.assertEqual(called.call_count, 1)
                self.assertIsNone(native._prefetched.get())

    def test_failed_batch_falls_back_without_restarting_per_file(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="auto"):
            paths = [Path(tmp) / f"{i}.xlsx" for i in range(2)]
            for path in paths:
                book = Workbook()
                book.active["B2"] = .125
                book.active["B2"].number_format = "0.00%"
                book.save(path)
            with patch.object(native, "call", side_effect=ValueError("unsupported workbook")) as called:
                with self.assertWarns(RuntimeWarning):
                    with native.prefetch_workbooks(paths):
                        for path in paths:
                            book = native.load_value_workbook(path)
                            self.assertEqual(book.active["B2"].value, .125)
                            self.assertEqual(book.active["B2"].number_format, "0.00%")
                            book.close()
                self.assertEqual(called.call_count, 1)
                self.assertIsNone(native._prefetched.get())

    def test_stale_dimensions_are_ignored_after_reset(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
            path = Path(tmp) / "input.xlsx"
            book = Workbook()
            book.active.append(["时刻", "电量", "结算电价"])
            book.active.append(["00:00", 1.125, 711.47])
            book.save(path)
            with zipfile.ZipFile(path) as source:
                files = {name: source.read(name) for name in source.namelist()}
            files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b'ref="A1:C2"', b'ref="A1:A1"')
            with zipfile.ZipFile(path, "w") as target:
                for name, data in files.items():
                    target.writestr(name, data)
            actual = native.load_value_workbook(path)
            self.assertEqual(actual.active.max_row, 1)
            actual.active.reset_dimensions()
            self.assertEqual(actual.active["C2"].value, 711.47)
            self.assertEqual(len(list(actual.active.iter_rows())), 2)

    def test_existing_styled_report_cache_is_reused(self):
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "南方分公司电力市场日报"))
        from gd_power_report import workbook_reader
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            Workbook().save(path)
            @workbook_reader.reuse_workbook_inputs
            def parse():
                full = workbook_reader.load_workbook(path, data_only=True, read_only=False)
                with patch.object(workbook_reader, "_load_value_workbook", side_effect=AssertionError("duplicate open")):
                    self.assertIs(workbook_reader.load_value_workbook(path, data_only=True, read_only=False), full)
            parse()

    def test_realtime_payload_and_validation_errors_match(self):
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "南方分公司电力市场日报"))
        from gd_power_report.realtime_unit_xlsx import EXPECTED_UNITS, parse_realtime_unit_workbook
        target = datetime(2026, 10, 18).date()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "units.xlsx"
            book = Workbook()
            book.active.title = "基础信息"
            book.active["A2"] = target
            for plant, units in EXPECTED_UNITS.items():
                sheet = book.create_sheet(plant)
                for index, unit in enumerate(units):
                    base = 1 + index * 30
                    sheet.cell(base, 1, unit)
                    for col, label in enumerate(("时刻", "日前出清", "实时出清", "实际上网电量"), 1):
                        sheet.cell(base + 1, col, label)
                    for hour in range(24):
                        row = base + 3 + hour
                        sheet.cell(row, 1, time(hour))
                        for col in range(2, 5):
                            sheet.cell(row, col, (hour + col) * .125)
            book.save(path)
            results = []
            for mode in ("python", "rust"):
                with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND=mode):
                    results.append(parse_realtime_unit_workbook(path, target))
            self.assertEqual(results[0], results[1])
            book["汕头"]["A23"] = time(18)  # duplicate hour in the first unit
            book.save(path)
            errors = []
            for mode in ("python", "rust"):
                with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND=mode):
                    with self.assertRaises(ValueError) as caught:
                        parse_realtime_unit_workbook(path, target)
                    errors.append(str(caught.exception))
            self.assertEqual(errors[0], errors[1])
            self.assertIn("重复小时", errors[0])
