from __future__ import annotations
import importlib

import gc
import subprocess
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
import runpy
import os
import shutil
from contextlib import ExitStack
from datetime import date
from datetime import timedelta
from pathlib import Path
from unittest import mock

def destroy_tk(root: tk.Misc) -> None:
    try:
        root.update_idletasks()
    except tk.TclError:
        pass
    root.destroy()


class GroupUploadRuntimeTests(unittest.TestCase):
    """Group upload runtime: locking, report naming, startup max and report creation."""
    def test_group_upload_process_is_alive_detects_current_process(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_process_alive_test", paths.group_upload)

        self.assertTrue(module.process_is_alive(os.getpid()))
        self.assertFalse(module.process_is_alive(-1))

    def test_group_upload_lock_cleanup_does_not_mask_body_exception(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_lock_cleanup_test", paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            lock_path = Path(directory) / "auth_state.json.lock"
            with mock.patch.object(module, "AUTH_LOCK_PATH", lock_path):
                with mock.patch.object(Path, "unlink", side_effect=PermissionError("locked")):
                    with self.assertRaisesRegex(ValueError, "original failure"):
                        with module.AuthStateLock():
                            raise ValueError("original failure")

    def test_group_upload_daily_report_name_uses_energy_date(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_daily_report_name_test", paths.group_upload)

        self.assertEqual(
            module.daily_report_name("2026-07-03"),
            "广东电力现货市场监测评估日报（20260703）",
        )

    def test_group_upload_startup_max_uses_day_ahead_when_realtime_is_missing(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_startup_fallback_test", paths.group_upload)

        class DummyResponse:
            status = 200

            def __init__(self, data):
                self._data = data

            def text(self):
                import json

                return json.dumps({"retCode": "T200", "data": self._data})

        class DummyRequest:
            def __init__(self):
                self.calls = 0

            def post(self, *_args, **_kwargs):
                self.calls += 1
                if self.calls == 2:
                    return DummyResponse([])
                return DummyResponse(
                    [
                        {
                            "date": "2026-07-18 00:00:00",
                            "dataItemGroup": "DAY_AHEAD_CLEARANCE_COAL",
                            "unitCount": [10, 12] + [0] * 22,
                        },
                        {
                            "date": "2026-07-18 00:00:00",
                            "dataItemGroup": "DAY_AHEAD_CLEARANCE_GAS",
                            "unitCount": [3, 4] + [0] * 22,
                        },
                    ]
                )

        context = mock.Mock(request=DummyRequest())
        with mock.patch.object(module, "log") as log:
            result = module.fetch_startup_max(context, "2026-07-18", "044")

        self.assertEqual(result, 16)
        self.assertTrue(any("实时出清电量未完整返回" in call.args[0] for call in log.call_args_list))

    def test_group_upload_startup_max_fails_when_neither_dataset_is_complete(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_startup_missing_test", paths.group_upload)
        response = mock.Mock(status=200)
        response.text.return_value = '{"retCode":"T200","data":[]}'
        context = mock.Mock()
        context.request.post.return_value = response

        with self.assertRaisesRegex(RuntimeError, "日前和实时出清电量都没有返回"):
            module.fetch_startup_max(context, "2026-07-18", "044")

    def test_group_upload_opens_day_report_from_authenticated_report_entry(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_day_report_entry_test", paths.group_upload)

        class DummyPage:
            def __init__(self) -> None:
                self.urls: list[str] = []

            def goto(self, url, **_kwargs):
                self.urls.append(url)

            def wait_for_load_state(self, *_args, **_kwargs):
                return None

        class DummyContext:
            def __init__(self) -> None:
                self.page = DummyPage()
                self.pages = [self.page]

            def new_page(self):
                return self.page

        context = DummyContext()
        with mock.patch.object(module.page_navigation, "page_needs_manual_navigation", return_value=False):
            page = module.open_day_report_page(context, allow_manual=False)

        self.assertIs(page, context.page)
        self.assertEqual(context.page.urls, [module.DAY_REPORT_URL])

    def test_group_upload_returns_to_report_entry_before_menu_fallback(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_day_report_fallback_test", paths.group_upload)

        class DummyText:
            @property
            def first(self):
                return self

            def click(self, **_kwargs):
                return None

        class DummyPage:
            url = "about:blank"

            def __init__(self) -> None:
                self.urls: list[str] = []

            def goto(self, url, **_kwargs):
                self.urls.append(url)
                self.url = url

            def wait_for_load_state(self, *_args, **_kwargs):
                return None

            def get_by_text(self, *_args, **_kwargs):
                return DummyText()

        class DummyContext:
            def __init__(self) -> None:
                self.page = DummyPage()
                self.pages = [self.page]

            def new_page(self):
                return self.page

        context = DummyContext()
        with (
            mock.patch.object(module.page_navigation, "page_needs_manual_navigation", side_effect=[True, False]),
            mock.patch.object(module.page_navigation, "enter_huaneng_marketing_platform", side_effect=lambda page: page.goto(module.USERCENTER_URL)),
            mock.patch.object(module.page_navigation, "click_button_by_text"),
        ):
            page = module.open_day_report_page(context, allow_manual=False)

        self.assertIs(page, context.page)
        self.assertEqual(
            context.page.urls,
            [module.DAY_REPORT_URL, module.USERCENTER_URL, module.DAY_REPORT_URL],
        )

    def test_group_upload_accepts_numeric_c2_when_rerunning_same_day_output(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_numeric_c2_rerun_test", paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            province_path = folder / "广东-省内数据-20260703.xlsx"
            energy_path = folder / "广东-能销数据-20260703.xlsx"
            energy_path.write_text("x", encoding="utf-8")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "省内数据"
            sheet["A2"] = date(2026, 7, 3)
            sheet["B2"] = 18
            sheet["C2"] = 108314.96
            workbook.save(province_path)

            uploads = [
                module.UploadFile(path=province_path, kind="省内", date="2026-07-03"),
                module.UploadFile(path=energy_path, kind="能销", date="2026-07-03"),
            ]
            with (
                mock.patch.object(module.upload_flow, "fetch_startup_max", return_value=20),
                mock.patch.object(module.upload_flow, "fetch_capacity_average", return_value=50000.0),
            ):
                updated_uploads = module.update_province_excel_for_upload(
                    object(),
                    "2026-07-03",
                    uploads,
                )

            updated = load_workbook(province_path, data_only=True)
            updated_sheet = updated["省内数据"]

        self.assertEqual(updated_uploads[0].path, province_path.resolve())
        self.assertEqual(updated_sheet["B2"].value, 20)
        self.assertEqual(updated_sheet["C2"].value, 108314.96)

    def test_group_upload_accepts_numeric_c2_from_nearest_template(self) -> None:
        from openpyxl import Workbook, load_workbook
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_numeric_c2_nearest_template_test", paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            province_path = folder / "广东-省内数据-20260704.xlsx"
            energy_path = folder / "广东-能销数据-20260705.xlsx"
            energy_path.write_text("x", encoding="utf-8")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "省内数据"
            sheet["A2"] = date(2026, 7, 4)
            sheet["B2"] = 18
            sheet["C2"] = 108314.96
            workbook.save(province_path)

            uploads = [
                module.UploadFile(path=province_path, kind="省内", date="2026-07-04", auto_template=True),
                module.UploadFile(path=energy_path, kind="能销", date="2026-07-05"),
            ]
            with (
                mock.patch.object(module.upload_flow, "fetch_startup_max", return_value=20),
                mock.patch.object(module.upload_flow, "fetch_capacity_average", return_value=50000.0),
            ):
                updated_uploads = module.update_province_excel_for_upload(
                    object(),
                    "2026-07-05",
                    uploads,
                )

            target_path = folder / "广东-省内数据-20260705.xlsx"
            updated = load_workbook(target_path, data_only=True)
            updated_sheet = updated["省内数据"]

        self.assertEqual(updated_uploads[0].path, target_path.resolve())
        self.assertEqual(updated_sheet["A2"].value.date(), date(2026, 7, 5))
        self.assertEqual(updated_sheet["B2"].value, 20)
        self.assertEqual(updated_sheet["C2"].value, 108314.96)

    def test_group_upload_prefers_guangdong_branch_tenant(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_tenant_selection_test", paths.group_upload)

        tenants = [
            {"tenantId": "plant", "name": "华能广东汕头海上风电有限责任公司"},
            {"tenantId": "branch", "name": "广东分公司"},
            {"tenantId": "other", "name": "华能集团"},
        ]

        self.assertEqual(module.select_day_report_tenant(tenants), tenants[1])

    def test_group_upload_switches_using_application_tenant(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_application_tenant_test", paths.group_upload)

        class FakeResponse:
            status = 200

            def __init__(self, body: dict[str, object]) -> None:
                self.body = body

            def text(self) -> str:
                import json

                return json.dumps(self.body, ensure_ascii=False)

        class FakeRequest:
            def __init__(self) -> None:
                self.calls: list[tuple[str, dict[str, object] | None]] = []

            def get(self, url, params=None, **_kwargs):
                self.calls.append((url, params))
                if url.endswith("/usercenter/web/pf/tenant/user/application"):
                    return FakeResponse(
                        {
                            "retCode": "T200",
                            "data": [
                                {"tenantId": "root", "name": "华能集团"},
                                {"tenantId": "branch", "name": "广东分公司"},
                            ],
                        }
                    )
                if url.endswith("/usercenter/web/switchTenant"):
                    return FakeResponse({"retCode": "T200", "retMsg": "切换租户成功"})
                raise AssertionError(f"unexpected request: {url}")

        class FakeContext:
            def __init__(self) -> None:
                self.request = FakeRequest()

        context = FakeContext()
        module.switch_to_day_report_tenant(context)

        self.assertEqual(
            context.request.calls,
            [
                (f"{module.BASE_URL}/usercenter/web/pf/tenant/user/application", None),
                (f"{module.BASE_URL}/usercenter/web/switchTenant", {"tenantId": "branch"}),
            ],
        )

    def test_group_upload_create_daily_report_posts_api_payload(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_create_report_api_test", paths.group_upload)

        class FakeResponse:
            status = 200

            def __init__(self, body: dict[str, object]) -> None:
                self.body = body

            def text(self) -> str:
                import json

                return json.dumps(self.body, ensure_ascii=False)

        class FakeRequest:
            def __init__(self) -> None:
                self.posts: list[tuple[str, dict[str, object]]] = []

            def get(self, url, params=None, **_kwargs):
                if url.endswith("/huaneng/group/api/report"):
                    return FakeResponse({"retCode": "T200", "data": {"datas": []}})
                if url.endswith("/huaneng/group/api/report/queryTemplates"):
                    return FakeResponse(
                        {
                            "retCode": "T200",
                            "data": [
                                {"name": "广东模板（修改）", "id": "template-id"},
                            ],
                        }
                    )
                raise AssertionError(f"unexpected request: {url}")

            def post(self, url, data=None, **_kwargs):
                self.posts.append((url, data or {}))
                return FakeResponse({"retCode": "T200", "data": "file-id"})

        class FakeContext:
            def __init__(self) -> None:
                self.request = FakeRequest()

        context = FakeContext()
        with (
            mock.patch.object(module.page_navigation, "open_day_report_page", return_value=object()),
        ):
            module.create_daily_report(context, "2026-07-04", headless=True)

        self.assertEqual(
            context.request.posts,
            [
                (
                    f"{module.BASE_URL}/huaneng/group/api/report",
                    {
                        "reportName": "广东电力现货市场监测评估日报（20260704）",
                        "startDate": "2026-07-04",
                        "endDate": "2026-07-04",
                        "templateId": "template-id",
                        "provinceId": None,
                    },
                )
            ],
        )

    def test_group_upload_create_daily_report_uses_template_fallback_after_timeout(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_template_timeout_fallback_test", paths.group_upload)

        class FakeResponse:
            status = 200

            def __init__(self, body: dict[str, object]) -> None:
                self.body = body

            def text(self) -> str:
                import json

                return json.dumps(self.body, ensure_ascii=False)

        class FakeRequest:
            def __init__(self) -> None:
                self.posts: list[tuple[str, dict[str, object]]] = []

            def get(self, url, params=None, **_kwargs):
                if url.endswith("/huaneng/group/api/report"):
                    return FakeResponse({"retCode": "T200", "data": {"datas": []}})
                if url.endswith("/huaneng/group/api/report/queryTemplates"):
                    raise module.PlaywrightTimeoutError("template endpoint timed out")
                raise AssertionError(f"unexpected request: {url}")

            def post(self, url, data=None, **_kwargs):
                self.posts.append((url, data or {}))
                return FakeResponse({"retCode": "T200", "data": "file-id"})

        class FakeContext:
            def __init__(self) -> None:
                self.request = FakeRequest()

        context = FakeContext()
        with (
            mock.patch.object(module.page_navigation, "open_day_report_page", return_value=object()),
            mock.patch.object(module.time, "sleep"),
        ):
            module.create_daily_report(context, "2026-07-06", headless=True)

        self.assertEqual(context.request.posts[0][1]["templateId"], module.DEFAULT_DAILY_REPORT_TEMPLATE_ID)

    def test_group_upload_smoke_report_cli_skips_file_selection(self) -> None:
        from toolbox.runtime import ToolPaths, load_module

        paths = ToolPaths(Path(__file__).resolve().parents[2])
        module = load_module("group_upload_smoke_report_cli_test", paths.group_upload)

        with (
            mock.patch.object(sys, "argv", ["upload_daily_report.py", "--smoke-report"]),
            mock.patch.object(module, "smoke_daily_report_page") as smoke,
            mock.patch.object(module, "choose_files_interactively") as choose_files,
        ):
            exit_code = module.main()

        self.assertEqual(exit_code, 0)
        smoke.assert_called_once()
        choose_files.assert_not_called()


class GroupUploadPageTests(unittest.TestCase):
    """Group upload page: staging, D-2 auto selection and template probing."""
    def setUp(self) -> None:
        try:
            self.root = tk.Tk()
            self.root.withdraw()
        except tk.TclError as exc:
            self.skipTest(f"Tk unavailable: {exc}")
        from toolbox.runtime import TaskRegistry, ToolPaths

        self.paths = ToolPaths(Path(__file__).resolve().parents[2])
        self.registry = TaskRegistry()

    def tearDown(self) -> None:
        if hasattr(self, "root"):
            # Finish asynchronous page imports before destroying their Tk owner.
            for child in self.root.winfo_children():
                loader = getattr(child, "_writer_loader", None)
                if loader is not None and loader.ident is not None:
                    loader.join(5)
                    self.assertFalse(loader.is_alive(), "WPS page import did not finish")
            destroy_tk(self.root)
            gc.collect()  # Tk variables must be finalized on the main thread.

    def test_group_upload_runs_in_toolbox_process_to_keep_macos_file_access(self) -> None:
        from toolbox.pages import GroupUploadPage

        page = GroupUploadPage(
            self.root,
            self.paths,
            self.registry,
            today=date(2000, 1, 1),
        )
        selected = [Path("/Desktop/广东-能销数据-20260718.xlsx")]
        page.selected_paths = selected
        page.force.set(True)
        module = mock.Mock()
        original_sink = object()
        module.LOG_SINK = original_sink

        with (
            mock.patch("toolbox.pages.messagebox.askokcancel", return_value=True),
            mock.patch("toolbox.pages.load_module", return_value=module),
            mock.patch("toolbox.pages.run_group_upload_with_staging") as staged_upload,
            mock.patch.object(page, "run_in_thread") as run_in_thread,
            mock.patch.object(page, "run_process") as run_process,
        ):
            page.confirm_upload()
            worker = run_in_thread.call_args.args[0]
            worker()

        run_process.assert_not_called()
        staged_upload.assert_called_once_with(
            module,
            selected,
            force=True,
            on_log=page.append_log,
        )
        self.assertIs(module.LOG_SINK, original_sink)

    def test_group_upload_stages_files_and_copies_generated_province_file_back(self) -> None:
        from toolbox.pages import run_group_upload_with_staging
        from toolbox.runtime import load_module

        module = load_module("group_upload_staging_test", self.paths.group_upload)
        with tempfile.TemporaryDirectory() as directory:
            source_dir = Path(directory)
            province = source_dir / "广东-省内数据-20260717.xlsx"
            energy = source_dir / "广东-能销数据-20260718.xlsx"
            province.write_bytes(b"province-template")
            energy.write_bytes(b"energy")

            def fake_upload(staged_paths, *, force):
                self.assertTrue(force)
                target_date, uploads = module.prepare_upload_files(staged_paths)
                staged_province = next(item for item in uploads if item.kind == "省内")
                output = module.province_output_path(staged_province.path, target_date)
                output.write_bytes(b"updated-province")

            logs: list[str] = []
            with mock.patch.object(module, "run_api_upload", side_effect=fake_upload):
                run_group_upload_with_staging(
                    module,
                    [province, energy],
                    force=True,
                    on_log=logs.append,
                )

            copied_output = source_dir / "广东-省内数据-20260718.xlsx"
            self.assertEqual(copied_output.read_bytes(), b"updated-province")
            self.assertTrue(any("复制回原目录" in line for line in logs))

    def test_group_upload_auto_selects_latest_d_minus_two_energy_file(self) -> None:
        from toolbox.pages import GroupUploadPage

        with tempfile.TemporaryDirectory() as directory:
            upload_dir = Path(directory)
            older = upload_dir / "广东-能销数据-20260615.xlsx"
            latest = upload_dir / "广东-能销数据-20260615-新版.xlsx"
            other_date = upload_dir / "广东-能销数据-20260616.xlsx"
            province_template = upload_dir / "广东-省内数据-20260614.xlsx"
            for index, path in enumerate(
                [older, latest, other_date, province_template], start=1
            ):
                path.write_text("x", encoding="utf-8")
                os.utime(path, (1_700_000_000 + index, 1_700_000_000 + index))
            os.utime(latest, (1_700_000_100, 1_700_000_100))

            match = GroupUploadPage.find_latest_energy_file_for_date(
                date(2026, 6, 17),
                upload_dir,
            )

        self.assertEqual(match, latest)

    def test_group_upload_page_initializes_with_d_minus_two_energy_file(self) -> None:
        from toolbox.pages import GroupUploadPage

        with tempfile.TemporaryDirectory() as directory:
            upload_dir = Path(directory)
            energy = upload_dir / "广东-能销数据-20260615.xlsx"
            province_template = upload_dir / "广东-省内数据-20260614.xlsx"
            energy.write_text("x", encoding="utf-8")
            province_template.write_text("x", encoding="utf-8")
            paths = type(
                "Paths",
                (),
                {
                    "group_upload": self.paths.group_upload,
                    "group_upload_dir": upload_dir,
                },
            )()

            page = GroupUploadPage(
                self.root,
                paths,  # type: ignore[arg-type]
                self.registry,
                today=date(2026, 6, 17),
            )

        self.assertEqual(page.selected_paths, [province_template.resolve(), energy.resolve()])
        self.assertIn("能销数据", page.selection_var.get())

    def test_group_upload_page_keeps_auto_province_template_in_selection(self) -> None:
        from toolbox.pages import GroupUploadPage

        with tempfile.TemporaryDirectory() as directory:
            upload_dir = Path(directory)
            energy = upload_dir / "广东-能销数据-20260705.xlsx"
            province_template = upload_dir / "广东-省内数据-20260704.xlsx"
            energy.write_text("x", encoding="utf-8")
            province_template.write_text("x", encoding="utf-8")
            paths = type(
                "Paths",
                (),
                {
                    "group_upload": self.paths.group_upload,
                    "group_upload_dir": upload_dir,
                },
            )()

            page = GroupUploadPage(
                self.root,
                paths,  # type: ignore[arg-type]
                self.registry,
                today=date(2026, 7, 7),
            )

        self.assertEqual(page.selected_paths, [province_template.resolve(), energy.resolve()])
        command = page.build_upload_command(page.selected_paths, force=False)
        self.assertIn(str(province_template.resolve()), command)
        self.assertIn(str(energy.resolve()), command)

    def test_group_upload_template_error_lists_visible_excel_files(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("group_upload_template_error_test", self.paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            (folder / "广东-能销数据-20260705.xlsx").write_text("x", encoding="utf-8")
            (folder / "广东-省内数据-没有日期.xlsx").write_text("x", encoding="utf-8")

            with (
                mock.patch.object(module, "BUILTIN_PROVINCE_TEMPLATE", folder / "missing-template.xlsx"),
                self.assertRaisesRegex(ValueError, "实际看到的 Excel 文件"),
            ):
                module.find_province_template(folder, "2026-07-05")

    def test_group_upload_page_accepts_energy_without_nearby_province_file(self) -> None:
        from toolbox.pages import GroupUploadPage
        from toolbox.runtime import load_module

        module = load_module("group_upload_builtin_page_test", self.paths.group_upload)
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            energy = folder / "广东-能销数据-20261023.xlsx"
            energy.write_bytes(b"energy")
            paths = type("Paths", (), {
                "group_upload": self.paths.group_upload, "group_upload_dir": folder,
            })()
            page = GroupUploadPage(self.root, paths, self.registry, today=date(2026, 10, 25))
            self.assertEqual(page.selected_paths, [module.BUILTIN_PROVINCE_TEMPLATE, energy.resolve()])
            self.assertIn("2026-10-23", page.summary_var.get())
            self.assertIn("省内数据模板.xlsx", page.selection_var.get())
            _, uploads = module.prepare_upload_files(page.selected_paths)
            self.assertTrue(uploads[0].builtin_template)

    def test_group_upload_auto_template_prefers_nearest_province_date(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("group_upload_nearest_template_test", self.paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            energy = folder / "广东-能销数据-20260705.xlsx"
            older = folder / "广东-省内数据-20260621.xlsx"
            nearest = folder / "广东-省内数据-20260704.xlsx"
            newer = folder / "广东-省内数据-20260710.xlsx"
            for path in [energy, older, nearest, newer]:
                path.write_text("x", encoding="utf-8")

            target_date, uploads = module.prepare_upload_files([energy])

        self.assertEqual(target_date, "2026-07-05")
        self.assertEqual(uploads[0].path, nearest.resolve())
        self.assertEqual(uploads[1].path, energy.resolve())

    def test_group_upload_auto_template_probes_standard_names_when_listing_is_empty(self) -> None:
        from toolbox.runtime import load_module

        module = load_module("group_upload_template_probe_test", self.paths.group_upload)

        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            energy = folder / "广东-能销数据-20260705.xlsx"
            nearest = folder / "广东-省内数据-20260704.xlsx"
            energy.write_text("x", encoding="utf-8")
            nearest.write_text("x", encoding="utf-8")
            with mock.patch.object(Path, "glob", return_value=iter(())):
                target_date, uploads = module.prepare_upload_files([energy])

        self.assertEqual(target_date, "2026-07-05")
        self.assertEqual(uploads[0].path, nearest.resolve())
        self.assertEqual(uploads[1].path, energy.resolve())


if __name__ == "__main__":
    unittest.main()
