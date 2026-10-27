"""Pure WPS value and range contracts, independent of Tk and browser state.

The legacy GUI re-exports these names for standalone callers.
"""
from __future__ import annotations
from datetime import date, datetime, time
from typing import Any
try:
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
except ImportError:
    get_column_letter = coordinate_from_string = column_index_from_string = None

def normalize_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return value


def normalize_for_compare(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if text.lower() in {"none", "null"}:
        return ""
    return text


def values_equivalent(expected: Any, actual: Any) -> bool:
    expected_value = normalize_for_compare(expected)
    actual_value = normalize_for_compare(actual)
    if expected_value == actual_value:
        return True
    if isinstance(expected_value, bool) or isinstance(actual_value, bool):
        return False
    try:
        expected_number = float(expected_value)
        actual_number = float(actual_value)
        return abs(expected_number - actual_number) <= 1e-9
    except (TypeError, ValueError):
        return False


def parse_start_cell(cell: str) -> tuple[int, int]:
    if coordinate_from_string is None or column_index_from_string is None:
        raise RuntimeError("缺少表格读取组件，请运行工具箱目录中的 setup_macos.command 完成初始化。")
    col_text, row = coordinate_from_string(cell.strip().upper())
    return row, column_index_from_string(col_text)


def normalize_cell_address(cell: str) -> str:
    row, col = parse_start_cell(cell)
    return f"{get_column_letter(col)}{row}"


def range_from_cells(start_cell: str, end_cell: str) -> str:
    start_row, start_col = parse_start_cell(start_cell)
    end_row, end_col = parse_start_cell(end_cell)
    if end_row < start_row or end_col < start_col:
        raise ValueError(f"结束单元格不能位于起始单元格的上方或左侧：{start_cell} → {end_cell}")
    return f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"


def range_shape(start_cell: str, end_cell: str) -> tuple[int, int]:
    start_row, start_col = parse_start_cell(start_cell)
    end_row, end_col = parse_start_cell(end_cell)
    if end_row < start_row or end_col < start_col:
        raise ValueError(f"结束单元格不能位于起始单元格的上方或左侧：{start_cell} → {end_cell}")
    return end_row - start_row + 1, end_col - start_col + 1


def split_range(range_text: str) -> tuple[str, str]:
    text = range_text.strip()
    if ":" not in text:
        cell = normalize_cell_address(text)
        return cell, cell
    start, end = text.split(":", 1)
    return normalize_cell_address(start), normalize_cell_address(end)


def target_range_from_start(start_cell: str, rows: int, cols: int) -> str:
    start_row, start_col = parse_start_cell(start_cell)
    end_row = start_row + rows - 1
    end_col = start_col + cols - 1
    return f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"


def target_rect(start_cell: str, rows: int, cols: int) -> tuple[int, int, int, int]:
    start_row, start_col = parse_start_cell(start_cell)
    return start_row, start_col, start_row + rows - 1, start_col + cols - 1


def rect_from_cells(start_cell: str, end_cell: str) -> tuple[int, int, int, int]:
    start_row, start_col = parse_start_cell(start_cell)
    end_row, end_col = parse_start_cell(end_cell)
    if end_row < start_row or end_col < start_col:
        raise ValueError(f"结束单元格不能位于起始单元格的上方或左侧：{start_cell} → {end_cell}")
    return start_row, start_col, end_row, end_col


def rects_overlap(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    a_r1, a_c1, a_r2, a_c2 = a
    b_r1, b_c1, b_r2, b_c2 = b
    return not (a_r2 < b_r1 or b_r2 < a_r1 or a_c2 < b_c1 or b_c2 < a_c1)


def coerce_wps_matrix(raw: Any, rows: int, cols: int) -> list[list[Any]]:
    if rows == 1 and cols == 1:
        if isinstance(raw, list):
            if raw and isinstance(raw[0], list):
                return [[raw[0][0] if raw[0] else ""]]
            return [[raw[0] if raw else ""]]
        return [[raw]]
    if isinstance(raw, list):
        if rows == 1 and (not raw or not isinstance(raw[0], list)):
            return [raw]
        matrix = []
        for row in raw:
            matrix.append(row if isinstance(row, list) else [row])
        return matrix
    return [[raw]]


def compare_matrix(expected: list[list[Any]], actual_raw: Any, rows: int, cols: int) -> list[str]:
    actual = coerce_wps_matrix(actual_raw, rows, cols)
    mismatches: list[str] = []
    for r in range(rows):
        for c in range(cols):
            raw_expected = expected[r][c] if r < len(expected) and c < len(expected[r]) else ""
            raw_actual = actual[r][c] if r < len(actual) and c < len(actual[r]) else ""
            if not values_equivalent(raw_expected, raw_actual):
                cell = f"区域内第 {r + 1} 行第 {c + 1} 列"
                expected_value = normalize_for_compare(raw_expected)
                actual_value = normalize_for_compare(raw_actual)
                mismatches.append(f"{cell}：来源值 {expected_value!r}，回读值 {actual_value!r}")
                if len(mismatches) >= 20:
                    return mismatches
    return mismatches
