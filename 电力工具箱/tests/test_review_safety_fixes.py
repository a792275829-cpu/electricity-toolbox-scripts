"""Regressions for the six maintenance review findings; no production I/O."""
from contextlib import nullcontext
from datetime import date
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from toolbox.runtime import load_module
from toolbox_shared.auth_lock import AuthFileLock

ROOT = Path(__file__).resolve().parents[2]
DAY = date(2026, 10, 28)


class DatedInputsTests(unittest.TestCase):
    def test_both_load_consumers_require_unique_matching_date(self):
        for relative in ('市场表更新/update_market_table.py', '交易系统信息更新/update_transaction_info.py'):
            module = load_module('dated_fix_' + relative.split('/')[0], ROOT / relative)
            def summarize(records):
                return module.summarize_load_type({'data': {'dataNetLoadDTOList': [
                    dict(loadType='1', forecastPeriodList=records)]}}, '1', DAY,
                    period_list_key='forecastPeriodList', label='日前')
            matching = dict(date='2026-10-28 00:00:00', max=10, min=0, avg=3, sum=4)
            stale = dict(matching, date='2026-10-27')
            self.assertEqual(summarize([stale, matching])[:3], (10, 0, 3) if '交易' in relative else [10, 0, 3])
            for records in ([stale], [dict(matching, date=None)], [matching, matching], []):
                with self.subTest(relative=relative, records=records), self.assertRaises(RuntimeError):
                    summarize(records)

    def export(self, rows, units, folder):
        module = load_module('dated_fix_online', ROOT / '上网电量抓取/export_online_energy.py')
        context = SimpleNamespace(close=Mock())
        with patch.multiple(module, sync_playwright=lambda: nullcontext(None),
            launch_context=lambda *a, **k: context, ensure_login=lambda *a: None,
            get_companies=lambda *a: [dict(tenantId='test', name='海门电厂')],
            switch_company=lambda *a: None, get_units=lambda *a: units,
            is_renewable_company=lambda *a: False, fetch_fire_rows=lambda *a: rows,
            log=lambda *a: None), patch.object(module.time, 'sleep'):
            try:
                return module.export_online_energy(str(DAY), module.Config(output_dir=folder))
            finally:
                context.close.assert_called_once()

    def test_bad_online_records_preserve_previous_export(self):
        valid = dict(unitId='u', runDate=str(DAY), onlineEle=[1]*24)
        cases = ([dict(valid, runDate='2026-10-27')], [dict(valid, runDate=None)], [],
                 [valid, valid], [dict(valid, unitId='unknown')])
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            target = folder / f'上网电量汇总_{DAY}.xlsx'
            target.write_bytes(b'previous workbook')
            for rows in cases:
                with self.subTest(rows=rows), self.assertRaises(ValueError):
                    self.export(rows, [dict(id='u', name='#1')], folder)
                self.assertEqual(target.read_bytes(), b'previous workbook')
            with self.assertRaisesRegex(ValueError, '缺少机组'):
                self.export([valid], [dict(id='u', name='#1'), dict(id='v', name='#2')], folder)

    def test_valid_online_export_reads_back_zero_negative_and_hours(self):
        with tempfile.TemporaryDirectory() as folder:
            target = self.export([dict(unitId='u', runDate=str(DAY), onlineEle=[0,-1.25]+[2]*22)],
                                 [dict(id='u', name='#1')], Path(folder))
            book = load_workbook(target)
            try:
                self.assertEqual([book.active[f'A{row}'].value for row in (3,4,26)], [0,-1.25,2])
                self.assertEqual(book.active['A4'].number_format, 'General')
            finally:
                book.close()


class ForecastCompletionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.module = load_module('forecast_completion_fix', ROOT / '广东电价预测/update_gdfire_data.py')

    def sample(self, day='2026-10-28'):
        frame = self.module.empty_time_frame(day,day).drop(columns=['时点序号'])
        frame['全省日前平均电价'] = np.arange(96, dtype=float)
        frame['日前统调负荷(MW)'] = 10000. + np.arange(96)
        return frame

    def test_missing_price_curve_invalid_slots_and_nonfinite_values_are_retried(self):
        complete = self.sample()
        self.assertEqual(self.module.missing_dates(complete, str(DAY)), [])
        for column, value in [('全省日前平均电价', np.nan), ('全省日前平均电价', np.inf),
                              ('日前统调负荷(MW)', np.nan), ('时刻', '00:07')]:
            broken = complete.copy()
            broken.loc[1, column] = value
            self.assertEqual(self.module.missing_dates(broken, str(DAY)), [DAY])
        self.assertEqual(self.module.missing_dates(complete.drop(columns='全省日前平均电价'), str(DAY)), [DAY])

    def update(self, folder, increment):
        master = folder / 'master.xlsx'
        client = Mock()
        client.__enter__ = Mock(return_value=client)
        client.__exit__ = Mock(return_value=False)
        with (patch.object(self.module, 'load_config', return_value=dict(username='test',password='test')),
             patch.object(self.module, 'resolve_paths', return_value=dict(masterWorkbook=master,initialWorkbook=folder/'absent.xlsx')),
             patch.object(self.module, 'GDFireClient', return_value=client),
             patch.object(self.module, 'build_increment', return_value=increment)):
            result = self.module.update_master_workbook(run_date=str(DAY))
        return master, result, client

    def test_cold_start_retains_business_columns_and_reopens_output(self):
        with tempfile.TemporaryDirectory() as folder:
            master, result, _ = self.update(Path(folder), self.sample())
            book = load_workbook(master)
            try:
                self.assertEqual([c.value for c in book.active[1]], list(self.sample().columns))
                self.assertEqual(book.active['C2'].value, 0)
                self.assertEqual(book.active['D97'].value, 10095)
                self.assertEqual(book.active['A2'].value.date(), DAY)
                self.assertIn('yy', book.active['A2'].number_format)
                self.assertEqual(book.active['C2'].number_format, 'General')
            finally:
                book.close()
            self.assertEqual(result['remaining_missing_dates'], [])
            self.assertFalse(list(Path(folder).glob('.*.tmp.xlsx')))

    def test_second_run_fills_unpublished_price_and_preserves_old_nonempty_values(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            incomplete = self.sample(); incomplete['全省日前平均电价'] = np.nan
            master, _, _ = self.update(folder, incomplete)
            published = self.sample(); published.loc[0,'日前统调负荷(MW)'] = np.nan
            _, result, client = self.update(folder, published)
            client.get.assert_called()
            self.assertEqual(result['requested_dates'], [str(DAY)])
            book = load_workbook(master)
            try:
                self.assertEqual(book.active['C2'].value, 0)
                self.assertEqual(book.active['D2'].value, 10000)
                self.assertEqual(book.active['C97'].value, 95)
            finally:
                book.close()

    def test_recent_refresh_accepts_revised_values_and_new_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            master, _, _ = self.update(folder, self.sample())
            book = load_workbook(master); book.active['C2'].number_format = '0.00'; book.save(master); book.close()
            revised = self.sample(); revised.loc[0,'全省日前平均电价'] = -25.
            revised['日前光伏出力(MW)'] = np.arange(96.) + 200
            _, result, client = self.update(folder, revised)
            client.get.assert_called()
            self.assertEqual(result['requested_dates'], [str(DAY)])
            book = load_workbook(master)
            try:
                self.assertEqual(book.active['C2'].value, -25)
                self.assertEqual(book.active['C2'].number_format, '0.00')
                self.assertEqual(book.active['E1'].value, '日前光伏出力(MW)')
                self.assertEqual(book.active['E97'].value, 295)
            finally:
                book.close()

    def test_equal_business_curves_are_distinct_fields(self):
        base = self.sample()
        increment = base.copy()
        increment['日前光伏出力(MW)'] = increment['全省日前平均电价']
        merged = self.module.merge_data_frames(base, increment)
        self.assertIn('日前光伏出力(MW)', merged.columns)
        self.assertEqual(merged['日前光伏出力(MW)'].tolist(), list(range(96)))

    def test_failed_serialization_keeps_existing_master_and_removes_temporary(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            master = folder / 'master.xlsx'; master.write_bytes(b'old master')
            def corrupt(frame, path, *args, **kwargs):
                Path(path).write_bytes(b'incomplete XLSX')
            with patch.object(self.module, '_write_large_xlsx', side_effect=corrupt), self.assertRaises(Exception):
                self.module.write_large_xlsx(self.sample(), master, '整合数据')
            self.assertEqual(master.read_bytes(), b'old master')
            self.assertFalse(list(folder.glob('.*.tmp.xlsx')))


class AuthLockTests(unittest.TestCase):
    def test_contender_timeout_stale_empty_recovery_and_release_identity(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'auth.lock'
            path.touch()  # Interrupted initialization must be recoverable.
            with AuthFileLock(path, timeout_seconds=0, log=lambda _: None):
                with self.assertRaisesRegex(RuntimeError, '超时'):
                    with AuthFileLock(path, timeout_seconds=0, log=lambda _: None):
                        self.fail('second owner')
                self.assertEqual(path.stat().st_mode & 0o777, 0o600)
                path.unlink(); path.write_text('replacement')
            self.assertEqual(path.read_text(), 'replacement')
            with AuthFileLock(path, timeout_seconds=0, log=lambda _: None):
                pass
            self.assertFalse(path.exists())

    def test_live_legacy_owner_is_never_reaped_by_age(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'auth.lock'
            content = json.dumps(dict(pid=os.getpid(), startedAt='2000-01-01'))
            path.write_text(content)
            with self.assertRaisesRegex(RuntimeError, '超时'):
                with AuthFileLock(path, timeout_seconds=0, log=lambda _: None):
                    self.fail('stole a live legacy lock')
            self.assertEqual(path.read_text(), content)

    def test_real_process_competition_and_crashed_owner(self):
        code = '''
from pathlib import Path
import os,sys,time
from toolbox_shared.auth_lock import AuthFileLock
p=Path(sys.argv[1]); marker=p.with_suffix('.active')
for i in range(12):
    with AuthFileLock(p, timeout_seconds=10, log=lambda _:None):
        fd=os.open(marker,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
        time.sleep(.003)
        os.close(fd);marker.unlink()
'''
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'auth.lock'
            path.write_text(json.dumps(dict(pid=999999999,startedAt='2000-01-01')))
            processes = [subprocess.Popen([sys.executable,'-c',code,str(path)],cwd=ROOT,
                          stdout=subprocess.PIPE,stderr=subprocess.PIPE) for _ in range(4)]
            try:
                for process in processes:
                    _, error = process.communicate(timeout=20)
                    self.assertEqual(process.returncode, 0, error.decode())
            finally:
                for process in processes:
                    if process.poll() is None: process.kill(); process.wait()
            self.assertFalse(path.exists())
            crash = 'from toolbox_shared.auth_lock import AuthFileLock; import os,sys; from pathlib import Path; lock=AuthFileLock(Path(sys.argv[1]),log=lambda _:None); lock.__enter__(); os._exit(9)'
            self.assertEqual(subprocess.run([sys.executable,'-c',crash,str(path)],cwd=ROOT).returncode,9)
            with AuthFileLock(path,timeout_seconds=0,log=lambda _:None):
                pass
            self.assertFalse(path.exists())

    def test_all_python_login_consumers_keep_independent_paths(self):
        from toolbox_shared.auth_loader import load_login_module
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, TOOLBOX_DATA_ROOT=folder,TOOLBOX_CODE_ROOT=str(ROOT)):
            for name in ('上网电量抓取','市场表更新','现货价格更新','交易系统信息更新','功率调整收益统计表更新','风力预测截图','广东电价预测','广东运行日报PDF解析','南方分公司电力市场日报/.state','每日生产经营情况汇报自动生成工具/scripts'):
                auth = load_login_module(ROOT,module_name='lock_regression_'+name.replace('/','_'),state_dir=ROOT/name)
                with auth.AuthStateLock(timeout_seconds=0):
                    self.assertTrue(auth.AUTH_LOCK_PATH.is_relative_to(Path(folder).resolve()))
                self.assertFalse(auth.AUTH_LOCK_PATH.exists())
            group=load_module('lock_regression_group',ROOT/'集团每日上传/upload_daily_report.py')
            with group.AuthStateLock(timeout_seconds=0):
                self.assertIsInstance(group.AuthStateLock(),AuthFileLock)
            self.assertFalse(group.AUTH_LOCK_PATH.exists())
