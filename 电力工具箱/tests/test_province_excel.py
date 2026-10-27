from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from toolbox_shared import province_excel as province


class ProvinceExcelTests(unittest.TestCase):
    def template(self, folder, value='=100000-20000-100', name='source.xlsx'):
        path = Path(folder)/name
        book = Workbook()
        sheet = book.active
        sheet.title = '省内数据'
        sheet['A2'] = datetime(2026, 10, 26)
        sheet['A2'].number_format = 'yyyy-mm-dd'
        sheet['B2'] = 10
        sheet['C2'] = value
        sheet['C2'].number_format = '0.00'
        sheet['G2'] = '100'
        sheet['H2'] = .125
        sheet['H2'].number_format = '0.0%'
        book.save(path)
        book.close()
        return path

    def test_serializes_once_and_reads_values_formats_and_unrelated_cells(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self.template(temp)
            before = source.read_bytes()
            target = Path(temp)/'output.xlsx'
            saved = []
            original = Workbook.save

            def save(book, path):
                saved.append(path)
                return original(book, path)

            with patch.object(Workbook, 'save', save):
                province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
            self.assertEqual(len(saved), 1)
            for data_only in (False, True):
                book = load_workbook(target, data_only=data_only)
                try:
                    sheet = book['省内数据']
                    self.assertEqual([sheet[c].value for c in ('A2','B2','C2')],
                                     [datetime(2026,10,27),20,49900])
                    self.assertEqual([sheet[c].number_format for c in ('A2','C2','H2')],
                                     ['yyyy-mm-dd','0.00','0.0%'])
                    self.assertEqual(sheet['G2'].value, '100')
                    self.assertEqual(sheet['H2'].value, .125)
                finally:
                    book.close()
            self.assertEqual(source.read_bytes(), before)
            self.assertFalse(list(Path(temp).glob('.province-*')))

    def test_in_place_numeric_value_is_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self.template(temp, value=108314.96)
            province.write_province_workbook(source, source, '2026-10-27', 162, 25219.7)
            book = load_workbook(source)
            try:
                self.assertEqual(book.active['C2'].value, 108314.96)
                self.assertEqual(book.active['C2'].number_format, '0.00')
                self.assertEqual(book.active['B2'].value, 162)
            finally:
                book.close()
            self.assertFalse(list(Path(temp).glob('.*.rollback.xlsx')))

    def test_invalid_formula_and_partial_save_preserve_existing_output(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self.template(temp, value='=SUM(A1:A2)-100-20')
            target = Path(temp)/'output.xlsx'
            target.write_bytes(b'previous output')
            before = source.read_bytes()
            with self.assertRaisesRegex(RuntimeError, '超出自动计算范围'):
                province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
            self.assertEqual(source.read_bytes(), before)
            source = self.template(temp)

            def fail_save(book, path):
                Path(path).write_bytes(b'partial zip')
                raise OSError('disk failure fixture')

            with patch.object(Workbook, 'save', fail_save), self.assertRaisesRegex(OSError, 'disk failure'):
                province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
            self.assertEqual(target.read_bytes(), b'previous output')
            self.assertFalse(list(Path(temp).glob('.province-*')))

    def test_post_commit_readback_failure_restores_existing_output(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self.template(temp)
            target = self.template(temp, value=42, name='output.xlsx')
            before = target.read_bytes()
            verify = province._verify_values

            def fail_target(path, *args):
                if path == target:
                    raise RuntimeError('readback failure fixture')
                return verify(path, *args)

            with patch.object(province, '_verify_values', fail_target), self.assertRaisesRegex(RuntimeError, 'readback failure'):
                province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
            self.assertEqual(target.read_bytes(), before)
            self.assertFalse(list(Path(temp).glob('.province-*')))

    def test_concurrent_new_output_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as temp:
            source = self.template(temp)
            target = Path(temp)/'output.xlsx'

            def race(temporary, destination):
                destination.write_bytes(b'external save')
                raise FileExistsError(str(destination))

            with patch.object(province, 'publish_new_file', race), self.assertRaisesRegex(RuntimeError, '停止覆盖'):
                province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
            self.assertEqual(target.read_bytes(), b'external save')
            self.assertFalse(list(Path(temp).glob('.province-*')))

    def test_external_source_or_destination_save_during_preparation_is_preserved(self):
        for changed in ('source', 'target'):
            with self.subTest(changed=changed), tempfile.TemporaryDirectory() as temp:
                source = self.template(temp)
                target = self.template(temp, value=42, name='output.xlsx')
                initial = target.read_bytes()
                save = Workbook.save

                def external_save(book, path):
                    save(book, path)
                    (source if changed == 'source' else target).write_bytes(b'external edit')

                with patch.object(Workbook, 'save', external_save), self.assertRaisesRegex(RuntimeError, '工作簿被修改'):
                    province.write_province_workbook(source, target, '2026-10-27', 20, 50000)
                self.assertEqual(target.read_bytes(), initial if changed == 'source' else b'external edit')
                self.assertFalse(list(Path(temp).glob('.province-*')))


if __name__ == '__main__':
    unittest.main()
