"""Reuse immutable workbook inputs during one report parse or pipeline run."""
from __future__ import annotations

from contextvars import ContextVar
from functools import wraps
from pathlib import Path
import sys

from openpyxl import load_workbook as _load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from toolbox_shared.native_excel import load_value_workbook as _load_value_workbook

_workbooks: ContextVar[dict | None] = ContextVar("report_workbooks", default=None)


def load_workbook(path, **kwargs):
    cache = _workbooks.get()
    if cache is None:
        return _load_workbook(path, **kwargs)
    source = Path(path).expanduser().resolve()
    state = source.stat()
    key = (source, state.st_mtime_ns, state.st_size, tuple(sorted(kwargs.items())))
    if key not in cache:
        cache[key] = _load_workbook(source, **kwargs)
    return cache[key]


def reuse_workbook_inputs(function):
    """Nested parsers share their caller's scope; all handles close on exit."""
    @wraps(function)
    def wrapped(*args, **kwargs):
        if _workbooks.get() is not None:
            return function(*args, **kwargs)
        cache = {}
        token = _workbooks.set(cache)
        try:
            return function(*args, **kwargs)
        finally:
            _workbooks.reset(token)
            for workbook in cache.values():
                workbook.close()
    return wrapped


def load_value_workbook(path, **kwargs):
    """Reuse an already loaded styled input, otherwise read only unit values."""
    cache = _workbooks.get()
    if cache is None:
        return _load_value_workbook(path, selected=("基础信息", "新", "汕头", "海门", "东莞"), **kwargs)
    source = Path(path).expanduser().resolve()
    state = source.stat()
    key = (source, state.st_mtime_ns, state.st_size, tuple(sorted(kwargs.items())))
    if key in cache:
        return cache[key]
    value_key = (*key, "values")
    if value_key not in cache:
        cache[value_key] = _load_value_workbook(source, selected=("基础信息", "新", "汕头", "海门", "东莞"), **kwargs)
    return cache[value_key]
