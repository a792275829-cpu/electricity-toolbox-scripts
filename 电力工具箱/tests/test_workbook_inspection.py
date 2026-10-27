from __future__ import annotations
from dataclasses import asdict
from datetime import datetime, date
import os
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from toolbox.runtime import load_module
from toolbox_shared.workbook_inspection import UnsupportedInspection

ROOT = Path(__file__).resolve().parents[2]


class WorkbookInspectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.module = load_module('stream_inspection_tests', ROOT / '功率调整收益统计表更新/update_power_benefit.py')

    def template(self, folder, epoch1904=False):
        path = Path(folder) / '功率预测调整收益统计-归湖.xlsx'
        book = Workbook()
        if epoch1904:
            book.epoch = CALENDAR_MAC_1904
        sheet = book.active; sheet.title = '10月'
        for column, label in self.module.HEADERS.items():
            sheet[f'{column}2'] = label
        for start, value in [(4, datetime(2026, 10, 1)), (29, '=$A$4 + 1')]:
            sheet.cell(start, 1, value).number_format = 'yyyy-mm-dd'
            sheet.merge_cells(start_row=start, start_column=1, end_row=start+23, end_column=1)
            for hour in range(24):
                row = start + hour
                sheet.cell(row, 2, f'{hour}:00')
                for column in self.module.FIELDS:
                    cell = sheet[f'{column}{row}']; cell.number_format = '0.00%'
                    cell.font = Font(name='Arial'); cell.fill = PatternFill('solid', fgColor='FFEEDD')
        sheet['C4'] = 0; sheet['J5'] = '=1+1'
        book.save(path); book.close()
        return path

    def inspect(self, path, mode):
        with patch.dict(os.environ, TOOLBOX_BENEFIT_INSPECT_BACKEND=mode):
            return self.module.inspect_workbook(path, end_date=date(2026,10,2))

    def test_stream_matches_full_for_dates_formulas_styles_and_merged_dates(self):
        for epoch in [False, True]:
            with self.subTest(epoch1904=epoch), tempfile.TemporaryDirectory() as folder:
                path = self.template(folder, epoch)
                before = path.read_bytes()
                self.assertEqual(asdict(self.inspect(path,'full')), asdict(self.inspect(path,'stream')))
                self.assertEqual(path.read_bytes(), before)

    def test_cached_date_and_unselected_shared_formula_master(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.template(folder, True)
            source = path.with_name('source.xlsx'); path.rename(source)
            with ZipFile(source) as archive, ZipFile(path,'w',ZIP_DEFLATED) as output:
                for info in archive.infolist():
                    value = archive.read(info.filename)
                    if info.filename == 'xl/worksheets/sheet1.xml':
                        text = value.decode()
                        serial = to_excel(datetime(2026,10,1), CALENDAR_MAC_1904)
                        text = re.sub(r'<c\b[^>]*r="A4"[^>]*>.*?</c>', f'<c r="A4" s="1"><f>DATE(2026,10,1)</f><v>{serial}</v></c>', text)
                        text = text.replace('<row r="4">','<row r="4"><c r="D4"><f t="shared" si="0" ref="D4:J4">1+1</f><v>2</v></c>')
                        text = re.sub(r'<c\b[^>]*r="J4"[^>]*(?:/>|>.*?</c>)','<c r="J4"><f t="shared" si="0"/><v>2</v></c>',text)
                        # The full loader must clear non-anchor merged values.
                        text = text.replace('<row r="5">','<row r="5"><c r="A5" s="1"><v>999</v></c>')
                        value = text.encode()
                    output.writestr(info,value)
            full = self.inspect(path,'full'); stream = self.inspect(path,'stream')
            self.assertEqual(asdict(full), asdict(stream))
            self.assertNotIn('J4', {cell.coordinate for cell in stream.cells})
            self.assertEqual({cell.run_date for cell in stream.cells}, {date(2026,10,1),date(2026,10,2)})

    def test_merged_input_cells_use_full_loader_in_auto(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path = self.template(folder)
            book = load_workbook(path); book.active.merge_cells('C5:C6'); book.save(path); book.close()
            full = self.inspect(path,'full')
            with self.assertRaises(UnsupportedInspection):
                self.inspect(path,'stream')
            with self.assertWarns(RuntimeWarning):
                automatic = self.inspect(path,'auto')
            self.assertEqual(asdict(full), asdict(automatic))

    def test_unknown_inspection_backend_fails_explicitly(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.template(folder)
            with self.assertRaisesRegex(ValueError,'未知收益检查后端'):
                self.inspect(path,'invalid')

    def test_border_propagation_uses_full_loader(self):
        from openpyxl import load_workbook
        from openpyxl.styles import Border, Side
        with tempfile.TemporaryDirectory() as folder:
            path = self.template(folder)
            book = load_workbook(path)
            book.active['A27'].border = Border(bottom=Side(style='thick'))
            book.save(path); book.close()
            full = self.inspect(path,'full')
            with self.assertRaisesRegex(UnsupportedInspection,'边框'):
                self.inspect(path,'stream')
            with self.assertWarns(RuntimeWarning):
                automatic = self.inspect(path,'auto')
            self.assertEqual(asdict(full),asdict(automatic))

    def test_missing_input_keeps_general_format_with_custom_style_zero(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            path = self.template(folder)
            book = load_workbook(path)
            book._cell_styles[0].numFmtId = 10
            book.save(path); book.close()
            source = path.with_name('source.xlsx'); path.rename(source)
            with ZipFile(source) as archive, ZipFile(path,'w',ZIP_DEFLATED) as output:
                for info in archive.infolist():
                    value = archive.read(info.filename)
                    if info.filename == 'xl/worksheets/sheet1.xml':
                        value = re.sub(rb'<c\b[^>]*r="C6"[^>]*(?:/>|>.*?</c>)',b'',value)
                    output.writestr(info,value)
            full = self.inspect(path,'full'); stream = self.inspect(path,'stream')
            self.assertEqual(asdict(full),asdict(stream))
            self.assertEqual(next(cell.number_format for cell in stream.cells if cell.coordinate=='C6'),'General')

    def test_sparse_verification_rejects_value_formula_format_and_style_changes(self):
        from datetime import date
        from openpyxl import load_workbook
        from openpyxl.styles import Font
        with tempfile.TemporaryDirectory() as folder:
            path=self.template(folder)
            plan=self.inspect(path,'stream')
            series={day:{col:tuple(index+h/100 for h in range(24)) for index,col in enumerate(self.module.FIELDS,1)} for day in {c.run_date for c in plan.cells}}
            prepared=self.module.prepare_workbook(plan,series)
            clean=prepared.temporary.read_bytes()
            for mode in ['full','stream']:
                with patch.dict(os.environ,TOOLBOX_BENEFIT_VERIFY_BACKEND=mode):
                    self.module.verify_workbook(prepared.temporary,prepared)
            cell,_=prepared.expectations[0]
            for kind in ['value','formula','format','font']:
                prepared.temporary.write_bytes(clean)
                book=load_workbook(prepared.temporary);actual=book[cell.sheet][cell.coordinate]
                if kind=='value':actual.value=999
                elif kind=='formula':actual.value='=1+1'
                elif kind=='format':actual.number_format='0.0000'
                else:actual.font=Font(name='Courier New',bold=True)
                book.save(prepared.temporary);book.close()
                for mode in ['full','stream']:
                    with patch.dict(os.environ,TOOLBOX_BENEFIT_VERIFY_BACKEND=mode),self.assertRaises(RuntimeError):
                        self.module.verify_workbook(prepared.temporary,prepared)
