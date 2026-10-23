"""Build a reviewable region/city/node hierarchy from a supplied city workbook."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.analyze_500kv_section_gaps import normalize_price_node_name, parse_topology_edges
from src.node_city_catalog import REGIONS, USER_CONFIRMED_CITIES

MODULE_ROOT = Path(__file__).resolve().parents[1]
CITY_REGION = {city: region for region, cities in REGIONS for city in cities}


def node_key(value: object) -> str:
    # Reuse voltage/bus cleanup; only normalize explicit facility suffix variants.
    # A/B/C unit names and provincial qualifiers remain distinct.
    name = normalize_price_node_name(value).rstrip(".。 ")
    name = re.sub(r"核电厂$", "核电", name)
    return re.sub(r"电厂$", "厂", name)


def labelled_rows(sheet, required: tuple[str, ...]):
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if all(label in row for label in required):
            columns = {label: row.index(label) for label in required}
            break
    else:
        raise ValueError(f"工作表 {sheet.title} 缺少表头：{required}")
    for number, row in enumerate(sheet.iter_rows(min_row=row_number + 1, values_only=True), row_number + 1):
        if any(value is not None for value in row):
            yield number, {label: row[index] if index < len(row) else None
                           for label, index in columns.items()}


def topology_inventory(path: Path):
    edges, _ = parse_topology_edges(path)
    connected = {name for edge in edges for name in (edge.a, edge.b)}
    declared = set()
    active = False
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if re.match(r"^##\s+1\.", line):
            active = True
        elif re.match(r"^##\s+2\.", line):
            break
        elif active and line and not line.startswith("#"):
            declared.update(name.strip() for name in re.split(r"[、，,]", line) if name.strip())
    return connected | declared, connected


def classify(workbook: Path, topology: Path, *, confirmed_cities=None) -> dict:
    confirmed_cities = USER_CONFIRMED_CITIES if confirmed_cities is None else confirmed_cities
    nodes, connected = topology_inventory(topology)
    book = load_workbook(workbook, read_only=True, data_only=True)
    accepted: dict[str, list[dict]] = {}
    pending: dict[str, list[dict]] = {}
    main_names: dict[str, set[str]] = {}
    warnings = []
    source_count = 0
    seen_cities = set()

    def add(index, alias, evidence):
        key = node_key(alias)
        if key and evidence not in index.setdefault(key, []):
            index[key].append(evidence)

    try:
        main = book["各市节点名单"]
        for row, fields in labelled_rows(main, ("地市", "去重节点数量", "节点名称（顿号分隔）")):
            city_raw = str(fields["地市"] or "").strip()
            city = city_raw.removesuffix("市")
            if city not in CITY_REGION:
                raise ValueError(f"来源存在不在指定21市中的地市：{city_raw}")
            seen_cities.add(city)
            names = [name.strip() for name in re.split(r"[、，,\n]", str(fields["节点名称（顿号分隔）"] or "")) if name.strip()]
            source_count += len(names)
            if fields["去重节点数量"] != len(set(names)):
                warnings.append(f"各市节点名单!B{row}：{city_raw}数量栏为 {fields['去重节点数量']}，实际名称 {len(set(names))} 个；按实际名单匹配，未修改原表。")
            for name in names:
                main_names.setdefault(name, set()).add(city)
                add(accepted, name, {"city": city, "source_city": city_raw, "source_node": name,
                                    "sheet": main.title, "row": row, "source_status": "地市主表"})
        if seen_cities != set(CITY_REGION):
            raise ValueError(f"来源主表缺少地市：{sorted(set(CITY_REGION) - seen_cities)}")
        supplement = book["补充节点待核验"]
        raw_label = "原始节点名称（逐条）"
        labels = ("识别状态", "匹配地市", "补充记录统一节点名", raw_label)
        for row, fields in labelled_rows(supplement, labels):
            source_node = str(fields["补充记录统一节点名"] or "").strip()
            city_raw = str(fields["匹配地市"] or "").strip()
            city = city_raw.removesuffix("市")
            status = str(fields["识别状态"] or "")
            evidence = {"city": city if city in CITY_REGION else None, "source_city": city_raw,
                        "source_node": source_node, "sheet": supplement.title, "row": row,
                        "source_status": status}
            verified = status == "已匹配地市主表" and city in main_names.get(source_node, set())
            index = accepted if verified else pending
            for alias in (source_node, *str(fields[raw_label] or "").splitlines()):
                add(index, alias, evidence)
    finally:
        book.close()

    records = []
    for name in sorted(nodes):
        evidence = accepted.get(node_key(name), [])
        cities = {item["city"] for item in evidence}
        city = next(iter(cities)) if len(cities) == 1 else None
        if name in confirmed_cities:
            city = confirmed_cities[name]
            if city not in CITY_REGION:
                raise ValueError(f"确认地市不在指定21市中：{city}")
            evidence = [{"city": city, "source_city": city + "市", "source_node": name,
                         "sheet": "用户确认", "row": None, "source_status": "2026-10-23用户确认"}]
        candidate = pending.get(node_key(name), [])
        reason = ("来源存在多个地市归属，需确认" if len(cities) > 1 else
                  "来源补充表标记待核验，暂不归类" if candidate else
                  "来源表未找到可确认的同名节点；不按相邻接线或相似名称推断")
        records.append({"name": name, "region": CITY_REGION.get(city), "city": city,
                        "status": "classified" if city else "unclassified",
                        "connected": name in connected,
                        "evidence": evidence if city else [*evidence, *candidate],
                        "reason": "" if city else reason})
    hierarchy = [{"name": region, "label": f"{i}. {region}（{len(cities)}市）",
                  "cities": [{"name": city, "nodes": [n["name"] for n in records if n["city"] == city]}
                             for city in cities]}
                 for i, (region, cities) in enumerate(REGIONS, 1)]
    unclassified = [record for record in records if record["status"] == "unclassified"]
    isolated = sorted(nodes - connected)
    if isolated:
        warnings.append("拓扑第一部分列出、但第二部分尚无接线边的节点：" + "、".join(isolated))
    return {"schema_version": 1,
            "source": {"workbook": workbook.name, "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
                       "topology": topology.name, "topology_sha256": hashlib.sha256(topology.read_bytes()).hexdigest()},
            "rules": ["仅包含拓扑名单及已确认接线中的节点，使用拓扑原名作稳定键",
                      "按用户指定的4大区21市排序；不从接线片区标题推断行政地市",
                      "站/变、厂/电厂及核电/核电厂写法规范化；A/B/C及省份后缀保留",
                      "徐闻开关站、玄陆升压站、铜鼓电厂、阳西A厂使用2026-10-23用户确认的地市",
                      "待核验与无法匹配节点单独保留，不纳入已分类菜单；不计算区域均价"],
            "summary": {"topology_nodes": len(nodes), "connected_nodes": len(connected),
                        "classified_nodes": len(records) - len(unclassified), "unclassified_nodes": len(unclassified),
                        "source_listed_names": source_count},
            "regions": hierarchy, "nodes": records, "unclassified": unclassified, "warnings": warnings}


def write_report(result: dict, path: Path) -> None:
    s = result["summary"]
    lines = ["# 500kV拓扑节点区域分类", "",
             f"拓扑名单 {s['topology_nodes']} 个节点，其中实际接线 {s['connected_nodes']} 个；"
             f"已分类 {s['classified_nodes']} 个，待确认 {s['unclassified_nodes']} 个。", "",
             "来源：" + result["source"]["workbook"] + "。原始文件未修改。", "", "## 待确认", ""]
    for node in result["unclassified"]:
        candidates = sorted({e["source_city"] for e in node["evidence"] if e["source_city"]})
        suffix = "；来源候选：" + "、".join(candidates) if candidates else ""
        lines.append(f"- {node['name']}：{node['reason']}{suffix}。")
    if not result["unclassified"]:
        lines.append("无；4个原待确认节点已按用户确认归类。")
    lines.extend(["", "## 来源及范围差异", "", *["- " + w for w in result["warnings"]], ""])
    for region in result["regions"]:
        lines.extend(["## " + region["label"], ""])
        for city in region["cities"]:
            lines.extend([f"### {city['name']}（{len(city['nodes'])}节点）", "", "、".join(city["nodes"]) or "当前拓扑无已分类节点", ""])
    lines.extend(["## 后续曲线界面", "", "一级勾选大区，二级勾选市，三级勾选节点；节点使用拓扑原名。待确认节点单列，不强行放入地市。区域／地市的电价汇总算法留待曲线功能实施时明确。", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--topology", type=Path, default=MODULE_ROOT / "500kV节点拓扑.md")
    parser.add_argument("--output-dir", type=Path, default=MODULE_ROOT / "local_data/区域节点分类")
    args = parser.parse_args()
    result = classify(args.workbook, args.topology)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "500kv_node_regions.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_report(result, args.output_dir / "500kV拓扑节点区域分类.md")
    print(json.dumps(result["summary"], ensure_ascii=False))
    print("待确认：" + ("、".join(node["name"] for node in result["unclassified"]) or "无"))


if __name__ == "__main__":
    main()
