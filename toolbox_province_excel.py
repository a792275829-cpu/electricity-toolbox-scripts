"""Province upload workbook preparation; independent of login, browser and GUI."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
from typing import Any
from types import SimpleNamespace
import os
import re
import shutil
import tempfile

from toolbox_workbook_commit import (assert_workbook_unchanged, commit_workbooks,
                                     workbook_commit_lock, workbook_fingerprint)


def number_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def format_formula_number(value: float) -> str:
    if abs(value - round(value)) < 0.000001:
        return str(int(round(value)))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def replace_capacity_in_formula(formula: str, capacity_value: float) -> str:
    match = re.match(r"^(=.*?-\s*)(\d+(?:\.\d+)?)(\s*-.*)$", formula)
    if not match:
        raise RuntimeError(f"C2 公式格式无法自动替换容量值：{formula}")
    return f"{match.group(1)}{format_formula_number(capacity_value)}{match.group(3)}"


def evaluate_simple_excel_formula(formula: str) -> float:
    expression = formula.strip()
    if expression.startswith("="):
        expression = expression[1:]
    if not re.fullmatch(r"\s*\d+(?:\.\d+)?(?:\s*[+-]\s*\d+(?:\.\d+)?)*\s*", expression):
        raise RuntimeError(f"C2 公式超出自动计算范围：{formula}")

    total = 0.0
    for sign, number_text in re.findall(r"([+-]?)\s*(\d+(?:\.\d+)?)", expression):
        value = float(number_text)
        total += -value if sign == "-" else value
    return total


def write_formula_result_as_value(workbook_path: Path, sheet_name: str, cell: str, value: float) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        worksheet[cell].value = int(round(value)) if abs(value - round(value)) < 0.000001 else value
        workbook.save(workbook_path)
    finally:
        workbook.close()


def _verify_values(path, sheet_name, values, formats):
    from openpyxl import load_workbook
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = book[sheet_name]
        for cell, expected in values.items():
            if sheet[cell].value != expected or sheet[cell].number_format != formats[cell]:
                raise RuntimeError(f"省内数据写入回读不一致：{cell}")
    finally:
        book.close()


def write_province_workbook(source, target, run_date, startup_max, capacity_average, *, log=lambda _: None):
    """Calculate in memory, serialize once, verify, then publish atomically.

    Existing numeric C2 values are preserved exactly; only supported formulas
    have the capacity term substituted and their numeric result written.
    """
    from openpyxl import load_workbook
    source, target = Path(source), Path(target)
    source_fingerprint = workbook_fingerprint(source)
    target_fingerprint = workbook_fingerprint(target) if target.exists() else None
    book = load_workbook(source, data_only=False)
    temporary = None
    try:
        sheet = book['省内数据'] if '省内数据' in book.sheetnames else book.active
        sheet['A2'] = datetime.strptime(run_date, '%Y-%m-%d')
        sheet['B2'] = startup_max
        current = sheet['C2'].value
        if isinstance(current, str) and current.startswith('='):
            formula = replace_capacity_in_formula(current, capacity_average)
            result = evaluate_simple_excel_formula(formula)
            sheet['C2'] = int(round(result)) if abs(result - round(result)) < .000001 else result
            if hasattr(book, 'calculation'):
                book.calculation.fullCalcOnLoad = True
                book.calculation.forceFullCalc = True
        else:
            if number_or_none(current) is None:
                raise RuntimeError(f"C2 不是公式，无法按模板更新：{current!r}")
            log(f"C2 已是数值，保留现有结果：{current}")
        sheet_name = sheet.title
        values = {cell: sheet[cell].value for cell in ('A2', 'B2', 'C2')}
        formats = {cell: sheet[cell].number_format for cell in values}
        descriptor, name = tempfile.mkstemp(prefix='.province-', suffix='.xlsx', dir=target.parent)
        os.close(descriptor)
        temporary = Path(name)
        book.save(temporary)
    except BaseException:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise
    finally:
        book.close()
    try:
        _verify_values(temporary, sheet_name, values, formats)
        assert_workbook_unchanged(source, source_fingerprint)
        if target.with_name(f'~${target.name}').exists():
            raise RuntimeError(f"省内数据目标文件正在被 Excel 打开，请先关闭：{target.name}")
        shutil.copymode(target if target_fingerprint is not None else source, temporary)
        if target_fingerprint is not None:
            prepared = SimpleNamespace(original=target, temporary=temporary,
                                       source_fingerprint=target_fingerprint)
            commit_workbooks((prepared,), lambda path, _: _verify_values(path, sheet_name, values, formats),
                             create_backup=False)
        else:
            with workbook_commit_lock((target,)):
                # link() publishes without overwriting a file created while we prepared.
                try:
                    os.link(temporary, target)
                except FileExistsError as exc:
                    raise RuntimeError(f"准备期间目标工作簿已出现，停止覆盖：{target}") from exc
        return target
    finally:
        temporary.unlink(missing_ok=True)
