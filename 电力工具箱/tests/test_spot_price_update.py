from __future__ import annotations

import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import openpyxl

from toolbox.runtime import ToolPaths, load_module


class DummyAuth:
    @staticmethod
    def unit_display_name(unit):
        return unit.get("unitName") or ""

    @staticmethod
    def is_renewable_unit(unit):
        return unit.get("subType") in {"4", "5"}


class SpotPriceUpdateTests(unittest.TestCase):
    def test_commit_rejects_external_save_without_overwriting_prices(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "归湖.xlsx"
            day = date(2026, 8, 30)
            self._build_fixture(path, day)
            target = self.module.default_targets(path, Path("鮀莲.xlsx"))[0]
            prices = self.module.TargetPrices(target, "company", "unit", "name",
                {day: (300.0,) * 24}, {day: (400.0,) * 24})
            prepared = self.module.prepare_workbook(prices, overwrite=True)
            book = openpyxl.load_workbook(path)
            book[self.module.DAY_AHEAD_SHEET]["B2"] = 999
            book.save(path)
            book.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(RuntimeError, "工作簿被修改"):
                self.module.commit_prepared_workbooks((prepared,), create_backup=False)
            self.assertEqual(path.read_bytes(), original)
            saved = openpyxl.load_workbook(path)
            try:
                self.assertEqual(saved[self.module.DAY_AHEAD_SHEET]["B2"].value, 999)
                self.assertEqual(saved[self.module.DAY_AHEAD_SHEET]["B2"].number_format, "0.0000")
                self.assertIsNone(saved[self.module.REAL_TIME_SHEET]["B2"].value)
            finally:
                saved.close()
            self.assertEqual(list(Path(directory).iterdir()), [path])

    @classmethod
    def setUpClass(cls) -> None:
        cls.paths = ToolPaths(Path(__file__).resolve().parents[2])
        cls.module = load_module("spot_price_update_tests", cls.paths.spot_price_update)

    def test_parse_price_payload_extracts_exact_24_points(self) -> None:
        points = [str(300 + index / 10) for index in range(24)]
        payload = {
            "retCode": "T200",
            "data": {
                "unitDateList": [
                    {
                        "date": "2026-08-30 00:00:00",
                        "unitId": "unit-1",
                        "acPrice": points,
                    }
                ]
            },
        }
        parsed = self.module.parse_price_payload(
            payload,
            unit_id="unit-1",
            expected_dates=(date(2026, 8, 30),),
            label="测试日前电价",
        )
        self.assertEqual(len(parsed[date(2026, 8, 30)]), 24)
        self.assertAlmostEqual(parsed[date(2026, 8, 30)][23], 302.3)

    def test_parse_price_payload_rejects_partial_day(self) -> None:
        payload = {
            "data": {
                "unitDateList": [
                    {
                        "date": "2026-08-30",
                        "unitId": "unit-1",
                        "acPrice": [300] * 23,
                    }
                ]
            }
        }
        with self.assertRaisesRegex(RuntimeError, "应有 24 个电价点"):
            self.module.parse_price_payload(
                payload,
                unit_id="unit-1",
                expected_dates=(date(2026, 8, 30),),
                label="测试实时电价",
            )

    def test_parse_price_payload_can_skip_incomplete_dates_by_price_type(self) -> None:
        complete_date = date(2026, 9, 1)
        incomplete_date = date(2026, 9, 2)
        payload = {
            "data": {
                "unitDateList": [
                    {
                        "date": complete_date.isoformat(),
                        "unitId": "unit-1",
                        "acPrice": [300] * 24,
                    },
                    {
                        "date": incomplete_date.isoformat(),
                        "unitId": "unit-1",
                        "acPrice": [],
                    },
                ]
            }
        }

        parsed = self.module.parse_price_payload(
            payload,
            unit_id="unit-1",
            expected_dates=(complete_date, incomplete_date),
            label="测试实时电价",
            skip_incomplete=True,
        )

        self.assertEqual(set(parsed), {complete_date})
        self.assertEqual(len(parsed[complete_date]), 24)

    def test_company_and_station_mapping(self) -> None:
        target = self.module.default_targets(Path("归湖.xlsx"), Path("鮀莲.xlsx"))[0]
        tenant = self.module.resolve_tenant(
            [
                {"tenantId": "a", "name": "华能（汕头金平）新能源有限责任公司"},
                {"tenantId": "b", "name": "华能（潮州潮安）新能源有限责任公司"},
            ],
            target,
        )
        unit = self.module.resolve_unit(
            [
                {"id": "x", "unitName": "其他风电场", "subType": "4"},
                {"id": "y", "unitName": "华归光伏电站", "subType": "5"},
            ],
            target,
            DummyAuth,
        )
        self.assertEqual(tenant["tenantId"], "b")
        self.assertEqual(unit["id"], "y")

    def test_haifeng_mapping_and_station_selection(self) -> None:
        targets = self.module.default_targets(
            Path("归湖.xlsx"), Path("鮀莲.xlsx"), Path("海风.xlsx"),
            selected_stations=("海风",),
        )
        self.assertEqual([target.key for target in targets], ["海风"])
        self.assertEqual(targets[0].workbook, Path("海风.xlsx"))
        tenant = self.module.resolve_tenant(
            [
                {"tenantId": "a", "name": "华能（汕头金平）新能源有限责任公司"},
                {"tenantId": "b", "name": "华能广东汕头海上风电有限责任公司"},
            ], targets[0],
        )
        unit = self.module.resolve_unit(
            [
                {"id": "x", "unitName": "其他风电场", "subType": "4"},
                {"id": "y", "unitName": "广澳海风场", "subType": "4"},
            ], targets[0], DummyAuth,
        )
        self.assertEqual(tenant["tenantId"], "b")
        self.assertEqual(unit["id"], "y")
        self.assertEqual(
            [target.key for target in self.module.default_targets()],
            ["归湖", "鮀莲", "海风"],
        )
        self.assertEqual(
            [target.key for target in self.module.default_targets(selected_stations=("归湖", "海风"))],
            ["归湖", "海风"],
        )

    def test_station_selection_rejects_empty_and_unknown(self) -> None:
        with self.assertRaisesRegex(ValueError, "至少选择一个"):
            self.module.default_targets(selected_stations=())
        with self.assertRaisesRegex(ValueError, "未知场站"):
            self.module.default_targets(selected_stations=("不存在",))

    def test_cli_dry_run_fetches_only_selected_station(self) -> None:
        with patch("sys.argv", [
            "update_spot_prices.py", "--stations", "海风", "--dry-run",
            "--haifeng-workbook", "/tmp/海风.xlsx",
            "--guihu-workbook", "/not-present/归湖.xlsx",
            "--tuolian-workbook", "/not-present/鮀莲.xlsx",
        ]), patch.object(self.module, "load_config", return_value=self.module.Config()), \
                patch.object(self.module, "collect_target_prices", return_value=()) as collect:
            self.assertEqual(self.module.main(), 0)
        targets = collect.call_args.args[2]
        self.assertEqual([target.key for target in targets], ["海风"])
        self.assertEqual(targets[0].workbook, Path("/tmp/海风.xlsx"))

    def _build_fixture(self, path: Path, target_date: date) -> None:
        workbook = openpyxl.Workbook()
        day_ahead = workbook.active
        day_ahead.title = self.module.DAY_AHEAD_SHEET
        real_time = workbook.create_sheet(self.module.REAL_TIME_SHEET)
        for sheet in (day_ahead, real_time):
            sheet.append(["日期"] + [f"{hour}时" for hour in range(24)])
            sheet.append([datetime.combine(target_date, datetime.min.time())] + [None] * 24)
            sheet["A2"].number_format = "m/d/yy"
            for column in range(2, 26):
                sheet.cell(row=2, column=column).number_format = "0.0000"
        workbook.save(path)
        workbook.close()

    def test_prepare_and_commit_preserve_formats_and_write_both_sheets(self) -> None:
        target_date = date(2026, 8, 30)
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "归湖.xlsx"
            self._build_fixture(workbook_path, target_date)
            target = self.module.TargetSpec(
                key="归湖",
                workbook=workbook_path,
                tenant_keywords=("潮州潮安",),
                unit_keywords=("华归",),
            )
            prices = self.module.TargetPrices(
                target=target,
                tenant_name="华能（潮州潮安）新能源有限责任公司",
                unit_id="unit-1",
                unit_name="华归光伏电站",
                day_ahead={target_date: tuple(300 + index for index in range(24))},
                real_time={target_date: tuple(400 + index for index in range(24))},
            )
            prepared = self.module.prepare_workbook(prices, overwrite=True)
            backups = self.module.commit_prepared_workbooks((prepared,), create_backup=True)

            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            try:
                self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["B2"].value, 300)
                self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["Y2"].value, 323)
                self.assertEqual(workbook[self.module.REAL_TIME_SHEET]["B2"].value, 400)
                self.assertEqual(workbook[self.module.REAL_TIME_SHEET]["Y2"].value, 423)
                self.assertEqual(
                    workbook[self.module.DAY_AHEAD_SHEET]["B2"].number_format,
                    "0.0000",
                )
            finally:
                workbook.close()
            self.assertTrue(backups[workbook_path.resolve()].is_file())

    def test_prepare_can_update_day_ahead_when_real_time_is_missing(self) -> None:
        target_date = date(2026, 9, 1)
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "归湖.xlsx"
            self._build_fixture(workbook_path, target_date)
            target = self.module.TargetSpec(
                key="归湖",
                workbook=workbook_path,
                tenant_keywords=("潮州潮安",),
                unit_keywords=("华归",),
            )
            prices = self.module.TargetPrices(
                target=target,
                tenant_name="华能（潮州潮安）新能源有限责任公司",
                unit_id="unit-1",
                unit_name="华归光伏电站",
                day_ahead={target_date: tuple(300 + index for index in range(24))},
                real_time={},
            )

            prepared = self.module.prepare_workbook(prices, overwrite=True)
            self.module.commit_prepared_workbooks((prepared,), create_backup=False)

            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            try:
                self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["B2"].value, 300)
                self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["Y2"].value, 323)
                self.assertIsNone(workbook[self.module.REAL_TIME_SHEET]["B2"].value)
                self.assertIsNone(workbook[self.module.REAL_TIME_SHEET]["Y2"].value)
            finally:
                workbook.close()

    def test_skip_existing_leaves_row_unchanged(self) -> None:
        target_date = date(2026, 8, 30)
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "归湖.xlsx"
            self._build_fixture(workbook_path, target_date)
            workbook = openpyxl.load_workbook(workbook_path)
            workbook[self.module.DAY_AHEAD_SHEET]["B2"] = 999
            workbook.save(workbook_path)
            workbook.close()

            workbook = openpyxl.load_workbook(workbook_path)
            expectations, updated, skipped, row_values = self.module.plan_price_sheet(
                workbook,
                self.module.DAY_AHEAD_SHEET,
                {target_date: tuple(300 + index for index in range(24))},
                overwrite=False,
            )
            try:
                self.assertEqual(expectations, [])
                self.assertEqual(updated, [])
                self.assertEqual(skipped, [target_date])
                self.assertEqual(row_values, {})
                self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["B2"].value, 999)
            finally:
                workbook.close()

    def test_fixture_dates_cover_requested_range(self) -> None:
        start = date(2026, 8, 30)
        self.assertEqual(
            self.module.date_range(start, start + timedelta(days=2)),
            (start, start + timedelta(days=1), start + timedelta(days=2)),
        )

    def test_extend_tail_preserves_formats_existing_values_and_package(self) -> None:
        start = date(2026, 10, 8)
        for styled_tail in (False, True):
            with self.subTest(styled_tail=styled_tail), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "归湖.xlsx"
                self._build_fixture(path, start)
                workbook = openpyxl.load_workbook(path)
                for sheet in workbook:
                    sheet["B2"] = 999
                    sheet.row_dimensions[2].height = 26
                    if styled_tail:
                        sheet["A50"].number_format = "m/d/yy"
                        sheet["B50"].number_format = "0.0000"
                workbook.save(path)
                workbook.close()
                with zipfile.ZipFile(path, "a") as archive:
                    archive.writestr("customXml/keep.xml", b"<keep/>")
                before = path.read_bytes()
                target = self.module.default_targets(path, Path("unused.xlsx"), selected_stations=("归湖",))[0]
                points = tuple(float(index - 12) for index in range(24))
                prices = self.module.TargetPrices(target, "company", "unit", "name",
                    {start: points, start + timedelta(days=3): points},
                    {start + timedelta(days=1): points})
                prepared = self.module.prepare_workbook(prices, overwrite=False)
                self.assertEqual(path.read_bytes(), before)
                self.module.commit_prepared_workbooks((prepared,), create_backup=False)
                workbook = openpyxl.load_workbook(path)
                try:
                    day = workbook[self.module.DAY_AHEAD_SHEET]
                    real = workbook[self.module.REAL_TIME_SHEET]
                    self.assertEqual(day["B2"].value, 999)
                    self.assertEqual([day.cell(row, 1).value.date() for row in range(2, 6)],
                        [start + timedelta(days=offset) for offset in range(4)])
                    self.assertEqual(tuple(day.cell(5, column).value for column in range(2, 26)), points)
                    self.assertIsNone(day["B4"].value)
                    self.assertEqual(day["B4"].number_format, "0.0000")
                    self.assertEqual(day["A5"].number_format, "m/d/yy")
                    self.assertEqual(day["Y5"].number_format, "0.0000")
                    self.assertEqual(day.row_dimensions[5].height, 26)
                    self.assertEqual(real["A3"].value.date(), start + timedelta(days=1))
                    self.assertIsNone(real["A4"].value)
                    self.assertEqual(real["Y3"].value, points[-1])
                finally:
                    workbook.close()
                with zipfile.ZipFile(path) as archive:
                    self.assertEqual(archive.read("customXml/keep.xml"), b"<keep/>")
                self.module.verify_workbook(path, prepared.expectations, prepared.date_expectations)

    def test_extension_respects_1904_epoch_and_date_readback(self) -> None:
        start = date(2026, 10, 8)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prices.xlsx"
            self._build_fixture(path, start)
            workbook = openpyxl.load_workbook(path)
            workbook.epoch = openpyxl.utils.datetime.MAC_EPOCH
            workbook.save(path)
            workbook.close()
            target = self.module.default_targets(path, Path("unused.xlsx"), selected_stations=("归湖",))[0]
            prices = self.module.TargetPrices(target, "company", "unit", "name",
                {start + timedelta(days=1): (0.0,) * 24}, {})
            prepared = self.module.prepare_workbook(prices, overwrite=True)
            self.module.commit_prepared_workbooks((prepared,), create_backup=False)
            workbook = openpyxl.load_workbook(path)
            self.assertEqual(workbook[self.module.DAY_AHEAD_SHEET]["A3"].value.date(), start + timedelta(days=1))
            workbook[self.module.DAY_AHEAD_SHEET]["A3"].number_format = "yyyy-mm"
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(RuntimeError, "日期回读失败"):
                self.module.verify_workbook(path, prepared.expectations, prepared.date_expectations)

    def test_extension_rejects_missing_history_duplicates_and_footer(self) -> None:
        for problem in ("gap", "duplicate", "footer"):
            with self.subTest(problem=problem), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "prices.xlsx"
                start = date(2026, 10, 8)
                self._build_fixture(path, start)
                workbook = openpyxl.load_workbook(path)
                sheet = workbook[self.module.DAY_AHEAD_SHEET]
                if problem == "gap":
                    sheet["A3"] = start + timedelta(days=2)
                elif problem == "duplicate":
                    sheet["A3"] = start
                else:
                    sheet["Z3"] = "合计"
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                target = self.module.default_targets(path, Path("unused.xlsx"), selected_stations=("归湖",))[0]
                prices = self.module.TargetPrices(target, "company", "unit", "name",
                    {start + timedelta(days=offset): (300.0,) * 24 for offset in (1, 3)}, {})
                with self.assertRaisesRegex(RuntimeError, "缺少日期|日期重复|其他内容"):
                    self.module.prepare_workbook(prices, overwrite=True)
                self.assertEqual(path.read_bytes(), before)
                self.assertEqual(list(path.parent.glob(".*.tmp.xlsx")), [])

    def test_extension_cleans_temporary_file_when_verification_fails(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prices.xlsx"
            start = date(2026, 10, 8)
            self._build_fixture(path, start)
            target = self.module.default_targets(path, Path("unused.xlsx"), selected_stations=("归湖",))[0]
            prices = self.module.TargetPrices(target, "company", "unit", "name",
                {start + timedelta(days=1): (300.0,) * 24}, {})
            before = path.read_bytes()
            with patch.object(self.module, "verify_workbook", side_effect=RuntimeError("readback failed")):
                with self.assertRaisesRegex(RuntimeError, "readback failed"):
                    self.module.prepare_workbook(prices, overwrite=True)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(path.parent.glob(".*.tmp.xlsx")), [])

    def test_xml_extension_handles_self_closing_rows_and_cells(self) -> None:
        source = b'<worksheet><dimension ref="A1:Z50"/><sheetData><row r="2" ht="26"><c r="A2" s="1"><v>46303</v></c><c r="B2" s="2"/></row><row r="3"/><row r="4"><c r="A4" s="1"/><c r="B4" s="2"/><c r="Z4" s="3"/></row><row r="50"/></sheetData></worksheet>'
        extended = self.module.extend_sheet_dates(source, {3: date(2026, 10, 9), 4: date(2026, 10, 10)},
            epoch=openpyxl.utils.datetime.WINDOWS_EPOCH)
        patched = self.module.patch_sheet_xml(extended, {3: (300.0,) * 24, 4: (400.0,) * 24})
        root = self.module.ElementTree.fromstring(patched)
        rows = root.findall("sheetData/row")
        self.assertEqual([int(row.attrib["r"]) for row in rows], [2, 3, 4, 50])
        self.assertEqual(root.find("dimension").attrib["ref"], "A1:Z50")
        self.assertIsNotNone(rows[2].find('c[@r="Z4"]'))
        self.assertEqual(rows[2].find('c[@r="B4"]/v').text, "400")


if __name__ == "__main__":
    unittest.main()
