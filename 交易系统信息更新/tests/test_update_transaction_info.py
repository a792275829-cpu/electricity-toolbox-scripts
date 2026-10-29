from __future__ import annotations

import importlib.util
import math
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

MODULE_PATH = Path(__file__).resolve().parents[1] / "update_transaction_info.py"
SPEC = importlib.util.spec_from_file_location("transaction_info_tests", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)

RUN_DATE = date(2026, 10, 10)


def period_row(day: date, *, max_value=1, min_value=2, avg_value=3, total=4) -> dict:
    return {
        "date": day.isoformat() + " 00:00:00",
        "max": max_value, "min": min_value, "avg": avg_value, "sum": total,
    }


def dataset(*, day: date = RUN_DATE, day_ahead: bool = True,
            labels=("1", "6", "4", "2", "3", "5")) -> dict:
    key = "forecastPeriodList" if day_ahead else "actualPeriodList"
    return {"data": {"dataNetLoadDTOList": [
        {"loadType": label, key: [period_row(day)]} for label in labels
    ]}}


class TempFolderTest(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.folder = Path(self._temp.name)

    def tearDown(self):
        self._temp.cleanup()


class DateRangeTests(unittest.TestCase):
    def test_range_is_inclusive(self):
        self.assertEqual(MODULE.date_range(date(2026, 10, 1), date(2026, 10, 2)),
                         (date(2026, 10, 1), date(2026, 10, 2)))

    def test_reversed_range_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "结束日期不能早于开始日期"):
            MODULE.date_range(date(2026, 10, 2), date(2026, 10, 1))


class NumberTests(unittest.TestCase):
    def test_integral_values_become_ints(self):
        self.assertEqual(MODULE._number("4", label="总电量"), 4)
        self.assertIsInstance(MODULE._number(4.0, label="总电量"), int)

    def test_fractional_values_keep_ten_digits(self):
        self.assertEqual(MODULE._number(1.234567890123, label="平均"), round(1.234567890123, 10))

    def test_rejects_placeholders_and_non_finite(self):
        for value in (None, "", "-", True):
            with self.subTest(value=value), self.assertRaisesRegex(RuntimeError, "为空"):
                MODULE._number(value, label="最大负荷")
        with self.assertRaisesRegex(RuntimeError, "不是有效数值"):
            MODULE._number("abc", label="最大负荷")
        for value in (math.inf, math.nan):
            with self.subTest(value=value), self.assertRaisesRegex(RuntimeError, "不是有限数值"):
                MODULE._number(value, label="最大负荷")


class PayloadRowsTests(unittest.TestCase):
    def test_dict_payload_rows(self):
        self.assertEqual(len(MODULE._payload_rows(dataset())), 6)

    def test_other_shapes_are_rejected(self):
        with self.assertRaisesRegex(RuntimeError, "返回格式异常"):
            MODULE._payload_rows({"data": {"dataNetLoadDTOList": "oops"}})
        self.assertEqual(MODULE._payload_rows({"data": []}), [])
        self.assertEqual(MODULE._payload_rows({}), [])

    def test_non_dict_rows_are_dropped(self):
        payload = {"data": {"dataNetLoadDTOList": [None, "x", {"loadType": "1"}]}}
        self.assertEqual(MODULE._payload_rows(payload), [{"loadType": "1"}])

    def test_period_date_parsing(self):
        self.assertEqual(MODULE._period_date("2026-10-10 00:00:00"), RUN_DATE)
        self.assertIsNone(MODULE._period_date("10/10/2026"))
        self.assertIsNone(MODULE._period_date(None))


class SummarizeLoadTypeTests(unittest.TestCase):
    def summarize(self, payload, load_type="1", *, day=RUN_DATE, key="forecastPeriodList"):
        return MODULE.summarize_load_type(
            payload, load_type, day, period_list_key=key, label="日前")

    def test_values_and_unit_conversion(self):
        result = self.summarize(dataset())
        self.assertEqual(result, (1, 2, 3, 400_000))

    def test_single_period_is_rejected_when_date_misses(self):
        payload = {"data": {"dataNetLoadDTOList": [
            {"loadType": "1", "forecastPeriodList": [period_row(date(2026, 10, 9))]}]}}
        with self.assertRaisesRegex(RuntimeError, "对应日期"):
            self.summarize(payload)

    def test_missing_load_type_lists_available_types(self):
        with self.assertRaisesRegex(RuntimeError, "可用类型：1、2"):
            self.summarize(dataset(labels=("1", "2")), load_type="6")

    def test_missing_period_is_reported(self):
        payload = {"data": {"dataNetLoadDTOList": [
            {"loadType": "1", "forecastPeriodList": [period_row(date(2026, 10, 9)),
                                                     period_row(date(2026, 10, 8))]}]}}
        with self.assertRaisesRegex(RuntimeError, "对应日期"):
            self.summarize(payload)


class ParseLoadDatasetTests(unittest.TestCase):
    def test_actual_requires_every_category(self):
        parsed = MODULE.parse_load_dataset(dataset(day_ahead=False), RUN_DATE, day_ahead=False)
        self.assertEqual(tuple(parsed), MODULE.CATEGORY_ORDER)

    def test_day_ahead_skips_the_western_power_category(self):
        parsed = MODULE.parse_load_dataset(dataset(labels=("1", "6", "2", "3", "5")),
                                           RUN_DATE, day_ahead=True)
        self.assertEqual(tuple(parsed), MODULE.DAY_AHEAD_REQUIRED)
        self.assertNotIn("西电", parsed)

    def test_day_ahead_rejects_missing_required_category(self):
        with self.assertRaisesRegex(RuntimeError, "loadType=6"):
            MODULE.parse_load_dataset(dataset(labels=("1", "2", "3", "5")),
                                      RUN_DATE, day_ahead=True)

    def test_actual_reads_actual_period_list(self):
        payload = dataset(day_ahead=False)
        parsed = MODULE.parse_load_dataset(payload, RUN_DATE, day_ahead=False)
        self.assertEqual(parsed["统调"], (1, 2, 3, 400_000))


class DailyPairTests(unittest.TestCase):
    def test_actual_date_is_day_ahead_minus_three(self):
        pair = MODULE.build_daily_load_pair(RUN_DATE, dataset(day_ahead=True),
                                            dataset(day=RUN_DATE - timedelta(days=3), day_ahead=False))
        self.assertEqual(pair.day_ahead_date, RUN_DATE)
        self.assertEqual(pair.actual_date, RUN_DATE - timedelta(days=MODULE.DATE_OFFSET_DAYS))
        self.assertEqual(tuple(pair.day_ahead), MODULE.DAY_AHEAD_REQUIRED)
        self.assertEqual(tuple(pair.actual), MODULE.CATEGORY_ORDER)


class CellDateTests(unittest.TestCase):
    def test_datetime_and_date_values(self):
        epoch = openpyxl.Workbook().epoch
        self.assertEqual(MODULE._as_date(datetime(2026, 10, 10, 8), epoch=epoch), RUN_DATE)
        self.assertEqual(MODULE._as_date(RUN_DATE, epoch=epoch), RUN_DATE)

    def test_excel_serial_and_string_forms(self):
        epoch = openpyxl.Workbook().epoch
        self.assertEqual(MODULE._as_date(46305, epoch=epoch), date(2026, 10, 10))
        self.assertEqual(MODULE._as_date("2026/10/10", epoch=epoch), RUN_DATE)

    def test_unparseable_values_return_none(self):
        epoch = openpyxl.Workbook().epoch
        self.assertIsNone(MODULE._as_date("10/10/2026", epoch=epoch))
        self.assertIsNone(MODULE._as_date(None, epoch=epoch))


class TailWorkbookMixin:
    def build(self, *, day_ahead=RUN_DATE, gap=MODULE.DATE_OFFSET_DAYS,
              order=MODULE.CATEGORY_ORDER, sheet=MODULE.SHEET_NAME, groups=1) -> Path:
        workbook = openpyxl.Workbook()
        sheet_object = workbook.active
        sheet_object.title = sheet
        sheet_object.append(["表头"] * 11)
        for group in range(groups):
            group_day = day_ahead + timedelta(days=group)
            actual = group_day - timedelta(days=gap)
            for index, label in enumerate(order):
                row = [None] * 11
                if index == 0:
                    row[2] = datetime(group_day.year, group_day.month, group_day.day)
                    row[9] = datetime(actual.year, actual.month, actual.day)
                row[3] = label
                row[10] = label
                sheet_object.append(row)
                current_row = sheet_object.max_row
                for column in (*(() if label == "西电" else range(5, 9)), *range(12, 16)):
                    sheet_object.cell(current_row, column, 1)
        path = self.folder / "交易系统信息.xlsx"
        workbook.save(path)
        return path


class WorkbookTailTests(TempFolderTest, TailWorkbookMixin):
    def test_tail_dates_and_row_are_read(self):
        tail = MODULE.inspect_workbook_tail(self.build())
        self.assertEqual(tail.start_row, 2)
        self.assertEqual(tail.day_ahead_date, RUN_DATE)
        self.assertEqual(tail.actual_date, RUN_DATE - timedelta(days=MODULE.DATE_OFFSET_DAYS))

    def test_last_group_wins(self):
        tail = MODULE.inspect_workbook_tail(self.build(groups=2))
        self.assertEqual(tail.start_row, 8)
        self.assertEqual(tail.day_ahead_date, RUN_DATE + timedelta(days=1))

    def test_wrong_offset_is_rejected(self):
        with self.assertRaisesRegex(RuntimeError, "不符合 D / D-3"):
            MODULE.inspect_workbook_tail(self.build(gap=2))

    def test_wrong_field_order_is_rejected(self):
        broken = ("统调", "西电", "B类", "A类", "地方电", "粤港联络线")
        with self.assertRaisesRegex(RuntimeError, "字段顺序异常"):
            MODULE.inspect_workbook_tail(self.build(order=broken))

    def test_missing_sheet_and_suffix_are_rejected(self):
        with self.assertRaisesRegex(RuntimeError, "没有工作表"):
            MODULE.inspect_workbook_tail(self.build(sheet="其他"))
        wrong_suffix = self.folder / "旧表.xls"
        wrong_suffix.write_bytes(b"x")
        with self.assertRaisesRegex(ValueError, "仅支持 .xlsx"):
            MODULE.inspect_workbook_tail(wrong_suffix)
        with self.assertRaisesRegex(FileNotFoundError, "不存在"):
            MODULE.inspect_workbook_tail(self.folder / "无.xlsx")


class StartDateTests(TempFolderTest, TailWorkbookMixin):
    def test_start_date_follows_the_workbook_tail(self):
        workbook = self.build()
        self.assertEqual(MODULE.resolve_start_date(workbook, None), RUN_DATE + timedelta(days=1))
        self.assertEqual(MODULE.resolve_start_date(workbook, RUN_DATE + timedelta(days=1)),
                         RUN_DATE + timedelta(days=1))

    def test_mismatched_input_is_rejected(self):
        workbook = self.build()
        with self.assertRaisesRegex(ValueError, "开始日期必须是 2026-10-11"):
            MODULE.resolve_start_date(workbook, RUN_DATE)


class ValueComparisonTests(unittest.TestCase):
    def test_near_equal_numbers_match(self):
        self.assertTrue(MODULE._values_equal("12.00000001", 12))
        self.assertFalse(MODULE._values_equal("12.1", 12))
        self.assertFalse(MODULE._values_equal(None, 12))
        self.assertFalse(MODULE._values_equal(True, 1))


class RowXmlTests(unittest.TestCase):
    XML = (
        '<worksheet><sheetData>'
        '<row r="7" spans="1:26"><c r="C7" s="3"><v>45000</v></c><c r="D7" t="s"><v>0</v></c></row>'
        '<row r="8" spans="1:26"><c r="C8" s="3"><v>1</v></c></row>'
        '</sheetData></worksheet>'
    )

    def test_row_is_extracted_exactly(self):
        row = MODULE._row_xml(self.XML, 7)
        self.assertIn('r="7"', row)
        self.assertNotIn('r="8"', row)

    def test_missing_row_is_reported(self):
        with self.assertRaisesRegex(RuntimeError, "未找到模板第 9 行"):
            MODULE._row_xml(self.XML, 9)

    def test_renumber_moves_row_and_cell_references(self):
        renumbered = MODULE._renumber_row(MODULE._row_xml(self.XML, 7), 7, 14)
        self.assertIn('r="14"', renumbered)
        self.assertIn('r="C14"', renumbered)
        self.assertNotIn('C7', renumbered)

    def test_numeric_cell_replacement_keeps_style(self):
        row = MODULE._row_xml(self.XML, 7)
        replaced = MODULE._replace_numeric_cell(row, "C7", 12.5)
        self.assertIn('<c r="C7" s="3"><v>12.5</v></c>', replaced)
        self.assertIn('r="D7"', replaced)

    def test_missing_cell_is_reported(self):
        with self.assertRaisesRegex(RuntimeError, "缺少单元格 E7"):
            MODULE._replace_numeric_cell(MODULE._row_xml(self.XML, 7), "E7", 1)


if __name__ == "__main__":
    unittest.main()
