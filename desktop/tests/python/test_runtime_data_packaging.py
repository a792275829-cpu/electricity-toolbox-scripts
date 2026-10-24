from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path[:0] = [str(ROOT), str(ROOT / "packaging/tauri")]
import prepare_payload
import prepare_workspace


class RuntimeDataPackagingTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.payload = self.base / "payload"
        (self.payload / "app").mkdir(parents=True)
        self.original = self.base / "总表.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "整合数据"
        sheet.append(["日期", "时刻", "全省日前平均电价"])
        sheet.append(["2026-10-11", "00:00", 123.45])
        sheet["C2"].number_format = "0.00"
        book.save(self.original)
        book.close()

    def write_manifest(self, build):
        name = prepare_payload.FORECAST_SEED.as_posix()
        (self.payload / "manifest.json").write_text(json.dumps({
            "files": {}, "build": build, "defaults": [name],
            "default_hashes": {name: prepare_workspace.digest(self.payload / "defaults" / name)},
        }))

    def install(self):
        prepare_workspace.install(
            self.payload, self.base / "user/app", data_root=self.base / "user/data",
            state_root=self.base / "user/state",
        )

    def test_seed_roundtrip_and_upgrade_preserve_updated_workbook(self):
        before = self.original.read_bytes()
        target = prepare_payload.prepare_forecast_seed(self.payload / "defaults", self.original)
        self.assertEqual(target.read_bytes(), before)
        self.write_manifest("first")
        self.install()
        installed = self.base / "user/data" / prepare_payload.FORECAST_SEED
        book = load_workbook(installed, data_only=True)
        self.assertEqual(book["整合数据"]["C2"].value, 123.45)
        self.assertEqual(book["整合数据"]["C2"].number_format, "0.00")
        book.close()
        installed.write_bytes(b"users updated master")
        self.write_manifest("next")
        self.install()
        self.assertEqual(installed.read_bytes(), b"users updated master")
        self.assertEqual(self.original.read_bytes(), before)

    def test_corrupt_seed_blocks_install_before_state_is_written(self):
        target = prepare_payload.prepare_forecast_seed(self.payload / "defaults", self.original)
        self.write_manifest("first")
        target.write_bytes(b"corrupt")
        with self.assertRaisesRegex(ValueError, "首次安装资源损坏"):
            self.install()
        self.assertFalse((self.base / "user/state/installed.json").exists())

    def test_missing_invalid_and_env_selected_workbook(self):
        with self.assertRaisesRegex(FileNotFoundError, "总数据表"):
            prepare_payload.prepare_forecast_seed(self.base, self.base / "missing.xlsx")
        with mock.patch.dict("os.environ", {"FORECAST_MASTER_WORKBOOK": str(self.original)}):
            target = prepare_payload.prepare_forecast_seed(self.payload / "defaults")
        self.assertEqual(target.read_bytes(), self.original.read_bytes())
        book = Workbook()
        book.save(self.original)
        book.close()
        with self.assertRaisesRegex(ValueError, "必要字段"):
            prepare_payload.prepare_forecast_seed(self.base, self.original)

    def test_each_required_runtime_asset_is_exported_and_enforced(self):
        source = self.base / "assets"
        for relative in prepare_payload.REQUIRED_RUNTIME_ASSETS:
            self.assertTrue(prepare_payload.runtime_source(relative), str(relative))
            path = source / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"resource")
        prepare_payload.validate_runtime_assets(source)
        for relative in prepare_payload.REQUIRED_RUNTIME_ASSETS:
            with self.subTest(relative=relative):
                path = source / relative
                path.unlink()
                with self.assertRaises(FileNotFoundError):
                    prepare_payload.validate_runtime_assets(source)
                path.write_bytes(b"resource")

    def test_installed_group_upload_uses_bundled_template_with_separate_data_root(self):
        source = self.payload / "app"
        prepare_payload.build_support.export_source(source, predicate=prepare_payload.runtime_source)
        relative = Path("集团每日上传/templates/省内数据模板.xlsx")
        self.assertIn(relative, prepare_payload.REQUIRED_RUNTIME_ASSETS)
        files = {p.relative_to(source).as_posix(): prepare_workspace.digest(p)
                 for p in source.rglob("*") if p.is_file()}
        (self.payload / "manifest.json").write_text(json.dumps({
            "layout": 2, "version": "test", "build": "group-template-test", "files": files,
        }))
        self.install()
        installed = self.base / "user/app"
        self.assertEqual((installed / relative).read_bytes(), (ROOT / relative).read_bytes())
        data_root = self.base / "user/data"
        result = subprocess.run([
            sys.executable, "-B", "-c",
            "import sys; from pathlib import Path; "
            "sys.path[:0] = ['.', '电力工具箱']; "
            "from toolbox.runtime import load_module; "
            "m = load_module('installed_group_template_test', Path('集团每日上传/upload_daily_report.py')); "
            "folder = Path(sys.argv[1]); folder.mkdir(parents=True, exist_ok=True); "
            "energy = folder / '广东-能销数据-20261023.xlsx'; energy.write_bytes(b'energy'); "
            "day, uploads = m.prepare_upload_files([energy]); "
            "assert day == '2026-10-23' and uploads[0].builtin_template; "
            "assert uploads[0].path.is_relative_to(Path.cwd()); "
            "assert not uploads[0].path.is_relative_to(folder); "
            "assert m.province_output_path(uploads[0].path, day, output_dir=folder).parent == folder; "
            "print('Installed builtin province template OK')",
            str(data_root / "energy"),
        ], cwd=installed, env={
            **os.environ, "TOOLBOX_CODE_ROOT": str(installed), "TOOLBOX_DATA_ROOT": str(data_root),
        }, capture_output=True, text=True, check=True)
        self.assertIn("Installed builtin province template OK", result.stdout)


if __name__ == "__main__":
    unittest.main()
