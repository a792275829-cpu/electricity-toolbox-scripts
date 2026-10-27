import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[3]
SCRIPT = ROOT / '电力工具脚本/电力交易分析工具/generate_electricity_report.py'


def make_workbook(path, realtime=False):
    book = Workbook()
    sheet = book.active
    sheet.title = '日前'
    for row in (4, 5, 8, 9, 10, 13, 16, 17, 18, 19, 21):
        for column, value in ((3, 12), (4, 250), (5, .5), (6, 24),
                              (7, 300), (8, 20), (9, 4), (10, 50), (11, .2)):
            sheet.cell(row, column, value)
    for name in ('运行补偿', '汕头', '海门', '东莞', '海风', '鮀莲', '归湖'):
        sheet = book.create_sheet(name)
        if name in ('海风', '鮀莲', '归湖'):
            for hour in range(24):
                sheet.cell(hour + 4, 1, f'{hour}:00')
                for column, value in ((2, 10), (3, 300), (4, 12 if realtime else 0),
                                      (5, 350 if realtime else 0), (12, 8), (13, 250)):
                    sheet.cell(hour + 4, column, value)
            sheet['AE28'] = 321.5
    book.save(path)
    book.close()


def legacy_module():
    spec = importlib.util.spec_from_file_location('trade_legacy_contract', SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TradeEngineTests(unittest.TestCase):
    def test_outputs_match_pre_refactor_contract_in_both_backends(self):
        from toolbox_trade_report import ReportBuilder, generate_report
        legacy = legacy_module()
        self.assertIs(legacy.ReportBuilder, ReportBuilder)
        cases = json.loads(Path(__file__).with_name('trade_render_contract.json').read_text())
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for case in cases:
                source = root / case['name']
                make_workbook(source, case['realtime'])
                before = source.read_bytes()
                for backend in ('python', 'rust'):
                    with self.subTest(case=case, backend=backend), patch.dict(os.environ, TOOLBOX_XLSX_BACKEND=backend):
                        output = generate_report(source, root/'new.html')
                        self.assertEqual(hashlib.sha256(output.read_bytes()).hexdigest(), case['sha256'])
                        self.assertEqual(legacy.generate_report(source, root/'legacy.html').read_bytes(), output.read_bytes())
                        self.assertEqual(source.read_bytes(), before)

    def test_service_cli_and_deviation_import_work_without_tk_or_dynamic_loader(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/'trade-2026-10-27.xlsx'
            make_workbook(source, realtime=True)
            script = '''
import sys, importlib.abc, runpy
from pathlib import Path
class NoTk(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'tkinter' or fullname.startswith('tkinter.'):
            raise AssertionError('unexpected Tk import: ' + fullname)
sys.meta_path.insert(0, NoTk())
from toolbox.runtime import ToolPaths
from toolbox import workflows
from unittest.mock import patch
source, output, entry, deviation = map(Path, sys.argv[1:])
with patch.object(workflows, 'load_module', side_effect=AssertionError('dynamic load')):
    p = workflows.validate(ToolPaths(Path.cwd()), 'trade-analysis', 'run',
                           dict(files=[str(source)], output_dir=str(output)))
    result = workflows.execute(ToolPaths(Path.cwd()), 'trade-analysis', 'run', p)
    assert len(result['outputs']) == 1
sys.argv = [str(entry), str(source), '--no-gui', '--output', str(output/'cli.html')]
runpy.run_path(str(entry), run_name='__main__')
assert (output/'cli.html').read_bytes() == Path(result['outputs'][0]).read_bytes()
runpy.run_path(str(deviation), run_name='import_check')
assert 'tkinter' not in sys.modules
'''
            result = subprocess.run([sys.executable, '-c', script, str(source), str(root/'out'),
                                     str(SCRIPT), str(SCRIPT.with_name('find_output_range_deviation_days.py'))],
                                    cwd=ROOT, env=dict(os.environ, PYTHONPATH=os.pathsep.join([str(ROOT), str(ROOT/'电力工具箱')])),
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_standalone_gui_opens_with_available_macos_theme(self):
        import tkinter as tk
        module = legacy_module()
        try:
            root = tk.Tk()
        except tk.TclError as exc:
            self.skipTest(str(exc))
        try:
            root.withdraw()
            with patch.object(tk, 'Tk', return_value=root), patch.object(root, 'mainloop') as loop:
                module.launch_gui()
                loop.assert_called_once()
        finally:
            root.destroy()


if __name__ == '__main__':
    unittest.main()
