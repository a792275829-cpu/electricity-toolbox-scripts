from __future__ import annotations

from datetime import datetime
import hashlib
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from toolbox.runtime import load_module
from toolbox.upload_service import run_group_upload_with_staging


class GroupUploadTemplateTests(unittest.TestCase):
    def setUp(self):
        self.module = load_module("group_upload_builtin_template_test", ROOT / "集团每日上传/upload_daily_report.py")
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name).resolve()
        self.energy = self.folder / "广东-能销数据-20261023.xlsx"
        self.energy.write_bytes(b"energy-input")
        self.template = self.module.BUILTIN_PROVINCE_TEMPLATE
        self.original_template = self.template.read_bytes()

    def prepare(self):
        return self.module.prepare_upload_files([self.energy])

    def update(self, paths):
        module = self.module
        run_date, uploads = module.prepare_upload_files(paths)
        with (
            mock.patch.object(module.upload_flow, "fetch_startup_max", return_value=162),
            mock.patch.object(module.upload_flow, "fetch_capacity_average", return_value=25219.7),
        ):
            return module.update_province_excel_for_upload(object(), run_date, uploads)

    def assert_output(self, path):
        self.assertEqual(path.name, "广东-省内数据-20261023.xlsx")
        for data_only in (False, True):
            workbook = load_workbook(path, data_only=data_only)
            source = load_workbook(self.template, data_only=data_only)
            try:
                self.assertEqual(workbook.sheetnames, ["省内数据"])
                sheet = workbook["省内数据"]
                self.assertEqual(sheet["A2"].value, datetime(2026, 10, 23))
                self.assertEqual(sheet["B2"].value, 162)
                self.assertEqual(sheet["C2"].value, 108314.96)
                for row in sheet.iter_rows(min_row=1, max_row=2, max_col=11):
                    for cell in row:
                        expected = source.active[cell.coordinate]
                        self.assertEqual(cell.number_format, expected.number_format)
                        if cell.coordinate not in ("A2", "B2"):
                            self.assertEqual(cell.value, expected.value)
                            self.assertEqual(cell.data_type, expected.data_type)
                self.assertEqual(sheet["A2"].number_format, "yyyy/m/d;@")
                self.assertEqual(sheet["G2"].value, "100")
                self.assertIsNone(sheet["J2"].value)
                self.assertIsNone(sheet["K2"].value)
            finally:
                workbook.close()
                source.close()
        self.assertEqual(self.template.read_bytes(), self.original_template)
        self.assertEqual(self.energy.read_bytes(), b"energy-input")

    def test_energy_only_uses_builtin_and_gui_selection_can_be_prepared_again(self):
        run_date, uploads = self.prepare()
        self.assertEqual(run_date, "2026-10-23")
        self.assertEqual([item.kind for item in uploads], ["省内", "能销"])
        self.assertTrue(uploads[0].builtin_template)
        self.assertTrue(uploads[0].auto_template)
        self.assertEqual(uploads[0].path, self.template)
        repeated_date, repeated = self.module.prepare_upload_files([item.path for item in uploads])
        self.assertEqual(repeated_date, run_date)
        self.assertEqual(repeated, uploads)

    def test_undated_province_and_excel_lock_file_do_not_block_fallback(self):
        (self.folder / "省内数据.xlsx").write_bytes(b"undated")
        (self.folder / "~$广东-省内数据-20261023.xlsx").write_bytes(b"lock")
        self.assertTrue(self.prepare()[1][0].builtin_template)

    def test_builtin_generates_real_workbook_next_to_energy(self):
        output = self.update([self.energy])[0].path
        self.assertEqual(output.parent, self.folder)
        self.assert_output(output)
        source = load_workbook(self.template)
        try:
            self.assertIsNone(source.active["A2"].value)
            self.assertIsNone(source.active["B2"].value)
        finally:
            source.close()

    def test_staging_copies_builtin_output_back_to_energy_folder(self):
        observed = []
        def upload(paths, *, force):
            self.assertFalse(force)
            observed.extend(paths)
            self.assertNotEqual(paths[1].parent, self.folder)
            self.assertEqual(paths[0], self.template)
            self.update(paths)
        with mock.patch.object(self.module, "run_api_upload", side_effect=upload):
            run_group_upload_with_staging(self.module, [self.energy], force=False, on_log=lambda _: None)
        self.assert_output(self.folder / "广东-省内数据-20261023.xlsx")
        self.assertFalse(observed[1].parent.exists())

    def test_same_directory_template_is_staged_even_when_only_energy_is_selected(self):
        from shutil import copy2
        province = self.folder / "广东-省内数据-20261022.xlsx"
        copy2(self.template, province)
        def upload(paths, *, force):
            self.assertEqual(paths[0].name, province.name)
            self.assertNotEqual(paths[0].parent, province.parent)
            self.assertFalse(self.module.prepare_upload_files(paths)[1][0].builtin_template)
            self.update(paths)
        with mock.patch.object(self.module, "run_api_upload", side_effect=upload):
            run_group_upload_with_staging(self.module, [self.energy], force=False, on_log=lambda _: None)
        self.assert_output(self.folder / "广东-省内数据-20261023.xlsx")
        self.assertEqual(province.read_bytes(), self.original_template)

    def test_locked_generated_file_is_rejected_without_touching_template(self):
        (self.folder / "~$广东-省内数据-20261023.xlsx").write_bytes(b"lock")
        with self.assertRaisesRegex(RuntimeError, "正在被 Excel 打开"):
            self.update([self.energy])
        self.assertEqual(self.template.read_bytes(), self.original_template)

    def test_builtin_output_cannot_be_written_to_resource_directory(self):
        with self.assertRaisesRegex(ValueError, "必须指定输出目录"):
            self.module.province_output_path(self.template, "2026-10-23")

    def test_desktop_data_root_does_not_redirect_template(self):
        with mock.patch.dict(os.environ, {"TOOLBOX_CODE_ROOT": str(ROOT), "TOOLBOX_DATA_ROOT": str(self.folder)}):
            module = load_module("group_upload_builtin_desktop_test", ROOT / "集团每日上传/upload_daily_report.py")
        self.assertEqual(module.BUILTIN_PROVINCE_TEMPLATE, self.template)
        self.assertEqual(module.AUTH_STATE_PATH, self.folder / "集团每日上传/auth_state.json")
        self.assertEqual(hashlib.sha256(self.template.read_bytes()).digest(), hashlib.sha256(self.original_template).digest())
