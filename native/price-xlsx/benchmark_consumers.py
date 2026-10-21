"""Benchmark summary and standalone market-unit parsing without remote I/O."""
from __future__ import annotations

import argparse
from datetime import date
import hashlib
import json
import os
from pathlib import Path
import statistics
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "电力工具箱"), str(ROOT / "南方分公司电力市场日报")]
from toolbox.runtime import load_module
from gd_power_report.realtime_unit_xlsx import parse_realtime_unit_workbook
from openpyxl import load_workbook


def workbook_snapshot(path):
    book = load_workbook(path, data_only=False)
    try:
        return [dict(name=s.title, cells=[[(c.value, c.data_type, c.number_format, list(c._style) if c.has_style else None)
                                         for c in row] for row in s.iter_rows()],
                     merges=sorted(str(r) for r in s.merged_cells.ranges), freeze=s.freeze_panes,
                     columns={k: (d.width, d.hidden, d.min, d.max) for k, d in s.column_dimensions.items()},
                     rows={k: (d.height, d.hidden) for k, d in s.row_dimensions.items()},
                     fonts=[str(f) for f in book._fonts], fills=[str(f) for f in book._fills],
                     borders=[str(b) for b in book._borders], alignments=[str(a) for a in book._alignments])
                for s in book.worksheets]
    finally:
        book.close()


def benchmark(summary_dir, market_file, market_date, repeats):
    module = load_module("summary_native_benchmark", ROOT / "电力工具脚本/电量汇总工具/summarize_511_excel.py")
    report = {"repeats": repeats, "summary_input": summary_dir.name, "market_input": market_file.name,
              "market_date": market_date.isoformat(), "backends": {},
              "scope": "summary read/full local export; standalone realtime-unit parse; no network or PDF rendering"}
    references = {}
    previous = os.environ.get("TOOLBOX_XLSX_BACKEND")
    timings = {m: {k: [] for k in ("summary_read", "summary_full", "market_units")} for m in ("python", "rust")}
    try:
        with tempfile.TemporaryDirectory(prefix="native-consumers-") as tmp:
            for repeat in range(repeats):
                for mode in (("python", "rust") if repeat % 2 == 0 else ("rust", "python")):
                    os.environ["TOOLBOX_XLSX_BACKEND"] = mode
                    start = time.perf_counter()
                    units = module.collect_units(summary_dir, lambda _: None)
                    timings[mode]["summary_read"].append(time.perf_counter() - start)
                    start = time.perf_counter()
                    output = module.run_summary(summary_dir, Path(tmp) / f"{mode}.xlsx")
                    timings[mode]["summary_full"].append(time.perf_counter() - start)
                    snap = workbook_snapshot(output)
                    start = time.perf_counter()
                    payload = parse_realtime_unit_workbook(market_file, market_date)
                    timings[mode]["market_units"].append(time.perf_counter() - start)
                    for label, value in (("units", units), ("workbook", snap), ("market", payload)):
                        if label not in references:
                            references[label] = value
                        elif references[label] != value:
                            raise AssertionError(f"{mode}: {label} parity failed")
            for mode, stages in timings.items():
                report["backends"][mode] = {stage: dict(seconds=values, median=statistics.median(values))
                                           for stage, values in stages.items()}
            report["speedups"] = {stage: report["backends"]["python"][stage]["median"] /
                                  report["backends"]["rust"][stage]["median"] for stage in timings["python"]}
            report["parity"] = "unit objects, all output cell values/types/formats/styles, merges/dimensions/freeze panes, complete realtime payload identical"
            report["summary_units"] = len(references["units"])
            report["realtime_units"] = sum(len(plant["units"]) for plant in references["market"]["plants"])
            report["realtime_hours"] = sum(len(unit["hours"]) for plant in references["market"]["plants"] for unit in plant["units"])
            report["payload_sha256"] = hashlib.sha256(json.dumps(references["market"], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    finally:
        if previous is None:
            os.environ.pop("TOOLBOX_XLSX_BACKEND", None)
        else:
            os.environ["TOOLBOX_XLSX_BACKEND"] = previous
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--summary-dir", type=Path, required=True)
    parser.add_argument("--market-file", type=Path, required=True)
    parser.add_argument("--market-date", type=date.fromisoformat, required=True)
    parser.add_argument("--repeats", type=int, default=5)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    report = benchmark(args.summary_dir.resolve(), args.market_file.resolve(), args.market_date, args.repeats)
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
