"""Same-machine offline comparison of full and trimmed macOS installer runtimes.

Before uses the pre-optimization Python policy (default headless shell); after
uses the after package's actual browser policy. Node and headed Chrome are control groups.
No production data, credentials, installed app, or network requests are used.
"""
from __future__ import annotations

import argparse
import ctypes
import hashlib
import json
import os
from pathlib import Path
import platform
import random
import signal
import statistics
import subprocess
import sys
import tempfile
import time

CASES = ("imports", "ui_start", "excel", "forecast", "browser_headless",
         "browser_headed", "browser_parallel", "pdf", "node")


class UsageInfo(ctypes.Structure):
    # macOS SDK sys/resource.h: struct rusage_info_v2 (16-byte UUID + 18 u64).
    _fields_ = [('uuid', ctypes.c_uint8 * 16)] + [(name, ctypes.c_uint64) for name in (
        'user_time', 'system_time', 'pkg_idle_wkups', 'interrupt_wkups', 'pageins',
        'wired_size', 'resident_size', 'phys_footprint', 'proc_start_abstime',
        'proc_exit_abstime', 'child_user_time', 'child_system_time',
        'child_pkg_idle_wkups', 'child_interrupt_wkups', 'child_pageins',
        'child_elapsed_abstime', 'diskio_bytesread', 'diskio_byteswritten')]


def footprint_sample(pids, libproc):
    total = 0
    sampled = 0
    for pid in pids:
        usage = UsageInfo()
        if libproc.proc_pid_rusage(pid, 2, ctypes.byref(usage)) == 0:
            total += usage.phys_footprint
            sampled += 1
    return total / 1048576, sampled


def cpu_seconds(value):
    parts = value.split(":")
    return sum(float(part) * 60 ** i for i, part in enumerate(reversed(parts)))


def tree_sample(rows, root):
    selected = {root}
    while True:
        children = {pid for pid, ppid, *_ in rows if ppid in selected}
        if children <= selected:
            break
        selected |= children
    members = [(pid, rss, cpu) for pid, _, rss, cpu in rows if pid in selected]
    return sum(rss for _, rss, _ in members) / 1024, {pid: cpu for pid, _, cpu in members}


def median_change(before, after):
    return (statistics.median(after) / statistics.median(before) - 1) * 100


def bootstrap_change(before, after):
    rng = random.Random(20261028)
    estimates = []
    for _ in range(4000):
        indices = [rng.randrange(len(before)) for _ in before]
        estimates.append(median_change([before[i] for i in indices], [after[i] for i in indices]))
    estimates.sort()
    return [estimates[99], estimates[3899]]


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()


def worker(args):
    source = args.app / "Contents/Resources/payload/app"
    sys.path[:0] = [str(source), str(source / "电力工具箱"), str(source / "广东电价预测")]
    metrics = {}
    backend = None
    def measure(name, callback):
        start = time.perf_counter()
        value = callback()
        metrics[name] = time.perf_counter() - start
        return value
    def phase(name):
        args.result.with_suffix('.phase').write_text(name)
    start = time.perf_counter()
    case = args.case
    before = args.worker == "before"
    if case == "imports":
        import openpyxl, pandas, numpy, sklearn, scipy, tkinter, PIL, pdfplumber, pypdfium2, docx
        metrics['operation_s'] = time.perf_counter() - start
        parity = [numpy.__version__, pandas.__version__, sklearn.__version__, scipy.__version__]
    elif case == "ui_start":
        from toolbox.app import ToolboxApp
        app = measure('operation_s', lambda: ToolboxApp(workspace=source))
        try:
            app.withdraw()
            app.update_idletasks()
            metrics['ready_s'] = time.perf_counter() - start
            parity = [app.current_page, len(app.nav_buttons)]
            phase('idle')
            time.sleep(1)
            phase('active')
        finally:
            app.destroy()
    elif case == "excel":
        from openpyxl import Workbook, load_workbook
        from datetime import datetime
        target = args.result.parent / 'synthetic.xlsx'
        def excel():
            book = Workbook(write_only=True)
            sheet = book.create_sheet('对比数据')
            sheet.append(['日期', *[f'值{i}' for i in range(11)]])
            for row in range(1, 10001):
                sheet.append([datetime(2026, 1, 1), *[row * 10 + col for col in range(11)]])
            book.save(target)
            book.close()
            loaded = load_workbook(target, read_only=True, data_only=True)
            try:
                values = list(loaded['对比数据'].iter_rows(values_only=True))
                assert len(values) == 10001 and values[-1][-1] == 100010
                assert isinstance(values[1][0], datetime)
                return [len(values), values[-1][-1], loaded['对比数据']['A2'].number_format]
            finally:
                loaded.close()
        parity = measure('operation_s', excel)
    elif case == "forecast":
        import numpy as np
        import pandas as pd
        from gd_price_forecast.modeling import fit_model, predict_model
        rng = np.random.default_rng(42)
        X = pd.DataFrame(rng.normal(size=(12000, 24)), columns=[f'x{i}' for i in range(24)])
        y = X['x0'] * 25 + X['x1'] ** 2 * 10 + 300
        X.iloc[::11, 3] = np.nan
        def forecast():
            model = fit_model(X.iloc[:-96], y.iloc[:-96], {'kind': 'hist_gradient_boosting',
                'params': {'max_iter': 60, 'max_leaf_nodes': 15}})
            return predict_model(model, X.iloc[-96:])
        prediction = measure('operation_s', forecast)
        assert prediction.shape == (96,) and np.isfinite(prediction).all()
        parity = prediction.round(8).tolist()
    elif case == "node":
        backend = "chromium"
        module = source / '电力工具脚本/private-data-uploader-tool/node_modules/playwright/index.mjs'
        script = '''
import fs from 'node:fs';
const {chromium} = await import(process.argv[1]);
const t = performance.now();
const context = await chromium.launchPersistentContext(process.argv[2],
  {headless:true, executablePath:chromium.executablePath(), viewport:{width:1280,height:900}});
const launch_s=(performance.now()-t)/1000;
try {
 const page = await context.newPage();
 await page.setContent('<h1>offline benchmark</h1><div id="rows"></div><script>rows.innerHTML=Array.from({length:96},(_,i)=>`<p>${i}</p>`).join("")</script>');
 if (await page.locator('p').count() !== 96) throw Error('row mismatch');
 fs.writeFileSync(process.argv[3], 'idle');
 await new Promise(resolve=>setTimeout(resolve,1000));
 fs.writeFileSync(process.argv[3], 'active');
 const s=performance.now();
 for(let i=0;i<3;i++) await page.screenshot({path:process.argv[2]+`/shot${i}.png`});
 console.log(JSON.stringify({launch_s,screenshot_s:(performance.now()-s)/1000}));
} finally {await context.close();}
'''
        output = subprocess.check_output([str(args.app / 'Contents/Resources/payload/runtime/node/bin/node'),
            '--input-type=module', '-e', script, module.as_uri(), str(args.result.parent / 'node-profile'),
            str(args.result.with_suffix('.phase'))], text=True, timeout=60)
        metrics.update(json.loads(output))
        parity = {'rows': 96, 'screenshots': 3}
    else:
        from playwright.sync_api import sync_playwright
        from toolbox_shared.browser_runtime import chromium_launch_options
        options = {} if before else chromium_launch_options()
        headed = case == 'browser_headed'
        backend = 'chromium' if headed or options.get('channel') == 'chromium' else 'chromium-headless-shell'
        contexts = []
        with sync_playwright() as sdk:
            t = time.perf_counter()
            if case == 'pdf':
                browser = sdk.chromium.launch(headless=True, **options)
                contexts.append(browser.new_context(viewport={'width':1280,'height':900}, device_scale_factor=1))
            else:
                browser = None
                for i in range(2 if case == 'browser_parallel' else 1):
                    contexts.append(sdk.chromium.launch_persistent_context(
                        str(args.result.parent / f'profile-{i}'), headless=not headed,
                        viewport={'width':1280,'height':900}, device_scale_factor=1, **options))
            metrics['launch_s'] = time.perf_counter() - t
            try:
                pages = []
                html = '<meta charset="utf-8"><style>body{font:14px Arial}table{border-collapse:collapse}td{border:1px solid #888;padding:4px}</style><h1>OFFLINE BENCHMARK</h1><table>'
                html += ''.join(f'<tr><td>{i}</td><td>{i*10}</td><td>{i*20}</td></tr>' for i in range(96)) + '</table>'
                for context in contexts:
                    page = context.new_page()
                    page.set_content(html)
                    assert page.locator('tr').count() == 96
                    pages.append(page)
                phase('idle')
                time.sleep(1)
                phase('active')
                def screenshots():
                    hashes = []
                    for index, page in enumerate(pages):
                        for j in range(3):
                            content = page.screenshot()
                            hashes.append(hashlib.sha256(content).hexdigest())
                    return hashes
                images = measure('screenshot_s', screenshots)
                parity = {'rows':96, 'contexts':len(contexts), 'screenshots':images}
                if case == 'pdf':
                    content = measure('pdf_s', lambda: pages[0].pdf(format='A4'))
                    import pypdfium2 as pdfium
                    with pdfium.PdfDocument(content) as document:
                        texts = []
                        for page in document:
                            textpage = page.get_textpage()
                            try:
                                texts.append(textpage.get_text_range())
                            finally:
                                textpage.close()
                                page.close()
                        text = ''.join(texts).replace('\r','')
                        assert 'OFFLINE BENCHMARK' in text and '1900' in text
                        parity['pdf_pages'] = len(document)
                        parity['pdf_text'] = text
            finally:
                for context in contexts:
                    context.close()
                if browser:
                    browser.close()
    metrics['worker_s'] = time.perf_counter() - start
    args.result.write_text(json.dumps({'metrics': metrics, 'parity': parity, 'browser_backend': backend}, ensure_ascii=False))


def execute(args, app, variant, case, folder):
    folder.mkdir(parents=True)
    result = folder / 'result.json'
    runtime = app / 'Contents/Resources/payload/runtime'
    env = {**os.environ, 'TOOLBOX_CODE_ROOT':str(app / 'Contents/Resources/payload/app'),
           'TOOLBOX_DATA_ROOT': str(folder / 'data'), 'TOOLBOX_PORTABLE':'1',
           'PLAYWRIGHT_BROWSERS_PATH':str(runtime / 'ms-playwright'),
           'PYTHONDONTWRITEBYTECODE':'1', 'PYTHONNOUSERSITE':'1',
           'OMP_NUM_THREADS':'4', 'OPENBLAS_NUM_THREADS':'4', 'VECLIB_MAXIMUM_THREADS':'4'}
    for key in ('PYTHONPATH','PYTHONHOME','__PYVENV_LAUNCHER__'):
        env.pop(key, None)
    command = [str(runtime / 'python/bin/python3'), '-I', '-B']
    if args.isolate_bytecode_cache:
        command += ['-X', f'pycache_prefix={folder / "empty-pycache"}']
    command += [str(Path(__file__).resolve()),
               '--worker',variant,'--app',str(app),'--case',case,'--result',str(result)]
    libproc = ctypes.CDLL('/usr/lib/libproc.dylib')
    libproc.proc_pid_rusage.argtypes = [ctypes.c_int, ctypes.c_int, ctypes.c_void_p]
    libproc.proc_pid_rusage.restype = ctypes.c_int
    peak_footprint = 0
    idle_footprint = []
    foot_ok = foot_queries = 0
    peak = 0
    idle = []
    cpus = {}
    start = time.perf_counter()
    with (folder / 'stdout.log').open('w') as stdout, (folder / 'stderr.log').open('w') as stderr:
        process = subprocess.Popen(command, stdout=stdout, stderr=stderr, cwd=folder,
                                   env=env, start_new_session=True)
        try:
            while process.poll() is None:
                sample_start = time.perf_counter()
                raw = subprocess.check_output(['/bin/ps','-axo','pid=,ppid=,rss=,time='], text=True)
                rows = [(int(p),int(pp),int(rss),cpu_seconds(cpu)) for p,pp,rss,cpu in
                        (line.split() for line in raw.splitlines())]
                rss, usage = tree_sample(rows, process.pid)
                peak = max(peak, rss)
                footprint, sampled = footprint_sample(usage, libproc)
                peak_footprint = max(peak_footprint, footprint)
                foot_ok += sampled
                foot_queries += len(usage)
                for pid, cpu in usage.items():
                    cpus[pid] = max(cpus.get(pid,0), cpu)
                if result.with_suffix('.phase').exists() and result.with_suffix('.phase').read_text() == 'idle':
                    idle.append(rss)
                    idle_footprint.append(footprint)
                if time.perf_counter() - start > args.timeout:
                    raise TimeoutError(f'{variant}/{case} exceeded {args.timeout}s')
                time.sleep(max(0, 0.05 - (time.perf_counter()-sample_start)))
            if process.returncode:
                raise RuntimeError(f'{variant}/{case} exited {process.returncode}: {(folder / "stderr.log").read_text()[-3000:]}')
        except BaseException:
            try:
                os.killpg(process.pid, signal.SIGTERM)
                process.wait(timeout=5)
            except (ProcessLookupError, subprocess.TimeoutExpired):
                if process.poll() is None:
                    os.killpg(process.pid, signal.SIGKILL)
                    process.wait()
            raise
    value = json.loads(result.read_text())
    value['metrics'].update(wall_s=time.perf_counter()-start, peak_tree_rss_mib=peak,
                            sampled_tree_cpu_s=sum(cpus.values()), peak_tree_footprint_mib=peak_footprint,
                            footprint_query_coverage=foot_ok / max(1, foot_queries))
    if idle:
        value['metrics']['idle_tree_rss_mib'] = statistics.median(idle)
        value['metrics']['idle_tree_footprint_mib'] = statistics.median(idle_footprint)
    return value


def summarize(records):
    summary = {}
    for case in sorted({r['case'] for r in records if not r['warmup']}):
        samples = {v: [r for r in records if r['case']==case and r['variant']==v and not r['warmup']]
                   for v in ('before','after')}
        reference = samples['before'][0]['parity']
        if any(r['parity'] != reference for runs in samples.values() for r in runs):
            raise ValueError(f'Output parity failed: {case}')
        metrics = {}
        for key in samples['before'][0]['metrics']:
            values = {v:[r['metrics'][key] for r in runs] for v,runs in samples.items()}
            metrics[key] = {v:dict(median=statistics.median(x), minimum=min(x), maximum=max(x)) for v,x in values.items()}
            metrics[key]['change_percent'] = median_change(values['before'], values['after'])
            metrics[key]['paired_bootstrap_95_percent'] = bootstrap_change(values['before'], values['after'])
        summary[case] = {'metrics':metrics, 'parity_sha256':digest(reference), 'runs_per_variant':len(samples['before'])}
    return summary


def run(args):
    args.output = args.output.resolve()
    apps = {'before':args.before.resolve(), 'after':args.after.resolve()}
    if args.repeats < 3:
        raise ValueError('At least 3 measured repeats required')
    for app in apps.values():
        if args.output.resolve().is_relative_to(app):
            raise ValueError('Output must be outside application bundles')
    args.output.mkdir(parents=True, exist_ok=False)
    manifests = [json.loads((app / 'Contents/Resources/payload/manifest.json').read_text()) for app in apps.values()]
    source_differences = [name for name in sorted(set(manifests[0]['files']) | set(manifests[1]['files']))
                          if manifests[0]['files'].get(name) != manifests[1]['files'].get(name)]
    if set(source_differences) - {'toolbox_shared/browser_runtime.py'}:
        raise ValueError('Compare the same business source files (only browser policy helper may differ)')
    report = {'metadata':{'platform':platform.platform(), 'python_controller':sys.version,
              'cpu':subprocess.check_output(['sysctl','-n','machdep.cpu.brand_string'],text=True).strip(),
              'memory_bytes':int(subprocess.check_output(['sysctl','-n','hw.memsize'],text=True)),
              'power':subprocess.check_output(['pmset','-g','batt'],text=True).strip(),
              'source_policy_differences':source_differences,
              'before':str(apps['before']), 'after':str(apps['after']), 'repeats':args.repeats,
              'warmups_per_case_variant':1, 'sample_interval_target_s':0.05, 'compute_threads':4,
              'isolate_bytecode_cache':args.isolate_bytecode_cache,
              'browser_policy':'before=Python default headless shell; after=actual packaged browser_runtime policy; see per-sample browser_backend; headed/Node=full Chromium',
              'timing_scope':'Fresh subprocess per sample, one discarded warmup, alternating before/after order. OS file cache is not purged; not a cold-boot benchmark.',
              'footprint_scope':'Sum of per-process ri_phys_footprint via macOS proc_pid_rusage v2; includes compressed/ledger charged memory, not whole-system memory pressure.',
              'rss_scope':'Sampled sum of worker and descendant RSS; shared pages can be counted more than once. Not physical memory pressure or total Tauri application memory.',
              'cpu_scope':'Sum of last observed cumulative CPU per PID; short-lived processes may be missed.',
              'limitations':'Synthetic offline workloads; no production website, true cold boot, battery/energy, swapping or low-RAM machine measured.'}, 'records':[]}
    def save():
        path = args.output / 'results.json'
        temporary = path.with_suffix('.tmp')
        temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2))
        temporary.replace(path)
    save()
    for repeat in range(-1,args.repeats):
        for case in args.cases:
            for variant in (('before','after') if repeat % 2 == 0 else ('after','before')):
                folder = args.output / f'{repeat+1:02d}-{case}-{variant}'
                load = os.getloadavg()
                value = execute(args, apps[variant], variant, case, folder)
                report['records'].append(dict(case=case,variant=variant,round=repeat,
                    warmup=repeat<0,load_before=load,**value))
                save()
                print(f'{repeat+1}/{args.repeats} {case} {variant}: {value["metrics"]["wall_s"]:.3f}s, {value["metrics"]["peak_tree_rss_mib"]:.1f}MiB',flush=True)
                # Keep only diagnostics and measured results, not test workbooks/profiles.
                for child in folder.iterdir():
                    if child.name not in ('result.json','result.phase','stdout.log','stderr.log'):
                        import shutil
                        if child.is_dir():
                            shutil.rmtree(child)
                        else:
                            child.unlink()
    report['summary'] = summarize(report['records'])
    save()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--before',type=Path)
    p.add_argument('--after',type=Path)
    p.add_argument('--output',type=Path)
    p.add_argument('--repeats',type=int,default=7)
    p.add_argument('--timeout',type=float,default=180)
    p.add_argument('--isolate-bytecode-cache',action='store_true',
                   help='Read bytecode only from a new empty per-worker prefix; writes remain disabled')
    p.add_argument('--cases',nargs='+',choices=CASES,default=CASES)
    p.add_argument('--worker',choices=('before','after'))
    p.add_argument('--app',type=Path)
    p.add_argument('--case',choices=CASES)
    p.add_argument('--result',type=Path)
    args = p.parse_args()
    if args.worker:
        worker(args)
    else:
        if not all((args.before,args.after,args.output)):
            p.error('--before, --after and --output are required')
        run(args)


if __name__ == '__main__':
    main()
