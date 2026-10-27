"""Optional bulk Rust XLSX I/O; no compiler or new dependency required to run.

TOOLBOX_500KV_BACKEND=auto|python|rust. Strict rust mode never falls back.
"""
from __future__ import annotations

import os
import sys
import subprocess
import tempfile
import warnings
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ERROR_CODES

ROOT = Path(__file__).resolve().parents[2]


def backend_mode() -> str:
    mode = os.environ.get("TOOLBOX_500KV_BACKEND", "auto")
    if mode not in {"auto", "python", "rust"}:
        raise ValueError(f"未知 500kV Excel 后端: {mode}")
    return mode


# Shared transport also serves summary and market-report readers.
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from toolbox_shared.native_excel import native_binary, call as _call


def _fallback_warning(operation, exc):
    detail = exc.stderr.strip() if isinstance(exc, subprocess.CalledProcessError) else str(exc)
    warnings.warn(f"500kV Rust Excel {operation}不可用，改用 Python: {detail}", RuntimeWarning)


class PriceSheet:
    def __init__(self, title, rows=None):
        self.title = title
        self.rows = [] if rows is None else rows

    def iter_rows(self, *, values_only=True):
        return iter(self.rows)

    def append(self, row):
        self.rows.append(list(row))


class PriceWorkbook:
    def __init__(self, binary, sheets=()):
        self.binary = binary
        self.worksheets = list(sheets)

    @property
    def sheetnames(self):
        return [sheet.title for sheet in self.worksheets]

    def create_sheet(self, title):
        sheet = PriceSheet(title)
        self.worksheets.append(sheet)
        return sheet

    def close(self):
        pass

    def save(self, path):
        path = Path(path)
        handle, temporary = tempfile.mkstemp(prefix=".500kv-", suffix=".xlsx", dir=path.parent)
        os.close(handle)
        try:
            # openpyxl treats these strings as formulas/error cells. Preserve
            # that behavior via the compatibility writer rather than text.
            compatible = all(
                not isinstance(value, str) or not (value.startswith("=") or value in ERROR_CODES)
                for sheet in self.worksheets for row in sheet.rows for value in row
            )
            try:
                if not compatible:
                    raise ValueError("公式或错误单元格需要 Python 兼容写入")
                response = _call(self.binary, "write", Path(temporary), sheets=[
                    # openpyxl serializes numbers with %.16g. Round before
                    # transport so readback and downstream thresholds match.
                    dict(name=sheet.title, rows=[
                        [float(format(value, ".16g")) if type(value) in (int, float) else value
                         for value in row] for row in sheet.rows
                    ]) for sheet in self.worksheets
                ])
                if response.get("written") is not True:
                    raise ValueError("Rust Excel 写入未确认")
            except (OSError, ValueError, subprocess.SubprocessError) as exc:
                if backend_mode() == "rust":
                    raise
                _fallback_warning("写入", exc)
                book = Workbook()
                book.remove(book.active)
                for sheet in self.worksheets:
                    ws = book.create_sheet(sheet.title)
                    for row in sheet.rows:
                        ws.append(row)
                book.save(temporary)
                book.close()
            os.replace(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)


def open_price_workbooks(path: Path):
    """Return matching input/output adapters, using Python for unsupported files."""
    mode = backend_mode()
    binary = native_binary() if mode != "python" else None
    if mode == "rust" and binary is None:
        raise RuntimeError("未安装 500kV Rust Excel 后端，请先运行 native/price-xlsx/build.py")
    if binary is not None:
        try:
            response = _call(binary, "read", path)
            if not isinstance(response.get("sheets"), list) or any(
                not isinstance(s, dict) or not isinstance(s.get("name"), str)
                or not isinstance(s.get("rows"), list)
                or any(not isinstance(row, list) for row in s["rows"])
                for s in response["sheets"]
            ):
                raise ValueError("500kV Rust Excel 返回格式不正确")
            source = PriceWorkbook(binary, [PriceSheet(s["name"], s["rows"]) for s in response["sheets"]])
            return source, PriceWorkbook(binary)
        except (OSError, ValueError, subprocess.SubprocessError) as exc:
            if mode == "rust":
                raise
            _fallback_warning("读取", exc)
    source = load_workbook(path, read_only=True, data_only=True)
    output = Workbook()
    output.remove(output.active)
    return source, output
