from __future__ import annotations
import os
from datetime import datetime, time, timedelta
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pandas.testing import assert_frame_equal
from toolbox.runtime import load_module
from toolbox_native_excel import native_binary
from toolbox_native_write import read_table_layout
ROOT = Path(__file__).resolve().parents[2]

@unittest.skipUnless(native_binary(), "optional Rust helper is not installed")
class NativeUpdateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.module = load_module("native_update_forecast", ROOT / "广东电价预测/update_gdfire_data.py")

    def snapshot(self, path):
        book = load_workbook(path)
        try:
            sheet = book.active
            return (sheet.title, sheet.freeze_panes, sheet.max_row, sheet.max_column,
                    tuple((key, value.width) for key, value in sheet.column_dimensions.items()),
                    tuple(tuple((cell.value, cell.data_type, cell.number_format,
                                 str(cell.font), str(cell.fill), str(cell.border),
                                 str(cell.alignment), str(cell.protection)) for cell in row)
                          for row in sheet))
        finally:
            book.close()

    def test_streaming_output_matches_python_for_excel_types_and_styles(self):
        frame = pd.DataFrame({
            "日期": [datetime(2026, 10, 20), None, datetime(2026, 10, 22)],
            "比例": [0.12345678901234567, float("nan"), -2.75],
            "文本": [" 空格\r\n<&>中文\x01", "=SUM(B2:B3)", "#DIV/0!"],
            "布尔": [True, False, None], "时间": [time(12, 15), timedelta(hours=6), None],
            "整数": [9007199254740993, 42, None],
        })
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory); source = folder / "template.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "其他名称"
            sheet.append(list(frame.columns)); sheet.append([datetime(2026, 1, 1), .25, "示例", True, time(0), 1])
            for row in sheet:
                for cell in row:
                    cell.font = Font(name="Arial", bold=cell.row == 1, color="123456")
                    cell.fill = PatternFill("solid", fgColor="FFEEDD")
                    cell.alignment = Alignment(horizontal="center", wrap_text=True)
                    cell.border = Border(bottom=Side(style="thin"))
            sheet["B2"].number_format = "0.00%"; sheet["A2"].number_format = "yyyy-mm-dd"
            sheet.column_dimensions["B"].width = 24.5
            book.save(source); book.close(); original = source.read_bytes(); outputs = []
            for mode in ("python", "rust"):
                target = folder / f"{mode}.xlsx"
                with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND=mode):
                    self.module.write_large_xlsx(frame, target, "整合数据", format_source=source)
                outputs.append(self.snapshot(target))
            self.assertEqual(outputs[0], outputs[1]); self.assertEqual(source.read_bytes(), original)
            self.assertEqual(outputs[1][1:4], ("A2", 4, 6))

    def test_native_reader_preserves_dataframe_contract(self):
        frame = pd.DataFrame({"日期": [datetime(2026, 10, 20)], "时刻": ["00:00"], "值": [.125]})
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"; frame.to_excel(path, sheet_name="备用表", index=False)
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="python"):
                expected = self.module.read_data_workbook(path)
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
                actual = self.module.read_data_workbook(path)
            assert_frame_equal(expected, actual, check_exact=True)

    def test_layout_only_scans_first_two_rows(self):
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"; book = Workbook(); sheet = book.active
            sheet["A1"] = "标题"; sheet["B1"] = "空列"; sheet["A5000"] = 42
            sheet.column_dimensions["A"].width = 31; book.save(path); book.close()
            original = ReadOnlyWorksheet.iter_rows
            with patch.object(ReadOnlyWorksheet, "iter_rows", autospec=True, side_effect=original) as scan:
                layout = read_table_layout(path)
            self.assertEqual(scan.call_count, 1)
            self.assertEqual(scan.call_args.kwargs, {"min_row": 1, "max_row": 2})
            self.assertEqual(layout["标题"]["width"], 31)
            self.assertEqual(layout["空列"]["data_style"]["number_format"], "General")

    def test_native_failure_preserves_destination_and_auto_falls_back(self):
        import toolbox_native_write as writer
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "output.xlsx"; path.write_bytes(b"existing-output")
            frame = pd.DataFrame({"值": [1.25]})
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"), patch.object(writer, "call", side_effect=ValueError("native failed")):
                with self.assertRaisesRegex(ValueError, "native failed"):
                    self.module.write_large_xlsx(frame, path, "整合数据")
            self.assertEqual(path.read_bytes(), b"existing-output")
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="auto"), patch.object(writer, "call", side_effect=ValueError("native failed")):
                with self.assertWarns(RuntimeWarning):
                    self.module.write_large_xlsx(frame, path, "整合数据")
            book = load_workbook(path); self.assertEqual(book.active["A2"].value, 1.25); book.close()

    def test_layout_handles_missing_dimensions_and_short_data_row(self):
        from zipfile import ZipFile, ZIP_DEFLATED
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            target = Path(directory) / "undimensioned.xlsx"
            book = Workbook(); sheet = book.active
            sheet.append(["日期", "价格", "无数据列"]); sheet.append([datetime(2026, 10, 20), 42])
            book.save(source); book.close()
            with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as output:
                for info in original.infolist():
                    value = original.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        import re
                        value = re.sub(rb"<dimension[^>]*/>", b"", value)
                    output.writestr(info, value)
            layout = read_table_layout(target)
            self.assertEqual(list(layout), ["日期", "价格", "无数据列"])
            self.assertEqual(layout["无数据列"]["data_style"]["number_format"], "General")

    def test_auto_preserves_faster_streaming_reader_when_requested(self):
        import toolbox_native_excel as reader
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            book = Workbook(); book.active["A1"] = .125; book.save(path); book.close()
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="auto"), patch.object(reader, "native_binary", side_effect=AssertionError("should retain Python streaming")):
                book = reader.load_value_workbook(path, native_in_auto=False)
                try:
                    self.assertEqual(book.active["A1"].value, .125)
                finally:
                    book.close()

    def test_layout_avoids_full_scan_when_dimensions_are_missing(self):
        from zipfile import ZipFile, ZipExtFile, ZIP_DEFLATED
        import re
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            target = Path(directory) / "no_dimensions.xlsx"
            book = Workbook(); sheet = book.active
            sheet.append(["日期", "值"])
            for i in range(5000):
                sheet.append([datetime(2026, 10, 20), i])
            book.save(source); book.close()
            with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as output:
                for info in original.infolist():
                    value = original.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        value = re.sub(rb"<dimension[^>]*/>", b"", value)
                        self.assertGreater(len(value), 200000)
                    output.writestr(info, value)
            consumed = []
            original_read = ZipExtFile.read
            def read(stream, *args, **kwargs):
                value = original_read(stream, *args, **kwargs)
                if stream.name == "xl/worksheets/sheet1.xml":
                    consumed.append(len(value))
                return value
            with patch.object(ZipExtFile, "read", autospec=True, side_effect=read):
                layout = read_table_layout(target)
            self.assertEqual(list(layout), ["日期", "值"])
            self.assertLess(sum(consumed), 65536)

    def test_prepared_frame_preserves_object_numeric_and_error_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "整合数据"
            sheet.append(["混合", "错误", "时间", "大整数"])
            sheet.append([1.0, "#DIV/0!", time(12, 15), 9007199254740993])
            sheet.append(["文本", "正常", timedelta(hours=27), -2.5])
            sheet.append([True, None, None, 0])
            sheet["C2"].number_format = "hh:mm:ss"
            sheet["C3"].number_format = "[h]:mm:ss"
            book.save(path); book.close()
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="python"):
                expected = self.module.read_data_workbook(path)
            with patch.dict(os.environ, TOOLBOX_XLSX_BACKEND="rust"):
                actual = self.module.read_data_workbook(path)
            assert_frame_equal(expected, actual, check_exact=True)
            self.assertEqual([type(v) for v in expected["混合"]], [type(v) for v in actual["混合"]])
