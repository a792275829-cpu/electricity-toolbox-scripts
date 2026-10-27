"""Offline coverage for the clearing workbook parser's header contract.

The workbook layout is the fragile part: metrics are located by merged parent
headers plus child labels, so a shifted column silently produces wrong numbers
instead of an error. These tests pin that contract with in-memory workbooks and
never touch real business files, credentials or the network.
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

TOOL_DIR = Path(__file__).resolve().parents[1]
ROOT = TOOL_DIR.parent
for entry in (str(ROOT), str(TOOL_DIR)):
    if entry not in sys.path:
        sys.path.insert(0, entry)

import openpyxl  # noqa: E402

from gd_power_report import company_clearing_xlsx as MODULE  # noqa: E402


def report_sheet(title: str = MODULE.SHEET_NAME):
    """A sheet whose shape passes the day-report detection check."""

    sheet = openpyxl.Workbook().active
    sheet.title = title
    sheet["A4"] = "日前市场简析：竞价空间回落，价格小幅下行。"
    sheet["A68"] = "实时市场简析：晚高峰价格上行。"
    sheet["Q4"] = "满足 17 列要求"
    return sheet


def recap_sheet(
    *,
    energy: float | None = 1000.0,
    fee: float | None = 400000.0,
    capacity: float | None = 20000.0,
    capacity_standalone: bool = False,
    summary_label: str = "分公司火电",
) -> "openpyxl.worksheet.worksheet.Worksheet":
    """Header rows 2/3 plus a company recap row, matching the real layout."""

    sheet = openpyxl.Workbook().active
    sheet.title = MODULE.SHEET_NAME
    sheet["A4"] = "日前市场简析"
    sheet["B2"] = "日前出清"
    sheet.merge_cells("B2:C2")
    sheet["B3"] = "电量"
    sheet["C3"] = "电价"
    sheet["D2"] = "当日合计"
    sheet.merge_cells("D2:E2")
    sheet["D3"] = "电能电费"
    if capacity_standalone:
        # A lone "容量电费" header without a parent group; the recap row still
        # carries the fee, so the fallback path must supply the capacity part.
        sheet["F3"] = "容量电费"
        sheet["F5"] = capacity
    else:
        sheet["E3"] = "容量电费"
        sheet["E5"] = capacity

    sheet["A5"] = summary_label
    sheet["B5"] = energy
    sheet["C5"] = 300.0
    sheet["D5"] = fee
    return sheet


class NumberHelperTests(unittest.TestCase):
    def test_numbers_accept_numeric_text_and_reject_noise(self):
        self.assertEqual(MODULE._number(3), 3.0)
        self.assertEqual(MODULE._number("2.5"), 2.5)
        for value in (None, True, False, "abc", "", "#N/A", "#DIV/0!"):
            self.assertIsNone(MODULE._number(value), value)

    def test_meaningful_number_drops_zero_and_near_zero(self):
        self.assertIsNone(MODULE._meaningful_number(0))
        self.assertIsNone(MODULE._meaningful_number(1e-12))
        self.assertEqual(MODULE._meaningful_number(-2), -2.0)

    def test_difference_needs_both_sides(self):
        self.assertEqual(MODULE._difference(5, 3), 2)
        self.assertIsNone(MODULE._difference(None, 3))
        self.assertIsNone(MODULE._difference(5, None))


class DateHelperTests(unittest.TestCase):
    def test_as_date_reads_the_three_supported_text_forms(self):
        self.assertEqual(MODULE._as_date(datetime(2026, 10, 11, 8, 0)), date(2026, 10, 11))
        self.assertEqual(MODULE._as_date(date(2026, 10, 11)), date(2026, 10, 11))
        for text in ("2026-10-11", "2026/10/11", "2026.10.11", " 2026-10-11 "):
            self.assertEqual(MODULE._as_date(text), date(2026, 10, 11), text)
        self.assertIsNone(MODULE._as_date("10月11日"))

    def test_date_label_names_the_weekday(self):
        self.assertEqual(MODULE._date_label(date(2026, 10, 11)), "10月11日（周日）")
        self.assertEqual(MODULE._date_label(date(2026, 1, 1)), "1月1日（周四）")


class SheetSelectionTests(unittest.TestCase):
    def test_preferred_sheet_must_exist(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = MODULE.SHEET_NAME
        with self.assertRaisesRegex(ValueError, "缺少工作表：不存在"):
            MODULE._select_report_sheet(workbook, "不存在")

    def test_named_sheet_wins_without_shape_check(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "无关"
        workbook.create_sheet(MODULE.SHEET_NAME)
        self.assertEqual(
            MODULE._select_report_sheet(workbook).title, MODULE.SHEET_NAME
        )

    def test_alias_sheet_needs_the_report_shape(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = MODULE.ENERGY_METRICS_SHEET_NAME
        workbook.active["A1"] = "只是汇总表"
        with self.assertRaisesRegex(ValueError, "缺少可识别的日报工作表"):
            MODULE._select_report_sheet(workbook)

    def test_fallback_picks_the_highest_versioned_snapshot(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "说明"
        for name in ("快报 (2)", "快报 (3)", "快报 (10)"):
            sheet = workbook.create_sheet(name)
            sheet["A4"] = "日前市场简析"
            sheet["A70"] = "实时市场简析"
            sheet["Q4"] = "占位"
        self.assertEqual(MODULE._select_report_sheet(workbook).title, "快报 (10)")

    def test_fallback_ignores_sheets_without_the_report_shape(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "说明"
        sheet = workbook.create_sheet("快报 (1)")
        sheet["A4"] = "别的表"
        with self.assertRaises(ValueError):
            MODULE._select_report_sheet(workbook)


class WorkbookFileTests(unittest.TestCase):
    def test_report_date_comes_from_a6_then_a2_then_filename(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = MODULE.SHEET_NAME
            sheet["A4"] = "日前市场简析"
            sheet["Q4"] = "占位"
            sheet["A6"] = datetime(2026, 10, 11)
            path = base / "出清情况.xlsx"
            workbook.save(path)
            self.assertEqual(MODULE.infer_report_date(path), date(2026, 10, 11))

            sheet["A6"] = None
            sheet["A2"] = "2026/10/12"
            workbook.save(path)
            self.assertEqual(MODULE.infer_report_date(path), date(2026, 10, 12))

            sheet["A2"] = None
            workbook.save(path)
            renamed = base / "分公司2026.10.13出清情况.xlsx"
            path.rename(renamed)
            self.assertEqual(MODULE.infer_report_date(renamed), date(2026, 10, 13))

    def test_unreadable_date_lists_no_guess(self):
        with tempfile.TemporaryDirectory() as folder:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = MODULE.SHEET_NAME
            sheet["A4"] = "日前市场简析"
            sheet["Q4"] = "占位"
            path = Path(folder) / "出清情况.xlsx"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "无法从“新快报”工作表识别报告日期"):
                MODULE.infer_report_date(path)

    def test_only_xlsx_files_are_accepted(self):
        with tempfile.TemporaryDirectory() as folder:
            wrong = Path(folder) / "出清情况.xls"
            wrong.write_bytes(b"x")
            with self.assertRaisesRegex(ValueError, "不存在或格式不正确"):
                MODULE.validate_report_workbook(wrong)
            with self.assertRaisesRegex(ValueError, "不存在或格式不正确"):
                MODULE.validate_report_workbook(Path(folder) / "无.xlsx")


class SummaryRowTests(unittest.TestCase):
    def test_company_labels_are_recognised(self):
        for label in ("分公司火电", "分公司", "火电机组合计"):
            sheet = openpyxl.Workbook().active
            sheet["A7"] = label
            self.assertEqual(MODULE._find_company_summary_row(sheet), 7, label)

    def test_new_energy_row_stops_the_scan(self):
        sheet = openpyxl.Workbook().active
        sheet["A3"] = "新能源"
        sheet["A9"] = "分公司"
        self.assertIsNone(MODULE._find_company_summary_row(sheet))

    def test_missing_company_row_returns_none(self):
        sheet = openpyxl.Workbook().active
        sheet["A3"] = "统调"
        self.assertIsNone(MODULE._find_company_summary_row(sheet))


class HeaderHierarchyTests(unittest.TestCase):
    def build(self):
        sheet = openpyxl.Workbook().active
        sheet["B2"] = "日前出清"
        sheet.merge_cells("B2:C2")
        sheet["B3"] = "电量"
        sheet["C3"] = "电价"
        sheet["D2"] = "日前出清合计"
        sheet["D3"] = "电量"
        return sheet

    def test_merged_parent_locates_the_child_column(self):
        sheet = self.build()
        self.assertEqual(
            MODULE._find_header_column_by_hierarchy(
                sheet, "日前出清", ("电量",), before_row=5
            ),
            2,
        )
        self.assertEqual(
            MODULE._find_header_column_by_hierarchy(
                sheet, "日前出清", ("电价",), before_row=5
            ),
            3,
        )

    def test_exact_parent_skips_similar_group_names(self):
        sheet = self.build()
        self.assertEqual(
            MODULE._find_header_column_by_hierarchy(
                sheet, "日前出清", ("电量",), before_row=5, exact_parent=True
            ),
            2,
        )
        self.assertEqual(
            MODULE._find_header_column_by_hierarchy(
                sheet, "日前出清合计", ("电量",), before_row=5, exact_parent=True
            ),
            4,
        )

    def test_unknown_labels_return_none(self):
        sheet = self.build()
        self.assertIsNone(
            MODULE._find_header_column_by_hierarchy(
                sheet, "不存在", ("电量",), before_row=5
            )
        )
        self.assertIsNone(
            MODULE._find_header_column_by_hierarchy(
                sheet, "日前出清", ("电量",), before_row=2
            )
        )


class DayAheadMetricsTests(unittest.TestCase):
    def test_price_is_derived_from_fee_and_clearing_energy(self):
        metrics = MODULE._day_ahead_summary_metrics(recap_sheet())
        self.assertEqual(metrics["energy"], 1000.0)
        self.assertEqual(metrics["price"], 300.0)
        self.assertEqual(metrics["electricity_fee"], 400000.0)
        self.assertEqual(metrics["capacity_fee"], 20000.0)
        self.assertEqual(metrics["long_term_energy_price"], 400000.0 / 1000.0 * 1000)
        self.assertEqual(
            metrics["capacity_energy_price"], (400000.0 + 20000.0) / 1000.0 * 1000
        )

    def test_zero_clearing_energy_never_divides(self):
        metrics = MODULE._day_ahead_summary_metrics(recap_sheet(energy=0))
        self.assertIsNone(metrics["long_term_energy_price"])
        self.assertIsNone(metrics["capacity_energy_price"])

    def test_missing_fee_leaves_prices_empty(self):
        metrics = MODULE._day_ahead_summary_metrics(recap_sheet(fee=None))
        self.assertIsNone(metrics["long_term_energy_price"])
        self.assertIsNone(metrics["capacity_energy_price"])

    def test_standalone_capacity_column_is_used_as_fallback(self):
        metrics = MODULE._day_ahead_summary_metrics(
            recap_sheet(capacity_standalone=True)
        )
        self.assertEqual(metrics["capacity_fee"], 20000.0)
        self.assertEqual(
            metrics["capacity_energy_price"], (400000.0 + 20000.0) / 1000.0 * 1000
        )

    def test_missing_company_row_returns_the_full_empty_shape(self):
        sheet = openpyxl.Workbook().active
        sheet["A2"] = "统调"
        metrics = MODULE._day_ahead_summary_metrics(sheet)
        self.assertEqual(
            set(metrics),
            {
                "energy", "price", "long_term_energy_price", "capacity_energy_price",
                "electricity_fee", "capacity_fee", "deviation", "margin",
                "long_term_energy", "long_term_margin", "day_total_energy_price",
                "day_total_capacity_price",
            },
        )
        self.assertTrue(all(value is None for value in metrics.values()))


class EnergyComparisonTests(unittest.TestCase):
    def build_realtime(
        self,
        price_one: float | None,
        price_two: float | None,
        *,
        energy_one: float = 100.0,
        energy_two: float = 300.0,
    ):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = MODULE.ENERGY_METRICS_SHEET_NAME
        sheet["F72"], sheet["P72"] = energy_one, price_one
        sheet["F81"], sheet["P81"] = energy_two, price_two
        return workbook

    def test_realtime_price_is_energy_weighted(self):
        metrics = MODULE._energy_comparison_metrics(self.build_realtime(200.0, 400.0))
        self.assertEqual(metrics["realtime_price"], 350.0)
        self.assertEqual(metrics["realtime_fee"], (100 * 200 + 300 * 400) / 1000)
        self.assertIsNone(metrics["day_ahead_price"])

    def test_missing_leg_price_drops_the_realtime_average(self):
        metrics = MODULE._energy_comparison_metrics(self.build_realtime(None, 400.0))
        self.assertIsNone(metrics["realtime_price"])
        self.assertIsNone(metrics["realtime_fee"])

    def test_zero_price_leg_still_counts_its_energy(self):
        # A 0 元/兆瓦时 leg contributes no fee but is real volume, so it lowers
        # the weighted average instead of being dropped.
        metrics = MODULE._energy_comparison_metrics(self.build_realtime(0.0, 400.0))
        self.assertEqual(metrics["realtime_price"], 300.0)

    def test_zero_energy_leg_is_skipped(self):
        metrics = MODULE._energy_comparison_metrics(
            self.build_realtime(None, 400.0, energy_one=0.0)
        )
        self.assertEqual(metrics["realtime_price"], 400.0)


class SeriesTests(unittest.TestCase):
    def build(self, hours: int) -> "openpyxl.worksheet.worksheet.Worksheet":
        sheet = openpyxl.Workbook().active
        for index in range(hours):
            sheet.cell(2 + index, 1).value = f"{index}:00"
            sheet.cell(2 + index, 2).value = 10.0 + index
            sheet.cell(2 + index, 3).value = 300.0 + index
        return sheet

    def test_every_hour_maps_to_its_row(self):
        series = MODULE._series(self.build(24), 2, 3)
        self.assertTrue(series["ready"])
        self.assertEqual(series["labels"][0], "0:00")
        self.assertEqual(series["labels"][23], "23:00")
        self.assertEqual(series["energy"][5], 15.0)
        self.assertEqual(series["price"][23], 323.0)

    def test_missing_hour_stays_empty(self):
        series = MODULE._series(self.build(10), 2, 3)
        self.assertEqual(series["energy"][9], 19.0)
        self.assertIsNone(series["energy"][10])
        self.assertIsNone(series["price"][23])

    def test_all_zero_energy_clears_the_series(self):
        sheet = self.build(24)
        for index in range(24):
            sheet.cell(2 + index, 2).value = 0
        series = MODULE._series(sheet, 2, 3)
        self.assertFalse(series["ready"])
        self.assertTrue(all(value is None for value in series["energy"]))
        self.assertTrue(all(value is None for value in series["price"]))


if __name__ == "__main__":
    unittest.main()
